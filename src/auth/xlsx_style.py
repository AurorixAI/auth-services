"""
ALGO - the house style of every Excel export (decision of the owner, 23.10.2026; spec ``artifacts/partner_doc_check/xlsx_house_style.md``).

One palette, one font, one title block, one header, one total row, one print setup for the statement, the payments journal, the payroll sheets,
the openings file and the import template.  The model of the whole family is the statement: «нежно-голубое выделение важного» - the header and
the totals are D9E2F3, a section (month / group / teacher) is the lighter EDF1F7.

What the helpers promise (and ``tests/test_xlsx_style.py`` checks):

* a text cell is ALWAYS a string (``data_type`` "s"): a student's note that starts with "=" is never a formula; characters XML cannot hold are dropped;
* money is a whole number with ``#,##0``, a date is a date with ``DD.MM.YYYY``, a moment is a datetime with ``DD.MM.YYYY HH:MM`` - never text;
* the title block is rows 1-3 (merged over the whole width), row 4 is empty, the header of the table is row 5 (``HEADER_ROW``);
* widths come from the content (the longest value + 2), a long text is capped at ``WRAP_CAP`` and wraps with the row height set from the number of lines,
  a header wraps by words: nothing is cut and no «####» appears;
* print: landscape, fit to the page width, the header row repeated on every page, narrow margins;
* a file name is Russian without technical ids; the header carries an ASCII fallback and ``filename*=UTF-8''...``.
"""
from __future__ import annotations

import math
import re
from copy import copy
from datetime import date, datetime
from typing import Any, Dict, Iterable, List, Mapping, Optional, Sequence
from urllib.parse import quote

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

# ------------------------------------------------------------------------------------------------------------ palette

FONT_NAME = "Calibri"
HEAD_FILL = "D9E2F3"          # the header of a table and the «Итого» row
SECTION_FILL = "EDF1F7"       # a month / a group / a teacher (lighter than the header)
SERVICE_FILL = "E7E6E6"       # the quiet blocks below a table («Следующие периоды», «Отменённые записи»)
PENDING_FILL = "FFF2B3"       # waits for a decision
WARN_FILL = "FDE2E1"          # a debt / a warning, point-wise
EXAMPLE_FILL = "F2F2F2"       # the example row of a template
TOTAL_LINE = "8EA9DB"         # the line above «Итого»
GRID = "D9D9D9"               # thin borders of an ordinary row
CANCELLED_COLOR = "808080"    # a cancelled record: grey and struck through (the reason is grey but not struck)
NOTE_COLOR = "6B7280"         # a quiet explanation (the legend line, the hints)
TITLE_COLOR = "1F2937"        # the name of the file (row 1 of the title block)
SUBTITLE_COLOR = "4B5563"     # rows 2-3 of the title block
PALETTE = frozenset({HEAD_FILL, SECTION_FILL, SERVICE_FILL, PENDING_FILL, WARN_FILL, EXAMPLE_FILL, "FFFFFF"})

# ------------------------------------------------------------------------------------------------------------ formats

MONEY = "#,##0"
PERCENT = '0"%"'
PERCENT_FRACTION = '0.0#"%"'
DATE = "DD.MM.YYYY"
DATE_SHORT = "DD.MM"
DATETIME = "DD.MM.YYYY HH:MM"

HEADER_ROW = 5                # rows 1-3 the title block, row 4 empty
TITLE_ROWS = 3
MIN_WIDTH = 10
WRAP_CAP = 60
HEADER_HEIGHT = 30            # the header row: two lines high, as in the owner's mock
LINE_HEIGHT = 15              # points of one line of Calibri 11
SPREADSHEET_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# not allowed in a worksheet: openpyxl refuses \x00-\x08, \x0b, \x0c, \x0e-\x1f; U+FFFE / U+FFFF are not XML characters (a line feed is fine)
_CONTROL = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f￾￿]")

# ------------------------------------------------------------------------------------------------------------- styles


def fill(rgb: str) -> PatternFill:
    return PatternFill("solid", start_color=rgb, end_color=rgb)


def font(*, bold: bool = False, italic: bool = False, size: float = 11, color: Optional[str] = None, strike: bool = False) -> Font:
    """Calibri always (the same look in Excel and LibreOffice: never ``None`` / a default font)."""
    return Font(name=FONT_NAME, size=size, bold=bold, italic=italic, color=color, strike=strike)


FONT_BODY = font()
FONT_BOLD = font(bold=True)
FONT_TITLE = font(bold=True, size=14, color=TITLE_COLOR)
FONT_SUBTITLE = font(color=SUBTITLE_COLOR)
FONT_NOTE = font(italic=True, color=NOTE_COLOR)
FONT_CANCELLED = font(color=CANCELLED_COLOR, strike=True)
FONT_CANCELLED_NOTE = font(color=CANCELLED_COLOR)
FONT_EXAMPLE = font(italic=True, color=CANCELLED_COLOR)
FONT_REFERENCE = font(color=CANCELLED_COLOR)

_THIN = Side(style="thin", color=GRID)
BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_TOTAL_TOP = Side(style="medium", color=TOTAL_LINE)
TOTAL_BORDER = Border(left=_THIN, right=_THIN, top=_TOTAL_TOP, bottom=_THIN)

ALIGN_LEFT_TOP = Alignment(horizontal="left", vertical="top", wrap_text=True)
ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_CENTER_TOP = Alignment(horizontal="center", vertical="top")
ALIGN_HEADER = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_HEADER_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def clean(value: Any) -> Any:
    """A text without the characters a worksheet refuses; any other value is returned as it is."""
    return _CONTROL.sub("", value) if isinstance(value, str) else value


def put(ws, row: int, col: int, value: Any, *, number_format: Optional[str] = None, font: Optional[Font] = None, fill: Optional[PatternFill] = None,
        alignment: Optional[Alignment] = None, border: Optional[Border] = None):
    """Write a cell.  A text stays text (``data_type`` "s"): a note starting with "=" is never a formula."""
    cell = ws.cell(row=row, column=col)
    cell.value = clean(value)
    if isinstance(value, str):
        cell.data_type = "s"
    if number_format:
        cell.number_format = number_format
    if font is not None:
        cell.font = font
    if fill is not None:
        cell.fill = fill
    if alignment is not None:
        cell.alignment = alignment
    if border is not None:
        cell.border = border
    return cell


# ------------------------------------------------------------------------------------------------------ the blocks


def title_block(ws, lines: Sequence[Optional[str]], last_col: int) -> int:
    """Rows 1-3, merged over ``last_col`` columns: the name of the file in bold 14, the as-of / filters line, the main totals in words.  A line that is
    ``None`` leaves an empty (still merged) row, so the header of the table is ALWAYS ``HEADER_ROW`` and the freeze ``A6``.  Returns ``HEADER_ROW``."""
    for number in range(1, TITLE_ROWS + 1):
        text = lines[number - 1] if number <= len(lines) else None
        if last_col > 1:
            ws.merge_cells(start_row=number, start_column=1, end_row=number, end_column=last_col)
        put(ws, number, 1, text or None, font=FONT_TITLE if number == 1 else FONT_SUBTITLE,
            alignment=Alignment(horizontal="left", vertical="center", wrap_text=True))
    return HEADER_ROW


def header_cell(ws, row: int, col: int, value: Any, *, number_format: Optional[str] = None, left: bool = False):
    """One cell of a table header: bold, wrapped, D9E2F3, thin border (a real date works too: pass ``number_format``)."""
    return put(ws, row, col, value, number_format=number_format, font=FONT_BOLD, fill=fill(HEAD_FILL), border=BORDER,
               alignment=ALIGN_HEADER_LEFT if left else ALIGN_HEADER)


def header_row(ws, row: int, headers: Sequence[str], *, start_col: int = 1, left: Iterable[int] = ()) -> int:
    """The header of a table: bold, centred (``left`` = the 1-based positions inside ``headers`` that are long texts and stay left), wrapped, D9E2F3,
    thin border.  Returns the row."""
    left = set(left)
    for i, text in enumerate(headers, start=1):
        header_cell(ws, row, start_col + i - 1, text, left=i in left)
    ws.row_dimensions[row].height = HEADER_HEIGHT
    return row


def section_row(ws, row: int, text: str, last_col: int, *, start_col: int = 1, fill_rgb: str = SECTION_FILL, font_: Optional[Font] = None) -> int:
    """A band over the whole width: a month, a group, a teacher (EDF1F7, bold)."""
    if last_col > start_col:
        ws.merge_cells(start_row=row, start_column=start_col, end_row=row, end_column=last_col)
    put(ws, row, start_col, text, font=font_ if font_ is not None else FONT_BOLD, fill=fill(fill_rgb), alignment=ALIGN_LEFT, border=BORDER)
    for col in range(start_col + 1, last_col + 1):
        ws.cell(row=row, column=col).fill = fill(fill_rgb)
        ws.cell(row=row, column=col).border = BORDER
    return row


def total_row(ws, row: int, label: str, values: Mapping[int, Any], last_col: int, *, start_col: int = 1, label_col: Optional[int] = None,
              formats: Optional[Mapping[int, str]] = None) -> int:
    """«Итого …»: D9E2F3, bold, a medium 8EA9DB line above.  ``values`` = {column: value}; a whole number gets ``#,##0``, ``formats`` overrides."""
    formats = formats or {}
    for col in range(start_col, last_col + 1):
        put(ws, row, col, None, fill=fill(HEAD_FILL), font=FONT_BOLD, border=TOTAL_BORDER)
    put(ws, row, label_col or start_col, label, font=FONT_BOLD, fill=fill(HEAD_FILL), border=TOTAL_BORDER, alignment=ALIGN_LEFT)
    for col, value in values.items():
        if value is None:
            continue
        shown = formats.get(col) or (MONEY if isinstance(value, int) and not isinstance(value, bool) else None)
        put(ws, row, col, value, number_format=shown, font=FONT_BOLD, fill=fill(HEAD_FILL), border=TOTAL_BORDER,
            alignment=ALIGN_LEFT if isinstance(value, str) else ALIGN_CENTER)
    return row


def example_row(ws, row: int, values: Sequence[Any], *, start_col: int = 1) -> int:
    """The example row of a template: F2F2F2, grey italic (the importer ignores a row that carries the mark ``EXAMPLE_MARK``)."""
    for i, value in enumerate(values):
        put(ws, row, start_col + i, value, font=FONT_EXAMPLE, fill=fill(EXAMPLE_FILL), border=BORDER, alignment=ALIGN_LEFT)
    return row


EXAMPLE_MARK = "пример — удалите"


# ----------------------------------------------------------------------------------------------------- measuring


def shown_length(value: Any, number_format: Optional[str] = None) -> int:
    """How many characters of width the cell needs (a date 10, a moment 16, a whole number with its thousands separators)."""
    if value is None or isinstance(value, bool):
        return 1 if isinstance(value, bool) else 0
    if isinstance(value, datetime):
        return len("00.00.0000 00:00")
    if isinstance(value, date):
        return len("00.00.0000")
    if isinstance(value, int):
        return len(f"{value:,}")
    if isinstance(value, float):
        return len(f"{value:,.2f}".rstrip("0").rstrip("."))
    return max((len(part) for part in str(value).split("\n")), default=0)


def lines_needed(text: Any, width: float) -> int:
    """The number of lines a wrapped text takes in a column of ``width`` (a word longer than the line is counted by characters)."""
    if not isinstance(text, str) or not text:
        return 1
    usable = max(1, int(width) - 2)
    total = 0
    for part in text.split("\n"):
        line, used = 1, 0
        for word in part.split(" "):
            n = len(word)
            if used == 0:
                used = n
            elif used + 1 + n <= usable:
                used += 1 + n
            else:
                line += 1
                used = n
            while used > usable:
                line += 1
                used -= usable
        total += line
    return total


def autosize(ws, *, header_row_number: int = HEADER_ROW, last_col: Optional[int] = None, min_width: float = MIN_WIDTH, wrap_cap: float = WRAP_CAP,
             wrap_cols: Iterable[int] = (), fixed: Optional[Mapping[int, float]] = None, first_col: int = 1) -> List[float]:
    """Widths from the content, heights from the wrapping.  Returns the widths (columns ``first_col`` ... ``last_col``).

    * a column is the longest value + 2, at least ``min_width``; a header that is longer than 18 characters wraps by words (the width follows the longest
      word, never below 18);
    * a text column that would pass ``wrap_cap`` (and every column of ``wrap_cols``) is capped there and its text cells wrap, aligned to the top;
    * cells inside a merged range of several columns (the title block, a band) do not widen anything: their ROW height follows the total width instead;
    * ``fixed`` = {column: width} sets a column as it is (an ID column that is hidden, a narrow date column of a matrix).
    """
    last_col = last_col or ws.max_column
    fixed = dict(fixed or {})
    wrap_cols = set(wrap_cols)
    inside_merged = set()
    for r in ws.merged_cells.ranges:
        if r.max_col > r.min_col:
            for row in range(r.min_row, r.max_row + 1):
                for col in range(r.min_col, r.max_col + 1):
                    inside_merged.add((row, col))
    need = {col: float(min_width) for col in range(first_col, last_col + 1)}
    for row in ws.iter_rows(min_row=header_row_number, max_row=ws.max_row, min_col=first_col, max_col=last_col):
        for cell in row:
            if cell.value is None or (cell.row, cell.column) in inside_merged:
                continue
            n = shown_length(cell.value, cell.number_format)
            if cell.font is not None and cell.font.bold:
                n = math.ceil(n * 1.1)
            if cell.row == header_row_number and isinstance(cell.value, str) and n > 18:
                n = max(18, max(len(w) for w in cell.value.split()) + 1)
            need[cell.column] = max(need[cell.column], n + 2)
    widths: List[float] = []
    for col in range(first_col, last_col + 1):
        width = fixed.get(col, need[col])
        if col not in fixed and (col in wrap_cols or width > wrap_cap):
            width = min(width, wrap_cap)
        widths.append(float(width))
        ws.column_dimensions[get_column_letter(col)].width = float(width)
    by_col = {first_col + i: w for i, w in enumerate(widths)}
    for row in ws.iter_rows(min_row=header_row_number, max_row=ws.max_row, min_col=first_col, max_col=last_col):
        lines = 1
        for cell in row:
            if not isinstance(cell.value, str) or (cell.row, cell.column) in inside_merged:
                continue
            width = by_col[cell.column]
            heading = cell.row == header_row_number
            wrapped = bool(cell.alignment is not None and cell.alignment.wrap_text)
            if not (heading or wrapped or cell.column in wrap_cols or width >= wrap_cap):
                continue
            if not wrapped:
                al = copy(cell.alignment) if cell.alignment is not None else Alignment()
                al.wrap_text = True
                if al.vertical is None:
                    al.vertical = "top"
                cell.alignment = al
            lines = max(lines, lines_needed(cell.value, width))
        if lines > 1:
            ws.row_dimensions[row[0].row].height = LINE_HEIGHT * lines
    for rng in ws.merged_cells.ranges:                           # the title block / a band / a paragraph: the height follows the whole width
        if rng.max_col == rng.min_col:
            continue
        cell = ws.cell(row=rng.min_row, column=rng.min_col)
        if not isinstance(cell.value, str):
            continue
        total = sum(by_col.get(c, MIN_WIDTH) for c in range(rng.min_col, rng.max_col + 1))
        lines = lines_needed(cell.value, total)
        big = (cell.font.size or 11) > 11 if cell.font is not None else False
        if lines > 1 or big:
            ws.row_dimensions[rng.min_row].height = max(LINE_HEIGHT * lines, 21 if big else LINE_HEIGHT)
    return widths


def print_setup(ws, *, header: Optional[int] = HEADER_ROW) -> None:
    """Landscape A4, fit to the page width (any number of pages high), narrow margins, the header row repeated on every page."""
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    if header:
        ws.print_title_rows = f"{header}:{header}"


# -------------------------------------------------------------------------------------------------- file names

_TRANSLIT = {
    "а": "a", "б": "b", "в": "v", "г": "g", "д": "d", "е": "e", "ё": "e", "ж": "zh", "з": "z", "и": "i", "й": "y", "к": "k", "л": "l", "м": "m",
    "н": "n", "о": "o", "п": "p", "р": "r", "с": "s", "т": "t", "у": "u", "ф": "f", "х": "kh", "ц": "ts", "ч": "ch", "ш": "sh", "щ": "shch",
    "ъ": "", "ы": "y", "ь": "", "э": "e", "ю": "yu", "я": "ya",
}
_BAD_NAME = re.compile(r'[\x00-\x1f\x7f"\\/:*?<>|;]')


def file_name(*parts: Optional[str], ext: str = "xlsx") -> str:
    """«Выписка_Алиев_Анвар_октябрь-2026.xlsx»: the parts joined by ``_``, a space inside a part becomes ``_``, no path characters, no control characters."""
    cleaned = []
    for part in parts:
        text = re.sub(r"\s+", "_", _BAD_NAME.sub(" ", str(part or "")).strip())
        text = text.strip("._")
        if text:
            cleaned.append(text)
    name = "_".join(cleaned)[:120].strip("._") or "Файл"
    return f"{name}.{ext}"


def ascii_fallback(filename: str) -> str:
    """The ASCII name for old clients: Cyrillic is transliterated, the rest of non-ASCII dropped, nothing that could break the header."""
    out = []
    for ch in filename:
        low = ch.lower()
        if low in _TRANSLIT:
            piece = _TRANSLIT[low]
            out.append(piece.capitalize() if ch != low and piece else piece)
        elif 0x20 <= ord(ch) < 0x7F and ch not in '"\\;':
            out.append(ch)
    name = re.sub(r"\s+", "_", "".join(out)).strip("._ ")
    ext = filename.rsplit(".", 1)[-1] if "." in filename else "xlsx"
    if not name or name.lower() == ext.lower():
        return f"file.{ext}"
    return name if name.lower().endswith("." + ext.lower()) else f"{name}.{ext}"


def content_disposition(filename: str) -> str:
    """``attachment; filename="Vypiska_Aliev.xlsx"; filename*=UTF-8''%D0%92...``: the ASCII fallback and the real Russian name (RFC 5987)."""
    safe = _BAD_NAME.sub(" ", filename).strip() or "file.xlsx"
    return f"attachment; filename=\"{ascii_fallback(safe)}\"; filename*=UTF-8''{quote(safe, safe='')}"


def attachment_headers(filename: str, **extra: str) -> Dict[str, str]:
    return {"Content-Disposition": content_disposition(filename), **extra}


# ---------------------------------------------------------------------------------------------------- words

MONTH_NAMES = ["январь", "февраль", "март", "апрель", "май", "июнь", "июль", "август", "сентябрь", "октябрь", "ноябрь", "декабрь"]


def month_slug(day: date) -> str:
    """``октябрь-2026`` (a part of a file name)."""
    return f"{MONTH_NAMES[day.month - 1]}-{day.year}"


def month_title(key: str) -> str:
    """``2026-10`` -> «Октябрь 2026»."""
    name = MONTH_NAMES[int(key[5:7]) - 1]
    return f"{name[:1].upper()}{name[1:]} {key[:4]}"


def date_ru(day: date) -> str:
    return day.strftime("%d.%m.%Y")


def words_join(parts: Iterable[Optional[str]]) -> str:
    """The lines of the title block: the non-empty parts joined by « · »."""
    return " · ".join(p for p in parts if p)
