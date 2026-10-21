"""
ALGO - PAY-2: the statement of a student, the receipts of the centre, the balances (design §4.5, §4.6).

Reads only: the numbers come from ``billing_ledger.run_ledger`` of every membership (one batch load of the whole centre).
Money is never recomputed here - sums of what the ledger distributed.
"""
from __future__ import annotations

import csv
import io
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from typing import Any, Dict, List, Optional, Sequence, Tuple
from uuid import UUID

from sqlalchemy.orm import Session

from src.models.auth import Group, Student, StudentGroup, StudentPayment, StudentRefund, StudentWriteoff, User

from . import billing_ledger as bl
from . import ledger_service as ls
from . import payment_status as ps
from . import payments_texts as tx
from . import payments_views as pv
from .payments_files import person_name, staff_name
from .payments_schemas import (
    BalanceRow,
    BalancesPage,
    ReceiptItem,
    ReceiptItems,
    ReceiptsByMethod,
    ReceiptsGroupRow,
    ReceiptsPosition,
    ReceiptsReport,
    ReceiptsSeriesPoint,
    ReceiptsTeacherRow,
    ReceiptsTotals,
    Reconciliation,
    ReconciliationCheck,
    StatementDetails,
    StatementMembership,
    StatementOut,
    StatementRow,
    StatementSummary,
)
from .payments_service import PaymentNotFound, PaymentService, RangeTooLong, PaymentInvalid

MAX_RANGE_DAYS = 366
MAX_DAY_SERIES_DAYS = 92
COLUMNS = ("received", "to_past_debt", "to_current", "to_advance", "written_off", "refunded")
BUCKET_COLUMN = {bl.PAST_DEBT: "to_past_debt", bl.CURRENT: "to_current", bl.ADVANCE: "to_advance"}


def _person(user: Optional[User]) -> Optional[str]:
    return staff_name(user.first_name, user.last_name) if user else None


def _iso(value: Optional[datetime]) -> Optional[str]:
    return value.isoformat() if value else None


class PaymentsReports:
    def __init__(self, db: Session) -> None:
        self.db = db
        self.service = PaymentService(db)

    # ============================================================== statement

    def statement(
        self, org_id: str, membership_id: UUID, *, date_from: Optional[date] = None, date_to: Optional[date] = None,
        include_voided: bool = False,
    ) -> StatementOut:
        self.service._membership(org_id, membership_id)
        today = ps.business_today()
        loader = ls.LedgerLoader(self.db, str(org_id), today)
        (entry,) = loader.run([membership_id])
        ledger = entry.ledger
        users = self.service._users(
            x for coll in (entry.payments, entry.writeoffs, entry.refunds) for r in coll for x in (r.created_by,))
        payments = {str(p.id): p for p in entry.payments}
        writeoffs = {str(w.id): w for w in entry.writeoffs}
        refunds = {str(r.id): r for r in entry.refunds}
        transfers = {f"tr:{t.id}": t for t in entry.transfers_in + entry.transfers_out}
        segments = {s.id: s for s in ledger.segments}

        def group_name(mid: UUID) -> str:
            found = loader.members.get(mid)
            return found[2].name if found else ""

        rows: List[StatementRow] = []
        for seq, e in enumerate(ledger.statement, 1):
            balance = e.balance_after
            common = dict(seq=seq, date=e.day, debit=e.debit, credit=e.credit, outflow=e.outflow, balance_after=balance,
                          balance_text=tx.balance_text(balance))
            if e.kind == "charge":
                seg = segments[e.ref_id]
                title = f"Начисление {tx.period_text(seg.period_start, seg.period_end)}" if seg.kind == "period" else (
                    f"Начисление: пакет №{seg.package_no} ({tx.period_text(seg.period_start, seg.period_end)})")
                rows.append(StatementRow(type="charge", title=title, details=StatementDetails(segment=pv.segment_out(seg)), **common))
            elif e.kind == bl.PAYMENT:
                p = payments[e.ref_id]
                rows.append(StatementRow(
                    type="payment", title=f"Оплата № {p.receipt_no}" if p.receipt_no else "Оплата",
                    details=StatementDetails(
                        payment_id=str(p.id), receipt_no=p.receipt_no, method=p.method or "unknown",
                        created_by_name=_person(users.get(p.created_by)), note=p.note,
                        allocations=pv.allocations_of(ledger, e.ref_id)), **common))
            elif e.kind == bl.WRITEOFF:
                w = writeoffs[e.ref_id]
                rows.append(StatementRow(type="writeoff", title="Списание долга", details=StatementDetails(
                    writeoff_id=str(w.id), reason=w.reason, created_by_name=_person(users.get(w.created_by))), **common))
            elif e.kind == bl.REFUND:
                r = refunds[e.ref_id]
                rows.append(StatementRow(type="refund", title="Возврат аванса", details=StatementDetails(
                    refund_id=str(r.id), method=r.method, reason=r.reason, note=r.note,
                    created_by_name=_person(users.get(r.created_by))), **common))
            elif e.kind == bl.TRANSFER_IN:
                t = transfers[e.ref_id]
                name = group_name(t.from_membership_id)
                rows.append(StatementRow(type="transfer_in", title=f"Перенос из «{name}»", details=StatementDetails(
                    from_group_name=name, from_membership_id=str(t.from_membership_id),
                    allocations=pv.allocations_of(ledger, e.ref_id)), **common))
            elif e.kind == bl.TRANSFER_OUT:
                t = transfers[e.ref_id]
                name = group_name(t.to_membership_id)
                rows.append(StatementRow(type="transfer_out", title=f"Перенос в «{name}»", details=StatementDetails(
                    to_group_name=name, to_membership_id=str(t.to_membership_id)), **common))
        if date_from or date_to:
            rows = [r for r in rows if (not date_from or r.date >= date_from) and (not date_to or r.date <= date_to)]
        voided: List[StatementRow] = []
        if include_voided:
            voided = self._voided_rows(org_id, membership_id)
        summary = StatementSummary(
            charged=ledger.charged, paid=ledger.paid, transferred_in=ledger.transferred_in, written_off=ledger.written_off,
            refunded=ledger.refunded, transferred_out=ledger.transferred_out, debt=ledger.debt, advance=ledger.advance,
            paid_until=ledger.status.paid_until, status=ledger.status.status, reconciles=ledger.reconciles)
        return StatementOut(
            membership=StatementMembership(
                membership_id=str(entry.membership.id), student_id=str(entry.student.id),
                student_name=person_name(entry.student.last_name, entry.student.first_name) or "",
                display_id=entry.student.display_id, group_id=str(entry.group.id), group_name=entry.group.name,
                teacher_name=_person(entry.teacher), is_active=bool(entry.membership.is_active)),
            as_of=today, mode=ledger.mode, accounting_start=loader.settings.accounting_start, summary=summary, rows=rows,
            voided=voided, upcoming=[pv.segment_out(s) for s in ledger.segments if not s.begun and (s.money or s.forgiven)])

    def _voided_rows(self, org_id: str, membership_id: UUID) -> List[StatementRow]:
        users_ids: List[Any] = []
        pay = self.db.query(StudentPayment).filter(
            StudentPayment.org_id == org_id, StudentPayment.membership_id == membership_id,
            StudentPayment.voided_at.isnot(None), StudentPayment.kind == "payment").all()
        wo = self.db.query(StudentWriteoff).filter(
            StudentWriteoff.org_id == org_id, StudentWriteoff.membership_id == membership_id,
            StudentWriteoff.voided_at.isnot(None)).all()
        rf = self.db.query(StudentRefund).filter(
            StudentRefund.org_id == org_id, StudentRefund.membership_id == membership_id,
            StudentRefund.voided_at.isnot(None)).all()
        users = self.service._users(x for r in (*pay, *wo, *rf) for x in (r.created_by, r.voided_by))
        out: List[tuple] = []
        for p in pay:
            out.append((p.paid_on, p.voided_at, StatementRow(
                seq=0, date=p.paid_on, type="payment", title=f"Оплата № {p.receipt_no}" if p.receipt_no else "Оплата",
                credit=int(p.amount), voided_at=_iso(p.voided_at), void_reason=p.void_reason,
                voided_by_name=_person(users.get(p.voided_by)),
                details=StatementDetails(payment_id=str(p.id), receipt_no=p.receipt_no, method=p.method or "unknown",
                                         created_by_name=_person(users.get(p.created_by)), note=p.note))))
        for w in wo:
            out.append((w.effective_on, w.voided_at, StatementRow(
                seq=0, date=w.effective_on, type="writeoff", title="Списание долга", credit=int(w.amount),
                voided_at=_iso(w.voided_at), void_reason=w.void_reason, voided_by_name=_person(users.get(w.voided_by)),
                details=StatementDetails(writeoff_id=str(w.id), reason=w.reason,
                                         created_by_name=_person(users.get(w.created_by))))))
        for r in rf:
            out.append((r.effective_on, r.voided_at, StatementRow(
                seq=0, date=r.effective_on, type="refund", title="Возврат аванса", outflow=int(r.amount),
                voided_at=_iso(r.voided_at), void_reason=r.void_reason, voided_by_name=_person(users.get(r.voided_by)),
                details=StatementDetails(refund_id=str(r.id), method=r.method, reason=r.reason, note=r.note,
                                         created_by_name=_person(users.get(r.created_by))))))
        out.sort(key=lambda x: (x[0], x[1]))
        rows = []
        for i, (_, _, row) in enumerate(out, 1):
            row.seq = i
            rows.append(row)
        return rows

    def statement_file(
        self, org_id: str, membership_id: UUID, *, fmt: str, date_from: Optional[date] = None,
        date_to: Optional[date] = None, include_voided: bool = False,
    ) -> Tuple[bytes, str, str]:
        """``(bytes, content type, filename)`` of the statement as xlsx / csv."""
        st = self.statement(org_id, membership_id, date_from=date_from, date_to=date_to, include_voided=include_voided)
        header = ["№", "Дата", "Операция", "Начислено", "Внесено / списано", "Возвращено / перенесено", "Баланс", "Подробности"]

        def detail(r: StatementRow) -> str:
            d = r.details
            parts: List[str] = []
            if r.type == "payment" and d.method:
                parts.append(f"способ: {d.method}")
            if d.allocations:
                parts.append("; ".join(f"{a.label}: {tx.num(a.amount)}" for a in d.allocations))
            if d.reason:
                parts.append(f"причина: {d.reason}")
            if r.voided_at:
                parts.append(f"ОТМЕНЕНО: {r.void_reason}")
            return " | ".join(parts)

        top = [
            ["Выписка", f"{st.membership.student_name} · {st.membership.group_name}", f"на {st.as_of.strftime('%d.%m.%Y')}"],
            ["Начислено", st.summary.charged, "Оплачено", st.summary.paid + st.summary.transferred_in],
            ["Списано", st.summary.written_off, "Возвращено / перенесено", st.summary.refunded + st.summary.transferred_out],
            ["Долг", st.summary.debt, "Аванс", st.summary.advance],
        ]
        body = [[r.seq, r.date.strftime("%d.%m.%Y"), r.title, r.debit or "", r.credit or "", r.outflow or "",
                 r.balance_after if r.balance_after is not None else "", detail(r)] for r in [*st.rows, *st.voided]]
        who = st.membership.display_id or str(membership_id)[:8]
        base = f"statement-{who}-{st.as_of.isoformat()}"
        if fmt == "csv":
            buf = io.StringIO()
            writer = csv.writer(buf, delimiter=";")
            for line in top:
                writer.writerow(line)
            writer.writerow([])
            writer.writerow(header)
            for line in body:
                writer.writerow(line)
            return ("﻿" + buf.getvalue()).encode("utf-8"), "text/csv; charset=utf-8", base + ".csv"
        from openpyxl import Workbook
        from openpyxl.styles import Font

        wb = Workbook()
        ws = wb.active
        ws.title = "Выписка"
        for line in top:
            ws.append(line)
        ws.append([])
        ws.append(header)
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
        for line in body:
            ws.append(line)
        for col, width in zip("ABCDEFGH", (6, 12, 34, 14, 18, 22, 14, 70)):
            ws.column_dimensions[col].width = width
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", base + ".xlsx"

    # ============================================================== the centre

    def _everything(self, org_id: str) -> Tuple[ls.LedgerLoader, Dict[UUID, ls.MembershipLedger]]:
        loader = ls.LedgerLoader(self.db, str(org_id), ps.business_today())
        entries = loader.load(scope="all")
        return loader, {e.membership.id: e for e in entries}

    @staticmethod
    def _validate_range(date_from: date, date_to: date, group_by: Optional[str] = None) -> None:
        if date_to < date_from:
            raise PaymentInvalid("Конец периода раньше начала")
        days = (date_to - date_from).days + 1
        if days > MAX_RANGE_DAYS:
            raise RangeTooLong(f"Период не длиннее {MAX_RANGE_DAYS} дней")
        if group_by == "day" and days > MAX_DAY_SERIES_DAYS:
            raise RangeTooLong(f"Разбивка по дням — не длиннее {MAX_DAY_SERIES_DAYS} дней")

    def _records(self, org_id: str, date_from: date, date_to: date, method: Optional[str],
                 group_id: Optional[UUID], teacher_id: Optional[UUID]):
        """Live payments / write-offs / refunds of the range with the filters (group / teacher of the membership)."""
        def scoped(model, query):
            query = query.join(Group, Group.id == model.group_id).filter(model.org_id == org_id, model.voided_at.is_(None))
            if group_id:
                query = query.filter(model.group_id == group_id)
            if teacher_id:
                query = query.filter(Group.teacher_id == teacher_id)
            return query

        payments = scoped(StudentPayment, self.db.query(StudentPayment, Group)).filter(
            StudentPayment.kind == "payment", StudentPayment.paid_on >= date_from, StudentPayment.paid_on <= date_to)
        if method:
            payments = payments.filter(StudentPayment.method == method)
        writeoffs = scoped(StudentWriteoff, self.db.query(StudentWriteoff, Group)).filter(
            StudentWriteoff.effective_on >= date_from, StudentWriteoff.effective_on <= date_to)
        refunds = scoped(StudentRefund, self.db.query(StudentRefund, Group)).filter(
            StudentRefund.effective_on >= date_from, StudentRefund.effective_on <= date_to)
        if method:
            refunds = refunds.filter(StudentRefund.method == method)
        return payments.all(), writeoffs.all(), refunds.all()

    @staticmethod
    def _bucket_key(day: date, group_by: str) -> str:
        return day.isoformat() if group_by == "day" else day.strftime("%Y-%m")

    @staticmethod
    def _split(ledger: bl.Ledger, payment_id: str) -> Dict[str, int]:
        out = {bl.PAST_DEBT: 0, bl.CURRENT: 0, bl.ADVANCE: 0}
        for a in ledger.allocations:
            if a.credit_id == payment_id:
                out[a.bucket] += a.amount
        return out

    def receipts_report(
        self, org_id: str, date_from: date, date_to: date, *, group_by: str = "day", method: Optional[str] = None,
        group_id: Optional[UUID] = None, teacher_id: Optional[UUID] = None,
    ) -> ReceiptsReport:
        self._validate_range(date_from, date_to, group_by)
        loader, entries = self._everything(org_id)
        payments, writeoffs, refunds = self._records(org_id, date_from, date_to, method, group_id, teacher_id)
        totals: Dict[str, int] = defaultdict(int)
        by_method: Dict[str, List[int]] = defaultdict(lambda: [0, 0])
        by_group: Dict[UUID, Dict[str, Any]] = {}
        by_teacher: Dict[Optional[UUID], Dict[str, Any]] = {}
        series: Dict[str, Dict[str, Any]] = {}

        def bucket_of(day: date) -> Dict[str, Any]:
            return series.setdefault(self._bucket_key(day, group_by), {
                "received": 0, "by_method": defaultdict(int), "to_past_debt": 0, "to_current": 0, "to_advance": 0,
                "written_off": 0, "refunded": 0, "count": 0})

        def group_row(group: Group) -> Dict[str, Any]:
            return by_group.setdefault(group.id, {
                "group": group, "received": 0, "to_past_debt": 0, "to_current": 0, "to_advance": 0, "written_off": 0,
                "refunded": 0, "count": 0})

        def teacher_row(group: Group) -> Dict[str, Any]:
            return by_teacher.setdefault(group.teacher_id, {
                "received": 0, "to_past_debt": 0, "to_current": 0, "to_advance": 0, "written_off": 0, "refunded": 0,
                "count": 0})

        def add(group: Group, day: date, **fields: int) -> None:
            for target in (group_row(group), teacher_row(group), bucket_of(day)):
                for k, v in fields.items():
                    target[k] += v

        allocated = refunded_from = 0
        for p, group in payments:
            entry = entries.get(p.membership_id)
            split = self._split(entry.ledger, str(p.id)) if entry else {bl.PAST_DEBT: 0, bl.CURRENT: 0, bl.ADVANCE: 0}
            cut = entry.ledger.refunded_from.get(str(p.id), 0) if entry else 0
            amount = int(p.amount)
            totals["received"] += amount
            totals["payments_count"] += 1
            for bucket, column in BUCKET_COLUMN.items():
                totals[column] += split[bucket]
            allocated += sum(split.values())
            refunded_from += cut
            by_method[p.method or "unknown"][0] += amount
            by_method[p.method or "unknown"][1] += 1
            add(group, p.paid_on, received=amount, count=1, to_past_debt=split[bl.PAST_DEBT], to_current=split[bl.CURRENT],
                to_advance=split[bl.ADVANCE])
            bucket_of(p.paid_on)["by_method"][p.method or "unknown"] += amount
        for w, group in writeoffs:
            totals["written_off"] += int(w.amount)
            add(group, w.effective_on, written_off=int(w.amount))
        for r, group in refunds:
            totals["refunded"] += int(r.amount)
            add(group, r.effective_on, refunded=int(r.amount))
        totals["refunded_from_received"] = refunded_from
        totals["net_received"] = totals["received"] - totals["refunded"]

        ledgers = [e.ledger for e in entries.values()]
        position = ReceiptsPosition(
            debt_total=sum(l.debt for l in ledgers), past_debt_total=sum(l.past_debt for l in ledgers),
            advance_total=sum(l.advance for l in ledgers), memberships_with_debt=sum(1 for l in ledgers if l.debt > 0),
            students_with_debt=len({e.student.id for e in entries.values() if e.ledger.debt > 0}),
            memberships_with_advance=sum(1 for l in ledgers if l.advance > 0))
        checks = [
            ReconciliationCheck(code="received_eq_allocated_plus_refunded", left=totals["received"],
                                right=allocated + refunded_from, diff=totals["received"] - allocated - refunded_from),
            ReconciliationCheck(
                code="charged_minus_paid_eq_debt_minus_advance",
                left=sum(l.charged - l.paid - l.transferred_in - l.written_off + l.refunded + l.transferred_out for l in ledgers),
                right=position.debt_total - position.advance_total,
                diff=sum(l.charged - l.paid - l.transferred_in - l.written_off + l.refunded + l.transferred_out for l in ledgers)
                - (position.debt_total - position.advance_total)),
            ReconciliationCheck(code="transfers_net_zero", left=sum(l.transferred_out for l in ledgers),
                                right=sum(l.transferred_in for l in ledgers),
                                diff=sum(l.transferred_out for l in ledgers) - sum(l.transferred_in for l in ledgers)),
        ]
        teachers = {t.id: t for t in self.db.query(User).filter(User.id.in_([k for k in by_teacher if k] or [None]))}

        def group_out(item: Dict[str, Any]) -> ReceiptsGroupRow:
            g: Group = item["group"]
            t = teachers.get(g.teacher_id)
            return ReceiptsGroupRow(
                group_id=str(g.id), group_name=g.name, teacher_id=str(g.teacher_id) if g.teacher_id else None,
                teacher_name=_person(t), **{k: v for k, v in item.items() if k != "group"})

        return ReceiptsReport(
            as_of=loader.today, from_=date_from, to=date_to, group_by=group_by,
            totals=ReceiptsTotals(**{k: totals.get(k, 0) for k in ReceiptsTotals.model_fields}),
            by_method=[ReceiptsByMethod(method=m, amount=v[0], count=v[1])
                       for m, v in sorted(by_method.items(), key=lambda kv: (-kv[1][0], kv[0]))],
            by_group=sorted((group_out(i) for i in by_group.values()), key=lambda r: (-r.received, r.group_name)),
            by_teacher=sorted((ReceiptsTeacherRow(
                teacher_id=str(tid) if tid else None, teacher_name=_person(teachers.get(tid)), **v)
                for tid, v in by_teacher.items()), key=lambda r: (-r.received, r.teacher_name or "")),
            series=[ReceiptsSeriesPoint(bucket=k, **{**v, "by_method": dict(v["by_method"])}) for k, v in sorted(series.items())],
            position=position,
            reconciliation=Reconciliation(ok=all(c.diff == 0 for c in checks), checks=checks))

    def receipts_items(
        self, org_id: str, date_from: date, date_to: date, *, column: str = "received", bucket: Optional[str] = None,
        method: Optional[str] = None, group_id: Optional[UUID] = None, teacher_id: Optional[UUID] = None, limit: int = 50,
        offset: int = 0,
    ) -> ReceiptItems:
        self._validate_range(date_from, date_to)
        loader, entries = self._everything(org_id)
        payments, writeoffs, refunds = self._records(org_id, date_from, date_to, method, group_id, teacher_id)

        def in_bucket(day: date) -> bool:
            if not bucket:
                return True
            return day.isoformat() == bucket if len(bucket) == 10 else day.strftime("%Y-%m") == bucket

        names = {}
        member_ids = {x.membership_id for x, _ in (*payments, *writeoffs, *refunds)}
        for mid in member_ids:
            found = loader.members.get(mid)
            if found:
                names[mid] = found
        items: List[Tuple[date, str, ReceiptItem]] = []

        def make(kind: str, rec, group: Group, day: date, amount: int, column_amount: int, **extra) -> None:
            _m, student, grp, teacher = names.get(rec.membership_id, (None, None, group, None))
            items.append((day, str(rec.id), ReceiptItem(
                type=kind, id=str(rec.id), date=day, amount=amount, column_amount=column_amount,
                student_id=str(rec.student_id), student_name=person_name(student.last_name, student.first_name) if student else "",
                membership_id=str(rec.membership_id), group_name=group.name, teacher_name=_person(teacher), **extra)))

        if column in ("received", "to_past_debt", "to_current", "to_advance"):
            for p, group in payments:
                if not in_bucket(p.paid_on):
                    continue
                entry = entries.get(p.membership_id)
                split = self._split(entry.ledger, str(p.id)) if entry else {bl.PAST_DEBT: 0, bl.CURRENT: 0, bl.ADVANCE: 0}
                value = int(p.amount) if column == "received" else split[next(b for b, c in BUCKET_COLUMN.items() if c == column)]
                if value <= 0:
                    continue
                make("payment", p, group, p.paid_on, int(p.amount), value, method=p.method or "unknown", receipt_no=p.receipt_no,
                     allocations=pv.allocations_of(entry.ledger, str(p.id)) if entry else [])
        elif column == "written_off":
            for w, group in writeoffs:
                if in_bucket(w.effective_on):
                    make("writeoff", w, group, w.effective_on, int(w.amount), int(w.amount))
        elif column == "refunded":
            for r, group in refunds:
                if in_bucket(r.effective_on):
                    make("refund", r, group, r.effective_on, int(r.amount), int(r.amount), method=r.method)
        items.sort(key=lambda x: (x[0], x[1]), reverse=True)
        page = [i[2] for i in items[offset : offset + limit]]
        return ReceiptItems(items=page, total=len(items), limit=limit, offset=offset)

    def balances(
        self, org_id: str, *, kind: str = "debt", group_id: Optional[UUID] = None, teacher_id: Optional[UUID] = None,
        search: Optional[str] = None, limit: int = 50, offset: int = 0,
    ) -> BalancesPage:
        loader = ls.LedgerLoader(self.db, str(org_id), ps.business_today())
        entries = loader.load(group_id=group_id, teacher_id=teacher_id, search=search, scope="all")
        field_name = "debt" if kind == "debt" else "advance"
        picked = [e for e in entries if getattr(e.ledger, field_name) > 0]
        picked.sort(key=lambda e: (-getattr(e.ledger, field_name), (person_name(e.student.last_name, e.student.first_name) or "").lower(),
                                   str(e.membership.id)))
        rows = []
        for e in picked[offset : offset + limit]:
            l = e.ledger
            rows.append(BalanceRow(
                membership_id=str(e.membership.id), student_id=str(e.student.id),
                student_name=person_name(e.student.last_name, e.student.first_name) or "", display_id=e.student.display_id,
                group_id=str(e.group.id), group_name=e.group.name, teacher_id=str(e.group.teacher_id) if e.group.teacher_id else None,
                teacher_name=_person(e.teacher), is_active=bool(e.membership.is_active),
                student_archived=e.student.status == "archived", status=l.status.status, paid_until=l.status.paid_until,
                debt=l.debt, past_debt=l.past_debt, current_due=l.current_due, advance=l.advance, refundable=l.refundable))
        return BalancesPage(items=rows, total=len(picked), limit=limit, offset=offset)
