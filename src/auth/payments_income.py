"""
ALGO - INCOME-2 («Поступления» v2): the income overview with the comparison, the month collection, the payments journal.

Design: ``artifacts/partner_doc_check/income2_spec.md``.  Reads only; money is never recomputed here - sums of what the ledger
distributed (``billing_ledger.run_ledger`` of every membership, one batch load) and of the payment rows.  Business dates are Tashkent
(``paid_on`` is already a business date; the time of entry is ``created_at`` shown in the business timezone).  A fixed number of
queries whatever the number of payments / memberships.
"""
from __future__ import annotations

import calendar
import csv
import io
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime, timedelta, timezone
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple
from uuid import UUID

from sqlalchemy import func, or_
from sqlalchemy.orm import Session

from src.models.auth import Group, Student, StudentPayment, StudentRefund, StudentWriteoff

from . import billing_ledger as bl
from . import ledger_service as ls
from . import payment_status as ps
from . import payments_texts as tx
from . import payments_views as pv
from .business_clock import BUSINESS_TIMEZONE
from .payments_files import person_name, staff_name
from .payments_income_schemas import (
    CollectionRow,
    IncomeCash,
    IncomeCheck,
    IncomeDay,
    IncomeMethod,
    IncomeOverview,
    IncomePeriod,
    IncomePosition,
    IncomeTotals,
    JournalAuthor,
    JournalRow,
    JournalTotals,
    MonthCollection,
    PaymentsJournal,
)
from .payments_reports import MAX_RANGE_DAYS, PaymentsReports
from .payments_service import PaymentInvalid, PaymentService, RangeTooLong

MAX_DAY_SERIES_DAYS = 92
MAX_JOURNAL_PAGE = 100
MAX_EXPORT_ROWS = 20_000
STATUS_TEXT = {"active": "Действует", "voided": "Отменена", "replaced": "Исправлена"}


class InvalidPeriod(PaymentInvalid):
    code = "INVALID_PERIOD"


class ExportTooLong(PaymentInvalid):
    code = "EXPORT_TOO_LONG"


def validate_range(first: date, last: date) -> None:
    if last < first:
        raise InvalidPeriod("Конец периода раньше начала")
    if (last - first).days + 1 > MAX_RANGE_DAYS:
        raise RangeTooLong(f"Период не длиннее {MAX_RANGE_DAYS} дней")


# ============================================================================================================
# Periods
# ============================================================================================================


@dataclass(frozen=True)
class Period:
    kind: str
    first: date
    last: date
    label: str
    in_progress: bool

    @property
    def days(self) -> int:
        return (self.last - self.first).days + 1


def month_end(day: date) -> date:
    return day.replace(day=calendar.monthrange(day.year, day.month)[1])


def _same_month(a: date, b: date) -> bool:
    return (a.year, a.month) == (b.year, b.month)


def _ddmm_span(first: date, last: date) -> str:
    year = first.year != last.year
    f = first.strftime("%d.%m.%Y" if year else "%d.%m")
    return f"{f}–{last.strftime('%d.%m.%Y' if year else '%d.%m')}"


def _label(kind: str, first: date, last: date, today: date, partial: bool = False) -> str:
    if partial:
        return tx.span_text(first, last)
    if kind == "day":
        if first == today:
            return "сегодня"
        if first == today - timedelta(days=1):
            return "вчера"
        return tx.day_genitive(first)
    if kind == "week":
        return "неделя " + tx.span_text(first, last)
    if kind == "month":
        name = tx.MONTHS_NOMINATIVE[first.month - 1]
        return name if first.year == today.year else f"{name} {first.year}"
    return "период " + _ddmm_span(first, last)


def resolve_period(first: date, last: date, today: date, kind: Optional[str] = None) -> Period:
    """The period of the page; the kind is told or inferred (one day / seven days from a Monday / a whole calendar month)."""
    if last < first:
        raise InvalidPeriod("Конец периода раньше начала")
    days = (last - first).days + 1
    if kind is None:
        if days == 1:
            kind = "day"
        elif first.day == 1 and _same_month(first, last) and last == month_end(first):
            kind = "month"
        elif days == 7 and first.isoweekday() == 1:
            kind = "week"
        else:
            kind = "custom"
    if kind == "day" and days != 1:
        raise InvalidPeriod("День — это один день")
    if kind == "week" and days != 7:
        raise InvalidPeriod("Неделя — это семь дней")
    if kind == "month" and not (first.day == 1 and _same_month(first, last)):
        raise InvalidPeriod("Месяц начинается с первого числа и не выходит за свой месяц")
    return Period(kind, first, last, _label(kind, first, last, today), first <= today < last)


def previous_period(p: Period, today: date) -> Period:
    """The same kind of period before it.  A day -> the day before; a week / a custom range -> the range of the same length right before;
    a month -> the previous calendar month.  While the period is still running (today inside it) the previous one is cut to the same
    number of days («1–21 сентября» against «1–21 октября»), so the comparison is like for like."""
    span = p.days
    if p.in_progress:
        span = (today - p.first).days + 1
    if p.kind == "month":
        prev_last_month = p.first.replace(day=1) - timedelta(days=1)
        first = prev_last_month.replace(day=1)
        whole = p.last == month_end(p.first) and not p.in_progress
        last = prev_last_month if whole else first + timedelta(days=min(prev_last_month.day, span) - 1)
    else:
        first = p.first - timedelta(days=p.days)
        last = first + timedelta(days=span - 1)
    partial = p.kind != "custom" and (p.in_progress or (p.kind == "month" and last != month_end(first)))
    return Period(p.kind, first, last, _label(p.kind, first, last, today, partial), False)


def _lead(p: Period, today: date) -> str:
    if p.kind == "day" and p.first == today:
        return "Сегодня"
    if p.kind == "day" and p.first == today - timedelta(days=1):
        return "Вчера"
    if p.kind == "week":
        return "За неделю " + tx.span_text(p.first, p.last)
    return "За " + p.label


def _previous_ref(cur: Period, prev: Period, today: date) -> str:
    if cur.kind == "day":
        return "вчера" if prev.first == today - timedelta(days=1) else "за " + prev.label
    if cur.kind != "custom" and (cur.in_progress or (cur.kind == "month" and prev.last != month_end(prev.first))):
        return "за " + tx.span_text(prev.first, prev.last)
    if cur.kind == "week":
        return "за предыдущую неделю"
    if cur.kind == "month":
        return "за " + tx.MONTHS_NOMINATIVE[prev.first.month - 1]
    n = prev.days
    return f"за предыдущие {n} {tx.days_word(n)}"


def _previous_dative(cur: Period, prev: Period, today: date) -> str:
    """«к сентябрю» / «к 1–21 сентября» / «к вчерашнему дню» / «к прошлой неделе» / «к предыдущему периоду»."""
    if cur.kind == "day":
        return "ко вчерашнему дню" if prev.first == today - timedelta(days=1) else "к " + tx.day_genitive(prev.first)
    if cur.kind != "custom" and (cur.in_progress or (cur.kind == "month" and prev.last != month_end(prev.first))):
        return "к " + tx.span_text(prev.first, prev.last)
    if cur.kind == "week":
        return "к прошлой неделе"
    if cur.kind == "month":
        return "к " + tx.MONTHS_DATIVE[prev.first.month - 1]
    return "к предыдущему периоду"


def _period_out(p: Period) -> IncomePeriod:
    return IncomePeriod(kind=p.kind, **{"from": p.first}, to=p.last, label=p.label, days=p.days, in_progress=p.in_progress)


# ============================================================================================================
# Service
# ============================================================================================================


def _business_time(moment: Optional[datetime]) -> Optional[str]:
    if moment is None:
        return None
    if moment.tzinfo is None:
        moment = moment.replace(tzinfo=timezone.utc)
    return moment.astimezone(BUSINESS_TIMEZONE).strftime("%H:%M")


def _iso(moment: Optional[datetime]) -> Optional[str]:
    return moment.isoformat() if moment else None


def _staff(user) -> Optional[str]:
    return staff_name(user.first_name, user.last_name) if user else None


class IncomeReports:
    def __init__(self, db: Session) -> None:
        self.db = db
        self.service = PaymentService(db)
        self.reports = PaymentsReports(db)

    # ------------------------------------------------------------------ overview

    def overview(
        self, org_id: str, first: date, last: date, *, kind: Optional[str] = None, method: Optional[str] = None,
        group_id: Optional[UUID] = None, teacher_id: Optional[UUID] = None,
    ) -> IncomeOverview:
        validate_range(first, last)
        today = ps.business_today()
        cur = resolve_period(first, last, today, kind)
        prev = previous_period(cur, today)
        loader, entries = self.reports._everything(org_id)
        lo = min(prev.first, cur.first)
        pay_rows, refund_rows, writeoff_rows = self._facts(org_id, lo, max(prev.last, cur.last), group_id, teacher_id)

        def totals_of(p: Period):
            return self._totals(entries, p, method, pay_rows, refund_rows, writeoff_rows)

        current, by_method, daily = totals_of(cur)
        previous, _, _ = totals_of(prev)

        collection = None
        if cur.kind == "month":
            collection = self._collection(self._shown(entries, group_id, teacher_id), cur.first, month_end(cur.first))
        pct = tx.percent_change(current.received, previous.received)
        direction = None if pct is None else ("same" if current.received == previous.received else ("up" if pct > 0 else "down"))
        if pct == 0 and direction != "same":
            direction = "same"          # rounds to 0%: shown as «0%», neither up nor down
        collected = collection.total.pct if collection and collection.total.charged > 0 else None
        text = tx.income_summary_text(
            lead=_lead(cur, today), received=current.received, previous_received=previous.received,
            previous_ref=_previous_ref(cur, prev, today), collected_pct=collected)

        shown = self._shown(entries, group_id, teacher_id)
        position = IncomePosition(
            debt_total=sum(e.ledger.debt for e in shown), past_debt_total=sum(e.ledger.past_debt for e in shown),
            advance_total=sum(e.ledger.advance for e in shown),
            memberships_with_debt=sum(1 for e in shown if e.ledger.debt > 0),
            students_with_debt=len({e.student.id for e in shown if e.ledger.debt > 0}),
            memberships_with_advance=sum(1 for e in shown if e.ledger.advance > 0),
            students_with_advance=len({e.student.id for e in shown if e.ledger.advance > 0}))

        series, series_by = self._series(cur, daily)
        journal_sum = self._journal_totals(org_id, cur.first, cur.last, method, group_id, teacher_id, None, None)[1]
        checks = [
            IncomeCheck(code="received_eq_strip", left=current.received, right=(
                current.to_past_debt + current.to_current + current.to_advance + current.refunded_from_received + current.transferred),
                diff=0),
            IncomeCheck(code="received_eq_methods", left=current.received, right=sum(m.amount for m in by_method), diff=0),
            IncomeCheck(code="received_eq_series", left=current.received, right=sum(d.received for d in series), diff=0),
            IncomeCheck(code="received_eq_journal", left=current.received, right=journal_sum, diff=0),
        ]
        if collection is not None:
            t = collection.total
            g = collection.groups
            checks.append(IncomeCheck(code="collection_total_eq_groups", left=t.charged, right=sum(r.charged for r in g), diff=0))
        for c in checks:
            c.diff = c.left - c.right
        return IncomeOverview(
            as_of=today, period=_period_out(cur), previous_period=_period_out(prev), current=current, previous=previous,
            delta_amount=current.received - previous.received, delta_pct=pct, delta_direction=direction,
            delta_text=tx.delta_text(pct, _previous_dative(cur, prev, today)) if pct is not None else None, summary_text=text, by_method=by_method,
            series=series, series_by=series_by, cash=self._cash(cur, today, pay_rows, refund_rows), collection=collection,
            position=position, checks=checks, ok=all(c.diff == 0 for c in checks))

    @staticmethod
    def _shown(entries: Dict[UUID, ls.MembershipLedger], group_id: Optional[UUID], teacher_id: Optional[UUID]):
        return [e for e in entries.values()
                if (not group_id or e.group.id == group_id) and (not teacher_id or e.group.teacher_id == teacher_id)]

    def _facts(self, org_id: str, first: date, last: date, group_id: Optional[UUID], teacher_id: Optional[UUID]):
        """The live payments / refunds / write-offs of the union of both periods, three queries, no method filter (cash needs them all)."""
        def scoped(model, columns, day):
            query = self.db.query(*columns)
            if teacher_id:
                query = query.join(Group, Group.id == model.group_id).filter(Group.teacher_id == teacher_id)
            query = query.filter(model.org_id == org_id, model.voided_at.is_(None), day >= first, day <= last)
            return query.filter(model.group_id == group_id) if group_id else query

        pay = scoped(StudentPayment, (StudentPayment.id, StudentPayment.membership_id, StudentPayment.paid_on,
                                      StudentPayment.amount, StudentPayment.method), StudentPayment.paid_on).filter(
            StudentPayment.kind == "payment").all()
        refunds = scoped(StudentRefund, (StudentRefund.effective_on, StudentRefund.amount, StudentRefund.method),
                         StudentRefund.effective_on).all()
        writeoffs = scoped(StudentWriteoff, (StudentWriteoff.effective_on, StudentWriteoff.amount),
                           StudentWriteoff.effective_on).all()
        return pay, refunds, writeoffs

    def _totals(self, entries, p: Period, method: Optional[str], pay_rows, refund_rows, writeoff_rows):
        t: Dict[str, int] = defaultdict(int)
        methods: Dict[str, List[int]] = defaultdict(lambda: [0, 0])
        daily: Dict[date, List[int]] = defaultdict(lambda: [0, 0])
        for pid, mid, paid_on, amount, pay_method in pay_rows:
            if not (p.first <= paid_on <= p.last) or (method and (pay_method or "unknown") != method):
                continue
            amount = int(amount)
            entry = entries.get(mid)
            split = PaymentsReports._split(entry.ledger, str(pid)) if entry else {bl.PAST_DEBT: 0, bl.CURRENT: 0, bl.ADVANCE: 0}
            t["received"] += amount
            t["payments_count"] += 1
            t["to_past_debt"] += split[bl.PAST_DEBT]
            t["to_current"] += split[bl.CURRENT]
            t["to_advance"] += split[bl.ADVANCE]
            if entry:
                t["refunded_from_received"] += entry.ledger.refunded_from.get(str(pid), 0)
                t["transferred"] += entry.ledger.transferred_from.get(str(pid), 0)
            key = pay_method or "unknown"
            methods[key][0] += amount
            methods[key][1] += 1
            daily[paid_on][0] += amount
            daily[paid_on][1] += 1
        for day, amount, refund_method in refund_rows:
            if p.first <= day <= p.last and (not method or refund_method == method):
                t["refunded"] += int(amount)
        for day, amount in writeoff_rows:
            if p.first <= day <= p.last:
                t["written_off"] += int(amount)
        t["net_received"] = t["received"] - t["refunded"]
        out_methods = [IncomeMethod(method=m, label=tx.method_label(m), amount=v[0], count=v[1])
                       for m, v in sorted(methods.items(), key=lambda kv: (-kv[1][0], kv[0])) if v[0] or v[1]]
        return IncomeTotals(**{k: t.get(k, 0) for k in IncomeTotals.model_fields}), out_methods, daily

    @staticmethod
    def _series(p: Period, daily: Dict[date, List[int]]):
        if p.days <= MAX_DAY_SERIES_DAYS:
            return [IncomeDay(bucket=(p.first + timedelta(days=i)).isoformat(),
                              received=daily.get(p.first + timedelta(days=i), [0, 0])[0],
                              count=daily.get(p.first + timedelta(days=i), [0, 0])[1]) for i in range(p.days)], "day"
        months: Dict[str, List[int]] = {}
        day = p.first.replace(day=1)
        while day <= p.last:
            months[day.strftime("%Y-%m")] = [0, 0]
            day = (day + timedelta(days=32)).replace(day=1)
        for d, v in daily.items():
            months[d.strftime("%Y-%m")][0] += v[0]
            months[d.strftime("%Y-%m")][1] += v[1]
        return [IncomeDay(bucket=k, received=v[0], count=v[1]) for k, v in months.items()], "month"

    @staticmethod
    def _cash(p: Period, today: date, pay_rows, refund_rows) -> Optional[IncomeCash]:
        if p.first != p.last:
            return None
        day = p.first
        got = [int(a) for _pid, _mid, d, a, m in pay_rows if d == day and (m or "unknown") == "cash"]
        out = [int(a) for d, a, m in refund_rows if d == day and m == "cash"]
        label = "Наличные в кассе за " + ("сегодня" if day == today else tx.day_genitive(day))
        expected = sum(got) - sum(out)
        text = f"{label}: {tx.money(expected)}"
        if out:
            text += f" (принято {tx.money(sum(got))}, возвращено {tx.money(sum(out))})"
        return IncomeCash(date=day, label=label, cash_received=sum(got), cash_refunded=sum(out), expected=expected,
                          payments_count=len(got), text=text)

    # ------------------------------------------------------------------ month collection

    @staticmethod
    def _collection(shown: Sequence[ls.MembershipLedger], first: date, last: date) -> MonthCollection:
        """«Собрано за месяц»: the periods (packages) that BEGAN in the month, from the ledger of every membership."""
        groups: Dict[Any, Dict[str, Any]] = {}
        all_students, debtors = set(), set()
        total = defaultdict(int)
        for e in shown:
            charged = paid = forgiven = debt = 0
            for seg in e.ledger.segments:
                if seg.kind not in ("period", "package") or not seg.begun or not (first <= seg.period_start <= last):
                    continue
                forg = min(seg.forgiven, seg.amount)
                unpaid = seg.unpaid
                charged += seg.amount
                forgiven += forg
                debt += unpaid
                paid += seg.amount - forg - unpaid
            if charged == 0:
                continue
            row = groups.setdefault(e.group.id, {
                "group": e.group, "teacher": e.teacher, "charged": 0, "paid": 0, "forgiven": 0, "debt": 0, "students": set(),
                "debtors": set()})
            row["charged"] += charged
            row["paid"] += paid
            row["forgiven"] += forgiven
            row["debt"] += debt
            row["students"].add(e.student.id)
            all_students.add(e.student.id)
            if debt > 0:
                row["debtors"].add(e.student.id)
                debtors.add(e.student.id)
            for k, v in (("charged", charged), ("paid", paid), ("forgiven", forgiven), ("debt", debt)):
                total[k] += v

        def pct(paid_: int, charged_: int) -> Optional[int]:
            return paid_ * 100 // charged_ if charged_ > 0 else None

        rows = []
        for g in groups.values():
            group: Group = g["group"]
            teacher = g["teacher"]
            rows.append(CollectionRow(
                group_id=str(group.id), group_name=group.name, teacher_id=str(group.teacher_id) if group.teacher_id else None,
                teacher_name=_staff(teacher), charged=g["charged"], paid=g["paid"], forgiven=g["forgiven"], debt=g["debt"],
                students=len(g["students"]), students_with_debt=len(g["debtors"]), pct=pct(g["paid"], g["charged"])))
        rows.sort(key=lambda r: (r.paid / r.charged, -r.debt, r.group_name or ""))
        name = tx.MONTHS_NOMINATIVE[first.month - 1]
        return MonthCollection(
            month=first.strftime("%Y-%m"), month_label=f"{name} {first.year}", **{"from": first}, to=month_end(first),
            text=tx.collection_text(total["paid"], total["charged"], pct(total["paid"], total["charged"])),
            total=CollectionRow(charged=total["charged"], paid=total["paid"], forgiven=total["forgiven"], debt=total["debt"],
                                students=len(all_students), students_with_debt=len(debtors),
                                pct=pct(total["paid"], total["charged"])),
            groups=rows)

    # ------------------------------------------------------------------ journal

    def _journal_filters(self, query, org_id, first, last, method, group_id, teacher_id, created_by, q):
        query = query.filter(
            StudentPayment.org_id == org_id, Student.org_id == org_id, Group.org_id == org_id,
            StudentPayment.kind == "payment", StudentPayment.paid_on >= first, StudentPayment.paid_on <= last)
        if method:
            query = query.filter(StudentPayment.method == method)
        if group_id:
            query = query.filter(StudentPayment.group_id == group_id)
        if teacher_id:
            query = query.filter(Group.teacher_id == teacher_id)
        if created_by:
            query = query.filter(StudentPayment.created_by == created_by)
        words = (q or "").split()
        for word in words:
            pattern = ls._like(word)
            conds = [Student.first_name.ilike(pattern, escape="\\"), Student.last_name.ilike(pattern, escape="\\"),
                     Student.display_id.ilike(pattern, escape="\\")]
            if len(words) == 1 and word.isdigit() and len(word) <= 9:
                conds.append(StudentPayment.receipt_no == int(word))      # «№ квитанции» typed into the search
            query = query.filter(or_(*conds))
        return query

    def _journal_base(self, *columns):
        return (self.db.query(*columns).select_from(StudentPayment)
                .join(Student, Student.id == StudentPayment.student_id).join(Group, Group.id == StudentPayment.group_id))

    def _journal_totals(self, org_id, first, last, method, group_id, teacher_id, created_by, q) -> Tuple[int, int, int]:
        """``(active count, active sum, voided count)`` of the filtered set: one aggregate query."""
        query = self._journal_filters(
            self._journal_base(
                func.count(StudentPayment.id).filter(StudentPayment.voided_at.is_(None)),
                func.coalesce(func.sum(StudentPayment.amount).filter(StudentPayment.voided_at.is_(None)), 0),
                func.count(StudentPayment.id).filter(StudentPayment.voided_at.isnot(None))),
            org_id, first, last, method, group_id, teacher_id, created_by, q)
        count, total, voided = query.one()
        return int(count), int(total), int(voided)

    def journal(
        self, org_id: str, first: date, last: date, *, page: int = 1, size: int = 25, method: Optional[str] = None,
        group_id: Optional[UUID] = None, teacher_id: Optional[UUID] = None, created_by: Optional[UUID] = None,
        q: Optional[str] = None, status: str = "all",
    ) -> PaymentsJournal:
        validate_range(first, last)
        size = max(1, min(size, MAX_JOURNAL_PAGE))
        page = max(1, page)
        count, total_sum, voided = self._journal_totals(org_id, first, last, method, group_id, teacher_id, created_by, q)
        rows_total = {"active": count, "voided": voided}.get(status, count + voided)
        query = self._journal_filters(
            self._journal_base(StudentPayment, Student, Group), org_id, first, last, method, group_id, teacher_id, created_by, q)
        query = self._status_filter(query, status)
        query = self._journal_order(query).offset((page - 1) * size).limit(size)
        rows = self._rows(org_id, query.all())
        authors = self._authors(org_id, first, last)
        return PaymentsJournal(
            as_of=ps.business_today(), **{"from": first}, to=last, page=page, size=size, total=rows_total,
            pages=(rows_total + size - 1) // size, totals=JournalTotals(count=count, sum=total_sum, voided_count=voided),
            rows=rows, authors=authors)

    @staticmethod
    def _status_filter(query, status: str):
        if status == "active":
            return query.filter(StudentPayment.voided_at.is_(None))
        if status in ("voided", "replaced"):
            return query.filter(StudentPayment.voided_at.isnot(None))
        return query

    @staticmethod
    def _journal_order(query):
        return query.order_by(StudentPayment.paid_on.desc(), StudentPayment.created_at.desc(),
                              StudentPayment.receipt_no.desc().nullslast(), StudentPayment.id.desc())

    def _authors(self, org_id: str, first: date, last: date) -> List[JournalAuthor]:
        found = (self.db.query(StudentPayment.created_by, func.count(StudentPayment.id))
                 .filter(StudentPayment.org_id == org_id, StudentPayment.kind == "payment", StudentPayment.voided_at.is_(None),
                         StudentPayment.paid_on >= first, StudentPayment.paid_on <= last)
                 .group_by(StudentPayment.created_by).all())
        users = self.service._users(uid for uid, _ in found)
        out = [JournalAuthor(user_id=str(uid), name=_staff(users.get(uid)), count=int(n)) for uid, n in found]
        out.sort(key=lambda a: (-a.count, a.name or ""))
        return out

    def _rows(self, org_id: str, triples: Sequence[Tuple[StudentPayment, Student, Group]]) -> List[JournalRow]:
        """Journal rows of a set of payments: the ledgers of their memberships (one batch), the correction chain (two queries), the names."""
        payments = [p for p, _s, _g in triples]
        live_members = list(dict.fromkeys(p.membership_id for p in payments if p.voided_at is None))
        ledgers: Dict[UUID, bl.Ledger] = {}
        if live_members:
            loader = ls.LedgerLoader(self.db, str(org_id), ps.business_today())
            ledgers = {e.membership.id: e.ledger for e in loader.run(live_members)}
        ids = [p.id for p in payments]
        was: Dict[Any, StudentPayment] = {}
        replaced_ids = [p.replaces_payment_id for p in payments if p.replaces_payment_id]
        if replaced_ids:
            was = {o.id: o for o in self.db.query(StudentPayment).filter(
                StudentPayment.org_id == org_id, StudentPayment.id.in_(replaced_ids))}
        becomes: Dict[Any, Optional[int]] = {}
        if ids:
            for new in self.db.query(StudentPayment).filter(
                    StudentPayment.org_id == org_id, StudentPayment.replaces_payment_id.in_(ids)).order_by(StudentPayment.created_at):
                becomes[new.replaces_payment_id] = new.receipt_no
        users = self.service._users(x for p, _s, g in triples for x in (p.created_by, p.voided_by, g.teacher_id))
        out: List[JournalRow] = []
        for p, student, group in triples:
            ledger = ledgers.get(p.membership_id) if p.voided_at is None else None
            replaced_by = becomes.get(p.id)
            status = "active" if p.voided_at is None else ("replaced" if p.id in becomes else "voided")
            old = was.get(p.replaces_payment_id)
            out.append(JournalRow(
                id=str(p.id), receipt_no=p.receipt_no, paid_on=p.paid_on, paid_time=_business_time(p.created_at),
                created_at=_iso(p.created_at), student_id=str(student.id),
                student_name=person_name(student.last_name, student.first_name) or "", display_id=student.display_id,
                membership_id=str(p.membership_id), group_id=str(group.id), group_name=group.name,
                teacher_name=_staff(users.get(group.teacher_id)), amount=int(p.amount), method=p.method or "unknown",
                method_label=tx.method_label(p.method), created_by_id=str(p.created_by) if p.created_by else None,
                created_by_name=_staff(users.get(p.created_by)), note=p.note,
                allocation_text=pv.allocation_text(ledger, str(p.id)) if ledger else None,
                allocations=pv.allocations_of(ledger, str(p.id)) if ledger else [],
                refunded=ledger.refunded_from.get(str(p.id), 0) if ledger else 0,
                transferred=ledger.transferred_from.get(str(p.id), 0) if ledger else 0, status=status,
                replaces_receipt_no=old.receipt_no if old else None, replaces_amount=int(old.amount) if old else None,
                replaced_by_receipt_no=replaced_by, void_reason=p.void_reason if p.voided_at else None,
                voided_at=_iso(p.voided_at), voided_by_name=_staff(users.get(p.voided_by))))
        return out

    # ------------------------------------------------------------------ journal file

    def journal_file(
        self, org_id: str, first: date, last: date, *, fmt: str = "xlsx", method: Optional[str] = None,
        group_id: Optional[UUID] = None, teacher_id: Optional[UUID] = None, created_by: Optional[UUID] = None,
        q: Optional[str] = None, status: str = "all",
    ) -> Tuple[bytes, str, str]:
        """``(bytes, content type, filename)`` of the whole filtered journal (xlsx / csv); a cell that starts with = + - @ never runs."""
        validate_range(first, last)
        count, total_sum, voided = self._journal_totals(org_id, first, last, method, group_id, teacher_id, created_by, q)
        if {"active": count, "voided": voided}.get(status, count + voided) > MAX_EXPORT_ROWS:
            raise ExportTooLong("Слишком много оплат для файла — сузьте период или фильтры")
        query = self._journal_filters(
            self._journal_base(StudentPayment, Student, Group), org_id, first, last, method, group_id, teacher_id, created_by, q)
        rows = self._rows(org_id, self._journal_order(self._status_filter(query, status)).all())
        header = ["№ квитанции", "Дата", "Время", "Ученик", "ID ученика", "Группа", "Преподаватель", "Сумма", "Способ", "Кто внёс",
                  "Куда пошли", "Статус", "Примечание"]
        top = [
            ["Журнал оплат", f"с {first.strftime('%d.%m.%Y')} по {last.strftime('%d.%m.%Y')}"],
            ["Оплат", count, "На сумму", total_sum] + (["Отменённых / исправленных", voided] if voided else []),
        ]
        body = []
        for r in rows:
            status_text = STATUS_TEXT[r.status]
            if r.status == "replaced" and r.replaced_by_receipt_no:
                status_text += f" (№ {r.replaced_by_receipt_no})"
            elif r.status == "active" and r.replaces_receipt_no:
                status_text += f" (исправляет № {r.replaces_receipt_no})"
            body.append([r.receipt_no if r.receipt_no is not None else "", r.paid_on.strftime("%d.%m.%Y"), r.paid_time or "",
                         r.student_name, r.display_id or "", r.group_name, r.teacher_name or "", r.amount, r.method_label,
                         r.created_by_name or "", r.allocation_text or "", status_text, r.note or ""])
        base = f"payments-{first.isoformat()}-{last.isoformat()}"
        if fmt == "csv":
            buf = io.StringIO()
            writer = csv.writer(buf, delimiter=";")

            def safe(row):
                return [("'" + v) if isinstance(v, str) and v[:1] in ("=", "+", "-", "@") else v for v in row]

            for line in top:
                writer.writerow(safe(line))
            writer.writerow([])
            writer.writerow(header)
            for line in body:
                writer.writerow(safe(line))
            return ("﻿" + buf.getvalue()).encode("utf-8"), "text/csv; charset=utf-8", base + ".csv"
        from openpyxl import Workbook
        from openpyxl.styles import Font

        wb = Workbook()
        ws = wb.active
        ws.title = "Журнал оплат"

        def put(line):
            ws.append(line)
            for cell in ws[ws.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"          # text stays text: a note starting with "=" is never a formula

        for line in top:
            put(line)
        ws.append([])
        put(header)
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
        for line in body:
            put(line)
        for col, width in zip("ABCDEFGHIJKLM", (12, 12, 8, 28, 14, 22, 22, 14, 16, 22, 52, 22, 36)):
            ws.column_dimensions[col].width = width
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", base + ".xlsx"
