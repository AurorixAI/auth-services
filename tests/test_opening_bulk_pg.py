"""OPENING-BULK on PostgreSQL: the table, the batch preview / apply (atomic, idempotent, versioned), the state of the checklist, rights, the query budget,
lessons mode, left students.  Spec: artifacts/partner_doc_check/opening_bulk_spec.md; frozen contract: tests/fixtures/openings."""
import json
import uuid
from datetime import date, datetime, timedelta
from pathlib import Path

import pytest

from src.auth import ledger_service as ls
from src.auth import payment_status as ps
from src.models.auth import (
    GroupAudit, MembershipOpeningBalance, MembershipTransfer, Organization, Student, StudentGroup,
)

from tests.pg_world import (  # noqa: F401  (fixtures)
    give_group_price, make_client, make_group, make_student, make_user, make_member, pg_session, world,
)
from tests.test_payment_carry_over_pg import engine, run_in_thread  # noqa: F401  (engine = fixture)
from tests.test_pay2_api_pg import (  # noqa: F401
    API, D, TODAY, fixed_today, membership, owner, post_payment, pytestmark, seed_payment, settings_block, stay,
)

FIX = Path(__file__).resolve().parent / "fixtures" / "openings"
BASE = f"{API}/org/opening-balances"


def fixture(name):
    return json.loads((FIX / f"{name}.json").read_text(encoding="utf-8"))


def assert_same_shape(live, golden, path="$"):
    """The live answer has exactly the keys of the golden example (recursively; a list is compared by its first element)."""
    if live is None or golden is None:
        return                                                                    # a nullable field: the example may show the other side
    if isinstance(golden, dict):
        assert isinstance(live, dict), path
        assert set(live) == set(golden), (path, sorted(set(live) ^ set(golden)))
        for key in golden:
            assert_same_shape(live[key], golden[key], f"{path}.{key}")
    elif isinstance(golden, list):
        assert isinstance(live, list), path
        if golden and live:
            assert_same_shape(live[0], golden[0], path + "[0]")


# ---- a centre ---------------------------------------------------------------------------------------------------------------------------------


def centre(w, price=600_000, start="2026-11-01"):
    w.group.schedule = []
    w.db.commit()
    settings_block(w, accounting_start_date=start, billing_anchor="fixed_day", billing_day=1)
    give_group_price(w.db, w.group, price, valid_from=D(2026, 1, 1))


def students(w, n, first=D(2026, 9, 1), group=None, prefix="Уч"):
    """``n`` more students in the group (all enrolled on ``first``); returns their memberships."""
    group = group or w.group
    out = []
    for i in range(n):
        s = make_student(w.db, w.org, group, f"{prefix}{i:03d}", f"Фам{i:03d}", f"B-{uuid.uuid4().hex[:8]}")
        out.append(s)
    w.db.commit()
    ms = []
    for s in out:
        m = membership(w, s, group)
        m.joined_at = datetime(first.year, first.month, first.day, 6)
        (period,) = m.periods
        period.joined_at = m.joined_at
        w.db.commit()
        ms.append(m)
    return ms


def enrolled(w, first=D(2026, 9, 1), student=None):
    m = membership(w, student)
    stay(w, m, first)
    return m


def lst(w, client=None, expect=200, **params):
    r = (client or owner(w)).get(BASE, params=params)
    assert r.status_code == expect, r.text
    return r.json()


def rows_of(data):
    return [row for g in data["groups"] for row in g["memberships"]]


def row_of(w, m, **params):
    for row in rows_of(lst(w, **params)):
        if row["membership_id"] == str(m.id):
            return row
    raise AssertionError("row not in the table")


def ver(w, m):
    return row_of(w, m)["version"]


def it(w, m, **values):
    return {"membership_id": str(m.id), "version": values.pop("version", None) or ver(w, m), **values}


def preview(w, items, reason=None, client=None, expect=200):
    r = (client or owner(w)).post(f"{BASE}/batch/preview", json={"items": items, "reason": reason})
    assert r.status_code == expect, r.text
    return r.json()


def apply(w, items, reason=None, key=None, digest=None, client=None, expect=200):
    if digest is None:
        digest = preview(w, items, reason)["digest"]
    body = {"items": items, "reason": reason, "idempotency_key": key or str(uuid.uuid4()), "digest": digest}
    r = (client or owner(w)).post(f"{BASE}/batch", json=body)
    assert r.status_code == expect, r.text
    return r.json()


def live_openings(w):
    w.db.expire_all()
    return w.db.query(MembershipOpeningBalance).filter_by(org_id=w.org.id, voided_at=None).all()


def audits(w, action=None):
    w.db.expire_all()
    q = w.db.query(GroupAudit).filter(GroupAudit.org_id == w.org.id, GroupAudit.action.in_(["opening_balance_set", "opening_balance_void"]))
    return q.filter_by(action=action).all() if action else q.all()


def ledger_of(w, m):
    w.db.expire_all()
    (e,) = ls.LedgerLoader(w.db, str(w.org.id), ps.business_today()).run([m.id])
    return e.ledger


def settings_of(w):
    w.db.expire_all()
    return dict(w.db.get(Organization, w.org.id).settings or {}).get("payments", {})


# ---- S1: the table ----------------------------------------------------------------------------------------------------------------------------


def test_table_shape_totals_and_the_rules_of_a_row(world, fixed_today):
    centre(world)
    a, b, c = (enrolled(world, D(2026, 9, 1), s) for s in world.students)
    seed_payment(world, c, 100_000, D(2026, 10, 5))
    data = lst(world)
    assert_same_shape(data, fixture("O01_list_pending"))
    assert data["accounting_start"] == "2026-11-01" and data["state"] == "none" and data["period_mode"] == "calendar"
    assert data["paid_until_max"] == (TODAY + timedelta(days=366)).isoformat()
    assert data["totals"] == {"memberships": 3, "with_opening": 0, "missing": 2, "left_available": 0, "shown": 3}
    (block,) = data["groups"]
    assert block["name"] == world.group.name and block["total"] == 3 and block["missing"] == 2 and block["teacher_name"]
    by = {r["membership_id"]: r for r in block["memberships"]}
    ra, rc = by[str(a.id)], by[str(c.id)]
    assert ra["missing"] and ra["version"] == "-" and ra["opening"] is None and ra["editable"] and ra["floor"] == "2026-10-31"
    assert ra["enrolled_on"] == "2026-09-01" and ra["mode"] == "calendar" and ra["price"]["amount"] == 600_000
    assert ra["student"]["name"] and ra["group"]["id"] == str(world.group.id)
    assert rc["has_money_records"] and not rc["missing"]                       # money: not «missing», but editable (with a reason)
    # the filters; the totals stay centre-wide
    only = lst(world, missing_only=True)
    assert {r["membership_id"] for r in rows_of(only)} == {str(a.id), str(b.id)} and only["totals"]["memberships"] == 3
    assert only["totals"]["shown"] == 2
    q = lst(world, q=world.students[1].last_name[:4].lower())
    assert [r["membership_id"] for r in rows_of(q)] == [str(b.id)]
    assert lst(world, group_id=str(world.group.id))["totals"]["shown"] == 3
    assert not lst(world, q="нет такого")["groups"]


def test_table_locked_rows_transfer_free_a_student_who_came_after_the_start_paused_and_archived(world, fixed_today):
    centre(world)
    a, b, c = (enrolled(world, D(2026, 9, 1), s) for s in world.students)
    # b: a transfer in; c: free by an individual price 0
    world.db.add(MembershipTransfer(org_id=world.org.id, student_id=b.student_id, from_membership_id=a.id, to_membership_id=b.id, at=datetime(2026, 10, 10, 5)))
    world.db.commit()
    from tests.pg_world import give_price

    give_price(world.db, c, 0)
    late = students(world, 2, first=D(2026, 11, 5))
    paused = students(world, 1)[0]
    from src.auth.logica_service import PauseService
    from src.auth.payments_schemas import PauseCreate

    PauseService(world.db).create(str(world.org.id), paused.id, str(world.owner.id), PauseCreate(first_day=TODAY - timedelta(days=3), last_day=TODAY + timedelta(days=3)))
    gone = students(world, 1)[0]
    world.db.get(Student, gone.student_id).status = "archived"
    world.db.commit()
    rows = {r["membership_id"]: r for r in rows_of(lst(world))}
    assert str(gone.id) not in rows                                              # archived: not in the table
    assert (rows[str(b.id)]["editable"], rows[str(b.id)]["reason_not_editable"], rows[str(b.id)]["missing"]) == (False, "transfer", False)
    assert (rows[str(c.id)]["editable"], rows[str(c.id)]["reason_not_editable"], rows[str(c.id)]["missing"]) == (False, "free", False)
    assert rows[str(late[0].id)]["missing"] is False and rows[str(late[0].id)]["floor"] == "2026-11-05"   # came after the start: L = first_day
    assert rows[str(paused.id)]["paused"] and rows[str(paused.id)]["missing"]
    assert lst(world)["totals"]["missing"] == 2                                  # a, paused


def test_table_after_a_batch_shows_the_openings_and_the_versions(world, fixed_today):
    centre(world)
    a, b = enrolled(world, D(2026, 9, 1), world.students[0]), enrolled(world, D(2026, 9, 1), world.students[1])
    apply(world, [it(world, a, paid_until="2026-11-30"), it(world, b, debt_amount=550_000, debt_on="2026-10-31")])
    data = lst(world)
    ra, rb = row_of(world, a), row_of(world, b)
    assert ra["opening"] == {"paid_until": "2026-11-30", "debt_amount": None, "debt_on": None} and ra["version"] != "-"
    assert rb["opening"] == {"paid_until": None, "debt_amount": 550_000, "debt_on": "2026-10-31"}
    assert data["totals"]["with_opening"] == 2 and data["totals"]["missing"] == 1
    block = data["groups"][0]
    assert block["with_opening"] == 2 and block["missing"] == 1
    assert ra["version"] == str(next(o.id for o in live_openings(world) if o.membership_id == a.id))


def test_summary_state_and_the_rights(world, fixed_today):
    centre(world)
    enrolled(world, D(2026, 9, 1), world.students[0])
    s = owner(world).get(f"{BASE}/summary").json()
    assert_same_shape(s, fixture("O03_summary_pending"))
    assert s["state"] == "none" and s["totals"]["memberships"] == 3 and s["totals"]["shown"] == 0
    # only done / skipped; pending is the release script's
    for bad in ("pending", "none", "x"):
        assert owner(world).post(f"{BASE}/state", json={"state": bad}).status_code == 422
    out = owner(world).post(f"{BASE}/state", json={"state": "skipped"}).json()
    assert out["state"] == "skipped" and settings_of(world)["openings_state"] == "skipped"
    assert owner(world).post(f"{BASE}/state", json={"state": "done"}).json()["state"] == "done"
    assert lst(world)["state"] == "done"
    # an admin: without the right to money 403, with it - the same as the owner
    admin = make_client(world, user=world.admin)
    assert admin.get(BASE).status_code == 403 and admin.get(f"{BASE}/summary").status_code == 403
    assert admin.post(f"{BASE}/state", json={"state": "done"}).status_code == 403
    assert admin.post(f"{BASE}/batch/preview", json={"items": []}).status_code == 403
    settings_block(world, admin_can_manage=True)
    assert admin.get(BASE).status_code == 200
    assert admin.post(f"{BASE}/state", json={"state": "skipped"}).json()["state"] == "skipped"
    teacher = make_client(world, user=world.teacher)
    for call in (lambda: teacher.get(BASE), lambda: teacher.get(f"{BASE}/summary"), lambda: teacher.post(f"{BASE}/state", json={"state": "done"}),
                 lambda: teacher.get(f"{BASE}/export"), lambda: teacher.post(f"{BASE}/batch/preview", json={"items": []})):
        assert call().status_code == 403
    # another centre: its own state and its own rows
    foreign = make_client(world, user=world.foreign_admin, org="other", org_role="admin")
    assert foreign.get(BASE).status_code == 403                                   # (no right to money there either)
    assert foreign.get(BASE, params={"group_id": str(world.group.id)}).status_code == 403


def test_a_foreign_group_is_not_found_and_foreign_memberships_are_never_listed(world, fixed_today):
    centre(world)
    other_owner = make_user(world.db, "owner")
    make_member(world.db, world.other_org, other_owner, "owner")
    world.db.commit()
    fo = make_client(world, user=other_owner, org="other", org_role="owner")
    assert fo.get(BASE, params={"group_id": str(world.group.id)}).status_code == 404
    mine = owner(world).get(BASE, params={"group_id": str(world.foreign_group.id)})
    assert mine.status_code == 404
    seen = {r["student"]["id"] for r in rows_of(fo.get(BASE).json())}
    assert str(world.foreign_student.id) in seen and not seen & {str(s.id) for s in world.students}
    assert str(world.foreign_student.id) not in {r["student"]["id"] for r in rows_of(lst(world))}
    # a foreign membership in a batch = not found (the same answer as a made-up id), nowhere written
    fm = world.db.query(StudentGroup).filter_by(student_id=world.foreign_student.id).one()
    for body in ({"items": [{"membership_id": str(fm.id), "version": "-", "paid_until": "2026-11-30"}]},
                 {"items": [{"membership_id": str(uuid.uuid4()), "version": "-", "paid_until": "2026-11-30"}]},
                 {"items": [{"membership_id": "not-a-uuid", "version": "-", "paid_until": "2026-11-30"}]}):
        r = owner(world).post(f"{BASE}/batch/preview", json=body)
        assert r.status_code == 404 and r.json()["detail"]["code"] == "NOT_FOUND", r.text
        r = owner(world).post(f"{BASE}/batch", json={**body, "idempotency_key": str(uuid.uuid4()), "digest": "x"})
        assert r.status_code == 404
    assert not live_openings(world)


# ---- S2: preview -------------------------------------------------------------------------------------------------------------------------------


def test_preview_writes_nothing_and_says_what_will_happen(world, fixed_today):
    centre(world)
    a, b, c = (enrolled(world, D(2026, 9, 1), s) for s in world.students)
    before_settings = settings_of(world)
    items = [it(world, a, paid_until="2026-11-30"), it(world, b, debt_amount=550_000, debt_on="2026-10-31"), it(world, c, paid_until="2026-10-31")]
    pv = preview(world, items)
    assert_same_shape(pv, fixture("O04_preview_ok"))
    assert pv["ok"] and len(pv["digest"]) == 64 and pv["accounting_start"] == "2026-11-01" and pv["needs_reason"] == []
    assert pv["summary"] == {"items": 3, "changed": 2, "paid_until": 1, "debt": 1, "cleared": 0, "unchanged": 1, "errors": 0, "debt_total_added": 550_000}
    by = {r["membership_id"]: r for r in pv["items"]}
    ra, rb, rc = by[str(a.id)], by[str(b.id)], by[str(c.id)]
    assert (ra["status"], ra["kind"], ra["changed"], ra["target"]["paid_until"]) == ("ok", "paid_until", True, "2026-11-30")
    assert ra["before"]["debt"] == 0 and ra["after"]["paid_until"] == "2026-11-30" and ra["after"]["opening"]["paid_until"] == "2026-11-30"
    assert (rb["kind"], rb["after"]["debt"], rb["after"]["opening"]["debt_amount"]) == ("debt", 550_000, 550_000)
    assert (rc["status"], rc["kind"], rc["changed"], rc["target"]) == ("noop", "unchanged", False, None)         # «оплачено до 31.10» = nothing owed
    assert [w["code"] for w in rc["warnings"]] == ["PAID_UNTIL_AT_START"]
    imp = pv["impact"]
    assert imp["memberships_changed"] == 2 and imp["memberships_with_new_debt"] == 1 and imp["debt_delta"] == imp["debt_total_after"] - imp["debt_total_before"]
    assert imp["debtors_after"] == imp["debtors_before"] + 1 and imp["debt_delta"] == 550_000 + (rb["after"]["debt"] - 550_000) - rb["before"]["debt"] + (ra["after"]["debt"] - ra["before"]["debt"])
    # nothing was written: no rows, no audit, no change of the settings (the ring of the keys is untouched)
    assert not live_openings(world) and not audits(world)
    assert settings_of(world) == before_settings
    assert preview(world, items)["digest"] == pv["digest"]                        # the same plan, the same digest
    assert preview(world, items[:2])["digest"] != pv["digest"]


def test_preview_item_errors_each_with_its_code(world, fixed_today):
    centre(world)
    a, b, c = (enrolled(world, D(2026, 9, 1), s) for s in world.students)
    seed_payment(world, c, 100_000, D(2026, 10, 5))
    items = [
        it(world, a, paid_until="2026-10-15"),                           # before L
        it(world, b, debt_amount=0, debt_on="2026-10-01"),               # amount
        {"membership_id": str(c.id), "version": "stale-version", "paid_until": "2026-11-30"},
    ]
    pv = preview(world, items)
    assert not pv["ok"] and pv["summary"]["errors"] == 3 and pv["summary"]["changed"] == 0
    got = {r["membership_id"]: r for r in pv["items"]}
    assert (got[str(a.id)]["code"], got[str(a.id)]["status"], got[str(a.id)]["kind"]) == ("OPENING_RANGE", "error", "error")
    assert "внесите как долг" in got[str(a.id)]["message"] and got[str(a.id)]["after"] is None and got[str(a.id)]["target"] is None
    assert got[str(b.id)]["code"] == "AMOUNT_RANGE"
    assert got[str(c.id)]["code"] == "STALE_VERSION"
    # the other codes
    pv = preview(world, [it(world, a), it(world, b, clear=True, paid_until="2026-11-30"), it(world, c, lessons_left=3)])
    codes = {r["membership_id"]: r["code"] for r in pv["items"]}
    assert codes == {str(a.id): "EMPTY_ITEM", str(b.id): "CLEAR_WITH_VALUES", str(c.id): "LESSONS_NOT_APPLICABLE"}
    pv = preview(world, [it(world, a, paid_until="2026-11-30"), it(world, a, debt_amount=5, debt_on="2026-10-31")])
    assert [r["code"] for r in pv["items"]] == ["DUPLICATE_ITEM", "DUPLICATE_ITEM"]
    r = owner(world).post(f"{BASE}/batch/preview", json={"items": []})
    assert r.status_code == 422 and r.json()["detail"]["code"] == "EMPTY_BATCH"
    r = owner(world).post(f"{BASE}/batch/preview", json={"items": [it(world, a, paid_until="2026-11-30")] * 2001})
    assert r.status_code == 422 and r.json()["detail"]["code"] == "TOO_MANY_ITEMS"


def test_debt_without_a_date_gets_the_day_before_the_start_and_money_needs_a_reason(world, fixed_today):
    centre(world)
    a, c = enrolled(world, D(2026, 9, 1), world.students[0]), enrolled(world, D(2026, 9, 1), world.students[2])
    seed_payment(world, c, 100_000, D(2026, 10, 5))
    pv = preview(world, [it(world, a, debt_amount=250_000), it(world, c, paid_until="2026-12-15")])
    got = {r["membership_id"]: r for r in pv["items"]}
    assert got[str(a.id)]["target"] == {"paid_until": None, "debt_amount": 250_000, "debt_on": "2026-10-31"}
    assert got[str(c.id)]["needs_reason"] and got[str(c.id)]["has_money_records"] and pv["needs_reason"] == [str(c.id)]
    assert [w["code"] for w in got[str(c.id)]["warnings"]] == ["HAS_MONEY"]
    assert_same_shape(preview(world, [it(world, c, paid_until="2026-12-15")]), fixture("O05_preview_errors_and_reason"))


# ---- S2: apply --------------------------------------------------------------------------------------------------------------------------------


def test_apply_equals_the_preview_and_the_openings_equal_the_statement(world, fixed_today):
    centre(world)
    a, b, c = (enrolled(world, D(2026, 9, 1), s) for s in world.students)
    items = [it(world, a, paid_until="2026-11-30"), it(world, b, debt_amount=550_000, debt_on="2026-10-31"), it(world, c, paid_until="2026-10-31")]
    pv = preview(world, items)
    out = apply(world, items, digest=pv["digest"])
    assert_same_shape(out, fixture("O07_applied"))
    assert out["replayed"] is False and out["changed"] == 2 and out["summary"] == pv["summary"] and out["impact"] == pv["impact"] and out["state"] == "none"
    # I1: the live ledger equals the "after" of the preview
    for res in pv["items"]:
        if res["status"] == "error":
            continue
        led = ledger_of(world, next(m for m in (a, b, c) if str(m.id) == res["membership_id"]))
        assert (led.debt, led.advance, led.status.status, led.status.paid_until) == (
            res["after"]["debt"], res["after"]["advance"], res["after"]["status"], (date.fromisoformat(res["after"]["paid_until"]) if res["after"]["paid_until"] else None))
    # the openings in the base
    rows = {o.membership_id: o for o in live_openings(world)}
    assert set(rows) == {a.id, b.id} and rows[a.id].paid_until == D(2026, 11, 30) and rows[b.id].debt_amount == 550_000 and rows[b.id].debt_on == D(2026, 10, 31)
    assert rows[a.id].source == "manual" and rows[a.id].created_by == world.owner.id
    # I2: the statement shows the debt and the paid-until; the books reconcile
    st = owner(world).get(f"{API}/org/memberships/{b.id}/statement").json()
    assert st["rows"][0]["type"] == "opening_debt" and st["rows"][0]["debit"] == 550_000 and st["summary"]["reconciles"]
    st = owner(world).get(f"{API}/org/memberships/{a.id}/statement").json()
    assert any(r["type"] == "opening_paid_until" for r in st["rows"]) and st["summary"]["reconciles"]
    # I7: one audit row per changed membership, with the batch in the snapshot; nothing for the unchanged one
    logs = audits(world)
    assert len(logs) == 2 and {x.action for x in logs} == {"opening_balance_set"}
    snap = {x.snapshot["membership_id"]: x.snapshot for x in logs}
    assert snap[str(b.id)]["debt_before"] == 0 and snap[str(b.id)]["debt_after"] == 550_000
    assert {s["batch_id"] for s in snap.values()} == {out["batch_id"]} and {s["via"] for s in snap.values()} == {"batch"}
    # I3: the money did not move
    assert ledger_of(world, a).paid == 0 and ledger_of(world, a).reconciles


def test_apply_is_idempotent_by_the_key_and_a_new_key_with_the_same_values_changes_nothing(world, fixed_today):
    centre(world)
    a, b = enrolled(world, D(2026, 9, 1), world.students[0]), enrolled(world, D(2026, 9, 1), world.students[1])
    items = [it(world, a, paid_until="2026-11-30"), it(world, b, debt_amount=550_000, debt_on="2026-10-31")]
    pv = preview(world, items)
    key = str(uuid.uuid4())
    first = apply(world, items, key=key, digest=pv["digest"])
    again = apply(world, items, key=key, digest=pv["digest"])
    assert_same_shape(again, fixture("O08_applied_replayed"))
    assert again["replayed"] is True and {k: v for k, v in again.items() if k != "replayed"} == {k: v for k, v in first.items() if k != "replayed"}
    assert len(audits(world)) == 2 and len(live_openings(world)) == 2                # nothing written twice
    # the same key with another body
    other = [it(world, a, paid_until="2026-12-15")]
    r = owner(world).post(f"{BASE}/batch", json={"items": other, "idempotency_key": key, "digest": preview(world, other)["digest"]})
    assert r.status_code == 409 and r.json()["detail"]["code"] == "IDEMPOTENCY_MISMATCH"
    # a new key, the same values (the versions are the fresh ones): all unchanged - no write, no audit
    same = [it(world, a, paid_until="2026-11-30"), it(world, b, debt_amount=550_000, debt_on="2026-10-31")]
    pv2 = preview(world, same)
    assert pv2["summary"]["changed"] == 0 and pv2["summary"]["unchanged"] == 2 and pv2["impact"]["memberships_changed"] == 0
    out = apply(world, same, digest=pv2["digest"])
    assert out["changed"] == 0 and len(audits(world)) == 2 and out["batch_id"] != first["batch_id"]
    ring = settings_of(world)["openings_batches"]
    assert len(ring) == 2 and "answer" in ring[0]


def test_stale_versions_409_nothing_written_and_the_preview_goes_stale_when_money_arrives(world, fixed_today):
    centre(world)
    a, b = enrolled(world, D(2026, 9, 1), world.students[0]), enrolled(world, D(2026, 9, 1), world.students[1])
    items = [it(world, a, paid_until="2026-11-30"), it(world, b, debt_amount=550_000, debt_on="2026-10-31")]
    pv = preview(world, items)
    # somebody else changes the row of «a» in between (through the card)
    r = owner(world).put(f"{API}/org/memberships/{a.id}/opening-balance", json={"paid_until": "2026-12-31"})
    assert r.status_code == 200
    body = {"items": items, "idempotency_key": str(uuid.uuid4()), "digest": pv["digest"]}
    r = owner(world).post(f"{BASE}/batch", json=body)
    assert r.status_code == 409 and r.json()["detail"]["code"] == "STALE_ITEMS", r.text
    detail = r.json()["detail"]
    assert detail["membership_ids"] == [str(a.id)] and detail["fresh"][0]["opening"]["paid_until"] == "2026-12-31"
    assert_same_shape(r.json(), fixture("errors")["STALE_ITEMS"]["body"])
    assert len(live_openings(world)) == 1 and len(audits(world)) == 1               # «b» was not written either
    # with the fresh versions and a new preview it goes through
    items = [it(world, a, paid_until="2026-11-30"), it(world, b, debt_amount=550_000, debt_on="2026-10-31")]
    pv = preview(world, items)
    # ... but a payment arrives between the preview and the apply: the plan is no longer the one the client saw
    seed_payment(world, b, 50_000, D(2026, 10, 20))
    r = owner(world).post(f"{BASE}/batch", json={"items": items, "reason": "проверка", "idempotency_key": str(uuid.uuid4()), "digest": pv["digest"]})
    assert r.status_code == 409 and r.json()["detail"]["code"] == "PREVIEW_STALE", r.text
    assert r.json()["detail"]["preview"]["digest"] != pv["digest"]
    assert_same_shape(r.json(), fixture("errors")["PREVIEW_STALE"]["body"])
    assert len(live_openings(world)) == 1


def test_atomic_rollback_one_bad_item_and_a_failure_in_the_middle_of_the_writes(world, fixed_today, monkeypatch):
    centre(world)
    a, b, c = (enrolled(world, D(2026, 9, 1), s) for s in world.students)
    bad = [it(world, a, paid_until="2026-11-30"), it(world, b, debt_amount=5, debt_on="2026-10-31"), it(world, c, paid_until="2026-10-01")]
    pv = preview(world, bad)
    assert not pv["ok"]
    r = owner(world).post(f"{BASE}/batch", json={"items": bad, "idempotency_key": str(uuid.uuid4()), "digest": pv["digest"]})
    assert r.status_code == 422 and r.json()["detail"]["code"] == "BATCH_INVALID", r.text
    assert [i["membership_id"] for i in r.json()["detail"]["items"]] == [str(c.id)] and r.json()["detail"]["items"][0]["code"] == "OPENING_RANGE"
    assert_same_shape(r.json(), fixture("errors")["BATCH_INVALID"]["body"])
    assert not live_openings(world) and not audits(world) and "openings_batches" not in settings_of(world)
    # I6: a failure of the last write rolls everything back (rows, audit, the ring of the keys)
    from src.auth import logica_service

    good = [it(world, a, paid_until="2026-11-30"), it(world, b, debt_amount=550_000, debt_on="2026-10-31"), it(world, c, paid_until="2026-12-15")]
    pv = preview(world, good)
    real, calls = logica_service.OpeningService.apply, {"n": 0}

    def flaky(self, *args, **kwargs):
        calls["n"] += 1
        if calls["n"] == 3:
            raise logica_service.PaymentConflict("boom")
        return real(self, *args, **kwargs)

    monkeypatch.setattr(logica_service.OpeningService, "apply", flaky)
    r = owner(world).post(f"{BASE}/batch", json={"items": good, "idempotency_key": str(uuid.uuid4()), "digest": pv["digest"]})
    assert r.status_code == 409 and calls["n"] == 3
    monkeypatch.setattr(logica_service.OpeningService, "apply", real)
    assert not live_openings(world) and not audits(world) and "openings_batches" not in settings_of(world)
    assert apply(world, good, digest=pv["digest"])["changed"] == 3                  # and the same batch goes through afterwards


def test_a_reason_is_needed_for_a_row_with_money_and_lands_in_the_audit(world, fixed_today):
    centre(world)
    a, c = enrolled(world, D(2026, 9, 1), world.students[0]), enrolled(world, D(2026, 9, 1), world.students[2])
    seed_payment(world, c, 100_000, D(2026, 10, 5))
    items = [it(world, a, paid_until="2026-11-30"), it(world, c, paid_until="2026-12-15")]
    pv = preview(world, items)
    r = owner(world).post(f"{BASE}/batch", json={"items": items, "idempotency_key": str(uuid.uuid4()), "digest": pv["digest"]})
    assert r.status_code == 422 and r.json()["detail"]["code"] == "REASON_REQUIRED" and r.json()["detail"]["membership_ids"] == [str(c.id)]
    assert_same_shape(r.json(), fixture("errors")["REASON_REQUIRED"]["body"])
    r = owner(world).post(f"{BASE}/batch", json={"items": items, "reason": "ab", "idempotency_key": str(uuid.uuid4()), "digest": pv["digest"]})
    assert r.status_code == 422 and r.json()["detail"]["code"] == "REASON_REQUIRED"
    assert not live_openings(world)
    out = apply(world, items, reason="сверка с таблицей центра", digest=pv["digest"])
    assert out["changed"] == 2
    snaps = {x.snapshot["membership_id"]: x.snapshot for x in audits(world)}
    assert snaps[str(c.id)]["reason"] == "сверка с таблицей центра" and snaps[str(a.id)]["reason"] is None
    # money is untouched (I3)
    assert ledger_of(world, c).paid == 100_000


def test_clear_voids_the_live_opening_and_a_batch_equals_n_single_puts(world, fixed_today):
    centre(world)
    ms = students(world, 4)
    a, b, c = (enrolled(world, D(2026, 9, 1), s) for s in world.students)
    # three memberships by the batch, three by single PUTs with the same values (I5)
    singles = ms[:3]
    values = [{"paid_until": "2026-11-30"}, {"debt_amount": 550_000, "debt_on": "2026-10-31"}, {"paid_until": "2026-12-15", "debt_amount": 80_000, "debt_on": "2026-10-31"}]
    apply(world, [it(world, m, **v) for m, v in zip((a, b, c), values)])
    for m, v in zip(singles, values):
        assert owner(world).put(f"{API}/org/memberships/{m.id}/opening-balance", json=v).status_code == 200
    for x, y in zip((a, b, c), singles):
        lx, ly = ledger_of(world, x), ledger_of(world, y)
        assert (lx.debt, lx.advance, lx.covered_through, lx.status.status, lx.charged) == (ly.debt, ly.advance, ly.covered_through, ly.status.status, ly.charged)
    ax = sorted((x.snapshot["debt_before"], x.snapshot["debt_after"], x.snapshot["paid_until"] is None) for x in audits(world) if x.snapshot["membership_id"] in {str(m.id) for m in (a, b, c)})
    ay = sorted((x.snapshot["debt_before"], x.snapshot["debt_after"], x.snapshot["paid_until"] is None) for x in audits(world) if x.snapshot["membership_id"] in {str(m.id) for m in singles})
    assert ax == ay
    # clear
    out = apply(world, [it(world, a, clear=True), it(world, b, paid_until="2026-10-31")])             # b: «оплачено до 31.10» voids the debt (= nothing owed)
    assert out["summary"]["cleared"] == 2
    assert {o.membership_id for o in live_openings(world)} == {c.id} | {m.id for m in singles}
    assert len(audits(world, "opening_balance_void")) == 2
    assert ledger_of(world, a).debt == 0


# ---- the rules of the cases -------------------------------------------------------------------------------------------------------------------


def test_transfer_free_archived_closed_paused_rules_in_the_batch(world, fixed_today):
    centre(world)
    a, b, c = (enrolled(world, D(2026, 9, 1), s) for s in world.students)
    world.db.add(MembershipTransfer(org_id=world.org.id, student_id=b.student_id, from_membership_id=a.id, to_membership_id=b.id, at=datetime(2026, 10, 10, 5)))
    world.db.commit()
    from tests.pg_world import give_price

    give_price(world.db, c, 0)
    p, gone, closed = students(world, 3)
    from src.auth.logica_service import PauseService
    from src.auth.payments_schemas import PauseCreate

    PauseService(world.db).create(str(world.org.id), p.id, str(world.owner.id), PauseCreate(first_day=TODAY - timedelta(days=3), last_day=TODAY + timedelta(days=3)))
    versions = {m.id: ver(world, m) for m in (a, b, c, p, gone, closed)}
    world.db.get(Student, gone.student_id).status = "archived"
    closed.is_active = False
    closed.left_at = datetime(2026, 10, 15, 6)
    (per,) = closed.periods
    per.left_at = closed.left_at
    world.db.commit()
    mk = lambda m, **v: {"membership_id": str(m.id), "version": versions[m.id], **v}
    pv = preview(world, [mk(a, paid_until="2026-11-30"), mk(b, paid_until="2026-11-30"), mk(c, paid_until="2026-11-30"), mk(p, paid_until="2026-11-30"),
                         mk(gone, paid_until="2026-11-30"), mk(closed, debt_amount=100_000)])
    got = {r["membership_id"]: r for r in pv["items"]}
    assert got[str(a.id)]["status"] == "ok"
    assert got[str(b.id)]["code"] == "OPENING_NOT_ALLOWED"
    assert got[str(c.id)]["code"] == "FREE_MEMBERSHIP"
    assert got[str(gone.id)]["code"] == "STUDENT_ARCHIVED"
    assert got[str(p.id)]["status"] == "ok" and got[str(p.id)]["after"]["status"] in ("paused", "active")      # a pause does not forbid an opening
    assert (got[str(closed.id)]["status"], got[str(closed.id)]["kind"]) == ("ok", "debt")                       # left BEFORE the start: only a debt (see the left test)
    # «clear» of a transfer / a free membership (nothing to clear) is simply unchanged
    pv = preview(world, [mk(b, clear=True), mk(c, clear=True)])
    assert [r["status"] for r in pv["items"]] == ["noop", "noop"]
    # nothing was written by any of it
    assert not live_openings(world)


def test_lessons_mode_lessons_left_becomes_the_day_of_the_kth_lesson(world, fixed_today):
    world.group.schedule = [{"id": f"d{d}", "dayOfWeek": d, "startTime": "10:00", "endTime": "11:00", "format": "offline"} for d in (1, 3, 5)]
    world.db.commit()
    settings_block(world, accounting_start_date="2026-11-01", period_mode="lessons", lessons_per_payment=8)
    give_group_price(world.db, world.group, 1_200_000, valid_from=D(2026, 1, 1))
    a, b = enrolled(world, D(2026, 9, 1), world.students[0]), enrolled(world, D(2026, 9, 1), world.students[1])
    assert lst(world)["period_mode"] == "lessons" and row_of(world, a)["mode"] == "lessons"
    pv = preview(world, [it(world, a, lessons_left=8), it(world, b, lessons_left=0)])
    ra, rb = pv["items"]
    # Mon 02.11, Wed 04.11, Fri 06.11, Mon 09.11, Wed 11.11, Fri 13.11, Mon 16.11, Wed 18.11
    assert ra["lessons"] == {"lessons_left": 8, "paid_until": "2026-11-18", "text": "осталось 8 уроков → оплачено до 18.11"}
    assert ra["target"]["paid_until"] == "2026-11-18" and ra["status"] == "ok" and ra["kind"] == "paid_until"
    assert (rb["status"], rb["kind"], rb["lessons"]["paid_until"]) == ("noop", "unchanged", None)                  # nothing is left = no opening
    assert_same_shape(pv, fixture("O06_preview_lessons"))
    out = apply(world, [it(world, a, lessons_left=8)])
    assert out["changed"] == 1
    (o,) = live_openings(world)
    assert o.membership_id == a.id and o.paid_until == D(2026, 11, 18)
    # the same number again = unchanged; «оплачено до» together with the lessons = an error; K out of the range
    pv = preview(world, [it(world, a, lessons_left=8), it(world, b, lessons_left=5, paid_until="2026-11-30"), it(world, b, lessons_left=1000)])
    assert pv["items"][0]["status"] == "noop"
    assert {r["code"] for r in pv["items"][1:]} == {"DUPLICATE_ITEM"}
    pv = preview(world, [it(world, b, lessons_left=5, paid_until="2026-11-30")])
    assert pv["items"][0]["code"] == "LESSONS_AND_PAID_UNTIL"
    pv = preview(world, [it(world, b, lessons_left=1000)])
    assert pv["items"][0]["code"] == "LESSONS_RANGE"
    # a schedule too short for K lessons in a year
    world.group.schedule = [{"id": "d1", "dayOfWeek": 1, "startTime": "10:00", "endTime": "11:00", "format": "offline"}]
    world.db.commit()
    assert preview(world, [it(world, b, lessons_left=399)])["items"][0]["code"] == "LESSONS_RANGE"


def test_left_students_only_a_debt(world, fixed_today):
    centre(world)
    a, b = enrolled(world, D(2026, 9, 1), world.students[0]), world.students[1]
    gone = membership(world, b)
    stay(world, gone, D(2026, 3, 1), D(2026, 6, 30))                                   # left before the start
    after = membership(world, world.students[2])
    stay(world, after, D(2026, 3, 1), D(2026, 10, 20))                                  # left before the start too
    base = lst(world)
    assert base["totals"]["left_available"] == 2 and base["totals"]["memberships"] == 1 and len(rows_of(base)) == 1
    data = lst(world, include_left=1)
    assert_same_shape(data, fixture("O02_list_include_left"))
    rows = {r["membership_id"]: r for r in rows_of(data)}
    rg = rows[str(gone.id)]
    assert (rg["status"], rg["left_on"], rg["only_debt"], rg["editable"], rg["missing"], rg["floor"]) == ("left", "2026-06-30", True, True, False, "2026-10-31")
    assert data["totals"]["missing"] == 1 and data["totals"]["shown"] == 3
    # only a debt: paid_until / lessons_left -> OPENING_RANGE «для ушедших - только долг»
    items = [{"membership_id": str(gone.id), "version": "-", "paid_until": "2026-11-30"}]
    bad = preview(world, items)["items"][0]
    assert bad["code"] == "OPENING_RANGE" and "только долг" in bad["message"]
    out = apply(world, [{"membership_id": str(gone.id), "version": "-", "debt_amount": 300_000}])
    (o,) = live_openings(world)
    assert (o.membership_id, o.debt_amount, o.debt_on, o.paid_until) == (gone.id, 300_000, D(2026, 10, 31), None)
    assert out["impact"]["debt_delta"] == 300_000 and ledger_of(world, gone).debt == 300_000
    assert row_of(world, gone, include_left=1)["opening"]["debt_amount"] == 300_000
    # a left membership of the base table is not editable through a batch when it left AFTER the start
    fixed_today(D(2026, 11, 20))
    late = membership(world, world.students[2])
    stay(world, late, D(2026, 3, 1), D(2026, 11, 10))
    pv = preview(world, [{"membership_id": str(late.id), "version": "-", "debt_amount": 1000}])
    assert pv["items"][0]["code"] == "MEMBERSHIP_CLOSED"
    # clear works for them
    out = apply(world, [{"membership_id": str(gone.id), "version": row_of(world, gone, include_left=1)["version"], "clear": True}])
    assert out["summary"]["cleared"] == 1 and not live_openings(world)


# ---- concurrency ------------------------------------------------------------------------------------------------------------------------------


def _service_apply(s, w, items, digest, reason=None):
    from src.auth.openings_schemas import OpeningBatchItem, OpeningBatchRequest
    from src.auth.openings_service import OpeningsService

    request = OpeningBatchRequest(items=[OpeningBatchItem(**i) for i in items], idempotency_key=uuid.uuid4(), digest=digest, reason=reason)
    return OpeningsService(s).apply(str(w.org.id), str(w.owner.id), "owner", request)


def test_two_admins_on_the_same_row_one_wins_the_other_gets_stale_items(world, fixed_today, engine):
    centre(world)
    a, b = enrolled(world, D(2026, 9, 1), world.students[0]), enrolled(world, D(2026, 9, 1), world.students[1])
    v = ver(world, a)
    one = [{"membership_id": str(a.id), "version": v, "paid_until": "2026-11-30"}]
    two = [{"membership_id": str(a.id), "version": v, "debt_amount": 70_000, "debt_on": "2026-10-31"}]
    d1, d2 = preview(world, one)["digest"], preview(world, two)["digest"]
    t1 = run_in_thread(engine, lambda s: _service_apply(s, world, one, d1))
    t2 = run_in_thread(engine, lambda s: _service_apply(s, world, two, d2))
    for t, _ in (t1, t2):
        t.join(timeout=60)
    outcomes = [t1[1], t2[1]]
    from src.auth.openings_service import StaleItems

    assert sum("value" in o for o in outcomes) == 1, outcomes
    loser = next(o for o in outcomes if "error" in o)["error"]
    assert isinstance(loser, StaleItems) and loser.extra["membership_ids"] == [str(a.id)]
    assert len(live_openings(world)) == 1 and len(audits(world)) == 1


def test_batches_of_different_rows_and_a_payment_in_parallel_all_land_without_a_deadlock(world, fixed_today, engine):
    from src.auth.payments_schemas import PaymentCreate
    from src.auth.payments_service import PaymentError, PaymentService, is_deadlock

    centre(world)
    ms = [enrolled(world, D(2026, 9, 1), s) for s in world.students]
    mk = lambda m: [{"membership_id": str(m.id), "version": "-", "paid_until": "2026-11-30"}]
    digests = [preview(world, mk(m))["digest"] for m in ms[:2]]
    t = [run_in_thread(engine, lambda s, m=m, d=d: _service_apply(s, world, mk(m), d)) for m, d in zip(ms[:2], digests)]
    t.append(run_in_thread(engine, lambda s: PaymentService(s).create_payment(str(world.org.id), ms[2].id, str(world.owner.id), PaymentCreate(amount=100_000, paid_on=TODAY))))
    for th, _ in t:
        th.join(timeout=60)
    for _, o in t:
        err = o.get("error")
        assert err is None or isinstance(err, PaymentError) or is_deadlock(err), err
    assert sum("value" in o for _, o in t[:2]) == 2
    assert len(live_openings(world)) == 2


# ---- the query budget -------------------------------------------------------------------------------------------------------------------------


class Counter:
    def __init__(self, db):
        from sqlalchemy import event

        self.engine, self.n, self.event = db.get_bind(), 0, event
        self.fn = lambda *a, **k: setattr(self, "n", self.n + 1)

    def __enter__(self):
        self.n = 0
        self.event.listen(self.engine, "before_cursor_execute", self.fn)
        return self

    def __exit__(self, *exc):
        self.event.remove(self.engine, "before_cursor_execute", self.fn)


def test_the_query_budget_does_not_depend_on_the_number_of_rows(world, fixed_today):
    centre(world)
    small = [membership(world, s) for s in world.students]
    small += students(world, 17, prefix="Мал")
    teacher2 = world.teacher2
    big_group = make_group(world.db, world.org, teacher2, "Большая")
    world.db.commit()
    give_group_price(world.db, big_group, 600_000, valid_from=D(2026, 1, 1))
    big = students(world, 400, group=big_group, prefix="Бол")
    for m in small:
        stay(world, m, D(2026, 9, 1))
    counts = {}
    for name, ms in (("small", small), ("big", big)):
        items = [{"membership_id": str(m.id), "version": "-", "paid_until": "2026-11-30"} for m in ms]
        with Counter(world.db) as c:
            assert owner(world).get(BASE, params={"group_id": str(ms[0].group_id)}).status_code == 200
        listing = c.n
        with Counter(world.db) as c:
            r = owner(world).post(f"{BASE}/batch/preview", json={"items": items})
        assert r.status_code == 200 and r.json()["ok"] and r.json()["summary"]["changed"] == len(ms)
        pv = c.n
        with Counter(world.db) as c:
            r = owner(world).post(f"{BASE}/batch", json={"items": items, "idempotency_key": str(uuid.uuid4()), "digest": r.json()["digest"]})
        assert r.status_code == 200 and r.json()["changed"] == len(ms), r.text
        counts[name] = (listing, pv, c.n)
    (ls_, ps_, as_), (lb, pb, ab) = counts["small"], counts["big"]
    assert abs(ls_ - lb) <= 2 and abs(ps_ - pb) <= 2 and abs(as_ - ab) <= 2, counts
    assert max(ls_, ps_, as_, lb, pb, ab) < 120, counts
    assert len(live_openings(world)) == len(small) + len(big)


# ---- the release script -----------------------------------------------------------------------------------------------------------------------


def test_the_release_script_dry_run_by_default_and_only_centres_without_money_before_the_start(world, fixed_today):
    import importlib.util

    spec = importlib.util.spec_from_file_location("set_openings_pending", Path(__file__).resolve().parents[1] / "scripts" / "set_openings_pending.py")
    script = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(script)
    # a second centre that has a payment before the start
    fm = world.db.query(StudentGroup).filter_by(student_id=world.foreign_student.id).one()
    from src.models.auth import StudentPayment

    world.db.add(StudentPayment(org_id=world.other_org.id, membership_id=fm.id, student_id=fm.student_id, group_id=fm.group_id, amount=100_000, paid_on=D(2026, 10, 5),
                                created_by=world.foreign_admin.id, method="cash", created_at=datetime(2026, 10, 5, 9)))
    world.db.commit()
    mine = {str(world.org.id), str(world.other_org.id)}
    settings_block(world, accounting_start_date="2026-09-01")
    # DRY RUN: nothing is written
    report = script.run(world.db, execute=False, only=mine)
    assert [o["org_id"] for o in report["changed"]] == [str(world.org.id)] and [o["org_id"] for o in report["skipped_money"]] == [str(world.other_org.id)]
    assert report["skipped_money"][0]["earliest_money_event"] == "2026-10-05" and report["skipped_money"][0]["kind"] == "payment"
    assert "openings_state" not in settings_of(world) and settings_of(world)["accounting_start_date"] == "2026-09-01"
    # EXECUTE: the centre without money is moved, the other is not touched and is listed
    report = script.run(world.db, execute=True, only=mine)
    st = settings_of(world)
    assert (st["accounting_start_date"], st["openings_state"], st["openings_pending_since"]) == ("2026-11-01", "pending", TODAY.isoformat())
    world.db.expire_all()
    other = dict(world.db.get(Organization, world.other_org.id).settings or {}).get("payments", {})
    assert "openings_state" not in other and "accounting_start_date" not in other
    assert lst(world)["state"] == "pending"
    # the second run leaves the centre in its state
    again = script.run(world.db, execute=True, only=mine)
    assert [o["org_id"] for o in again["untouched"]] == [str(world.org.id)] and not again["changed"]
    # «Готово» ends it; the script never brings it back
    assert owner(world).post(f"{BASE}/state", json={"state": "done"}).json()["state"] == "done"
    script.run(world.db, execute=True, only=mine)
    assert settings_of(world)["openings_state"] == "done"


# ---- Excel round trip -------------------------------------------------------------------------------------------------------------------------

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
COL = {"last": 1, "first": 2, "class": 3, "phone": 4, "group": 5, "comment": 6, "paid": 7, "debt": 8, "debt_on": 9, "student": 10, "membership": 11, "none": 12,
       "lessons": 13, "price": 14, "mode": 15, "note": 16}


def export_xlsx(w, expect=200, **params):
    r = owner(w).get(f"{BASE}/export", params=params)
    assert r.status_code == expect, r.text[:200]
    return r


def sheet_of(data):
    import io

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data))
    return wb, wb["Остатки"]


def to_bytes(wb):
    import io

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def row_for(ws, m):
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=COL["membership"]).value == str(m.id):
            return r
    raise AssertionError("no such row")


def import_xlsx(w, data, expect=200, client=None):
    r = (client or owner(w)).post(f"{BASE}/import/preview", files={"file": ("ostatki.xlsx", data, XLSX)})
    assert r.status_code == expect, r.text[:400]
    return r.json()


def build_file(rows, header=None):
    """A file as a user could have made it: the header of ours and ``rows`` = dicts by COL keys."""
    from openpyxl import Workbook

    from src.auth.openings_excel import HEADERS

    wb = Workbook()
    ws = wb.active
    ws.title = "Остатки"
    ws.append(header or HEADERS)
    for row in rows:
        line = [None] * 16
        for key, value in row.items():
            line[COL[key] - 1] = value
        ws.append(line)
    return wb


def test_export_prefilled_columns_types_filters_and_text_cells(world, fixed_today):
    centre(world)
    a, b, c = (enrolled(world, D(2026, 9, 1), s) for s in world.students)
    world.students[0].last_name = "=HYPERLINK(\"http://evil\",\"x\")"
    world.db.commit()
    apply(world, [it(world, a, paid_until="2026-11-30"), it(world, b, debt_amount=550_000, debt_on="2026-10-31")])
    r = export_xlsx(world)
    assert r.headers["content-type"].startswith(XLSX) and "ostatki_na_nachalo_2026-10-21.xlsx" in r.headers["content-disposition"]
    wb, ws = sheet_of(r.content)
    assert wb.sheetnames == ["Остатки", "Инструкция"]
    assert [c.value for c in ws[1]] == [x["header"] for x in fixture("excel")["columns"]]
    ra, rb, rc = row_for(ws, a), row_for(ws, b), row_for(ws, c)
    assert ws.cell(row=ra, column=COL["paid"]).value == datetime(2026, 11, 30) and ws.cell(row=ra, column=COL["paid"]).number_format == "DD.MM.YYYY"
    assert ws.cell(row=rb, column=COL["debt"]).value == 550_000 and ws.cell(row=rb, column=COL["debt_on"]).value == datetime(2026, 10, 31)
    assert ws.cell(row=rc, column=COL["paid"]).value is None and ws.cell(row=rc, column=COL["debt"]).value is None
    assert ws.cell(row=ra, column=COL["student"]).value == str(a.student_id) and ws.cell(row=ra, column=COL["price"]).value == 600_000
    # a surname that looks like a formula stays text (data_type "s"), in the file and when read back
    cell = ws.cell(row=ra, column=COL["last"])
    assert cell.data_type == "s" and cell.value.startswith("=HYPERLINK")
    # the filters
    _, ws2 = sheet_of(export_xlsx(world, missing_only=True).content)
    assert ws2.max_row == 2 and ws2.cell(row=2, column=COL["membership"]).value == str(c.id)
    _, ws3 = sheet_of(export_xlsx(world, group_id=str(world.group.id), q=world.students[1].last_name[:4]).content)
    assert ws3.max_row == 2
    assert export_xlsx(world, group_id=str(world.foreign_group.id), expect=404)
    assert make_client(world, user=world.teacher).get(f"{BASE}/export").status_code == 403


def test_round_trip_unchanged_is_zero_changes_and_one_edit_is_one_change(world, fixed_today):
    centre(world)
    a, b, c = (enrolled(world, D(2026, 9, 1), s) for s in world.students)
    apply(world, [it(world, a, paid_until="2026-11-30"), it(world, b, debt_amount=550_000, debt_on="2026-10-31")])
    data = export_xlsx(world).content
    pv = import_xlsx(world, data)
    assert pv["ok"] and pv["batch_items"] == [] and pv["summary"]["items"] == 0
    assert pv["file"]["changed"] == 0 and pv["file"]["rows_read"] == 3 and pv["file"]["errors"] == 0 and pv["file"]["unchanged"] == 3
    assert {r["status"] for r in pv["rows"]} == {"noop"} and {r["matched_by"] for r in pv["rows"]} == {"membership_id"}
    # a sorted file (the rows in another order) matches the same
    wb, ws = sheet_of(data)
    lines = [[ws.cell(row=r, column=k).value for k in range(1, 17)] for r in range(2, ws.max_row + 1)][::-1]
    shuffled = build_file([dict(zip(COL, line)) for line in lines])
    assert import_xlsx(world, to_bytes(shuffled))["file"]["changed"] == 0
    # one edit = exactly one change
    wb, ws = sheet_of(data)
    ws.cell(row=row_for(ws, c), column=COL["paid"]).value = datetime(2026, 12, 20)
    pv = import_xlsx(world, to_bytes(wb))
    assert pv["file"]["changed"] == 1 and [i["membership_id"] for i in pv["batch_items"]] == [str(c.id)] and pv["summary"]["changed"] == 1
    assert_same_shape(pv, fixture("O09_import_preview"))
    assert pv["batch_items"][0] == {"membership_id": str(c.id), "version": "-", "paid_until": "2026-12-20", "debt_amount": None, "debt_on": None,
                                    "lessons_left": None, "clear": False}
    assert len(live_openings(world)) == 2                                                                      # the preview wrote nothing


def test_excel_apply_equals_the_direct_batch_and_an_emptied_row_clears(world, fixed_today):
    centre(world)
    a, b, c = (enrolled(world, D(2026, 9, 1), s) for s in world.students)
    twins = students(world, 3)
    wb, ws = sheet_of(export_xlsx(world).content)
    ws.cell(row=row_for(ws, a), column=COL["paid"]).value = "30.11.2026"                              # a text date is read too
    ws.cell(row=row_for(ws, b), column=COL["debt"]).value = "550 000"
    ws.cell(row=row_for(ws, c), column=COL["paid"]).value = datetime(2026, 12, 15)
    ws.cell(row=row_for(ws, c), column=COL["debt"]).value = 80_000
    ws.cell(row=row_for(ws, c), column=COL["debt_on"]).value = datetime(2026, 10, 31)
    pv = import_xlsx(world, to_bytes(wb))
    assert pv["ok"] and pv["file"]["changed"] == 3
    out = apply(world, pv["batch_items"], digest=pv["digest"])                                         # the Excel apply = the atomic batch, no re-upload
    values = [{"paid_until": "2026-11-30"}, {"debt_amount": 550_000, "debt_on": "2026-10-31"}, {"paid_until": "2026-12-15", "debt_amount": 80_000, "debt_on": "2026-10-31"}]
    apply(world, [it(world, m, **v) for m, v in zip(twins, values)])
    for x, y in zip((a, b, c), twins):
        lx, ly = ledger_of(world, x), ledger_of(world, y)
        assert (lx.debt, lx.advance, lx.covered_through, lx.status.status) == (ly.debt, ly.advance, ly.covered_through, ly.status.status)
    assert out["summary"] == pv["summary"] and out["impact"] == pv["impact"]
    # emptied cells of a row that has an opening = remove it («уберёт остаток»); «да» in the column does the same
    wb, ws = sheet_of(export_xlsx(world).content)
    for k in ("paid", "debt", "debt_on"):
        ws.cell(row=row_for(ws, c), column=COL[k]).value = None
    ws.cell(row=row_for(ws, b), column=COL["none"]).value = "да"
    ws.cell(row=row_for(ws, b), column=COL["debt"]).value = None
    ws.cell(row=row_for(ws, b), column=COL["debt_on"]).value = None
    pv = import_xlsx(world, to_bytes(wb))
    kinds = {r["membership_id"]: r["kind"] for r in pv["rows"] if r["membership_id"]}
    assert kinds[str(c.id)] == "cleared" and kinds[str(b.id)] == "cleared" and pv["summary"]["cleared"] == 2
    assert apply(world, pv["batch_items"], digest=pv["digest"])["summary"]["cleared"] == 2


def test_excel_matching_and_the_errors_of_the_rows(world, fixed_today):
    centre(world)
    a, b, c = (enrolled(world, D(2026, 9, 1), s) for s in world.students)
    gone = students(world, 1)[0]
    world.db.get(Student, gone.student_id).status = "archived"
    world.db.commit()
    fm = world.db.query(StudentGroup).filter_by(student_id=world.foreign_student.id).one()
    sa, sb, sc = world.students
    rows = [
        {"membership": str(a.id), "last": sa.last_name, "paid": datetime(2026, 11, 30)},                                              # 2: by the membership id
        {"student": str(sb.id), "group": world.group.name.upper(), "paid": datetime(2026, 12, 1)},                                    # 3: by the student id + the group (no case)
        {"last": sc.last_name, "first": sc.first_name, "class": 9, "group": world.group.name, "debt": 100_000},                       # 4: by the name
        {"last": "Несуществующий", "first": "Ученик", "group": world.group.name, "paid": datetime(2026, 11, 30)},                      # 5: unknown
        {"membership": str(a.id), "last": sb.last_name, "paid": datetime(2026, 11, 30)},                                              # 6: ID_MISMATCH (+ the same membership is not applied)
        {"membership": str(gone.id), "paid": datetime(2026, 11, 30)},                                                                  # 7: archived = not in the table
        {"membership": str(fm.id), "paid": datetime(2026, 11, 30)},                                                                    # 8: a foreign centre
        {"membership": str(b.id), "paid": "31.13.2026"},                                                                               # 9: BAD_DATE
        {"membership": str(c.id), "debt": "12abc"},                                                                                    # 10: BAD_AMOUNT
        {"membership": str(c.id), "none": "да", "paid": datetime(2026, 11, 30)},                                                       # 11: ROW_CONFLICT
    ]
    pv = import_xlsx(world, to_bytes(build_file(rows)))
    by = {r["row"]: r for r in pv["rows"]}
    assert (by[2]["status"], by[2]["matched_by"]) == ("ok", "membership_id") and (by[3]["status"], by[3]["matched_by"]) == ("ok", "student_group")
    assert (by[4]["status"], by[4]["matched_by"], by[4]["kind"]) == ("ok", "name", "debt")
    assert (by[5]["code"], by[5]["status"]) == ("UNKNOWN", "error") and "добавьте его" in by[5]["message"]
    assert by[6]["code"] == "ID_MISMATCH" and by[6]["membership_id"] == str(a.id)
    assert by[7]["code"] == "UNKNOWN" and by[8]["code"] == "UNKNOWN" and by[7]["message"] == by[8]["message"]          # a foreign id = a made-up id
    assert by[9]["code"] == "BAD_DATE" and by[10]["code"] == "BAD_AMOUNT" and by[11]["code"] == "ROW_CONFLICT"
    assert all(by[n]["status"] == "error" for n in range(5, 12))
    assert [i["membership_id"] for i in pv["batch_items"]] == [str(a.id), str(b.id), str(c.id)] and pv["ok"]
    f = pv["file"]
    assert (f["rows_read"], f["changed"], f["errors"], f["unknown"]) == (10, 3, 7, 3) and f["missing_in_file"] == 0
    # duplicates: the same membership in two rows - neither is applied
    rows = [{"membership": str(a.id), "paid": datetime(2026, 11, 30)}, {"membership": str(a.id), "paid": datetime(2026, 12, 30)},
            {"membership": str(b.id), "paid": datetime(2026, 11, 30)}]
    pv = import_xlsx(world, to_bytes(build_file(rows)))
    assert [(r["row"], r["code"]) for r in pv["rows"] if r["status"] == "error"] == [(2, "DUPLICATE"), (3, "DUPLICATE")] and "строки 2 и 3" in pv["rows"][0]["message"]
    assert [i["membership_id"] for i in pv["batch_items"]] == [str(b.id)]
    assert pv["file"]["missing_in_file"] == 1                                                                              # c is not mentioned: stays as it was
    # a row whose values the planner refuses: the error of the planner is the error of the row
    pv = import_xlsx(world, to_bytes(build_file([{"membership": str(a.id), "paid": datetime(2026, 10, 1)}])))
    assert pv["rows"][0]["code"] == "OPENING_RANGE" and pv["batch_items"] == [] and pv["summary"]["items"] == 0 and pv["ok"]
    # a row of a student who left: only a debt
    gone2 = membership(world, students(world, 1)[0].student)
    stay(world, gone2, D(2026, 3, 1), D(2026, 6, 30))
    pv = import_xlsx(world, to_bytes(build_file([{"membership": str(gone2.id), "paid": datetime(2026, 11, 30)}, {"membership": str(gone2.id).upper()[:0] + str(gone2.id), "debt": 5}])))
    assert {r["code"] for r in pv["rows"]} == {"DUPLICATE"}


def test_excel_formulas_are_never_executed_and_a_formula_without_a_value_is_an_error(world, fixed_today):
    centre(world)
    a, b = enrolled(world, D(2026, 9, 1), world.students[0]), enrolled(world, D(2026, 9, 1), world.students[1])
    wb = build_file([{"membership": str(a.id), "paid": "=1+1"}, {"membership": str(b.id), "debt": "=SUM(1,2)*100000"}])
    ws = wb["Остатки"]
    assert ws["G2"].data_type == "f"                                                                  # written as a formula (no cached value)
    pv = import_xlsx(world, to_bytes(wb))
    assert [r["code"] for r in pv["rows"]] == ["FORMULA", "FORMULA"] and pv["batch_items"] == []
    assert "формула без значения" in pv["rows"][0]["message"]
    # the same text stored as TEXT is just a text: not a date, never evaluated
    wb = build_file([{"membership": str(a.id)}])
    cell = wb["Остатки"].cell(row=2, column=COL["paid"], value="=cmd|' /C calc'!A0")
    cell.data_type = "s"
    wb["Остатки"].cell(row=2, column=COL["comment"], value="=1+1").data_type = "s"
    pv = import_xlsx(world, to_bytes(wb))
    assert pv["rows"][0]["code"] == "BAD_DATE" and not live_openings(world)


def test_excel_limits_and_unreadable_files(world, fixed_today, monkeypatch):
    import time

    centre(world)
    a, b, c = (enrolled(world, D(2026, 9, 1), s) for s in world.students)
    for data, code in ((b"not a workbook", "FILE_UNREADABLE"), (b"\0" * (5 * 1024 * 1024 + 10), "FILE_TOO_LARGE")):
        r = owner(world).post(f"{BASE}/import/preview", files={"file": ("x.xlsx", data, XLSX)})
        assert r.status_code == 422 and r.json()["detail"]["code"] == code, r.text[:200]
    from openpyxl import Workbook

    other = Workbook()
    other.active.title = "Ученики"
    r = owner(world).post(f"{BASE}/import/preview", files={"file": ("x.xlsx", to_bytes(other), XLSX)})
    assert r.status_code == 422 and r.json()["detail"]["code"] == "SHEET_MISSING"
    assert_same_shape(r.json(), fixture("errors")["SHEET_MISSING"]["body"])
    bare = build_file([], header=["Колонка", "Другая"])
    r = owner(world).post(f"{BASE}/import/preview", files={"file": ("x.xlsx", to_bytes(bare), XLSX)})
    assert r.status_code == 422 and r.json()["detail"]["code"] == "FILE_UNREADABLE"
    # 5 001 rows: refused; 5 000 rows (none of them a student of the centre) are read in seconds
    many = build_file([{"last": f"Ф{i}", "first": "И", "group": "Г", "paid": datetime(2026, 11, 30)} for i in range(5001)])
    r = owner(world).post(f"{BASE}/import/preview", files={"file": ("x.xlsx", to_bytes(many), XLSX)})
    assert r.status_code == 422 and r.json()["detail"]["code"] == "TOO_MANY_ROWS"
    ok = build_file([{"last": f"Ф{i}", "first": "И", "group": "Г", "paid": datetime(2026, 11, 30)} for i in range(5000)])
    started = time.time()
    pv = import_xlsx(world, to_bytes(ok))
    assert time.time() - started < 10 and pv["file"]["rows_read"] == 5000 and pv["file"]["unknown"] == 5000 and pv["batch_items"] == []
    # more than 2 000 changes (the limit lowered for the test)
    from src.auth import openings_excel

    monkeypatch.setattr(openings_excel, "MAX_CHANGES", 2)
    rows = [{"membership": str(m.id), "paid": datetime(2026, 11, 30)} for m in (a, b, c)]
    r = owner(world).post(f"{BASE}/import/preview", files={"file": ("x.xlsx", to_bytes(build_file(rows)), XLSX)})
    assert r.status_code == 422 and r.json()["detail"]["code"] == "TOO_MANY_CHANGES", r.text[:200]
    assert_same_shape(r.json(), fixture("errors")["TOO_MANY_CHANGES"]["body"])
    # the rights
    assert make_client(world, user=world.teacher).post(f"{BASE}/import/preview", files={"file": ("x.xlsx", to_bytes(build_file(rows)), XLSX)}).status_code == 403


def test_excel_lessons_column_left_students_and_the_rights_of_the_batch_after_a_file(world, fixed_today):
    world.group.schedule = [{"id": f"d{d}", "dayOfWeek": d, "startTime": "10:00", "endTime": "11:00", "format": "offline"} for d in (1, 3, 5)]
    world.db.commit()
    settings_block(world, accounting_start_date="2026-11-01", period_mode="lessons", lessons_per_payment=8)
    give_group_price(world.db, world.group, 1_200_000, valid_from=D(2026, 1, 1))
    a = enrolled(world, D(2026, 9, 1), world.students[0])
    gone = membership(world, world.students[1])
    stay(world, gone, D(2026, 3, 1), D(2026, 6, 30))
    wb, ws = sheet_of(export_xlsx(world, include_left=1).content)
    assert ws.cell(row=row_for(ws, a), column=COL["mode"]).value == "уроки" and ws.cell(row=row_for(ws, gone), column=COL["note"]).value == "ушёл: только долг"
    ws.cell(row=row_for(ws, a), column=COL["lessons"]).value = 8
    ws.cell(row=row_for(ws, gone), column=COL["paid"]).value = datetime(2026, 11, 30)
    pv = import_xlsx(world, to_bytes(wb))
    by = {r["membership_id"]: r for r in pv["rows"]}
    assert by[str(a.id)]["status"] == "ok" and by[str(gone.id)]["code"] == "OPENING_RANGE" and "только долг" in by[str(gone.id)]["message"]
    (item,) = pv["batch_items"]
    assert item["lessons_left"] == 8 and pv["items"][0]["lessons"]["text"] == "осталось 8 уроков → оплачено до 18.11"
    assert apply(world, pv["batch_items"], digest=pv["digest"])["changed"] == 1
    (o,) = live_openings(world)
    assert o.paid_until == D(2026, 11, 18)
