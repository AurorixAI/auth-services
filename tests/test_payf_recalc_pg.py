"""PAY-F part 3 on PostgreSQL: the recalculation for the lessons the CENTRE cancelled (scenarios 7-17 of ``PAY_F_SPEC.md``) through the real tables and endpoints.

The kinds are stored on the cancellation (``group_attendance.cancellation_kind``): ``centre`` (the default of a usual cancellation) is recalculated when the
setting is on, ``holiday`` never, a move and an old cancellation (NULL) never.  The deduction is derived by ``run_ledger``: no money row is written for it."""
from datetime import datetime

import pytest
from sqlalchemy import text

from src.auth.payments_service import PaymentService
from src.models.auth import GroupAttendance, StudentGroup, StudentPayment

from tests.ops_helpers import clock, history_row, needs_pg, pg_session  # noqa: F401  (fixtures)
from tests.payf_helpers import API, D, build_org, cancel, centre, client, lesson_id_on, pause, pay, rows, set_payments, statement

pytestmark = needs_pg
LINE = "−46 000 сум: урок 21.10 отменён центром"      # the text of the deduction (a real minus sign)


def month_world(w, clock, *, paid=0, switch_on="2026-10-06", **block):
    """«Свой месяц у каждого ученика»: joined 06.10, the price 600 000, the period 06.10-05.11 has 13 lessons; the setting is switched on on ``switch_on``."""
    clock.set(f"{switch_on or '2026-10-06'} 09:00")
    (g,), (s,), (m,) = centre(w, **block)
    if switch_on:
        r = client(w, "owner").put(f"{API}/org/settings/payments", json={"centre_cancel_recalc": True})
        assert r.status_code == 200 and r.json()["centre_cancel_recalc"] is True, r.text
    if paid:
        pay(w, m, paid, D(2026, 10, 6))
    clock.set("2026-10-25 12:00")
    return g, s, m


def row_of(w, s):
    (row,) = rows(w, s)
    return row


def kind_of(w, g, day, lesson):
    w.db.expire_all()
    return w.db.query(GroupAttendance).filter_by(group_id=g.id, lesson_date=day, lesson_id=lesson).one().cancellation_kind


# ---- the setting: its own history, acts from the day it was made --------------------------------------------------------------------------------------------


def test_settings_default_is_off_and_the_change_is_versioned_from_today(pg_session, clock):
    w = build_org(pg_session)
    clock.set("2026-10-25 12:00")
    centre(w)
    c = client(w, "owner")
    assert c.get(f"{API}/org/settings/payments").json()["centre_cancel_recalc"] is False
    assert c.put(f"{API}/org/settings/payments", json={"centre_cancel_recalc": True}).json()["centre_cancel_recalc"] is True
    w.db.expire_all()
    block = w.db.execute(text("SELECT settings->'payments' FROM organizations WHERE id=:o"), {"o": w.org.id}).scalar_one()
    assert block["recalc_history"] == [{"from": "2026-10-25", "on": True}] and "centre_cancel_recalc" not in block
    assert "billing_history" not in block                               # it is not a billing regime: the periods are not cut
    clock.set("2026-10-27 12:00")
    assert c.put(f"{API}/org/settings/payments", json={"centre_cancel_recalc": False}).json()["centre_cancel_recalc"] is False
    w.db.expire_all()
    assert w.db.execute(text("SELECT settings->'payments'->'recalc_history' FROM organizations WHERE id=:o"), {"o": w.org.id}).scalar_one() == [
        {"from": "2026-10-25", "on": True}, {"from": "2026-10-27", "on": False}]
    assert client(w, "admin").put(f"{API}/org/settings/payments", json={"centre_cancel_recalc": True}).status_code == 403          # the owner's alone


def test_a_cancellation_has_a_kind_centre_by_default_and_holiday_on_request(pg_session, clock):
    w = build_org(pg_session)
    g, s, m = month_world(w, clock)
    cancel(w, g, "lesson-wed", "2026-10-21")
    cancel(w, g, "lesson-fri", "2026-10-23", kind="holiday", reason="Хайит")
    assert kind_of(w, g, "2026-10-21", "lesson-wed") == "centre" and kind_of(w, g, "2026-10-23", "lesson-fri") == "holiday"
    bad = client(w, "owner").post(f"{API}/org/groups/{g.id}/attendance", json={
        "lessonDate": "2026-10-26", "lessonId": "lesson-mon", "status": "cancelled", "records": {}, "cancellationReason": "Нет света", "cancellationKind": "party"})
    assert bad.status_code == 422                                          # only « centre » / « holiday »
    only_cancel = client(w, "owner").post(f"{API}/org/groups/{g.id}/attendance/reconcile/cancel", json={
        "lessonDate": "2026-10-19", "lessonId": "lesson-mon", "cancellationReason": "Преподаватель заболел", "cancellationKind": "holiday"})
    assert only_cancel.status_code == 200, only_cancel.text
    assert kind_of(w, g, "2026-10-19", "lesson-mon") == "holiday"


# ---- 7. «Вычитать»: the lesson of 21.10 cancelled by the centre -------------------------------------------------------------------------------------------


def test_7_the_deduction_is_46000_in_the_row_the_statement_and_the_payment_dialog(pg_session, clock):
    w = build_org(pg_session)
    g, s, m = month_world(w, clock)
    assert row_of(w, s)["debt"] == 600_000
    cancel(w, g, "lesson-wed", "2026-10-21")                               # the usual dialog: the default kind is the centre's fault
    row = row_of(w, s)
    assert row["debt"] == 554_000 and row["advance"] == 0
    st = statement(w, m)
    (line,) = [r for r in st["rows"] if r["type"] == "recalc"]
    assert line["title"] == "Перерасчёт: урок 21.10 отменён центром" and line["credit"] == 46_000 and line["date"] == "2026-10-21"
    assert line["details"]["recalc"]["text"] == LINE and line["balance_after"] == -554_000
    assert st["summary"]["recalculated"] == 46_000 and st["summary"]["reconciles"] is True and st["summary"]["debt"] == 554_000
    preview = client(w, "owner").get(f"{API}/org/memberships/{m.id}/payments/preview").json()
    assert preview["suggested_amount"] == 554_000                              # the next payment is smaller
    assert [r["text"] for r in preview["recalc"]] == [LINE]
    # nothing is stored as money: the same rows of the tables as before
    assert w.db.query(StudentPayment).filter_by(membership_id=m.id).count() == 0
    # the statement FILES carry the line too (xlsx: its own column; csv: a row)
    from io import BytesIO

    from openpyxl import load_workbook

    sheet = load_workbook(BytesIO(client(w, "owner").get(f"{API}/org/memberships/{m.id}/statement/export", params={"format": "xlsx"}).content)).active
    cells = [str(c.value) for row_ in sheet.iter_rows() for c in row_ if c.value is not None]
    assert "Перерасчёт" in cells and any("Перерасчёт: урок 21.10 отменён центром" in c for c in cells)
    assert "Перерасчёт: урок 21.10 отменён центром" in client(w, "owner").get(f"{API}/org/memberships/{m.id}/statement/export", params={"format": "csv"}).text


def test_7_paid_in_full_the_deduction_is_an_advance_and_the_receipt_of_the_next_payment_shows_it_once(pg_session, clock):
    w = build_org(pg_session)
    g, s, m = month_world(w, clock, paid=600_000)
    cancel(w, g, "lesson-wed", "2026-10-21")
    row = row_of(w, s)
    assert row["debt"] == 0 and row["advance"] == 46_000
    clock.set("2026-11-06 12:00")                                              # the second period has begun: 600 000 minus the advance
    owner = client(w, "owner")
    due = owner.get(f"{API}/org/memberships/{m.id}/payments/preview").json()
    assert due["suggested_amount"] == 554_000 and [r["text"] for r in due["recalc"]] == [LINE]
    first = owner.post(f"{API}/org/memberships/{m.id}/payments", json={"amount": 554_000, "paid_on": "2026-11-06", "method": "cash"})
    assert first.status_code == 201, first.text
    receipt = owner.get(f"{API}/org/payments/{first.json()['payment']['id']}/receipt").json()
    assert [r["text"] for r in receipt["recalc"]] == [LINE] and receipt["debt_after"] == 0
    assert row_of(w, s)["debt"] == 0 and statement(w, m)["summary"]["reconciles"] is True
    clock.set("2026-12-04 12:00")
    second = owner.post(f"{API}/org/memberships/{m.id}/payments", json={"amount": 600_000, "paid_on": "2026-12-04", "method": "cash"}).json()
    again = owner.get(f"{API}/org/payments/{second['payment']['id']}/receipt").json()
    assert again["recalc"] == []                                               # the line was already on the previous receipt


# ---- 8. a holiday, 9. «Не пересчитывать», 10. a move ------------------------------------------------------------------------------------------------------


def test_8_a_holiday_is_never_recalculated(pg_session, clock):
    w = build_org(pg_session)
    g, s, m = month_world(w, clock)
    cancel(w, g, "lesson-wed", "2026-10-21", kind="holiday", reason="Хайит")
    assert row_of(w, s)["debt"] == 600_000 and statement(w, m)["summary"]["recalculated"] == 0
    assert not [r for r in statement(w, m)["rows"] if r["type"] == "recalc"]


def test_9_with_the_setting_off_nothing_is_recalculated(pg_session, clock):
    w = build_org(pg_session)
    g, s, m = month_world(w, clock, switch_on=None)
    cancel(w, g, "lesson-wed", "2026-10-21")
    assert kind_of(w, g, "2026-10-21", "lesson-wed") == "centre"              # the fact is recorded; the setting decides
    assert row_of(w, s)["debt"] == 600_000


def test_10_a_moved_lesson_is_neutral(pg_session, clock):
    w = build_org(pg_session)
    g, s, m = month_world(w, clock)
    r = client(w, "admin").post(f"{API}/org/groups/{g.id}/lesson-move", json={
        "lessonId": "lesson-wed", "lessonDate": "2026-10-28", "newDate": "2026-10-29", "startTime": "10:00", "endTime": "11:00", "reason": "Перенос по просьбе центра"})
    assert r.status_code == 200, r.text
    assert kind_of(w, g, "2026-10-28", "lesson-wed") is None                  # the original is cancelled with no kind: a move
    assert row_of(w, s)["debt"] == 600_000 and statement(w, m)["summary"]["recalculated"] == 0


def test_10_an_old_cancellation_without_a_kind_is_never_recalculated(pg_session, clock):
    w = build_org(pg_session)
    g, s, m = month_world(w, clock)
    cancel(w, g, "lesson-wed", "2026-10-21")
    w.db.execute(text("UPDATE group_attendance SET cancellation_kind = NULL WHERE group_id=:g"), {"g": g.id})        # what a cancellation before the release looks like
    w.db.commit()
    assert row_of(w, s)["debt"] == 600_000


# ---- 11. a student who was not charged that day ---------------------------------------------------------------------------------------------------------


def test_11_paused_that_day_or_not_yet_in_the_group_gets_no_line(pg_session, clock):
    w = build_org(pg_session)
    g, s, m = month_world(w, clock)
    pause(w, m, D(2026, 10, 19), D(2026, 10, 23))
    cancel(w, g, "lesson-wed", "2026-10-21")
    cancel(w, g, "lesson-mon", "2026-10-05")                                  # the day before he came
    assert row_of(w, s)["debt"] == 600_000 and statement(w, m)["summary"]["recalculated"] == 0


# ---- 12. the cancellation is taken back -------------------------------------------------------------------------------------------------------------------


def test_12_when_the_lesson_is_given_back_the_line_disappears(pg_session, clock):
    w = build_org(pg_session)
    g, s, m = month_world(w, clock)
    cancel(w, g, "lesson-wed", "2026-10-21")
    assert row_of(w, s)["debt"] == 554_000
    w.db.query(GroupAttendance).filter_by(group_id=g.id, lesson_date="2026-10-21").delete()
    w.db.commit()
    assert row_of(w, s)["debt"] == 600_000 and not [r for r in statement(w, m)["rows"] if r["type"] == "recalc"]


# ---- 13. switched on on 25.10: the cancellation of 19.10 stays as it was -----------------------------------------------------------------------------------


def test_13_the_setting_does_not_rewrite_the_past(pg_session, clock):
    w = build_org(pg_session)
    g, s, m = month_world(w, clock, switch_on=None)
    cancel(w, g, "lesson-mon", "2026-10-19")                                  # before the setting was switched on
    assert row_of(w, s)["debt"] == 600_000
    assert client(w, "owner").put(f"{API}/org/settings/payments", json={"centre_cancel_recalc": True}).status_code == 200      # 25.10
    assert row_of(w, s)["debt"] == 600_000                                    # switching on does not rewrite the past
    cancel(w, g, "lesson-mon", "2026-10-26")
    assert row_of(w, s)["debt"] == 554_000                                    # a lesson from 25.10 on is recalculated
    assert [r["date"] for r in statement(w, m)["rows"] if r["type"] == "recalc"] == ["2026-10-26"]


# ---- 14. an incomplete first period of the «Общий месяц» ------------------------------------------------------------------------------------------------


def test_14_the_partial_first_month_is_price_x_not_cancelled_over_the_template(pg_session, clock):
    """October 2026 has 13 Mon / Wed / Fri lessons by the schedule; the student came on 14.10 and has 8 of them: the charge is price x 8 / 13 (the lessons the centre
    cancels still count in it, LOGIC-A S-4 decision 3); two of his lessons cancelled by the centre cost the price of a lesson each: the total is price x 6 / 13."""
    w = build_org(pg_session)
    clock.set("2026-10-06 09:00")
    (g,), (s,), (m,) = centre(w, price=650_000, joined=datetime(2026, 10, 14, 6, 0), billing_anchor="fixed_day", billing_day=1)
    assert client(w, "owner").put(f"{API}/org/settings/payments", json={"centre_cancel_recalc": True}).status_code == 200
    clock.set("2026-10-25 12:00")
    assert row_of(w, s)["debt"] == 650_000 * 8 // 13 == 400_000
    cancel(w, g, "lesson-wed", "2026-10-21")
    cancel(w, g, "lesson-fri", "2026-10-23")
    cancel(w, g, "lesson-mon", "2026-10-12")                                  # before he came: not his
    row = row_of(w, s)
    assert row["debt"] == 650_000 * 6 // 13 == 300_000
    st = statement(w, m)
    assert [r["credit"] for r in st["rows"] if r["type"] == "recalc"] == [50_000, 50_000] and st["summary"]["reconciles"] is True
    (charge,) = [r for r in st["rows"] if r["type"] == "charge"]
    assert charge["debit"] == 400_000                                         # the charge itself did not move: the cancelled lessons are still in the numerator


# ---- 15. a package ---------------------------------------------------------------------------------------------------------------------------------------


def test_15_a_package_is_not_affected_by_the_setting(pg_session, clock):
    w = build_org(pg_session)
    g, s, m = month_world(w, clock, paid=600_000, period_mode="lessons", lessons_per_payment=12)
    before = row_of(w, s)
    cancel(w, g, "lesson-wed", "2026-10-21")                                  # not held: not charged, not counted - as it always was
    after = row_of(w, s)
    for key in ("debt", "advance", "lessons_total", "lessons_used", "lessons_left", "status"):
        assert before[key] == after[key], key
    assert statement(w, m)["summary"]["recalculated"] == 0


# ---- 16. a transfer in the middle of the period ----------------------------------------------------------------------------------------------------------------


def test_16_after_a_transfer_each_group_recalculates_its_own_lessons(pg_session, clock):
    w = build_org(pg_session)
    g, s, m = month_world(w, clock)
    from tests.ops_helpers import new_group
    from tests.payf_helpers import give_group_price, slots

    other = new_group(w.db, w.org, w.ids, "GB", "Группа B", w.A, slots())
    give_group_price(w.db, other, 600_000, valid_from=D(2026, 1, 1))
    w.db.commit()
    clock.set("2026-10-20 12:00")
    w.svc.add_student_to_group(str(w.org.id), str(other.id), str(s.id), mode="transfer", actor_id=str(w.owner.id))
    w.db.commit()
    clock.set("2026-10-25 12:00")
    cancel(w, g, "lesson-fri", "2026-10-16")           # group A, while he was there
    cancel(w, g, "lesson-wed", "2026-10-21")           # group A, after he left it
    cancel(w, other, "lesson-fri", "2026-10-16")       # group B, before he came
    cancel(w, other, "lesson-fri", "2026-10-23")       # group B, while he is there
    w.db.expire_all()
    ms = {mm.group_id: mm for mm in w.db.query(StudentGroup).filter_by(student_id=s.id)}
    old, new = statement(w, ms[g.id]), statement(w, ms[other.id])
    assert [r["date"] for r in old["rows"] if r["type"] == "recalc"] == ["2026-10-16"]
    assert [r["date"] for r in new["rows"] if r["type"] == "recalc"] == ["2026-10-23"]
    assert old["summary"]["reconciles"] and new["summary"]["reconciles"]


# ---- 17. payroll ------------------------------------------------------------------------------------------------------------------------------------------


def test_17_payroll_reads_the_students_money_as_it_was(pg_session, clock):
    """The deduction pushes «paid until» of the student two days into the next period; the teacher is paid by the lessons the students PAID for, so payroll
    does not see the deduction: its answer is the same with the setting on and off."""
    w = build_org(pg_session)
    g, s, m = month_world(w, clock, paid=600_000)
    cancel(w, g, "lesson-wed", "2026-10-21")
    clock.set("2026-11-06 12:00")
    history_row(w.db, w.org, g, "lesson-fri", "2026-11-06", w.A, [s], start="10:00", end="10:55")
    w.db.commit()
    own = statement(w, m)
    assert own["summary"]["paid_until"] == "2026-11-07"                       # the student's own account: two days into the second period
    on = PaymentService(w.db).lesson_coverage(str(w.org.id), [m.id], D(2026, 11, 6), D(2026, 11, 6))
    assert [(c.paid, c.cover_source) for c in on.values()] == [(False, "unpaid")]
    set_payments(w, recalc_history=[])                                       # the same facts, the setting off
    off = PaymentService(w.db).lesson_coverage(str(w.org.id), [m.id], D(2026, 11, 6), D(2026, 11, 6))
    assert on == off
