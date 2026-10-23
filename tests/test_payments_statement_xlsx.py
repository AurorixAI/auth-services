"""The statement Excel file reads like the page: a golden file built from a hand-made statement (no database), cell by cell."""
import io
from datetime import date

import pytest
from openpyxl import load_workbook

from src.auth.payments_schemas import (
    Allocation, PriceOut, Segment, StatementDetails, StatementMembership, StatementOut, StatementRow, StatementSummary,
)
from src.auth.payments_statement_xlsx import HEADER, build_statement_xlsx

D = date


def seg(start, end, charge, *, days_text=None, basis="full", written_off=0):
    return Segment(id=f"seg:{start}", kind="period", period_from=start, period_to=end, billable_days=1, days_text=days_text,
                   price=PriceOut(amount=500_000, source="override"), charge=charge, basis=basis, written_off=written_off)


def alloc(start, end, amount, *, bucket="current", closes=False, after=None):
    return Allocation(segment_id=f"seg:{start}", period_from=start, period_to=end, amount=amount, days=0, bucket=bucket,
                      closes=closes, label="x", segment_remaining_after=after)


AUG, SEP, OCT, NOV = ((D(2026, m, 1), D(2026, m, 28 if m == 2 else 30 if m in (4, 6, 9, 11) else 31)) for m in (8, 9, 10, 11))


def golden(*, note=None, voided=True) -> StatementOut:
    """Aug (a joiner on the 18th, individual price) - Sep - Oct - the advance for Nov; a payment with three lines, a write-off,
    a refund; a voided payment and a voided refund.  The balance before October is an advance of 74 194."""
    aug = seg(D(2026, 8, 18), AUG[1], 225_806, days_text="14 дней из 31", basis="days")
    sep, oct_, nov = seg(*SEP, 500_000), seg(*OCT, 500_000, written_off=100_000), seg(*NOV, 500_000)
    rows = [
        StatementRow(seq=1, date=D(2026, 8, 18), type="charge", title="s", debit=225_806, balance_after=-225_806, details=StatementDetails(segment=aug)),
        StatementRow(seq=2, date=D(2026, 8, 20), type="payment", title="s", credit=200_000, balance_after=-25_806, details=StatementDetails(
            receipt_no=1, method="cash", created_by_name="Айбек О.", note=note,
            allocations=[alloc(D(2026, 8, 18), AUG[1], 200_000, after=25_806)])),
        StatementRow(seq=3, date=D(2026, 9, 1), type="charge", title="s", debit=500_000, balance_after=-525_806, details=StatementDetails(segment=sep)),
        StatementRow(seq=4, date=D(2026, 9, 5), type="payment", title="s", credit=600_000, balance_after=74_194, details=StatementDetails(
            receipt_no=2, method="card", created_by_name="Айбек О.", allocations=[
                alloc(D(2026, 8, 18), AUG[1], 25_806, bucket="past_debt", closes=True, after=0),
                alloc(*SEP, 500_000, closes=True, after=0),
                alloc(*OCT, 74_194, bucket="advance")])),
        StatementRow(seq=5, date=D(2026, 10, 1), type="charge", title="s", debit=500_000, balance_after=-425_806, details=StatementDetails(segment=oct_)),
        StatementRow(seq=6, date=D(2026, 10, 5), type="writeoff", title="s", credit=100_000, balance_after=-325_806,
                     details=StatementDetails(reason="болел", created_by_name="Директор")),
        StatementRow(seq=7, date=D(2026, 10, 10), type="payment", title="s", credit=325_806, balance_after=0, details=StatementDetails(
            receipt_no=3, method="click", allocations=[alloc(*OCT, 325_806, closes=True, after=0)])),
        StatementRow(seq=8, date=D(2026, 10, 15), type="payment", title="s", credit=100_000, balance_after=100_000, details=StatementDetails(
            receipt_no=4, method="bank_transfer", created_by_name="Айбек О.", allocations=[alloc(*NOV, 100_000, bucket="advance")])),
        StatementRow(seq=9, date=D(2026, 10, 18), type="refund", title="s", outflow=40_000, balance_after=60_000,
                     details=StatementDetails(method="card", reason="переехал", created_by_name="Айбек О.")),
    ]
    cancelled = [
        StatementRow(seq=1, date=D(2026, 9, 3), type="payment", title="Оплата № 99", credit=50_000, voided_at="2026-09-04T10:00:00",
                     void_reason="ошибка ввода", voided_by_name="Директор",
                     details=StatementDetails(receipt_no=99, method="cash", replaced_by_receipt_no=100)),
        StatementRow(seq=2, date=D(2026, 10, 20), type="refund", title="Возврат аванса", outflow=5_000, voided_at="2026-10-21T10:00:00",
                     void_reason="передумали", details=StatementDetails(method="card")),
    ] if voided else []
    return StatementOut(
        membership=StatementMembership(membership_id="m", student_id="s", student_name="Иванов Иван", display_id="A-0042",
                                       group_id="g", group_name="Математика 9", teacher_name="Петрова Анна"),
        as_of=D(2026, 10, 21), accounting_start=D(2026, 8, 1),
        summary=StatementSummary(charged=1_225_806, paid=1_225_806, written_off=100_000, refunded=40_000, advance=60_000),
        rows=rows, voided=cancelled, upcoming=[nov])


def sheet(statement, **kw):
    return load_workbook(io.BytesIO(build_statement_xlsx(statement, **kw))).active


def title_block(ws):
    out = {}
    for r in range(1, 40):
        label = ws.cell(r, 1).value
        if label == HEADER[0]:
            return out, r
        if label and ws.cell(r, 2).value is not None:
            out[label] = ws.cell(r, 2)
    raise AssertionError("no header row")


def grid(ws, first):
    return [[ws.cell(r, c).value for c in range(1, 7)] for r in range(first, ws.max_row + 1)]


def test_title_block_header_freeze_and_widths():
    ws = sheet(golden())
    block, head = title_block(ws)
    assert ws.title == "Выписка" and ws["A1"].value == "Выписка" and ws["A1"].font.bold
    assert {k: v.value for k, v in block.items() if k not in ("По состоянию на", "Учёт ведётся с")} == {
        "Ученик": "Иванов Иван", "ID": "A-0042", "Группа": "Математика 9", "Преподаватель": "Петрова Анна", "Период": "Все месяцы",
        "Стоимость обучения": 1_225_806, "Оплачено": 1_225_806, "Прощено": 100_000, "Возвращено": 40_000, "Баланс": "Аванс 60 000"}
    for label, day in (("По состоянию на", D(2026, 10, 21)), ("Учёт ведётся с", D(2026, 8, 1))):
        assert block[label].value.date() == day and block[label].number_format == "DD.MM.YYYY"
    assert block["Стоимость обучения"].number_format == "#,##0" and block["Баланс"].data_type == "s"
    assert [ws.cell(head, c).value for c in range(1, 7)] == HEADER == ["Дата", "Операция", "Стоимость", "Оплачено", "Баланс", "Подробности"]
    assert all(ws.cell(head, c).font.bold and ws.cell(head, c).fill.fgColor.rgb not in (None, "00000000") for c in range(1, 7))
    assert ws.freeze_panes == f"A{head + 1}"
    widths = [ws.column_dimensions[c].width for c in "ABCDEF"]
    assert all(w >= 12 for w in widths) and widths[1] <= 60 and widths[5] <= 60
    # no #### : a date needs 10, the biggest sum 9 characters with the separators, the balance text 13
    assert widths[0] >= 12 and widths[2] >= 12 and widths[3] >= 12 and widths[4] >= len("Долг 425 806") + 2
    assert ws.column_dimensions["B"].width >= len("Оплата № 4 · перевод · внёс Айбек О.")
    assert ws.cell(head, 1).alignment.horizontal == "center" and ws.cell(head, 4).alignment.horizontal == "center"


def test_golden_all_months():
    ws = sheet(golden(note="за двоих"), include_voided=True)
    _, head = title_block(ws)
    got = grid(ws, head + 1)
    d = lambda y, m, dd: __import__("datetime").datetime(y, m, dd)
    assert got == [
        ["Август 2026", None, None, None, None, None],
        [d(2026, 8, 18), "Август 2026 — неполный месяц (14 дней из 31)", 225_806, None, "Долг 225 806",
         "цена 500 000 сум в месяц · индивидуальная цена"],
        [d(2026, 8, 20), "Оплата № 1 · наличные · внёс Айбек О.", None, 200_000, "Долг 25 806",
         "комментарий: за двоих\nза август — 200 000 (осталось 25 806)"],
        ["Сентябрь 2026", None, None, None, None, None],
        [d(2026, 9, 1), "Сентябрь 2026", 500_000, None, "Долг 525 806", "цена 500 000 сум в месяц · индивидуальная цена"],
        [d(2026, 9, 5), "Оплата № 2 · карта · внёс Айбек О.", None, 600_000, "Аванс 74 194",
         "за август — 25 806 (август закрыт)\nза сентябрь — 500 000 (сентябрь закрыт)\nаванс на октябрь — 74 194"],
        ["Октябрь 2026", None, None, None, None, None],
        [d(2026, 10, 1), "Октябрь 2026", 500_000, None, "Долг 425 806", "цена 500 000 сум в месяц · индивидуальная цена"],
        [d(2026, 10, 5), "Долг прощён — причина: болел · Директор", None, 100_000, "Долг 325 806", None],
        [d(2026, 10, 10), "Оплата № 3 · click", None, 325_806, "0", "за октябрь — 325 806 (октябрь закрыт)"],
        [d(2026, 10, 15), "Оплата № 4 · перевод · внёс Айбек О.", None, 100_000, "Аванс 100 000", "аванс на ноябрь — 100 000"],
        [d(2026, 10, 18), "Возврат аванса — карта · Айбек О.", None, -40_000, "Аванс 60 000", "причина: переехал"],
        [None] * 6,
        ["Отменённые записи", None, None, None, None, None],
        [d(2026, 9, 3), "Оплата № 99", None, 50_000, None, "Отменено: ошибка ввода · Директор · заменена оплатой № 100"],
        [d(2026, 10, 20), "Возврат аванса", None, -5_000, None, "Отменено: передумали"],
    ]


def test_cell_types_formats_alignment_and_merges():
    ws = sheet(golden())
    _, head = title_block(ws)
    first = head + 1
    merged = {str(r) for r in ws.merged_cells.ranges}
    assert {f"A{first}:F{first}", f"A{first + 3}:F{first + 3}", f"A{first + 6}:F{first + 6}"} <= merged     # a band per month
    band = ws.cell(first, 1)
    assert band.font.bold and band.fill.fgColor.rgb not in (None, "00000000") and band.alignment.horizontal == "left"
    # rows: the second row of the table is the August charge
    r = first + 1
    date_cell, op, charged, paid, balance, details = (ws.cell(r, c) for c in range(1, 7))
    assert date_cell.is_date and date_cell.number_format == "DD.MM.YYYY" and date_cell.alignment.horizontal == "center"
    assert isinstance(charged.value, int) and charged.number_format == "#,##0" and charged.alignment.horizontal == "center"
    assert paid.value is None and balance.data_type == "s" and balance.value == "Долг 225 806" and balance.alignment.horizontal == "center"
    assert op.data_type == "s" and op.alignment.wrap_text and op.alignment.vertical == "top"
    assert details.alignment.wrap_text and details.alignment.vertical == "top"
    pay = first + 2
    assert isinstance(ws.cell(pay, 4).value, int) and ws.cell(pay, 4).number_format == "#,##0" and ws.cell(pay, 3).value is None
    refund = next(rr for rr in range(first, ws.max_row + 1) if str(ws.cell(rr, 2).value).startswith("Возврат аванса —"))
    assert ws.cell(refund, 4).value == -40_000 and ws.cell(refund, 4).number_format == "#,##0"
    # the multi-line details are one cell with a taller row
    alloc_row = next(rr for rr in range(first, ws.max_row + 1) if str(ws.cell(rr, 2).value).startswith("Оплата № 2"))
    assert ws.cell(alloc_row, 6).value.count("\n") == 2 and ws.row_dimensions[alloc_row].height >= 45
    # no per-month totals, no old words
    flat = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
    assert not any("Начисл" in x for x in flat)
    assert not any(x in flat for x in ("cash", "card", "bank_transfer", "click", "payme", "№", "Возвращено / перенесено", "Оплачено / прощено"))
    assert not any(x.startswith("Итого") for x in flat)
    assert "ID" in {ws.cell(rr, 1).value for rr in range(1, head)}


def test_the_cancelled_block_is_below_in_grey_and_struck():
    ws = sheet(golden(), include_voided=True)
    rows = list(range(1, ws.max_row + 1))
    head = next(r for r in rows if ws.cell(r, 1).value == "Отменённые записи")
    assert head > max(r for r in rows if ws.cell(r, 2).value == "Возврат аванса — карта · Айбек О.")
    assert any(str(rg) == f"A{head}:F{head}" for rg in ws.merged_cells.ranges) and ws.cell(head, 1).font.bold
    for r in (head + 1, head + 2):
        for c in range(1, 5):
            cell = ws.cell(r, c)
            if cell.value is not None:
                assert cell.font.strike and cell.font.color.rgb.endswith("808080")
        assert ws.cell(r, 6).font.color.rgb.endswith("808080") and ws.cell(r, 6).value.startswith("Отменено:")
    # the cancelled rows are never mixed into the months: the month bands carry no «Отменено» text
    assert not any("Отменено" in str(ws.cell(r, 6).value) for r in range(1, head))
    # without include_voided there is no block at all
    ws2 = sheet(golden(voided=False), include_voided=False)
    assert not any(ws2.cell(r, 1).value == "Отменённые записи" for r in range(1, ws2.max_row + 1))
    ws3 = sheet(golden(voided=False), include_voided=True)
    assert [ws3.cell(r, 1).value for r in range(ws3.max_row - 1, ws3.max_row + 1)] == ["Отменённые записи", "Отменённых записей нет"]


def test_a_single_month_has_its_carry_line_and_the_period_in_the_title():
    ws = sheet(golden(), date_from=D(2026, 10, 1), date_to=D(2026, 10, 31), include_voided=True)
    block, head = title_block(ws)
    assert block["Период"].value == "Октябрь 2026"
    # the title totals stay for the whole membership
    assert block["Стоимость обучения"].value == 1_225_806 and block["Баланс"].value == "Аванс 60 000"
    got = grid(ws, head + 1)
    assert got[0][0] == "Октябрь 2026" and got[1][0] == "С сентября перенесён аванс 74 194"
    assert [x for x in ws.merged_cells.ranges if x.min_row == head + 2] and ws.cell(head + 2, 1).font.italic
    body = [g for g in got[2:] if g[1] and not str(g[0]).startswith("Отменённые")]
    assert [g[1] for g in body[:6]] == ["Октябрь 2026", "Долг прощён — причина: болел · Директор", "Оплата № 3 · click",
                                         "Оплата № 4 · перевод · внёс Айбек О.", "Возврат аванса — карта · Айбек О.", "Возврат аванса"]
    # months outside the range are gone, also from the cancelled block (the September payment)
    assert not any(isinstance(g[0], str) and g[0] in ("Август 2026", "Сентябрь 2026") for g in got)
    assert not any(g[1] == "Оплата № 99" for g in got)
    # the balance is the ledger's, not recomputed over the cut rows: a charge row keeps its running debt
    assert got[2][4] == "Долг 425 806"


def test_range_label_and_a_month_without_records_or_carry():
    ws = sheet(golden(), date_from=D(2026, 10, 1), date_to=D(2026, 10, 15))
    block, head = title_block(ws)
    assert block["Период"].value == "01.10.2026–15.10.2026"
    assert ws.cell(head + 1, 1).value == "Октябрь 2026" and not str(ws.cell(head + 2, 1).value).startswith("С сентября")   # not a month: no carry
    # a month before the first record: no carry line; an empty month after: carry and a quiet «Записей нет»
    ws = sheet(golden(), date_from=D(2026, 8, 1), date_to=D(2026, 8, 31))
    _, head = title_block(ws)
    assert [ws.cell(head + 1, 1).value, str(ws.cell(head + 2, 1).value)[:4]] == ["Август 2026", "2026"]
    ws = sheet(golden(), date_from=D(2026, 12, 1), date_to=D(2026, 12, 31))
    _, head = title_block(ws)
    assert [ws.cell(r, 1).value for r in (head + 1, head + 2, head + 3)] == ["Декабрь 2026", "С ноября перенесён аванс 60 000", "Записей нет"]
    ws = sheet(golden(), date_from=D(2026, 7, 1), date_to=D(2026, 7, 31))
    _, head = title_block(ws)
    assert [ws.cell(r, 1).value for r in (head + 1, head + 2)] == ["Июль 2026", "Записей нет"]


def test_a_hostile_text_stays_text_everywhere():
    st = golden(note='=HYPERLINK("http://evil")')
    st.membership.student_name = "=1+1"
    st.membership.teacher_name = "@SUM(1)"
    st.rows[5].details.reason = "=cmd|' /C calc'!A0"
    ws = sheet(st)
    hot = [c for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value[:1] in "=+-@" or
           (isinstance(c.value, str) and "=HYPERLINK" in c.value)]
    assert hot and all(c.data_type == "s" for c in hot)
    assert all(c.data_type != "f" for row in ws.iter_rows() for c in row)


def test_an_empty_statement():
    st = golden()
    st.rows, st.voided, st.upcoming = [], [], []
    st.summary = StatementSummary()
    ws = sheet(st)
    block, head = title_block(ws)
    assert block["Баланс"].value == "0" and ws.cell(head + 1, 1).value == "В выписке пока нет записей"
