"""Pack A2, fix round after review: a zip whose directory LIES about sizes, a read-only sheet with a wrong <dimension>, 503 for the display-id
prefix claim, X-Request-ID on the upload limit's own answers.  No database (the import itself is covered in test_a2_import_limits_pg.py)."""
import io
import os
import re
import struct
import tracemalloc
import uuid
import zipfile
import zlib
from unittest.mock import MagicMock

import pytest
from openpyxl import Workbook, load_workbook

os.environ.setdefault("SECRET_KEY", "test-secret-key-must-be-at-least-32-chars-ok")
os.environ.setdefault("JWT_SECRET_KEY", "test-jwt-secret-key-must-be-at-least-32")
os.environ.setdefault("APP_ENV", "testing")
os.environ.setdefault("DATABASE_URL", "postgresql://algo:algo@localhost:5432/algo_test")

from fastapi.testclient import TestClient  # noqa: E402

import src.api.internal_router  # noqa: E402,F401  (import order: avoids a pre-existing router cycle)
from src.api import main_v2  # noqa: E402
from src.auth import import_limits as limits  # noqa: E402
from src.auth import import_service as legacy  # noqa: E402
from src.auth import student_id_service  # noqa: E402
from src.auth.dependencies import get_current_active_user_payload  # noqa: E402
from src.auth.service import AuthService  # noqa: E402
from src.core.database import get_db  # noqa: E402

APP = main_v2.app
GH = ["Название группы*", "Класс (1-11)", "Почта учителя", "Макс. учеников"]
SH = ["Фамилия*", "Имя*", "Класс (1-11)*", "Телефон", "Группа*", "Комментарий"]
MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def book(n_groups=2, n_students=5):
    wb = Workbook()
    wb.remove(wb.active)
    g = wb.create_sheet("Группы")
    g.append(GH)
    for i in range(n_groups):
        g.append([f"Г-{i + 1}", 5, None, 30])
    s = wb.create_sheet("Ученики")
    s.append(SH)
    for i in range(n_students):
        s.append([f"Фам{i:03d}", f"Имя{i:03d}", 5, None, f"Г-{i % n_groups + 1}", None])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def rewrite(xlsx, edit, compression=zipfile.ZIP_DEFLATED):
    """The workbook with ``edit(name, bytes) -> bytes`` applied to each part."""
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(xlsx)) as src, zipfile.ZipFile(out, "w", compression) as dst:
        for info in src.infolist():
            dst.writestr(info.filename, edit(info.filename, src.read(info.filename)))
    return out.getvalue()


def lying_bomb(xlsx, real_extra: int, declared: int = 500, fix_crc: bool = True) -> bytes:
    """A valid-looking xlsx whose sheet part inflates to ``real_extra`` more bytes than its central directory (size and CRC) admits."""
    blob = rewrite(xlsx, lambda n, d: d + b" " * real_extra if n == "xl/worksheets/sheet2.xml" else d)
    raw = bytearray(blob)
    pos = 0
    while (pos := raw.find(b"PK\x01\x02", pos)) >= 0:
        name_len = struct.unpack("<H", raw[pos + 28:pos + 30])[0]
        if bytes(raw[pos + 46:pos + 46 + name_len]) == b"xl/worksheets/sheet2.xml":
            raw[pos + 24:pos + 28] = struct.pack("<I", declared)
            if fix_crc:                                                   # the attacker also forges the CRC of the first ``declared`` bytes
                offset = struct.unpack("<I", raw[pos + 42:pos + 46])[0]
                name2, extra2 = struct.unpack("<HH", raw[offset + 26:offset + 30])
                csize = struct.unpack("<I", raw[pos + 20:pos + 24])[0]
                start = offset + 30 + name2 + extra2
                head = zlib.decompressobj(-15).decompress(bytes(raw[start:start + csize]), declared)
                raw[pos + 16:pos + 20] = struct.pack("<I", zlib.crc32(head))
        pos += 4
    return bytes(raw)


# ------------------------------------------------------------------------------------------------------ 1. a lying zip directory
def test_a_zip_that_lies_about_its_size_is_counted_by_its_real_bytes():
    bomb = lying_bomb(book(), real_extra=limits.MAX_UNPACKED_BYTES + 5 * 1024 * 1024)
    zf = zipfile.ZipFile(io.BytesIO(bomb))
    assert zf.getinfo("xl/worksheets/sheet2.xml").file_size == 500                               # what the directory claims ...
    assert len(bomb) < 200 * 1024                                                       # ... of a file that is small on the wire
    with pytest.raises(limits.ImportFileRejected) as refused:
        limits.check_zip(bomb)
    assert refused.value.args[0] == limits.TOO_BIG_UNPACKED


def test_the_guard_holds_at_most_a_chunk_while_it_counts():
    bomb = lying_bomb(book(), real_extra=limits.MAX_UNPACKED_BYTES * 3)                 # 120 MB of inflated spaces behind a ~1 KB claim
    tracemalloc.start()
    try:
        with pytest.raises(limits.ImportFileRejected):
            limits.check_zip(bomb)
        peak = tracemalloc.get_traced_memory()[1]
    finally:
        tracemalloc.stop()
    assert peak < 2 * 1024 * 1024, peak                                                 # a few 64 KB chunks, not 40 MB


@pytest.mark.parametrize("declared,fix_crc", [(500, False), (10**7, True)])
def test_a_part_that_is_not_what_the_directory_says_is_unreadable_not_a_crash(declared, fix_crc):
    small = lying_bomb(book(), real_extra=2000, declared=declared, fix_crc=fix_crc)       # real size ~2.5 KB: no bomb, just a wrong claim
    with pytest.raises(limits.ImportFileRejected) as refused:
        limits.check_zip(small)
    assert refused.value.args[0] == limits.UNREADABLE


def test_the_http_answer_for_a_lying_bomb_is_422():
    from src.auth.router import import_preview  # noqa: F401  (the route exists)

    bomb = lying_bomb(book(), real_extra=limits.MAX_UNPACKED_BYTES + 1024 * 1024)
    payload = {"sub": str(uuid.uuid4()), "role": "owner", "type": "access", "org_id": str(uuid.uuid4()), "org_role": "owner"}
    APP.dependency_overrides[get_current_active_user_payload] = lambda: payload
    APP.dependency_overrides[get_db] = lambda: MagicMock()
    try:
        r = TestClient(APP, raise_server_exceptions=False).post("/api/v1/auth/org/import/preview", files={"file": ("x.xlsx", bomb, MIME)})
    finally:
        APP.dependency_overrides.clear()
    assert r.status_code == 422 and "после распаковки" in r.json()["detail"]


def test_honest_files_still_pass_whatever_the_compression():
    for compression in (zipfile.ZIP_DEFLATED, zipfile.ZIP_STORED):
        blob = rewrite(book(3, 12), lambda n, d: d, compression)
        limits.check_zip(blob)
        parsed = legacy.parse_workbook(blob)
        assert len(parsed.groups) == 3 and len(parsed.students) == 12
    limits.check_zip(book())                                                            # as openpyxl writes it


# ------------------------------------------------------------------------------------------------------ 2. a wrong <dimension>
def reference_parse(data):
    """The old, normal-mode reading (every cell, the real row number of the first cell) - what the importer did before the streaming reader."""
    wb = load_workbook(io.BytesIO(data), data_only=True)
    out = {}
    for sheet, cols in (("Группы", legacy.COLS_GROUPS), ("Ученики", legacy.COLS_STUDENTS)):
        ws = wb[sheet]
        rows_iter = ws.iter_rows(values_only=False)
        header = next(rows_iter)
        wanted = {legacy._norm_header(c): c for c in cols}
        col_map = {i: legacy.HEADER_ALIASES.get(legacy._norm_header(c.value), legacy._norm_header(c.value)) for i, c in enumerate(header)}
        col_map = {i: k for i, k in col_map.items() if k in wanted}
        rows = []
        for cells in rows_iter:
            if legacy._is_example(cells):
                continue
            values = {k: legacy._cell(cells[i].value if i < len(cells) else None) for i, k in col_map.items()}
            if any(values.values()):
                rows.append((cells[0].row, values))
        out[sheet] = rows
    return out


def with_dimension(data, ref):
    def edit(name, blob):
        if name.startswith("xl/worksheets/sheet"):
            blob = re.sub(rb"<dimension[^>]*/>", b"", blob)
            if ref is not None:
                blob = re.sub(rb"(<worksheet[^>]*>)", lambda m: m.group(1) + b'<dimension ref="' + ref.encode() + b'"/>', blob, count=1)
        return blob

    return rewrite(data, edit)


@pytest.mark.parametrize("ref", ["A1", "A1:C3", None, "A1:H1048576", "A1:A1", "Z9:Z9"])
def test_a_wrong_dimension_loses_no_rows_or_columns(ref):
    data = with_dimension(book(4, 30), ref)
    expected = reference_parse(data)
    parsed = legacy.parse_workbook(data)
    assert [(r.row, r.data) for r in parsed.groups] == expected["Группы"] and len(parsed.groups) == 4
    assert [(r.row, r.data) for r in parsed.students] == expected["Ученики"] and len(parsed.students) == 30


def test_the_dimension_test_really_breaks_the_unfixed_reader():
    """Guard for the test itself: without ``reset_dimensions`` a read-only sheet with dimension="A1" yields nothing."""
    data = with_dimension(book(4, 30), "A1")
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    assert len([r for r in wb["Ученики"].iter_rows(values_only=True)]) <= 1


# ------------------------------------------------------------------------------------------------------ 3. RuntimeError -> 503
PREFIX_TEXT = "Не удалось выделить уникальный префикс ID учеников: слишком много параллельных попыток"


@pytest.fixture
def owner_http():
    payload = {"sub": str(uuid.uuid4()), "role": "owner", "type": "access", "org_id": str(uuid.uuid4()), "org_role": "owner"}
    APP.dependency_overrides[get_current_active_user_payload] = lambda: payload
    APP.dependency_overrides[get_db] = lambda: MagicMock()
    yield TestClient(APP, raise_server_exceptions=False)
    APP.dependency_overrides.clear()


def boom(*a, **k):
    raise RuntimeError(PREFIX_TEXT)


ROUTES = [
    ("get", "/api/v1/auth/org/students/s1/profile", None, "get_student_profile"),
    ("put", "/api/v1/auth/org/students/s1", {"first_name": "Али"}, "update_student"),
    ("get", "/api/v1/auth/org/groups/g1", None, "get_group_detail"),
    ("post", "/api/v1/auth/org/groups/g1/students/bulk", {"student_ids": ["s1"]}, "bulk_add_students"),
    ("post", "/api/v1/auth/org/groups/g1/students/s1", None, "add_student_to_group"),
]


@pytest.mark.parametrize("method,path,body,service_method", ROUTES)
def test_the_display_id_claim_failure_is_a_503_with_its_russian_text(owner_http, monkeypatch, method, path, body, service_method):
    monkeypatch.setattr(AuthService, service_method, boom)
    r = owner_http.request(method, path, json=body)
    assert r.status_code == 503, (path, r.text)
    assert r.json()["detail"] == PREFIX_TEXT


def test_the_real_prefix_claim_failure_reaches_the_client_as_503(owner_http, monkeypatch):
    """Not a stub of the service: ``ensure_org_prefix`` runs out of attempts, a route that assigns display ids answers 503."""
    monkeypatch.setattr(student_id_service, "PREFIX_CLAIM_ATTEMPTS", 0)
    org = MagicMock(student_id_prefix=None)
    org.name = "Центр"
    with pytest.raises(RuntimeError, match="префикс ID учеников"):
        student_id_service.ensure_org_prefix(MagicMock(), org)

    def assign_fails(self, *a, **k):
        student_id_service.ensure_org_prefix(MagicMock(), org)

    monkeypatch.setattr(AuthService, "get_student_profile", assign_fails)
    r = owner_http.get("/api/v1/auth/org/students/s1/profile")
    assert r.status_code == 503 and "префикс ID учеников" in r.json()["detail"]


# ------------------------------------------------------------------------------------------------------ 4. X-Request-ID on 411 / 413
def test_the_upload_limit_answers_carry_the_request_id():
    http = TestClient(APP, raise_server_exceptions=False)
    too_big = http.post("/api/v1/auth/org/import/preview", files={"file": ("x.xlsx", b"\0" * (6 * 1024 * 1024), MIME)}, headers={"X-Request-ID": "upload-413"})
    assert too_big.status_code == 413 and too_big.headers["X-Request-ID"] == "upload-413"
    no_length = http.post("/api/v1/auth/org/import/commit", content=iter([b"--x\r\n", b"data"]), headers={"Content-Type": "multipart/form-data; boundary=x"})
    assert no_length.status_code == 411 and uuid.UUID(no_length.headers["X-Request-ID"])
