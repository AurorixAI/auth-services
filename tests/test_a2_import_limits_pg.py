"""Pack A2 - DM-001 / DM-054 / DM-050: the xlsx import under load and abuse, through the real app (``src.api.main_v2``: its middleware and handlers).

Matrix: a normal file (7 groups, 41 students) goes through exactly as before; 6 MB -> 413 BEFORE the body is read; no Content-Length -> 411;
2001 rows -> 422; 11 sheets -> 422; a broken / truncated / non-zip file -> 422; a zip bomb (unpacks to more than the limit) -> 422;
a sheet that "starts" a hundred thousand rows away -> 422; the plan costs a constant number of queries, not a query per row;
a failure inside the commit never reaches the client as SQL text.
"""
import io
import logging
import re
import time
import zipfile
from unittest.mock import MagicMock

import pytest
from openpyxl import Workbook
from sqlalchemy import event

from src.api import main_v2
from src.auth import import_limits as limits
from src.auth import import_ops, import_service as legacy
from src.auth import router as auth_router
from src.auth.dependencies import get_current_active_user_payload
from src.auth.service import AuthService
from src.core.database import get_db
from src.models.auth import Group, Student, StudentGroup

from tests.pg_world import needs_pg, pg_session, world  # noqa: F401  (fixtures)

pytestmark = needs_pg
IMP = "/api/v1/auth/org/import"
MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
GH = ["Название группы*", "Класс (1-11)", "Почта учителя", "Макс. учеников"]
SH = ["Фамилия*", "Имя*", "Класс (1-11)*", "Телефон", "Группа*", "Комментарий"]


@pytest.fixture
def owner(world):
    """A TestClient on the real application, signed in as the owner of ``world.org``."""
    from fastapi.testclient import TestClient

    payload = {"sub": str(world.owner.id), "role": "owner", "type": "access", "org_id": str(world.org.id), "org_role": "owner"}
    main_v2.app.dependency_overrides[get_db] = lambda: world.db
    main_v2.app.dependency_overrides[get_current_active_user_payload] = lambda: payload
    yield TestClient(main_v2.app, raise_server_exceptions=False)
    main_v2.app.dependency_overrides.pop(get_db, None)
    main_v2.app.dependency_overrides.pop(get_current_active_user_payload, None)


def book(groups, students, *, extra_sheets=0):
    wb = Workbook()
    wb.remove(wb.active)
    g = wb.create_sheet("Группы")
    g.append(GH)
    for row in groups:
        g.append(row)
    s = wb.create_sheet("Ученики")
    s.append(SH)
    for row in students:
        s.append(row)
    for i in range(extra_sheets):
        wb.create_sheet(f"Лист{i}")
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def school(groups=7, students=41, prefix="Тест"):
    """The shape of the TestSchool file: ``groups`` groups, ``students`` students spread over them (unique names)."""
    names = [f"{prefix}-{i + 1}" for i in range(groups)]
    rows = [[f"{prefix}{i:03d}ов", f"Имя{i:03d}", 5 + i % 7, None, names[i % groups], None] for i in range(students)]
    return book([[n, 5 + i, None, 30] for i, n in enumerate(names)], rows)


def post(client, kind, data, name="import.xlsx"):
    return client.post(f"{IMP}/{kind}", files={"file": (name, data, MIME)})


def zip_with_extra_part(xlsx: bytes, size: int) -> bytes:
    """The workbook plus one more zip part of ``size`` zero bytes (deflates to a few dozen KB): a zip bomb that is still a valid xlsx."""
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(xlsx)) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            dst.writestr(info, src.read(info.filename))
        dst.writestr("xl/media/bomb.bin", b"\0" * size)
    return out.getvalue()


# ------------------------------------------------------------------------------------------------------ a normal file, as before
def test_a_normal_school_file_imports_exactly_as_before(owner, world):
    data = school(7, 41)
    pv = post(owner, "preview", data)
    assert pv.status_code == 200, pv.text
    p = pv.json()
    assert p["ok"] is True and p["counts"]["groups_new"] == 7 and p["counts"]["students_enrolled"] == 41 and p["counts"]["students_new"] == 41
    assert p["counts"]["students_skipped"] == 0 and len(p["groups_preview"]) == 7 and len(p["students_preview"]) == 41
    assert world.db.query(Student).filter_by(org_id=world.org.id).count() == 3                                        # a preview writes nothing

    done = post(owner, "commit", data)
    assert done.status_code == 200, done.text
    out = done.json()
    assert out["ok"] is True and len(out["created_groups"]) == 7 and out["counts"]["students_enrolled"] == 41 and out["counts"]["students_new"] == 41
    assert world.db.query(Student).filter_by(org_id=world.org.id).count() == 3 + 41
    assert world.db.query(Group).filter(Group.org_id == world.org.id, Group.name.like("Тест-%")).count() == 7
    seats = world.db.query(StudentGroup).join(Group, Group.id == StudentGroup.group_id).filter(Group.name.like("Тест-%"), Group.org_id == world.org.id).count()
    assert seats == 41

    again = post(owner, "commit", data).json()                                                                         # a repeat changes nothing
    assert again["ok"] is True and again["counts"]["students_enrolled"] == 0 and again["counts"]["students_unchanged"] == 41
    assert world.db.query(Student).filter_by(org_id=world.org.id).count() == 3 + 41


def test_an_example_row_and_row_numbers_survive_the_streaming_reader():
    """The reader is read-only now: the Excel row of a record is still its real row, blanks and gaps included, and the template example is skipped."""
    wb = Workbook()
    wb.remove(wb.active)
    g = wb.create_sheet("Группы")
    g.append(GH)
    g.append(["5-А", 5, None, 30])
    s = wb.create_sheet("Ученики")
    s.append(SH)
    s.append([None] * 6)
    s.append(["Иванов", "Иван", 5, None, "5-А", None])
    s["A6"], s["B6"], s["C6"], s["E6"] = "Петров", "Пётр", 5, "5-А"
    s["A8"], s["F8"] = "Образец", "пример — удалите"
    out = io.BytesIO()
    wb.save(out)
    parsed = legacy.parse_workbook(out.getvalue())
    assert parsed.fatal is None and [r.row for r in parsed.groups] == [2] and [r.row for r in parsed.students] == [3, 6]
    assert parsed.students[1].data[legacy.K_LAST] == "Петров"


def test_a_workbook_without_dimension_info_is_read(tmp_path):
    """Some tools write no <dimension> tag; a read-only openpyxl then knows no sheet size - the rows must still be found."""
    data = school(2, 4)
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            blob = src.read(info.filename)
            if info.filename.startswith("xl/worksheets/"):
                blob = re.sub(rb"<dimension[^>]*/>", b"", blob)
            dst.writestr(info, blob)
    parsed = legacy.parse_workbook(out.getvalue())
    assert parsed.fatal is None and len(parsed.groups) == 2 and len(parsed.students) == 4


def test_a_readable_file_without_the_sheets_is_still_a_200_with_the_old_message(owner):
    wb = Workbook()
    wb.active.title = "Чужой лист"
    out = io.BytesIO()
    wb.save(out)
    r = post(owner, "preview", out.getvalue())
    assert r.status_code == 200 and r.json()["ok"] is False and "нет обязательных вкладок" in r.json()["issues"][0]["message"]


# ------------------------------------------------------------------------------------------------------ size: before the body is read
@pytest.mark.parametrize("kind", ["preview", "commit"])
def test_a_6_mb_file_is_413_before_anything_reads_it(owner, monkeypatch, kind):
    async def must_not_run(file):
        raise AssertionError("the body was read")

    monkeypatch.setattr(auth_router, "_read_import_upload", must_not_run)
    monkeypatch.setattr(import_ops, kind, MagicMock(side_effect=AssertionError("the importer ran")))
    r = post(owner, kind, b"\0" * (6 * 1024 * 1024))
    assert r.status_code == 413 and "макс. 5 МБ" in r.json()["detail"]


@pytest.mark.parametrize("kind", ["preview", "commit"])
def test_a_file_just_over_the_limit_inside_the_multipart_slack_is_413_from_the_route(owner, kind):
    r = post(owner, kind, b"\0" * (limits.MAX_IMPORT_BYTES + 100))
    assert r.status_code == 413 and "макс. 5 МБ" in r.json()["detail"]


def test_a_body_without_a_declared_length_is_411(owner):
    r = owner.post(f"{IMP}/preview", content=iter([b"--x\r\n", b"data"]), headers={"Content-Type": "multipart/form-data; boundary=x"})
    assert r.status_code == 411


def test_the_413_carries_cors_headers_like_any_other_error(owner):
    origin = main_v2._allowed_origins[0]
    r = owner.post(f"{IMP}/preview", files={"file": ("x.xlsx", b"\0" * (6 * 1024 * 1024), MIME)}, headers={"Origin": origin})
    assert r.status_code == 413 and r.headers.get("access-control-allow-origin") == origin


# ------------------------------------------------------------------------------------------------------ rows, sheets, broken files, bombs
def test_2000_rows_are_read_and_2001_are_refused_with_422(owner):
    def students(n):
        return [[f"Фам{i:04d}", f"Имя{i:04d}", 5, None, "Г-1", None] for i in range(n)]

    assert len(legacy.parse_workbook(book([["Г-1", 5, None, 30]], students(2000))).students) == 2000
    for kind in ("preview", "commit"):
        r = post(owner, kind, book([["Г-1", 5, None, 30]], students(2001)))
        assert r.status_code == 422, kind
        assert "«Ученики»" in r.json()["detail"] and "2000" in r.json()["detail"]
    r = post(owner, "preview", book([[f"Г-{i}", 5, None, 30] for i in range(2001)], []))
    assert r.status_code == 422 and "«Группы»" in r.json()["detail"]


def test_blank_rows_do_not_count_towards_the_limit():
    rows = [[None] * 6 for _ in range(3000)] + [["Иванов", "Иван", 5, None, "Г-1", None]]
    parsed = legacy.parse_workbook(book([["Г-1", 5, None, 30]], rows))
    assert len(parsed.students) == 1 and parsed.students[0].row == 3002


def test_a_sheet_that_starts_far_below_is_422(owner):
    wb = Workbook()
    wb.remove(wb.active)
    wb.create_sheet("Группы").append(GH)
    s = wb.create_sheet("Ученики")
    s.append(SH)
    s["A100000"] = "Иванов"
    out = io.BytesIO()
    wb.save(out)
    r = post(owner, "preview", out.getvalue())
    assert r.status_code == 422 and "2000" in r.json()["detail"]


def test_11_sheets_are_422_and_10_are_fine(owner):
    assert post(owner, "preview", book([["Г-1", 5, None, 30]], [], extra_sheets=8)).status_code == 200            # 2 + 8 = 10
    r = post(owner, "preview", book([["Г-1", 5, None, 30]], [], extra_sheets=9))                                  # 11
    assert r.status_code == 422 and "10 листов" in r.json()["detail"]


@pytest.mark.parametrize("kind", ["preview", "commit"])
def test_broken_files_are_422_with_a_clear_text(owner, kind):
    good = school(2, 3)
    for blob in (b"this is not a zip file at all" * 20, b"PK\x03\x04" + b"\x00garbage" * 50, good[: len(good) // 2], good[:-30] + b"\0" * 30):
        r = post(owner, kind, blob)
        assert r.status_code == 422, (kind, blob[:12])
        assert r.json()["detail"] == "Не удалось прочитать файл. Загрузите .xlsx по нашему шаблону."


def test_a_zip_with_xml_that_is_not_a_workbook_is_422(owner):
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        z.writestr("[Content_Types].xml", "<broken")
        z.writestr("xl/workbook.xml", "<also broken")
    assert post(owner, "preview", out.getvalue()).status_code == 422


@pytest.mark.parametrize("kind", ["preview", "commit"])
def test_a_zip_bomb_is_422_and_cheap(owner, kind):
    bomb = zip_with_extra_part(school(2, 3), limits.MAX_UNPACKED_BYTES + 1024 * 1024)
    assert len(bomb) < limits.MAX_IMPORT_BYTES                                       # it passes the size check: the unpacked size is what is wrong
    started = time.monotonic()
    r = post(owner, kind, bomb)
    assert r.status_code == 422 and "после распаковки" in r.json()["detail"]
    assert time.monotonic() - started < 10


def test_a_zip_just_under_the_unpacked_limit_is_read():
    blob = zip_with_extra_part(school(2, 3), 5 * 1024 * 1024)
    assert len(legacy.parse_workbook(blob).students) == 3


def test_a_zip_with_too_many_parts_is_refused():
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        for i in range(limits.MAX_ZIP_ENTRIES + 1):
            z.writestr(f"p{i}.txt", "x")
    with pytest.raises(limits.ImportFileRejected):
        limits.check_zip(out.getvalue())


# ------------------------------------------------------------------------------------------------------ DM-001: queries are dictionaries, not per row
def count_queries(world, fn):
    engine = world.db.get_bind()
    seen = []

    def on_execute(conn, cursor, statement, parameters, context, executemany):
        seen.append(statement)

    event.listen(engine, "before_cursor_execute", on_execute)
    try:
        fn()
    finally:
        event.remove(engine, "before_cursor_execute", on_execute)
    return len(seen)


def test_the_plan_costs_a_constant_number_of_queries(owner, world):
    small = school(7, 14, prefix="Малый")
    large = school(7, 700, prefix="Большой")
    # warm-up: first-use queries (settings, caches) are not the plan's
    post(owner, "preview", small)
    n_small = count_queries(world, lambda: post(owner, "preview", small))
    n_large = count_queries(world, lambda: post(owner, "preview", large))
    assert n_large <= n_small + 3, (n_small, n_large)                                # 50 times the rows, the same queries
    assert n_large < 60


def test_an_existing_student_and_group_are_still_recognised(owner, world):
    """The lookups are dictionaries now: the same case-insensitive match as the old per-row queries (a homonym, an existing group, a full group)."""
    existing = world.students[0]
    data = book(
        [[world.group.name.upper(), 9, None, None], ["Новая-1", 5, None, 1]],
        [[existing.last_name.upper(), existing.first_name.lower(), 9, None, "Новая-1", None],
         ["Один", "Раз", 5, None, "новая-1", None], ["Два", "Раза", 5, None, "НОВАЯ-1", None]],
    )
    p = post(owner, "preview", data).json()
    assert p["counts"]["groups_existing"] == 1 and p["counts"]["groups_new"] == 1 and p["counts"]["students_existing"] == 1
    assert sorted(s["code"] for s in p["skipped"]) == ["GROUP_FULL", "IN_OTHER_GROUP"]                                  # limit 1 in «Новая-1»; Каримова is already in her group
    assert p["counts"]["students_enrolled"] == 1 and p["issues"] == []


# ------------------------------------------------------------------------------------------------------ DM-050
def test_a_failing_group_creation_does_not_leak_the_sql_text(owner, monkeypatch, caplog):
    def broken(self, *a, **k):
        raise RuntimeError('(psycopg2.errors.UndefinedTable) relation "groups_secret" does not exist [SQL: INSERT INTO groups_secret ...]')

    monkeypatch.setattr(AuthService, "create_group", broken)
    with caplog.at_level(logging.ERROR):
        r = post(owner, "commit", school(2, 4, prefix="Сбой"))
    assert r.status_code == 200 and r.json()["ok"] is False
    message = r.json()["issues"][0]["message"]
    assert "groups_secret" not in r.text and "SQL" not in r.text and "psycopg2" not in r.text and "внутренняя ошибка" in message
    assert "groups_secret" in caplog.text                                            # the details are in the log


def test_a_domain_refusal_of_the_group_service_is_still_shown(owner, monkeypatch):
    def refuse(self, *a, **k):
        raise ValueError("Некорректный преподаватель")

    monkeypatch.setattr(AuthService, "create_group", refuse)
    r = post(owner, "commit", school(1, 1, prefix="Отказ"))
    assert r.json()["ok"] is False and "Некорректный преподаватель" in r.json()["issues"][0]["message"]


def test_the_legacy_commit_does_not_leak_either(caplog):
    plan = MagicMock(has_errors=False)
    plan.groups = [MagicMock(action="create", name="X-1", class_level=5, teacher_email="", max_students=3)]
    plan.students = []
    db = MagicMock()
    db.add.side_effect = RuntimeError("INSERT INTO groups_secret VALUES (:p)  [parameters: {'p': 'hunter2'}]")
    original = legacy.build_plan
    legacy.build_plan = lambda *a, **k: plan
    try:
        with caplog.at_level(logging.ERROR):
            result = legacy.commit("org-1", "user-1", school(1, 1), db)
    finally:
        legacy.build_plan = original
    assert result.ok is False and "groups_secret" not in result.issues[0].message and "hunter2" not in result.issues[0].message
    assert "groups_secret" in caplog.text and db.rollback.called
