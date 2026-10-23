"""The import template in the house style: «Инструкция» first as ordinary text, a grey example row per sheet that the importer IGNORES, format hints, «Почта учителя»
(the old «Email учителя» is read too).  No database: ``build_template_xlsx`` and ``parse_workbook``."""
import io
from datetime import datetime

from openpyxl import load_workbook

from src.auth import import_service as imp
from src.auth import xlsx_style as xs


def template():
    return load_workbook(io.BytesIO(imp.build_template_xlsx()))


def saved(wb):
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def test_the_instruction_is_first_and_ordinary_text():
    wb = template()
    assert wb.sheetnames == ["Инструкция", "Группы", "Ученики"]
    info = wb["Инструкция"]
    texts = [info.cell(r, 1).value for r in range(1, info.max_row + 1) if info.cell(r, 1).value]
    assert texts[0] == "Импорт данных в ALGO — как заполнять" and info["A1"].font.bold and info["A1"].font.size == 14
    assert len(texts) == 1 + len(imp.INSTRUCTION_PARAGRAPHS) and texts[1:] == list(imp.INSTRUCTION_PARAGRAPHS)
    assert not any("\n" in t or t.startswith("   ") or "——" in t or "|" in t for t in texts)          # no manual breaks, no indents, no drawn separators
    assert all(info.cell(r, 1).alignment.wrap_text and info.cell(r, 1).alignment.vertical == "top" for r in range(3, 3 + len(texts) - 1))
    assert info.column_dimensions["A"].width == 110 and all(info.row_dimensions[r].height >= 30 for r in range(3, 3 + len(texts) - 1) if len(info.cell(r, 1).value) > 100)
    assert any("Почта учителя" in t for t in texts) and any("31.10.2026" in t for t in texts) and any("+998901234567" in t for t in texts)


def test_headers_house_colour_example_rows_and_hints():
    wb = template()
    groups, students = wb["Группы"], wb["Ученики"]
    assert [c.value for c in groups[1]] == ["Название группы*", "Класс (1-11)", "Почта учителя", "Макс. учеников", "Пометка"]
    assert [c.value for c in students[1]] == ["Фамилия*", "Имя*", "Класс (1-11)*", "Телефон", "Группа*", "Комментарий", "Оплачено до (дата)", "Долг, сум",
                                              "Долг на дату", "Пометка"]
    for ws in (groups, students):
        for c in ws[1]:
            assert c.fill.fgColor.rgb.endswith("D9E2F3") and c.font.bold and c.font.name == "Calibri" and not (c.font.color and c.font.color.rgb.endswith("FFFFFF"))
        assert ws.freeze_panes == "A2" and ws.page_setup.orientation == "landscape" and ws.print_title_rows == "$1:$1"
        assert all("2F6FED" not in str(c.fill.fgColor.rgb) for row in ws.iter_rows() for c in row)
        # one grey example row, marked
        assert ws.cell(2, ws.max_column).value == "пример — удалите" == xs.EXAMPLE_MARK
        assert all(ws.cell(2, c).fill.fgColor.rgb.endswith("F2F2F2") and ws.cell(2, c).font.italic and ws.cell(2, c).font.color.rgb.endswith("808080")
                   for c in range(1, ws.max_column + 1))
        assert ws.max_row == 2
    assert [groups.cell(2, c).value for c in range(1, 5)] == ["Алгебра-9А", 9, "a.karimov@example.com", 25]
    assert [students.cell(2, c).value for c in range(1, 6)] == ["Усмонова", "Малика", 9, "+998901234567", "Алгебра-9А"]
    example_date = students.cell(2, 7).value
    assert isinstance(example_date, datetime) and students.cell(2, 7).number_format == "DD.MM.YYYY" and students.cell(2, 4).number_format == "@"
    # the hints: a comment on the header, the number format of the column, a prompt on the cells
    assert "ДД.ММ.ГГГГ" in students["G1"].comment.text and "ДД.ММ.ГГГГ" in students["I1"].comment.text and "+998901234567" in students["D1"].comment.text
    assert students.column_dimensions["G"].number_format == "DD.MM.YYYY" and students.column_dimensions["D"].number_format == "@"
    prompts = {(dv.type, dv.promptTitle) for dv in students.data_validations.dataValidation if dv.showInputMessage}
    assert prompts == {("date", "Дата"), ("textLength", "Телефон")}
    # widths from the content: nothing is cut (the mark fits, the long headers fit)
    for ws in (groups, students):
        for c in range(1, ws.max_column + 1):
            assert ws.column_dimensions[chr(64 + c)].width >= len(str(ws.cell(1, c).value)) - 4
        assert ws.column_dimensions[chr(64 + ws.max_column)].width >= len(xs.EXAMPLE_MARK)
    # the dropdown of the groups and the class limits are kept (they now start under the example row)
    assert {dv.type for dv in students.data_validations.dataValidation} >= {"list", "whole"}
    assert "Email" not in {c.value for ws in (groups, students) for c in ws[1]}


def test_the_importer_ignores_the_example_rows():
    wb = template()
    parsed = imp.parse_workbook(saved(wb))
    assert parsed.fatal is None and parsed.groups == [] and parsed.students == []                      # the examples are not data
    # a real row under the example row is read, the example stays out; the row numbers are the Excel rows
    wb["Группы"].append(["Физика-8Б", 8, None, 20])
    wb["Ученики"].append(["Каримов", "Тимур", 8, "+998901112233", "Физика-8Б", None, None, None, None])
    parsed = imp.parse_workbook(saved(wb))
    assert [(r.row, r.data[imp.K_GROUP_NAME]) for r in parsed.groups] == [(3, "Физика-8Б")]
    assert [(r.row, r.data[imp.K_LAST], r.data[imp.K_GROUP]) for r in parsed.students] == [(3, "Каримов", "Физика-8Б")]
    # the mark is found in any cell, any dash, any case - and only the marked row is dropped
    wb = template()
    wb["Ученики"].append(["Иванов", "Иван", 9, None, "Физика-8Б", "ПРИМЕР - УДАЛИТЕ", None, None, None])
    wb["Ученики"].append(["Петров", "Пётр", 9, None, "Физика-8Б", "пример удалите", None, None, None])    # no dash: NOT the mark, a real row
    parsed = imp.parse_workbook(saved(wb))
    assert [r.data[imp.K_LAST] for r in parsed.students] == ["Петров"]
    assert imp._is_example([type("C", (), {"value": "Пример – удалите"})()]) and not imp._is_example([type("C", (), {"value": "пример"})()])


def test_both_teacher_headers_are_accepted():
    for header in ("Почта учителя", "Email учителя", "E-mail учителя"):
        wb = template()
        wb["Группы"]["C1"] = header
        wb["Группы"].append(["Физика-8Б", 8, "T@Example.com", 20])
        parsed = imp.parse_workbook(saved(wb))
        assert [r.data[imp.K_TEACHER_EMAIL] for r in parsed.groups] == ["T@Example.com"], header
    assert imp.K_TEACHER_EMAIL == "почта учителя" and imp.COLS_GROUPS[2] == "Почта учителя"


def test_the_file_name_is_russian():
    assert imp.TEMPLATE_FILE_NAME == "Шаблон_импорта.xlsx"
    assert xs.content_disposition(imp.TEMPLATE_FILE_NAME) == (
        "attachment; filename=\"Shablon_importa.xlsx\"; filename*=UTF-8''%D0%A8%D0%B0%D0%B1%D0%BB%D0%BE%D0%BD_%D0%B8%D0%BC%D0%BF%D0%BE%D1%80%D1%82%D0%B0.xlsx")
