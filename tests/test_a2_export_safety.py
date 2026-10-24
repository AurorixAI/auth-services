"""Pack A4 (auth part) - MN-10 / MN-43: formula injection in every export, one rule.

xlsx: a text that starts with = + - @ (TAB, CR) is a STRING cell with Excel's quote-prefix flag; the text itself is exactly what was typed (no
apostrophe: «=Тест», «-Анвар» read correctly in Excel).  CSV: the cell gets a leading apostrophe, the file follows RFC 4180 quoting.  Numbers and
dates are never touched.  The files are opened with openpyxl / the csv module, not grepped.
"""
import csv
import io
from datetime import date, datetime

import pytest
from openpyxl import Workbook, load_workbook

from src.auth import export_safety as es
from src.auth import xlsx_style as xs

HOSTILE = ["=1+1", "+998901112233", "-Анвар", "@SUM(A1)", "=Тест", "\t=cmd", "\r=cmd", " =pad", "=HYPERLINK(\"http://evil\",\"x\")", "=cmd|' /C calc'!A0"]
SAFE_TEXT = ["Анвар", "a=b", "5-А", "Иванов - Петров", "", "'quoted", "1+1", "Алгебра-9А", "—"]


def roundtrip(values):
    wb = Workbook()
    ws = wb.active
    for i, v in enumerate(values, start=1):
        xs.put(ws, i, 1, v)
    out = io.BytesIO()
    wb.save(out)
    return load_workbook(io.BytesIO(out.getvalue())).active


# ------------------------------------------------------------------------------------------------------ xlsx
def test_hostile_text_is_a_string_cell_with_quote_prefix_and_reads_back_exactly():
    ws = roundtrip(HOSTILE)
    for i, original in enumerate(HOSTILE, start=1):
        cell = ws.cell(i, 1)
        assert cell.value == original.replace("\r", "\n"), (i, cell.value)    # no apostrophe in the text: «-Анвар» is «-Анвар» (XML turns a lone CR into LF)
        assert cell.data_type == "s" and cell.quotePrefix is True, original


def test_ordinary_text_has_no_quote_prefix_and_numbers_and_dates_stay_typed():
    ws = roundtrip([*SAFE_TEXT, -5, 7, 12.5, -0.5, date(2026, 10, 26), datetime(2026, 10, 26, 9, 30), True, None])
    for i, original in enumerate(SAFE_TEXT, start=1):
        cell = ws.cell(i, 1)
        assert cell.value == (original or None) and not cell.quotePrefix, original
    n = len(SAFE_TEXT)
    assert [ws.cell(n + k, 1).value for k in (1, 2, 3, 4)] == [-5, 7, 12.5, -0.5]
    assert all(ws.cell(n + k, 1).data_type == "n" and not ws.cell(n + k, 1).quotePrefix for k in (1, 2, 3, 4))
    assert ws.cell(n + 5, 1).value == datetime(2026, 10, 26) and ws.cell(n + 6, 1).value == datetime(2026, 10, 26, 9, 30)
    assert ws.cell(n + 7, 1).value is True


def test_a_formula_never_survives_as_a_formula():
    ws = roundtrip(HOSTILE)
    assert all(ws.cell(i, 1).data_type != "f" for i in range(1, len(HOSTILE) + 1))


# ------------------------------------------------------------------------------------------------------ CSV
def test_csv_prefixes_text_only_and_keeps_numbers_typed():
    row = ["=1+1", "+7", "-Анвар", "@x", "\t=cmd", "\r=cmd", " =pad", "ok", "a=b", "", 5, -5, 0.5, None, "-", date(2026, 10, 26)]
    assert es.csv_safe_row(row) == ["'=1+1", "'+7", "'-Анвар", "'@x", "'\t=cmd", "'\r=cmd", "' =pad", "ok", "a=b", "", 5, -5, 0.5, None, "'-", date(2026, 10, 26)]


def test_csv_follows_rfc4180_and_round_trips_through_a_parser():
    rows = [["=Тест", 'say "hi"', "a;b", "line1\nline2", "x,y", -3, ""], ["-Анвар", "@home", 10, "tab\there", "\r=1", "", "ok"]]
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=";")
    for r in rows:
        writer.writerow(es.csv_safe_row(r))
    text = buf.getvalue()
    assert text.endswith("\r\n") and '"say ""hi"""' in text and '"a;b"' in text and '"line1\nline2"' in text          # quotes doubled, fields quoted
    parsed = list(csv.reader(io.StringIO(text), delimiter=";"))
    expected = [[("'" + v if es.formula_like(v) else v) if isinstance(v, str) else str(v) for v in r] for r in rows]
    assert parsed == expected


def test_one_rule_for_both_csv_exports():
    """The journal and the statement used two copies of the rule (the statement missed TAB / CR): both use ``export_safety`` now."""
    from src.auth import payments_income, payments_reports

    assert payments_income.csv_safe_row is es.csv_safe_row and payments_reports.csv_safe_row is es.csv_safe_row
    assert payments_income.FORMULA_STARTS == es.FORMULA_STARTS


# ------------------------------------------------------------------------------------------------------ the real xlsx builders
def test_the_statement_xlsx_shows_hostile_names_as_typed():
    from tests.test_payments_statement_xlsx import golden
    from src.auth.payments_statement_xlsx import build_statement_xlsx

    st = golden(note='=HYPERLINK("http://evil")')
    st.membership.student_name = "=Тест"
    st.membership.teacher_name = "-Анвар"
    st.rows[5].details.reason = "@cmd"
    ws = load_workbook(io.BytesIO(build_statement_xlsx(st))).active
    cells = [c for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    hot = [c for c in cells if c.value.lstrip()[:1] in ("=", "+", "-", "@") or "=HYPERLINK" in c.value]
    assert hot and all(c.data_type == "s" and not c.value.startswith("'") for c in hot)
    assert all(c.quotePrefix for c in hot if c.value.lstrip()[:1] in ("=", "+", "-", "@"))
    assert any(c.value == "=Тест" or "=Тест" in c.value for c in cells)
    assert not [c for c in cells if c.value.startswith("'=") or c.value.startswith("'-") or c.value.startswith("'@")]


def test_the_import_template_has_no_apostrophes_and_no_formulas():
    from src.auth import import_service

    wb = load_workbook(io.BytesIO(import_service.build_template_xlsx()))
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                assert c.data_type != "f" and not (isinstance(c.value, str) and c.value.startswith("'")), (ws.title, c.coordinate)
