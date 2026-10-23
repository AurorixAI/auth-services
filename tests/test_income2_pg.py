"""INCOME-2 («Поступления» v2) on PostgreSQL: the overview with the comparison, the month collection, the payments journal and its Excel.

The figures of ``income2_world`` are hand-calculated (see tests/income2_helpers.py); the identities hold for any data (random test below)."""
import io
import random
from datetime import datetime, timedelta

import pytest
from openpyxl import load_workbook

from src.models.auth import StudentGroup, StudentPayment, StudentRefund, StudentWriteoff

from tests.income2_helpers import income2_world, uid
from tests.pg_world import give_group_price, make_client, make_group, make_student, needs_pg, pg_session, world  # noqa: F401
from tests.test_payment_carry_over_pg import engine  # noqa: F401
from tests.test_pay2_api_pg import API, D, TODAY, fixed_today, membership, owner, seed_payment, settings_block, stay  # noqa: F401
from tests.test_pay2_perf_concurrency_pg import count_statements, grow

pytestmark = needs_pg


def income(w, first, last, client=None, expect=200, **params):
    r = (client or owner(w)).get(f"{API}/org/payments/income", params={"from": first, "to": last, **params})
    assert r.status_code == expect, r.text
    return r.json()


def journal(w, first="2026-10-01", last="2026-10-31", client=None, expect=200, **params):
    r = (client or owner(w)).get(f"{API}/org/payments/journal", params={"from": first, "to": last, **params})
    assert r.status_code == expect, r.text
    return r.json()


def receipts(w, first, last, **params):
    r = owner(w).get(f"{API}/org/payments/receipts", params={"from": first, "to": last, "group_by": "month", **params})
    assert r.status_code == 200, r.text
    return r.json()


def journal_xlsx(w, client=None, **params):
    """``(response, worksheet)`` of the journal file of October (or of the given range)."""
    params = {"from": "2026-10-01", "to": "2026-10-31", **params}
    r = (client or owner(w)).get(f"{API}/org/payments/journal/export", params=params)
    assert r.status_code == 200, r.text
    return r, load_workbook(io.BytesIO(r.content)).active


def table_of(ws):
    """``(header row, {header: column}, [row as {header: value}])``: the table ends at the blue «Итого» row (or the note «Оплат за период нет»)."""
    head = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "№ квитанции")
    cols = {ws.cell(head, c).value: c for c in range(1, ws.max_column + 1) if ws.cell(head, c).value}
    body = []
    for r in range(head + 1, ws.max_row + 1):
        if not isinstance(ws.cell(r, 1).value, int):
            break
        body.append({h: ws.cell(r, c).value for h, c in cols.items()})
    return head, cols, body


def journal_csv(w, **params):
    import csv as csv_module

    r = owner(w).get(f"{API}/org/payments/journal/export", params={"from": "2026-10-01", "to": "2026-10-31", "format": "csv", **params})
    assert r.status_code == 200, r.text
    rows = list(csv_module.reader(io.StringIO(r.text.lstrip("\ufeff")), delimiter=";"))
    head = next(i for i, row in enumerate(rows) if row and row[0] == "№ квитанции")
    return rows[:head], rows[head], [row for row in rows[head + 1:] if row]


@pytest.fixture
def seeded(world, fixed_today):
    return income2_world(world)


# ---- the overview: hand-calculated figures ----------------------------------------------------------------------


def test_october_overview_hand_calculated(world, seeded):
    o = income(world, "2026-10-01", "2026-10-31")
    assert o["period"] == {"kind": "month", "from": "2026-10-01", "to": "2026-10-31", "label": "октябрь", "days": 31, "in_progress": True}
    assert o["previous_period"]["label"] == "1–21 сентября" and o["previous_period"]["to"] == "2026-09-21"
    cur = o["current"]
    # 500+200+100+1250+900+150+350 (the voided 777 and the replaced 400 are not money)
    assert (cur["received"], cur["payments_count"]) == (3_450_000, 7)
    assert (cur["to_past_debt"], cur["to_current"], cur["to_advance"]) == (1_000_000, 1_900_000, 450_000)
    assert (cur["refunded"], cur["refunded_from_received"], cur["transferred"], cur["written_off"], cur["net_received"]) == (
        100_000, 100_000, 0, 100_000, 3_350_000)
    assert cur["to_forward"] == 0 and cur["to_past_debt"] + cur["to_current"] + cur["to_advance"] + cur["refunded_from_received"] == cur["received"]
    assert [(x["kind"], x["label"], x["amount"]) for x in o["strip"]] == [
        ("old_debt", "на старые долги", 1_000_000), ("current", "за текущий период", 1_900_000), ("forward", "вперёд (за будущие периоды)", 0),
        ("advance", "пока не потрачено (аванс)", 450_000), ("refunded", "возвращено", 100_000), ("transferred", "перенесено в другую группу", 0)]
    assert o["previous_label"] == "те же дни сентября" and o["previous_hint"] == "1–21 сентября"
    assert [c["code"] for c in o["checks"] if c["code"].startswith("strip_")] == [
        "strip_old_debt_eq_journal", "strip_current_eq_journal", "strip_advance_eq_journal", "strip_refunded_eq_journal"]
    prev = o["previous"]
    assert (prev["received"], prev["payments_count"], prev["to_current"], prev["to_past_debt"]) == (1_500_000, 4, 1_500_000, 0)
    assert (o["delta_amount"], o["delta_pct"]) == (1_950_000, 130)
    assert o["summary_text"] == "За октябрь (по 21-е) пришло 3 450 000 сум — на 130% больше, чем за те же дни сентября. Собрано 76% стоимости месяца."
    assert [(m["method"], m["amount"], m["count"]) for m in o["by_method"]] == [
        ("click", 1_250_000, 1), ("cash", 950_000, 4), ("bank_transfer", 900_000, 1), ("card", 350_000, 1)]
    assert [m["label"] for m in o["by_method"]] == ["Click", "Наличные", "Перевод", "Карта"]
    assert o["position"] == {"debt_total": 500_000, "past_debt_total": 0, "advance_total": 450_000, "memberships_with_debt": 3,
                             "students_with_debt": 3, "memberships_with_advance": 2, "students_with_advance": 2}
    assert (o["delta_direction"], o["delta_text"]) == ("up", "130% к 1–21 сентября")
    assert o["cash"] is None and o["ok"] and all(c["diff"] == 0 for c in o["checks"])


def test_collection_hand_calculated_and_identities(world, seeded):
    col = income(world, "2026-10-01", "2026-10-31")["collection"]
    assert (col["month"], col["month_label"], col["from"], col["to"]) == ("2026-10", "октябрь 2026", "2026-10-01", "2026-10-31")
    rows = {r["group_name"]: r for r in col["groups"]}
    assert [r["group_name"] for r in col["groups"]] == ["Математика 9А", "Английский 7Б", "Физика 10В"]      # worst first
    a, b, c = rows["Математика 9А"], rows["Английский 7Б"], rows["Физика 10В"]
    assert (a["charged"], a["paid"], a["forgiven"], a["debt"], a["students"], a["students_with_debt"], a["pct"]) == (
        1_500_000, 1_100_000, 100_000, 300_000, 3, 1, 73)
    assert (b["charged"], b["paid"], b["forgiven"], b["debt"], b["students_with_debt"], b["pct"]) == (600_000, 450_000, 0, 150_000, 1, 75)
    assert (c["charged"], c["paid"], c["debt"], c["pct"]) == (400_000, 350_000, 50_000, 87)
    assert (a["teacher_name"], b["teacher_name"]) == ("Пётр Иванов", "Мария Ким")
    t = col["total"]
    assert col["text"] == "Оплачено 1 900 000 из 2 500 000 сум (76%)"
    assert (t["charged"], t["paid"], t["forgiven"], t["debt"], t["students"], t["students_with_debt"], t["pct"]) == (
        2_500_000, 1_900_000, 100_000, 500_000, 6, 3, 76)
    for key in ("charged", "paid", "forgiven", "debt", "students", "students_with_debt"):
        assert t[key] == sum(r[key] for r in col["groups"]), key
    for r in [t, *col["groups"]]:
        assert r["paid"] + r["forgiven"] + r["debt"] == r["charged"]


def test_collection_equals_the_statements_of_every_membership(world, seeded):
    """Independent path: the charge rows of each membership's statement whose period begins in the month."""
    col = income(world, "2026-10-01", "2026-10-31")["collection"]
    seen = {}
    for key, m in seeded["memberships"].items():
        st = owner(world).get(f"{API}/org/memberships/{m.id}/statement").json()
        for row in st["rows"]:
            seg = row["details"]["segment"]
            if row["type"] == "charge" and "2026-10-01" <= seg["period_from"] <= "2026-10-31":
                g = seen.setdefault(st["membership"]["group_name"], dict(charged=0, paid=0, forgiven=0, debt=0))
                g["charged"] += seg["charge"]
                g["paid"] += seg["paid"]
                g["forgiven"] += seg["written_off"]
                g["debt"] += seg["remaining"]
    assert seen == {r["group_name"]: {k: r[k] for k in ("charged", "paid", "forgiven", "debt")} for r in col["groups"]}


def test_a_closed_month_is_collected_by_later_payments_and_advances(world, seeded):
    """September was paid partly in October (late money, an advance of September): the share is as of today."""
    o = income(world, "2026-09-01", "2026-09-30")
    assert o["period"]["in_progress"] is False and o["previous_period"]["label"] == "август"
    assert o["current"]["received"] == 1_500_000 and o["previous"]["received"] == 0 and o["delta_pct"] is None
    t = o["collection"]["total"]
    assert (t["charged"], t["paid"], t["debt"], t["pct"], t["students_with_debt"]) == (2_500_000, 2_500_000, 0, 100, 0)
    assert o["summary_text"] == "За сентябрь пришло 1 500 000 сум. Собрано 100% стоимости месяца."
    assert income(world, "2026-08-01", "2026-08-31")["collection"]["groups"] == []        # the accounting starts on 01.09: nothing charged


def test_overview_agrees_with_the_old_receipts_report(world, seeded):
    for first, last in (("2026-10-01", "2026-10-31"), ("2026-09-01", "2026-09-30"), ("2026-10-05", "2026-10-15")):
        old = receipts(world, first, last)["totals"]
        new = income(world, first, last)["current"]
        for key in ("received", "payments_count", "refunded", "refunded_from_received", "to_past_debt", "to_current", "to_advance",
                    "written_off", "net_received"):
            assert new[key] == old[key], (first, key)
        assert new["transferred"] == old["transferred_from_received"]


def test_filters_follow_the_receipts_report(world, seeded):
    gb, ga = seeded["groups"]["B"], seeded["groups"]["A"]
    for params in ({"group_id": str(gb.id)}, {"teacher_id": str(world.teacher.id)}, {"method": "cash"},
                   {"group_id": str(ga.id), "method": "cash"}):
        old = receipts(world, "2026-10-01", "2026-10-31", **params)
        new = income(world, "2026-10-01", "2026-10-31", **params)
        assert new["current"]["received"] == old["totals"]["received"], params
        assert new["current"]["to_current"] == old["totals"]["to_current"], params
        assert {k: new["position"][k] for k in old["position"]} == old["position"], params
        assert new["ok"], params
    only_b = income(world, "2026-10-01", "2026-10-31", group_id=str(gb.id))
    assert only_b["current"]["received"] == 1_050_000 and [g["group_name"] for g in only_b["collection"]["groups"]] == ["Английский 7Б"]
    by_teacher = income(world, "2026-10-01", "2026-10-31", teacher_id=str(world.teacher2.id))
    assert by_teacher["collection"]["total"]["charged"] == 600_000


# ---- periods ---------------------------------------------------------------------------------------------------


def test_day_today_with_the_cash_hint(world, seeded):
    o = income(world, "2026-10-21", "2026-10-21")
    assert o["period"]["kind"] == "day" and o["period"]["label"] == "сегодня" and o["previous_period"]["label"] == "вчера"
    assert (o["current"]["received"], o["current"]["payments_count"]) == (250_000, 2)       # a2 100 000 + b2 150 000, both cash
    assert o["previous"]["received"] == 0                                          # 20.10: only a voided payment
    assert o["summary_text"] == "Сегодня пришло 250 000 сум." and o["delta_direction"] is None and o["delta_text"] is None
    assert o["collection"] is None and [d["bucket"] for d in o["series"]] == ["2026-10-21"]
    assert o["cash"] == {
        "date": "2026-10-21", "label": "Наличные в кассе за сегодня", "cash_received": 250_000, "cash_refunded": 100_000,
        "expected": 150_000, "payments_count": 2,
        "text": "Наличные в кассе за сегодня: 150 000 сум (принято 250 000 сум, возвращено 100 000 сум)",
        "scope": "centre", "filtered_view": False, "scope_text": None}
    # R3 #7: the cash hint is the physical till of the WHOLE centre: the method, group and teacher filters never narrow it
    card = income(world, "2026-10-21", "2026-10-21", method="card")
    assert card["current"]["received"] == 0 and card["cash"]["expected"] == 150_000 and card["cash"]["filtered_view"] is False
    ga, gb = seeded["groups"]["A"], seeded["groups"]["B"]
    only_a = income(world, "2026-10-21", "2026-10-21", group_id=str(ga.id))
    assert only_a["current"]["received"] == 100_000                                       # the page itself IS filtered ...
    assert (only_a["cash"]["cash_received"], only_a["cash"]["cash_refunded"], only_a["cash"]["expected"], only_a["cash"]["payments_count"]) == (
        250_000, 100_000, 150_000, 2)                                                     # ... the till is not
    assert only_a["cash"]["scope"] == "centre" and only_a["cash"]["filtered_view"] is True
    assert only_a["cash"]["scope_text"] == "в кассе центра за сегодня (все группы)"
    assert only_a["cash"]["text"] == o["cash"]["text"] and only_a["cash"]["label"] == o["cash"]["label"]
    only_b = income(world, "2026-10-21", "2026-10-21", group_id=str(gb.id), method="card")
    assert only_b["current"]["received"] == 0 and only_b["cash"]["expected"] == 150_000 and only_b["cash"]["filtered_view"] is True
    by_teacher = income(world, "2026-10-21", "2026-10-21", teacher_id=str(world.teacher2.id))
    assert by_teacher["cash"]["expected"] == 150_000 and by_teacher["cash"]["filtered_view"] is True and by_teacher["cash"]["cash_received"] == 250_000
    past_filtered = income(world, "2026-10-15", "2026-10-15", group_id=str(ga.id))
    assert past_filtered["cash"]["scope_text"] == "в кассе центра за 15 октября (все группы)"
    past = income(world, "2026-10-15", "2026-10-15")
    assert past["cash"]["label"] == "Наличные в кассе за 15 октября" and past["cash"]["expected"] == 0 and past["cash"]["text"].endswith("0 сум")


def test_week_and_custom_ranges(world, seeded):
    o = income(world, "2026-10-12", "2026-10-18")
    assert o["period"]["kind"] == "week" and o["previous_period"]["from"] == "2026-10-05" and o["previous_period"]["to"] == "2026-10-11"
    assert o["current"]["received"] == 550_000 and o["previous"]["received"] == 1_750_000     # 200+350 vs 500+1250
    assert o["delta_pct"] == -69 and (o["delta_direction"], o["delta_text"]) == ("down", "69% к прошлой неделе")
    assert o["summary_text"] == "За неделю 12–18 октября пришло 550 000 сум — на 69% меньше, чем за предыдущую неделю."
    assert o["cash"] is None and o["collection"] is None and len(o["series"]) == 7
    c = income(world, "2026-10-02", "2026-10-08")                                        # a Friday to a Thursday: custom, 7 days
    assert c["period"]["kind"] == "custom" and c["previous_period"]["from"] == "2026-09-25" and c["period"]["label"] == "период 02.10–08.10"
    assert c["current"]["received"] == 900_000 + 500_000 + 1_250_000 and c["previous"]["received"] == 0
    assert c["summary_text"] == "За период 02.10–08.10 пришло 2 650 000 сум." and c["previous_label"] == "25 сент. – 1 окт."
    forced = income(world, "2026-10-01", "2026-10-21", period="month")
    assert forced["period"]["kind"] == "month" and forced["previous_period"]["to"] == "2026-09-21" and forced["collection"] is not None
    assert forced["current"]["received"] == 3_450_000


def test_series_has_every_day_and_long_ranges_go_by_month(world, seeded):
    o = income(world, "2026-10-01", "2026-10-31")
    assert o["series_by"] == "day" and [d["bucket"] for d in o["series"]] == [f"2026-10-{i:02d}" for i in range(1, 32)]
    by_day = {d["bucket"]: (d["received"], d["count"]) for d in o["series"]}
    assert by_day["2026-10-21"] == (250_000, 2) and by_day["2026-10-02"] == (900_000, 1) and by_day["2026-10-03"] == (0, 0)
    assert sum(v[0] for v in by_day.values()) == o["current"]["received"] and sum(v[1] for v in by_day.values()) == 7
    long = income(world, "2026-07-01", "2026-10-31")
    assert long["series_by"] == "month" and [d["bucket"] for d in long["series"]] == ["2026-07", "2026-08", "2026-09", "2026-10"]
    assert [d["received"] for d in long["series"]] == [0, 0, 1_500_000, 3_450_000] and long["ok"]


def test_period_errors(world, seeded):
    assert income(world, "2025-10-01", "2026-10-21", expect=422)["detail"]["code"] == "RANGE_TOO_LONG"
    assert income(world, "2026-10-05", "2026-10-04", expect=422)["detail"] == {"code": "INVALID_PERIOD", "message": "Конец периода раньше начала"}
    assert journal(world, "2026-10-05", "2026-10-04", expect=422)["detail"]["code"] == "INVALID_PERIOD"
    assert journal(world, "2025-10-01", "2026-10-21", expect=422)["detail"]["code"] == "RANGE_TOO_LONG"
    assert income(world, "2026-10-01", "2026-10-02", expect=422, period="day")["detail"]["code"] == "INVALID_PERIOD"
    assert income(world, "2026-10-02", "2026-10-31", expect=422, period="month")["detail"]["code"] == "INVALID_PERIOD"
    journal(world, expect=200, size=100)
    journal(world, expect=422, size=101)
    journal(world, expect=422, page=0)


def test_transfers_are_counted_apart(world, fixed_today):
    """A student moved to another group takes his money along: the part of the payment that travelled is «перенесено»."""
    db = world.db
    settings_block(world, accounting_start_date="2026-09-01", billing_anchor="fixed_day", billing_day=1)
    group_b = make_group(db, world.org, world.teacher2, "Группа B")
    group_b.schedule = world.group.schedule = []
    give_group_price(db, world.group, 500_000, valid_from=D(2026, 1, 1))
    give_group_price(db, group_b, 400_000, valid_from=D(2026, 1, 1))
    db.commit()
    a = membership(world)
    stay(world, a, D(2026, 10, 1))
    seed_payment(world, a, 800_000, D(2026, 10, 1), receipt_no=1)
    world.svc.add_student_to_group(str(world.org.id), str(group_b.id), str(world.students[0].id), mode="transfer",
                                   actor_id=str(world.owner.id), payments_detail=True)
    old = receipts(world, "2026-10-01", "2026-10-31")["totals"]
    o = income(world, "2026-10-01", "2026-10-31")
    assert old["transferred_from_received"] > 0
    assert o["current"]["transferred"] == old["transferred_from_received"] and o["current"]["received"] == 800_000
    cur = o["current"]
    assert cur["to_past_debt"] + cur["to_current"] + cur["to_forward"] + cur["to_advance"] + cur["refunded_from_received"] + cur["transferred"] == 800_000
    assert o["ok"]
    row = journal(world)["rows"][0]
    assert row["transferred"] == cur["transferred"] and f"перенесено в другую группу — {tx_num(cur['transferred'])} сум" in row["allocation_text"]


def tx_num(n):
    return f"{n:,}".replace(",", " ")


def test_lessons_mode_collects_by_the_package_that_began_in_the_month(world, fixed_today):
    """Lessons packages: a package «begins» with its first held lesson; the cost of a package that has held lessons is charged then."""
    from src.models.auth import GroupAttendance

    settings_block(world, period_mode="lessons", lessons_per_payment=12, accounting_start_date="2026-09-01",
                   billing_history=[{"from": "min", "period_mode": "lessons", "lessons_per_payment": 12}])
    give_group_price(world.db, world.group, 600_000, valid_from=D(2026, 1, 1))
    m = membership(world)
    stay(world, m, D(2026, 9, 25))
    s1 = world.students[0]
    for d in (D(2026, 9, 25), D(2026, 9, 28), D(2026, 9, 30), D(2026, 10, 2), D(2026, 10, 5), D(2026, 10, 7), D(2026, 10, 9),
              D(2026, 10, 12), D(2026, 10, 14), D(2026, 10, 16), D(2026, 10, 19), D(2026, 10, 21)):
        world.db.add(GroupAttendance(org_id=world.org.id, group_id=world.group.id, lesson_date=d.isoformat(), lesson_id=f"s{d.isoweekday()}",
                                     status="completed", records={str(s1.id): "present"}, roster_snapshot=[{"student_id": str(s1.id)}]))
    world.db.commit()
    seed_payment(world, m, 300_000, D(2026, 10, 1), receipt_no=1)
    sept = income(world, "2026-09-01", "2026-09-30")["collection"]
    assert [(r["group_name"], r["charged"], r["paid"], r["debt"], r["pct"]) for r in sept["groups"]] == [("Группа", 600_000, 300_000, 300_000, 50)]
    assert (sept["total"]["students_with_debt"], sept["total"]["students"]) == (1, 1)
    octo = income(world, "2026-10-01", "2026-10-31")
    assert octo["collection"]["groups"] == [] and octo["collection"]["total"]["pct"] is None
    assert octo["collection"]["text"] == "За этот месяц стоимость не считалась"
    assert "Собрано" not in octo["summary_text"] and octo["current"]["received"] == 300_000 and octo["ok"]


def test_previous_label_is_always_there_and_human(world, seeded, fixed_today):
    running = income(world, "2026-10-01", "2026-10-31")                                  # S2: a running month is compared day for day
    assert (running["previous_label"], running["previous_hint"]) == ("те же дни сентября", "1–21 сентября")
    assert income(world, "2026-10-21", "2026-10-21")["previous_label"] == "вчера"
    assert income(world, "2026-10-15", "2026-10-15")["previous_label"] == "14 окт."
    assert income(world, "2026-10-12", "2026-10-18")["previous_label"] == "5–11 окт."
    assert income(world, "2026-10-02", "2026-10-08")["previous_label"] == "25 сент. – 1 окт."
    sept = income(world, "2026-09-01", "2026-09-30")
    assert sept["previous"]["received"] == 0 and sept["previous_label"] == "август"            # there was nothing in August: the label is still there
    assert income(world, "2026-08-01", "2026-08-31")["previous_label"] == "июль"
    assert income(world, "2026-01-01", "2026-01-31")["previous_label"] == "декабрь 2025"
    fixed_today(D(2026, 10, 25))                                                               # Sunday: the week 19-25.10 is over, the one before is «прошлая»
    assert income(world, "2026-10-19", "2026-10-25")["previous_label"] == "прошлая неделя"


def test_method_unknown_is_the_same_everywhere(world, seeded):
    """«Способ не указан» (method=unknown, also a missing method): Σ journal = the till line of the overview = the filtered overview = the file."""
    from sqlalchemy.dialects import postgresql

    from src.auth.payments_reports import method_clause
    from src.models.auth import StudentPayment as SP

    assert "IS NULL" in str(method_clause(SP.method, "unknown").compile(dialect=postgresql.dialect()))
    assert str(method_clause(SP.method, "cash").compile(dialect=postgresql.dialect())).count("IS NULL") == 0
    m = seeded["memberships"]["a1"]
    seed_payment(world, m, 40_000, D(2026, 10, 8), receipt_no=21, method="unknown")
    seed_payment(world, m, 60_000, D(2026, 10, 9), receipt_no=22, method="unknown")
    o = income(world, "2026-10-01", "2026-10-31")
    till = next(x for x in o["by_method"] if x["method"] == "unknown")
    assert (till["amount"], till["count"], till["label"]) == (100_000, 2, "Способ не указан")
    j = journal(world, method="unknown")
    assert (j["totals"]["count"], j["totals"]["sum"]) == (till["count"], till["amount"]) and {r["method"] for r in j["rows"]} == {"unknown"}
    only = income(world, "2026-10-01", "2026-10-31", method="unknown")
    assert only["current"]["received"] == till["amount"] and [x["method"] for x in only["by_method"]] == ["unknown"] and only["ok"]
    old = receipts(world, "2026-10-01", "2026-10-31", method="unknown")["totals"]
    assert (old["received"], old["payments_count"]) == (100_000, 2)
    _, ws = journal_xlsx(world, method="unknown")
    _, _, body = table_of(ws)
    assert ws["A3"].value == "Оплат 2 · Оплачено 100 000" and "Способ: Способ не указан" in ws["A2"].value
    assert [row["Способ"] for row in body] == ["Способ не указан"] * 2


def test_journal_authors_ignore_created_by_but_follow_the_other_filters(world, seeded):
    """«Кто внёс» keeps all its options after one is chosen; the other filters narrow them."""
    both = {a["name"]: a["count"] for a in journal(world)["authors"]}
    assert both == {"Азиз Хакимов": 4, "Нодира Саидова": 3}
    chosen = journal(world, created_by=str(world.admin.id))
    assert {a["name"]: a["count"] for a in chosen["authors"]} == both and len(chosen["rows"]) == 4        # the rows are narrowed, the options are not
    assert {a["name"]: a["count"] for a in journal(world, method="cash")["authors"]} == {"Нодира Саидова": 3, "Азиз Хакимов": 1}
    assert {a["name"]: a["count"] for a in journal(world, method="click", created_by=str(world.admin.id))["authors"]} == {"Азиз Хакимов": 1}
    assert {a["name"]: a["count"] for a in journal(world, q="алиев")["authors"]} == {"Азиз Хакимов": 1}
    assert {a["name"]: a["count"] for a in journal(world, group_id=str(seeded["groups"]["A"].id))["authors"]} == {"Нодира Саидова": 3, "Азиз Хакимов": 1}
    assert {a["name"]: a["count"] for a in journal(world, status="voided")["authors"]} == both              # the status does not narrow them either


# ---- the journal ------------------------------------------------------------------------------------------------


def test_journal_rows_order_and_texts(world, seeded):
    j = journal(world, size=50)
    assert (j["total"], j["pages"], j["page"], j["size"]) == (9, 1, 1, 50)
    assert j["totals"] == {"count": 7, "sum": 3_450_000, "voided_count": 2, "replaced_count": 1, "kind_sum": None}
    assert [r["receipt_no"] for r in j["rows"]] == [13, 12, 10, 11, 9, 8, 7, 6, 5]                  # date desc, then time desc
    by = {r["receipt_no"]: r for r in j["rows"]}
    top = j["rows"][0]
    assert (top["student_name"], top["display_id"], top["group_name"], top["teacher_name"], top["amount"], top["method"],
            top["method_label"], top["created_by_name"], top["paid_on"], top["paid_time"], top["note"]) == (
        "Алиев Тимур", "INC-005", "Английский 7Б", "Мария Ким", 150_000, "cash", "Наличные", "Азиз Хакимов", "2026-10-21", "11:40",
        "=SUM(A1)")                                                                              # 06:40 UTC = 11:40 Tashkent
    assert top["allocation_text"] == "за октябрь — 150 000 сум" and top["status"] == "active"
    assert by[7]["allocation_text"] == "за сентябрь — 500 000 сум · за октябрь — 500 000 сум · аванс — 150 000 сум · возвращено — 100 000 сум"
    assert (by[7]["refunded"], by[7]["transferred"]) == (100_000, 0)
    assert by[5]["allocation_text"] == "за сентябрь — 300 000 сум · за октябрь — 300 000 сум · аванс — 300 000 сум"
    assert by[8]["allocation_text"] == "за сентябрь — 200 000 сум"
    assert by[6]["allocation_text"] == "за октябрь — 500 000 сум"
    assert [a["amount"] for a in by[7]["allocations"]] == [500_000, 500_000, 150_000]
    # corrections: shown, grey, linked
    assert (by[9]["status"], by[9]["replaced_by_receipt_no"], by[9]["allocation_text"], by[9]["allocations"], by[9]["void_reason"]) == (
        "replaced", 11, None, [], "Исправлено")
    assert (by[11]["status"], by[11]["replaces_receipt_no"], by[11]["replaces_amount"], by[11]["allocation_text"]) == (
        "active", 9, 400_000, "за октябрь — 350 000 сум")
    assert (by[10]["status"], by[10]["void_reason"], by[10]["voided_by_name"], by[10]["replaced_by_receipt_no"]) == (
        "voided", "Ошибочный ввод", "Нодира Саидова", None)
    assert [a["name"] for a in j["authors"]] == ["Азиз Хакимов", "Нодира Саидова"] and [a["count"] for a in j["authors"]] == [4, 3]
    assert all(r["entered_at_text"] is None and r["paid_time"] for r in j["rows"])                  # every seeded payment was entered on its own day


def test_journal_pagination_walks_every_row_once(world, seeded):
    seen, sums = [], 0
    for page in (1, 2, 3):
        j = journal(world, size=4, page=page)
        assert j["total"] == 9 and j["pages"] == 3 and j["totals"]["sum"] == 3_450_000
        seen += [r["id"] for r in j["rows"]]
        sums += sum(r["amount"] for r in j["rows"] if r["status"] == "active")
    assert len(seen) == 9 and len(set(seen)) == 9 and sums == 3_450_000
    assert journal(world, size=4, page=4)["rows"] == []


def test_journal_filters_and_the_total_equals_the_period_received(world, seeded):
    ga, gb = seeded["groups"]["A"], seeded["groups"]["B"]
    cases = [
        ({}, 9), ({"method": "cash"}, 5), ({"method": "click"}, 1), ({"group_id": str(gb.id)}, 3),
        ({"teacher_id": str(world.teacher2.id)}, 3), ({"teacher_id": str(world.teacher.id)}, 6),
        ({"created_by": str(world.admin.id)}, 4), ({"created_by": str(world.owner.id)}, 5), ({"q": "Алиев"}, 2), ({"q": "inc-005"}, 2),
        ({"q": "каримова алия"}, 1), ({"q": "N13"}, 4), ({"q": "7"}, 1), ({"group_id": str(ga.id), "method": "cash"}, 3),
        ({"status": "voided"}, 2), ({"status": "active"}, 7), ({"q": "нет такого"}, 0),
    ]
    for params, rows in cases:
        j = journal(world, **params)
        assert j["total"] == len(j["rows"]) == rows, params
    for params in ({}, {"method": "cash"}, {"group_id": str(gb.id)}, {"teacher_id": str(world.teacher.id)},
                   {"group_id": str(ga.id), "method": "cash"}, {"method": "bank_transfer"}):
        j = journal(world, **params)
        r = receipts(world, "2026-10-01", "2026-10-31", **params)["totals"]
        assert (j["totals"]["count"], j["totals"]["sum"]) == (r["payments_count"], r["received"]), params
        o = income(world, "2026-10-01", "2026-10-31", **params)
        assert o["current"]["received"] == j["totals"]["sum"] and o["ok"], params
    # the status filter does not change the totals (they are the active money of the set)
    assert journal(world, status="voided")["totals"] == {"count": 7, "sum": 3_450_000, "voided_count": 2, "replaced_count": 1, "kind_sum": None}
    only_day = journal(world, "2026-10-21", "2026-10-21")
    assert [r["receipt_no"] for r in only_day["rows"]] == [13, 12] and only_day["totals"]["sum"] == 250_000
    assert journal(world, "2026-09-01", "2026-09-30")["totals"] == {"count": 4, "sum": 1_500_000, "voided_count": 0, "replaced_count": 0, "kind_sum": None}


def test_journal_is_limited_to_the_callers_organisation_and_to_money_access(world, seeded):
    from src.models.auth import Organization

    for path in ("income", "journal", "journal/export"):
        for user in (world.teacher, world.teacher2, world.admin):          # an admin has no right to payments until the owner gives it
            r = make_client(world, user=user).get(f"{API}/org/payments/{path}", params={"from": "2026-10-01", "to": "2026-10-31"})
            assert r.status_code == 403, (path, user.role, r.status_code)
    settings_block(world, admin_can_manage=True)
    admin = make_client(world, user=world.admin)
    assert journal(world, client=admin)["totals"]["sum"] == 3_450_000
    assert income(world, "2026-10-01", "2026-10-31", client=admin)["current"]["received"] == 3_450_000
    world.db.query(Organization).filter_by(id=world.other_org.id).update({"settings": {"payments": {"admin_can_manage": True}}})
    world.db.commit()
    foreign = make_client(world, user=world.foreign_admin, org="other")
    assert journal(world, client=foreign)["rows"] == [] and journal(world, client=foreign)["totals"]["sum"] == 0
    assert journal(world, client=foreign, group_id=str(seeded["groups"]["A"].id))["rows"] == []
    assert journal(world, client=foreign, created_by=str(world.owner.id))["authors"] == []
    o = income(world, "2026-10-01", "2026-10-31", client=foreign)
    assert o["current"]["received"] == 0 and o["collection"]["groups"] == [] and o["position"]["debt_total"] == 0
    assert o["cash"] is None and o["summary_text"].startswith("За октябрь (по 21-е) оплат не было")
    _, ws = journal_xlsx(world, client=foreign)                                       # nothing of the other centre in the file
    assert table_of(ws)[2] == [] and ws["A3"].value == "Оплат 0 · Оплачено 0" and any(
        ws.cell(r, 1).value == "Оплат за период нет" for r in range(1, ws.max_row + 1))


# ---- Excel ------------------------------------------------------------------------------------------------------


JOURNAL_HEADER = ["№ квитанции", "Дата", "Ученик", "ID ученика", "Группа", "Преподаватель", "Оплачено", "Способ", "Кто внёс", "Внесена", "Статус",
                  "Примечание", "Подробности"]


def test_journal_excel_golden_cells_types_formats_and_fills(world, seeded):
    r, ws = journal_xlsx(world)
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert r.headers["content-disposition"] == ("attachment; filename=\"Zhurnal_oplat_oktyabr-2026.xlsx\"; "
                                                "filename*=UTF-8''%D0%96%D1%83%D1%80%D0%BD%D0%B0%D0%BB_%D0%BE%D0%BF%D0%BB%D0%B0%D1%82_%D0%BE%D0%BA%D1%82%D1%8F%D0%B1%D1%80%D1%8C-2026.xlsx")
    assert ws.title == "Журнал оплат"
    # the title block: three merged rows over the whole width, the as-of / period line, the totals in the words of the screen
    assert ws["A1"].value == "Журнал оплат — N13 org" and ws["A1"].font.bold and ws["A1"].font.size == 14
    assert ws["A2"].value == "По состоянию на 21.10.2026 · Период: Октябрь 2026"
    assert ws["A3"].value == "Оплат 7 · Оплачено 3 450 000 · Отменённых 1"                  # voided 2 - replaced 1
    assert {str(x) for x in ws.merged_cells.ranges} >= {"A1:M1", "A2:M2", "A3:M3"} and ws.cell(4, 1).value is None
    head, cols, body = table_of(ws)
    assert head == 5 and list(cols) == JOURNAL_HEADER and ws.freeze_panes == "D6" and ws.auto_filter.ref == "A5:M13"
    for c in range(1, 14):
        cell = ws.cell(5, c)
        assert cell.fill.fgColor.rgb.endswith("D9E2F3") and cell.font.bold and cell.font.name == "Calibri" and cell.alignment.wrap_text
    # ONE payment = ONE row; a replaced payment (№ 9) is never a row
    assert [row["№ квитанции"] for row in body] == [13, 12, 10, 11, 8, 7, 6, 5]
    from datetime import datetime as dt

    assert body[0] == {
        "№ квитанции": 13, "Дата": dt(2026, 10, 21), "Ученик": "Алиев Тимур", "ID ученика": "INC-005", "Группа": "Английский 7Б", "Преподаватель": "Мария Ким",
        "Оплачено": 150_000, "Способ": "Наличные", "Кто внёс": "Азиз Хакимов", "Внесена": dt(2026, 10, 21, 11, 40), "Статус": "Действует",
        "Примечание": "=SUM(A1)", "Подробности": "за октябрь — 150 000 сум"}
    # the types and formats of a row
    n = 6
    assert ws.cell(n, 1).data_type == "n" and ws.cell(n, 2).is_date and ws.cell(n, 2).number_format == "DD.MM.YYYY"
    assert ws.cell(n, 7).data_type == "n" and ws.cell(n, 7).number_format == "#,##0"
    assert ws.cell(n, 10).is_date and ws.cell(n, 10).number_format == "DD.MM.YYYY HH:MM"
    assert ws.cell(n, 12).data_type == "s" and ws.cell(n, 12).value.startswith("=")          # the note starts with "=": text, never a formula
    assert all(c.data_type != "f" for row in ws.iter_rows() for c in row)
    # a cancelled payment: grey, struck through, the reason in «Подробности» (grey, not struck); its amount is not in the totals
    by = {row["№ квитанции"]: i for i, row in enumerate(body, start=6)}
    gone = by[10]
    assert body[by[10] - 6]["Статус"] == "Отменена" and body[by[10] - 6]["Оплачено"] == 777_000
    assert body[by[10] - 6]["Подробности"] == "отменена 20.10 · Нодира Саидова · причина: Ошибочный ввод"
    assert ws.cell(gone, 7).font.strike and ws.cell(gone, 7).font.color.rgb.endswith("808080") and ws.cell(gone, 3).font.strike
    assert not ws.cell(gone, 13).font.strike and ws.cell(gone, 13).font.color.rgb.endswith("808080")
    # a corrected payment: the final amount, the version history under the allocation (the default reason «Исправлено» is not repeated)
    fixed = body[by[11] - 6]
    assert fixed["Оплачено"] == 350_000 and fixed["Статус"] == "Действует (исправлена)"
    assert fixed["Подробности"] == "за октябрь — 350 000 сум\nисправлена 15.10: было 400 000 сум (№ 9), способ: Карта · исправил Азиз Хакимов"
    assert ws.cell(by[11], 13).alignment.wrap_text and ws.row_dimensions[by[11]].height >= 30
    # the blue «Итого»: the active payments only
    total = 14
    assert ws.cell(total, 1).value == "Итого" and ws.cell(total, 3).value == "Оплат: 7" and ws.cell(total, 7).value == 3_450_000
    assert ws.cell(total, 7).number_format == "#,##0"
    assert all(ws.cell(total, c).fill.fgColor.rgb.endswith("D9E2F3") and ws.cell(total, c).font.bold for c in range(1, 14))
    assert ws.cell(total, 1).border.top.style == "medium"
    # «Итого по способам»: a small block under the table, the biggest first
    assert ws.cell(16, 1).value == "Итого по способам" and ws.cell(16, 1).fill.fgColor.rgb.endswith("EDF1F7") and "A16:M16" in {str(x) for x in ws.merged_cells.ranges}
    assert [ws.cell(17, c).value for c in (3, 4, 5)] == ["Способ", "Оплат", "Оплачено"] and ws.cell(17, 3).fill.fgColor.rgb.endswith("D9E2F3")
    assert [[ws.cell(r, c).value for c in (3, 4, 5)] for r in range(18, 22)] == [
        ["Click", 1, 1_250_000], ["Наличные", 4, 950_000], ["Перевод", 1, 900_000], ["Карта", 1, 350_000]]
    assert sum(ws.cell(r, 5).value for r in range(18, 22)) == ws.cell(total, 7).value
    # widths from the content: the date, the moment and the biggest sum fit; the long texts wrap
    width = {name: ws.column_dimensions[chr(64 + c)].width for name, c in cols.items()}
    assert width["Дата"] >= 12 and width["Внесена"] >= 18 and width["Оплачено"] >= len("1,250,000") + 2 and width["№ квитанции"] >= len("№ квитанции") - 1
    assert width["Статус"] >= len("Действует (исправлена)") and width["Подробности"] == 60
    assert ws.page_setup.orientation == "landscape" and ws.page_setup.fitToWidth == 1 and ws.print_title_rows == "$5:$5"
    # the old words and the old blue are gone
    flat = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
    assert not any(x in flat for x in ("Сумма", "Куда пошли", "Время", "На сумму", "Исправлена (№ 11)")) and not any("Начисл" in x for x in flat)
    assert all("2F6FED" not in str(c.fill.fgColor.rgb) for row in ws.iter_rows() for c in row)


def test_journal_file_never_holds_a_replaced_payment_whatever_the_status(world, seeded):
    want = {"all": [13, 12, 10, 11, 8, 7, 6, 5], "active": [13, 12, 11, 8, 7, 6, 5], "voided": [10], "cancelled": [10], "replaced": []}
    for status, receipts_ in want.items():
        _, ws = journal_xlsx(world, status=status)
        assert [row["№ квитанции"] for row in table_of(ws)[2]] == receipts_, status
        _, _, body = journal_csv(world, status=status)
        assert [int(row[0]) for row in body] == receipts_, status
    _, ws = journal_xlsx(world, status="replaced")
    assert "заменённые версии в файл не выводятся" in ws["A2"].value and any(ws.cell(r, 1).value == "Оплат за период нет" for r in range(1, ws.max_row + 1))
    _, ws = journal_xlsx(world, status="cancelled")
    assert ws["A3"].value == "Отменённых 1" and "Только отменённые" in ws["A2"].value and table_of(ws)[2][0]["Статус"] == "Отменена"
    assert not any(ws.cell(r, 1).value == "Итого по способам" for r in range(1, ws.max_row + 1))          # no active payment: no block
    assert not any(c.value == 9 for row in journal_xlsx(world)[1].iter_rows(min_col=1, max_col=1) for c in row)


def test_journal_csv_has_the_same_table_without_a_replaced_row(world, seeded):
    top, header, body = journal_csv(world)
    assert header == JOURNAL_HEADER and [int(row[0]) for row in body] == [13, 12, 10, 11, 8, 7, 6, 5]
    assert top[0] == ["Журнал оплат — N13 org"] and top[2] == ["Оплат 7 · Оплачено 3 450 000 · Отменённых 1"]
    assert body[0][:2] == ["13", "21.10.2026"] and body[0][9] == "21.10.2026 11:40" and body[0][11] == "'=SUM(A1)"
    assert body[3][12] == "за октябрь — 350 000 сум | исправлена 15.10: было 400 000 сум (№ 9), способ: Карта · исправил Азиз Хакимов"
    rows_, _, _ = journal_csv(world)
    assert not any(p.strip().startswith("=") for r_ in (rows_, body) for line in r_ for p in line)


def test_a_filtered_file_says_its_filters_in_words_and_totals_what_it_shows(world, seeded):
    _, ws = journal_xlsx(world, method="click")
    assert ws["A2"].value == "По состоянию на 21.10.2026 · Период: Октябрь 2026 · Способ: Click" and ws["A3"].value == "Оплат 1 · Оплачено 1 250 000"
    assert [row["№ квитанции"] for row in table_of(ws)[2]] == [7]
    _, ws = journal_xlsx(world, group_id=str(seeded["groups"]["B"].id), created_by=str(world.owner.id), q="алиев")
    assert ws["A2"].value == "По состоянию на 21.10.2026 · Период: Октябрь 2026 · Группа: Английский 7Б · Кто внёс: Азиз Хакимов · Поиск: «алиев»"
    assert [row["№ квитанции"] for row in table_of(ws)[2]] == [13]
    _, ws = journal_xlsx(world, **{"from": "2026-10-21", "to": "2026-10-21"})
    assert ws["A2"].value == "По состоянию на 21.10.2026 · Период: 21.10.2026–21.10.2026"


def test_the_journal_file_name_is_russian_for_a_month_a_day_and_a_range(world, seeded):
    def name(**params):
        r = owner(world).get(f"{API}/org/payments/journal/export", params=params)
        from urllib.parse import unquote

        return unquote(r.headers["content-disposition"].split("filename*=UTF-8''")[1])

    assert name(**{"from": "2026-10-01", "to": "2026-10-31"}) == "Журнал_оплат_октябрь-2026.xlsx"
    assert name(**{"from": "2026-10-21", "to": "2026-10-21"}) == "Журнал_оплат_21.10.2026.xlsx"
    assert name(**{"from": "2026-10-05", "to": "2026-10-21"}) == "Журнал_оплат_05.10-21.10.2026.xlsx"
    assert name(**{"from": "2026-09-20", "to": "2026-10-21"}) == "Журнал_оплат_20.09-21.10.2026.xlsx"
    assert name(**{"from": "2025-12-20", "to": "2026-01-10"}) == "Журнал_оплат_20.12.2025-10.01.2026.xlsx"


# ---- history[] and the new status values of the API --------------------------------------------------------------


def test_the_journal_api_has_history_on_an_active_corrected_row_oldest_first(world, seeded):
    rows = {r["receipt_no"]: r for r in journal(world, size=100)["rows"]}
    assert rows[11]["history"] == [{"receipt_no": 9, "amount": 400_000, "method_label": "Карта", "replaced_at": "2026-10-15T06:30:00",
                                    "replaced_by_name": "Азиз Хакимов", "reason": "Исправлено"}]
    assert all(r["history"] == [] for n, r in rows.items() if n != 11)                          # never corrected, replaced (9) or cancelled (10): empty
    assert (rows[11]["replaces_receipt_no"], rows[11]["replaces_amount"]) == (9, 400_000)         # the older fields stay
    # a correction of a correction: the OLDEST version first
    db, o = world.db, world.owner
    p11, c1 = seeded["payments"]["c1-oct15"], seeded["memberships"]["c1"]
    p11.voided_at, p11.voided_by, p11.void_reason = datetime(2026, 10, 16, 8, 0), world.admin.id, "Сумма не та"
    db.add(StudentPayment(id=uid(world.org, "pay:c1-oct15-again"), org_id=world.org.id, membership_id=c1.id, student_id=c1.student_id, group_id=c1.group_id,
                          amount=360_000, paid_on=D(2026, 10, 15), method="cash", receipt_no=14, created_by=o.id, created_at=datetime(2026, 10, 16, 8, 0),
                          replaces_payment_id=p11.id))
    db.commit()
    rows = {r["receipt_no"]: r for r in journal(world, size=100)["rows"]}
    assert [(h["receipt_no"], h["amount"], h["method_label"], h["reason"]) for h in rows[14]["history"]] == [
        (9, 400_000, "Карта", "Исправлено"), (11, 350_000, "Карта", "Сумма не та")]
    assert rows[14]["history"][1]["replaced_by_name"] == "Нодира Саидова" and rows[11]["history"] == [] and rows[11]["status"] == "replaced"
    _, ws = journal_xlsx(world)
    text = next(row["Подробности"] for row in table_of(ws)[2] if row["№ квитанции"] == 14)
    assert text.splitlines()[1:] == ["исправлена 15.10: было 400 000 сум (№ 9), способ: Карта · исправил Азиз Хакимов",
                                     "исправлена 16.10: было 350 000 сум (№ 11), способ: Карта · исправил Нодира Саидова · причина: Сумма не та"]
    assert [row["№ квитанции"] for row in table_of(ws)[2]] == [13, 12, 10, 14, 8, 7, 6, 5]      # neither 9 nor 11 is a row any more


def test_journal_status_cancelled_and_replaced(world, seeded):
    cancelled, replaced, voided = (journal(world, status=s) for s in ("cancelled", "replaced", "voided"))
    assert [r["receipt_no"] for r in cancelled["rows"]] == [10] and [r["status"] for r in cancelled["rows"]] == ["voided"]
    assert [r["receipt_no"] for r in replaced["rows"]] == [9] and [r["status"] for r in replaced["rows"]] == ["replaced"]
    assert sorted(r["receipt_no"] for r in voided["rows"]) == [9, 10]                                # «voided» is both, as before
    assert (cancelled["total"], replaced["total"], voided["total"]) == (1, 1, 2) and (cancelled["pages"], replaced["pages"]) == (1, 1)
    for j in (cancelled, replaced, voided):                                                          # the counters do not depend on the status
        assert (j["totals"]["voided_count"], j["totals"]["replaced_count"], j["totals"]["count"], j["totals"]["sum"]) == (2, 1, 7, 3_450_000)
    assert journal(world, status="cancelled", size=1, page=2)["rows"] == []
    assert journal(world, status="cancelled", kind="advance")["rows"] == [] and journal(world, status="replaced", kind="advance")["total"] == 0
    assert owner(world).get(f"{API}/org/payments/journal", params={"from": "2026-10-01", "to": "2026-10-31", "status": "bogus"}).status_code == 422


# ---- the budget of queries --------------------------------------------------------------------------------------


def test_queries_do_not_grow_with_the_number_of_payments_and_memberships(world, fixed_today):
    settings_block(world, accounting_start_date="2026-09-01", billing_anchor="fixed_day", billing_day=1)
    give_group_price(world.db, world.group, 500_000)
    world.group.schedule = []
    world.db.commit()
    c = owner(world)
    probes = {
        "income_month": lambda: c.get(f"{API}/org/payments/income", params={"from": "2026-10-01", "to": "2026-10-31"}).json(),
        "income_day": lambda: c.get(f"{API}/org/payments/income", params={"from": "2026-10-21", "to": "2026-10-21"}).json(),
        "income_filtered": lambda: c.get(f"{API}/org/payments/income", params={"from": "2026-09-01", "to": "2026-10-31", "method": "cash"}).json(),
        "journal": lambda: c.get(f"{API}/org/payments/journal", params={"from": "2026-09-01", "to": "2026-10-31", "size": 100}).json(),
        "journal_q": lambda: c.get(f"{API}/org/payments/journal", params={"from": "2026-09-01", "to": "2026-10-31", "q": "Имя", "size": 100}).json(),
        "export": lambda: c.get(f"{API}/org/payments/journal/export", params={"from": "2026-09-01", "to": "2026-10-31"}).content,
    }
    grow(world, 12)
    small = {name: count_statements(world, fn)[1] for name, fn in probes.items()}
    grow(world, 190, start=12)
    big = {name: count_statements(world, fn)[1] for name, fn in probes.items()}
    print("QUERIES small/big:", small, big)
    assert all(n <= 32 for n in big.values()), big
    for name in big:
        assert big[name] <= small[name] + 2, (name, small, big)
    full = c.get(f"{API}/org/payments/journal", params={"from": "2026-09-01", "to": "2026-10-31", "size": 100}).json()
    assert full["total"] >= 200 and len(full["rows"]) == 100 and full["totals"]["count"] == full["total"]


# ---- identities on random data ----------------------------------------------------------------------------------


@pytest.mark.parametrize("seed", [1, 2, 3, 4, 5, 6])
def test_identities_hold_on_random_data(world, fixed_today, seed):
    rnd = random.Random(seed)
    db = world.db
    settings_block(world, accounting_start_date="2026-08-01", billing_anchor=rnd.choice(["fixed_day", "payment_date"]), billing_day=rnd.choice([1, 5, 15]))
    groups = [world.group] + [make_group(db, world.org, rnd.choice([world.teacher, world.teacher2]), f"Г{i}") for i in range(3)]
    for g in groups:
        g.schedule = []
    db.commit()
    for g in groups:
        give_group_price(db, g, rnd.choice([200_000, 350_000, 500_000]), valid_from=D(2026, 1, 1))
    members = []
    for i in range(14):
        g = rnd.choice(groups)
        s = make_student(db, world.org, g, f"Имя{i}", f"Фам{i}", f"R-{seed}-{i:02d}")
        db.commit()
        m = membership(world, s, g)
        first = D(2026, 8, 1) + timedelta(days=rnd.randint(0, 70))
        stay(world, m, first, None if rnd.random() < 0.8 else first + timedelta(days=rnd.randint(10, 40)))
        members.append(m)
    receipt = 0
    for m in members:
        for _ in range(rnd.randint(0, 4)):
            receipt += 1
            day = D(2026, 8, 1) + timedelta(days=rnd.randint(0, 80))
            p = seed_payment(world, m, rnd.choice([50_000, 120_000, 200_000, 350_000, 500_000, 900_000]), day, receipt_no=receipt,
                             method=rnd.choice(["cash", "card", "click", "bank_transfer", "payme", "other", "unknown"]))
            if rnd.random() < 0.12:
                p.voided_at, p.voided_by, p.void_reason = datetime(2026, 10, 1, 8), world.owner.id, "Ошибка"
        if rnd.random() < 0.3:
            db.add(StudentWriteoff(org_id=world.org.id, membership_id=m.id, student_id=m.student_id, group_id=m.group_id,
                                   amount=rnd.choice([10_000, 50_000]), reason="льгота", effective_on=D(2026, 10, rnd.randint(1, 20)),
                                   created_by=world.owner.id))
    db.commit()
    for first, last in (("2026-10-01", "2026-10-31"), ("2026-09-01", "2026-09-30"), ("2026-08-01", "2026-10-21"), ("2026-10-10", "2026-10-21")):
        for params in ({}, {"method": "cash"}, {"group_id": str(groups[1].id)}, {"teacher_id": str(world.teacher2.id)}):
            o = income(world, first, last, **params)
            j = journal(world, first, last, size=100, **params)
            old = receipts(world, first, last, **params)["totals"]
            assert o["ok"], (seed, first, last, params, o["checks"])
            assert o["current"]["received"] == j["totals"]["sum"] == old["received"], (seed, first, last, params)
            assert j["totals"]["count"] == old["payments_count"]
            assert sum(m["amount"] for m in o["by_method"]) == o["current"]["received"]
            col = o["collection"]
            if col:
                for r in [col["total"], *col["groups"]]:
                    assert r["paid"] + r["forgiven"] + r["debt"] == r["charged"], (seed, r)
                    assert 0 <= (r["pct"] or 0) <= 100
                for key in ("charged", "paid", "forgiven", "debt"):
                    assert col["total"][key] == sum(r[key] for r in col["groups"])
                assert col["total"]["students_with_debt"] <= sum(r["students_with_debt"] for r in col["groups"])
                pcts = [r["paid"] / r["charged"] for r in col["groups"]]
                assert pcts == sorted(pcts)


# ---- R3 fixes -----------------------------------------------------------------------------------------------------


def add_payments(world, seeded, *rows):
    """``(key, paid_on, created_at UTC, amount, note)`` for the membership a1; returns ``{key: payment id string}``."""
    m = seeded["memberships"]["a1"]
    out = {}
    for n, (key, paid_on, created_at, amount, note) in enumerate(rows):
        p = StudentPayment(id=uid(world.org, "r3:" + key), org_id=world.org.id, membership_id=m.id, student_id=m.student_id, group_id=m.group_id, amount=amount,
                           paid_on=paid_on, method="cash", receipt_no=900 + n, created_by=world.owner.id, created_at=created_at, note=note)
        world.db.add(p)
        out[key] = str(p.id)
    world.db.commit()
    return out


def test_paid_time_is_the_entry_time_only_for_a_payment_entered_on_its_own_day(world, seeded):
    """R3 #2: a back-dated payment has no time of its own (null) and says when it was entered; the Tashkent midnight (19:30 UTC = 00:30 next day) decides, not UTC."""
    ids = add_payments(
        world, seeded,
        ("backdated", D(2026, 10, 15), datetime(2026, 10, 21, 6, 40), 10_000, None),             # entered 21.10 11:40 Tashkent for 15.10
        ("after-midnight-own-day", D(2026, 10, 21), datetime(2026, 10, 20, 19, 30), 20_000, None),   # 00:30 on 21.10 Tashkent, for 21.10
        ("after-midnight-utc-day", D(2026, 10, 20), datetime(2026, 10, 20, 19, 30), 30_000, None))   # 00:30 on 21.10 Tashkent, but for 20.10: back-dated
    rows = {r["id"]: r for r in journal(world, "2026-10-01", "2026-10-31", size=100)["rows"]}
    back = rows[ids["backdated"]]
    assert (back["paid_time"], back["entered_at_text"], back["created_at"]) == (None, "внесена 21.10 в 11:40", "2026-10-21T06:40:00")
    own = rows[ids["after-midnight-own-day"]]
    assert (own["paid_time"], own["entered_at_text"]) == ("00:30", None)
    utc_day = rows[ids["after-midnight-utc-day"]]
    assert (utc_day["paid_time"], utc_day["entered_at_text"], utc_day["created_at"]) == (None, "внесена 21.10 в 00:30", "2026-10-20T19:30:00")
    # Excel / CSV: «Внесена» is the real moment of entry (Tashkent) of EVERY payment - one column, no gaps: the old «Время» / «Внесена» pair is gone
    _, ws = journal_xlsx(world)
    _, _, body = table_of(ws)
    assert "Время" not in [ws.cell(5, c).value for c in range(1, 14)]
    by_receipt = {row["№ квитанции"]: row for row in body}
    from datetime import datetime as dt

    assert (by_receipt[900]["Дата"], by_receipt[900]["Внесена"]) == (dt(2026, 10, 15), dt(2026, 10, 21, 11, 40))
    assert (by_receipt[901]["Дата"], by_receipt[901]["Внесена"]) == (dt(2026, 10, 21), dt(2026, 10, 21, 0, 30))
    assert (by_receipt[902]["Дата"], by_receipt[902]["Внесена"]) == (dt(2026, 10, 20), dt(2026, 10, 21, 0, 30))
    assert by_receipt[13]["Внесена"] == dt(2026, 10, 21, 11, 40) and all(row["Внесена"] is not None for row in body)
    _, _, csv_body = journal_csv(world)
    assert {int(row[0]): row[9] for row in csv_body}[900] == "21.10.2026 11:40"


def test_a_non_ascii_digit_in_the_search_is_text_not_a_receipt_number(world, seeded):
    """R3 #3: ``"²".isdigit()`` is True but ``int("²")`` raises - the receipt number is looked up for ASCII digits only; any other glyph is plain text (200, no rows)."""
    for q in ("²", "٣", "①", "１２"):                                                   # superscript, Arabic-Indic, circled, full-width
        j = journal(world, q=q)
        assert j["total"] == 0 and j["rows"] == [], q
    assert journal(world, q="7")["total"] == 1                                       # the receipt number still works for ASCII digits


def test_csv_cells_starting_with_tab_or_cr_are_neutralised(world, seeded):
    """R3 #4 (OWASP): not only = + - @ but a leading TAB / CR makes a spreadsheet read a formula."""
    add_payments(world, seeded, ("tab", D(2026, 10, 21), datetime(2026, 10, 21, 7, 0), 1_000, "\t=HYPERLINK(1)"),
                 ("cr", D(2026, 10, 21), datetime(2026, 10, 21, 7, 1), 2_000, "\r=HYPERLINK(2)"),
                 ("minus", D(2026, 10, 21), datetime(2026, 10, 21, 7, 2), 3_000, "-1+1"))
    r = owner(world).get(f"{API}/org/payments/journal/export", params={"from": "2026-10-21", "to": "2026-10-21", "format": "csv"})
    import csv as csv_module

    cells = [c for line in csv_module.reader(io.StringIO(r.text.lstrip("\ufeff")), delimiter=";") for c in line]
    assert "'\t=HYPERLINK(1)" in cells and "'\r=HYPERLINK(2)" in cells and "'-1+1" in cells and "'=SUM(A1)" in cells
    assert not [c for c in cells if c[:1] in ("=", "+", "-", "@", "\t", "\r")]


def test_export_too_long_cap_and_the_voided_export(world, seeded, monkeypatch):
    """R3 #10: the 20 000-row cap (monkeypatched low) answers 422 EXPORT_TOO_LONG for both formats and counts the SELECTED status; a voided-only export holds only the voided rows."""
    from src.auth import payments_income

    params = {"from": "2026-10-01", "to": "2026-10-31"}
    monkeypatch.setattr(payments_income, "MAX_EXPORT_ROWS", 3)
    for fmt in ("xlsx", "csv"):
        r = owner(world).get(f"{API}/org/payments/journal/export", params={**params, "format": fmt})
        assert r.status_code == 422 and r.json()["detail"]["code"] == "EXPORT_TOO_LONG" and "Слишком много оплат" in r.json()["detail"]["message"], fmt
    assert owner(world).get(f"{API}/org/payments/journal/export", params={**params, "status": "active"}).status_code == 422      # 7 live rows > 3
    ok = owner(world).get(f"{API}/org/payments/journal/export", params={**params, "status": "voided"})                       # 1 emitted row (the cancelled one) <= 3
    assert ok.status_code == 200
    monkeypatch.setattr(payments_income, "MAX_EXPORT_ROWS", 1)
    assert owner(world).get(f"{API}/org/payments/journal/export", params={**params, "status": "voided"}).status_code == 200       # exactly the cap passes: a replaced row is no row
    assert owner(world).get(f"{API}/org/payments/journal/export", params={**params, "status": "all"}).status_code == 422          # 8 rows
    monkeypatch.setattr(payments_income, "MAX_EXPORT_ROWS", 0)
    assert owner(world).get(f"{API}/org/payments/journal/export", params={**params, "status": "voided"}).status_code == 422
    assert owner(world).get(f"{API}/org/payments/journal/export", params={**params, "status": "replaced"}).status_code == 200      # an empty file is always fine
    monkeypatch.setattr(payments_income, "MAX_EXPORT_ROWS", 20_000)
    voided = journal(world, status="voided", size=100)
    assert voided["total"] == 2 and sorted(r["receipt_no"] for r in voided["rows"]) == [9, 10]                                      # the page still shows both
    for fmt in ("xlsx", "csv"):
        if fmt == "xlsx":
            body = [(row["№ квитанции"], row["Статус"]) for row in table_of(journal_xlsx(world, status="voided")[1])[2]]
        else:
            body = [(int(row[0]), row[10]) for row in journal_csv(world, status="voided")[2]]
        assert body == [(10, "Отменена")], fmt                                                                                      # the file never holds the replaced № 9



# ---- S1: «Куда ушли деньги» (the current allocation, each piece relative to the payment's own date) -------------------------


def split_identities(world, first, last, **filters):
    """The strip sums to ``received``; the rows of the SAME filters (their ``split``) sum to the strip kind by kind; every row sums to its amount (allocations +
    refunded + transferred); the advance part never exceeds the position's advance."""
    o = income(world, first, last, **filters)
    strip = {x["kind"]: x["amount"] for x in o["strip"]}
    assert list(strip) == ["old_debt", "current", "forward", "advance", "refunded", "transferred"]
    cur = o["current"]
    assert strip == {"old_debt": cur["to_past_debt"], "current": cur["to_current"], "forward": cur["to_forward"], "advance": cur["to_advance"],
                     "refunded": cur["refunded_from_received"], "transferred": cur["transferred"]}
    assert sum(strip.values()) == cur["received"], (filters, strip)
    rows = [r for r in journal(world, first, last, size=100, status="active", **filters)["rows"]]
    by_kind = {}
    for r in rows:
        assert sum(x["amount"] for x in r["split"]) == r["amount"], r["receipt_no"]
        assert sum(a["amount"] for a in r["allocations"]) + r["refunded"] + r["transferred"] == r["amount"], r["receipt_no"]
        for x in r["split"]:
            by_kind[x["kind"]] = by_kind.get(x["kind"], 0) + x["amount"]
    assert by_kind == {k: v for k, v in strip.items() if v}, (filters, by_kind, strip)
    assert strip["advance"] <= o["position"]["advance_total"], (filters, strip["advance"], o["position"]["advance_total"])
    return o


def test_where_the_money_went_identities(world, seeded):
    ga, gb = seeded["groups"]["A"], seeded["groups"]["B"]
    for first, last in (("2026-10-01", "2026-10-31"), ("2026-09-01", "2026-09-30"), ("2026-09-01", "2026-10-31"), ("2026-10-12", "2026-10-18"),
                        ("2026-10-21", "2026-10-21")):
        for filters in ({}, {"method": "cash"}, {"group_id": str(ga.id)}, {"group_id": str(gb.id), "method": "card"}, {"teacher_id": str(world.teacher2.id)}):
            split_identities(world, first, last, **filters)


def test_a_back_dated_payment_is_forward_once_its_period_has_begun(world, seeded, fixed_today):
    """S1: a1 has paid September and October; a payment entered on 21.10 for 13.10 pays the NOVEMBER period.  While November has not begun the money is «аванс» (not consumed);
    on 5.11 the period has begun: it is «вперёд (за будущие периоды)», NOT advance - the period starts after the payment's date but the money is spent."""
    ids = add_payments(world, seeded, ("back", D(2026, 10, 13), datetime(2026, 10, 21, 6, 0), 500_000, None))
    row = {r["id"]: r for r in journal(world, "2026-10-01", "2026-10-31", size=100)["rows"]}[ids["back"]]
    assert [(x["kind"], x["amount"]) for x in row["split"]] == [("advance", 500_000)] and row["allocation_text"] == "аванс — 500 000 сум"
    before = split_identities(world, "2026-10-01", "2026-10-31")
    assert {x["kind"]: x["amount"] for x in before["strip"]}["forward"] == 0
    fixed_today(D(2026, 11, 5))
    after = split_identities(world, "2026-10-01", "2026-10-31")
    row = {r["id"]: r for r in journal(world, "2026-10-01", "2026-10-31", size=100)["rows"]}[ids["back"]]
    assert [(x["kind"], x["label"], x["amount"]) for x in row["split"]] == [("forward", "вперёд (за будущие периоды)", 500_000)]
    assert row["allocation_text"] == "за ноябрь — 500 000 сум" and row["allocations"][0]["bucket"] == "advance"          # relative to ITS date: it is the future period
    strip = {x["kind"]: x["amount"] for x in after["strip"]}
    assert strip["forward"] >= 500_000
    split_identities(world, "2026-10-12", "2026-10-18")                                                                   # the week of the payment
    split_identities(world, "2026-09-01", "2026-11-30")


# ---- S2: a running week ---------------------------------------------------------------------------------------------


def test_a_running_week_is_worded_and_compared_day_for_day(world, seeded):
    o = income(world, "2026-10-19", "2026-10-25")                                      # Mon 19 - Sun 25, today Wed 21
    assert o["period"]["kind"] == "week" and o["period"]["in_progress"] is True
    assert (o["previous_label"], o["previous_hint"]) == ("те же дни прошлой недели", "тем же дням прошлой недели (12–14 окт.)")
    assert o["current"]["received"] == 250_000 and o["previous"]["received"] == 200_000                     # 12.10: a2 200 000
    assert o["summary_text"] == "За эту неделю (пн–ср) пришло 250 000 сум — на 25% больше, чем за те же дни прошлой недели."
    month = income(world, "2026-10-01", "2026-10-31")
    assert month["summary_text"].startswith("За октябрь (по 21-е) пришло 3 450 000 сум — на 130% больше, чем за те же дни сентября.")
    closed = income(world, "2026-09-01", "2026-09-30")
    assert closed["previous_hint"] == "августу" and income(world, "2026-10-21", "2026-10-21")["previous_hint"] == "вчерашнему дню"
    assert o["delta_text"] == "25% к тем же дням прошлой недели (12–14 окт.)"


# ---- S3: the journal by kind, the correction details, status=active -----------------------------------------------------------

KINDS = ("old_debt", "current", "forward", "advance", "refunded", "transferred")
PERIODS = (("2026-10-01", "2026-10-31"), ("2026-09-01", "2026-09-30"), ("2026-09-01", "2026-10-31"), ("2026-10-12", "2026-10-18"), ("2026-10-21", "2026-10-21"))


def kind_part(row, kind):
    return sum(x["amount"] for x in row["split"] if x["kind"] == kind)


def check_kind(world, first, last, **filters):
    """One kind at a time: the rows are the ACTIVE payments that have the kind (their whole split stays), ``kind_sum`` is the kind's part of those rows and equals the
    strip's amount of the same period and filters; the kinds add up to the period's «Поступило»."""
    o = income(world, first, last, **{k: v for k, v in filters.items() if k not in ("q", "created_by")})      # the overview has no search / author filter
    strip = {x["kind"]: x["amount"] for x in o["strip"]}
    active = journal(world, first, last, size=100, status="active", **filters)
    if not ({"q", "created_by"} & set(filters)):
        assert active["totals"]["sum"] == o["current"]["received"], (first, last, filters)
    total_kinds = 0
    for kind in KINDS:
        j = journal(world, first, last, size=100, kind=kind, **filters)
        rows = j["rows"]
        assert j["totals"]["kind_sum"] == sum(kind_part(r, kind) for r in rows), (first, last, filters, kind)
        if not ({"q", "created_by"} & set(filters)):
            assert j["totals"]["kind_sum"] == strip[kind], (first, last, filters, kind)
        assert {r["id"] for r in rows} == {r["id"] for r in active["rows"] if kind_part(r, kind) > 0}, (first, last, filters, kind)
        assert all(r["status"] == "active" and r["split"] for r in rows)                                    # the row keeps its whole split
        assert (j["total"], j["totals"]["count"], j["pages"] == (1 if rows else 0)) == (len(rows), len(rows), True), (kind, j["total"])
        assert j["totals"]["sum"] == sum(r["amount"] for r in rows)                                          # whole amounts of the shown rows
        assert (j["totals"]["voided_count"], j["totals"]["replaced_count"]) == (0, 0)
        total_kinds += j["totals"]["kind_sum"]
    assert total_kinds == active["totals"]["sum"], (first, last, filters)                                    # every split adds up to its payment
    return strip


def test_journal_kind_filter_adds_up_to_the_strip_over_periods_and_filters(world, seeded):
    ga, gb = seeded["groups"]["A"], seeded["groups"]["B"]
    seen = set()
    for first, last in PERIODS:
        for filters in ({}, {"method": "cash"}, {"group_id": str(ga.id)}, {"group_id": str(gb.id), "method": "card"}, {"teacher_id": str(world.teacher2.id)},
                        {"method": "cash", "q": "алиев"}, {"created_by": str(world.admin.id)}):
            strip = check_kind(world, first, last, **filters)
            seen |= {k for k, v in strip.items() if v}
    assert {"old_debt", "current", "advance", "refunded"} <= seen, seen                                        # the world really exercises the kinds


def test_journal_kind_hand_calculated_october(world, seeded):
    """October: a3 paid 1 250 000 on 07.10 (500k Sep + 500k Oct + 150k advance, 100k of it refunded); the kind filter shows the payment and its whole split."""
    j = journal(world, kind="advance", size=100)
    by = {r["receipt_no"]: r for r in j["rows"]}
    assert 7 in by and j["totals"]["kind_sum"] == sum(kind_part(r, "advance") for r in j["rows"]) > 0
    assert [(x["kind"], x["amount"]) for x in by[7]["split"]] == [(x["kind"], x["amount"]) for x in
                                                                    journal(world, size=100)["rows"][[r["receipt_no"] for r in journal(world, size=100)["rows"]].index(7)]["split"]]
    assert kind_part(by[7], "refunded") == 100_000
    r = journal(world, kind="refunded", size=100)
    assert [x["receipt_no"] for x in r["rows"]] == [7] and r["totals"] == {"count": 1, "sum": 1_250_000, "voided_count": 0, "replaced_count": 0, "kind_sum": 100_000}
    # the cancelled / replaced payments have no split: a kind never shows them, whatever the status
    for status in ("all", "active", "voided"):
        rows = journal(world, kind="current", size=100, status=status)["rows"]
        assert {r["status"] for r in rows} <= {"active"} and not ({9, 10} & {r["receipt_no"] for r in rows}), status
    voided = journal(world, kind="current", status="voided")
    assert (voided["rows"], voided["total"], voided["pages"], voided["totals"]["kind_sum"], voided["totals"]["count"]) == ([], 0, 0, 0, 0)


def test_journal_kind_with_method_and_a_day(world, seeded):
    day = journal(world, "2026-10-21", "2026-10-21", kind="current", method="cash", size=100)
    assert sorted(r["receipt_no"] for r in day["rows"]) == [12, 13] and day["totals"]["kind_sum"] == 250_000 == day["totals"]["sum"]
    strip = {x["kind"]: x["amount"] for x in income(world, "2026-10-21", "2026-10-21", method="cash")["strip"]}
    assert strip["current"] == 250_000
    assert journal(world, "2026-10-21", "2026-10-21", kind="current", method="card", size=100)["rows"] == []
    assert journal(world, "2026-10-21", "2026-10-21", kind="old_debt", method="cash", size=100)["totals"]["kind_sum"] == strip["old_debt"] == 0
    both = journal(world, kind="current", method="cash", group_id=str(seeded["groups"]["A"].id), size=100)
    assert both["totals"]["kind_sum"] == {x["kind"]: x["amount"] for x in income(
        world, "2026-10-01", "2026-10-31", method="cash", group_id=str(seeded["groups"]["A"].id))["strip"]}["current"]
    assert all(r["method"] == "cash" and r["group_name"] == "Математика 9А" for r in both["rows"])


def test_journal_kind_pagination_walks_every_matching_row_once(world, seeded):
    for kind in ("current", "advance", "old_debt"):
        whole = journal(world, kind=kind, size=100)
        n = whole["total"]
        assert n >= 2, kind
        for size in (1, 2, 3):
            seen, kind_sum, order = [], 0, []
            pages = (n + size - 1) // size
            for page in range(1, pages + 1):
                j = journal(world, kind=kind, size=size, page=page)
                assert (j["total"], j["pages"], j["page"], j["size"]) == (n, pages, page, size)
                assert j["totals"] == whole["totals"]                                                 # the totals are the set's, not the page's
                assert 1 <= len(j["rows"]) <= size
                seen += [r["id"] for r in j["rows"]]
                kind_sum += sum(kind_part(r, kind) for r in j["rows"])
            assert seen == [r["id"] for r in whole["rows"]] and len(set(seen)) == n, (kind, size)         # journal order, each row once
            assert kind_sum == whole["totals"]["kind_sum"]
            assert journal(world, kind=kind, size=size, page=pages + 1)["rows"] == []


def test_journal_export_honours_the_kind(world, seeded):
    for kind in ("advance", "current", "refunded"):
        j = journal(world, kind=kind, size=100)
        want = sorted(r["receipt_no"] for r in j["rows"])
        top, header, body = journal_csv(world, kind=kind)
        assert sorted(int(row[0]) for row in body) == want, kind
        label = {"advance": "пока не потрачено (аванс)", "current": "за текущий период", "refunded": "возвращено"}[kind]
        assert any(line == [f"Оплат {len(want)} · Оплачено {sum(r['amount'] for r in j['rows']):,} · Сумма «{label}»: {j['totals']['kind_sum']:,}".replace(",", " ")]
                   for line in top), kind
        assert any(f"Только оплаты, где есть: {label}" in line[0] for line in top), kind
        _, ws = journal_xlsx(world, kind=kind)
        assert sorted(row["№ квитанции"] for row in table_of(ws)[2]) == want, kind
        assert f"Сумма «{label}»: {j['totals']['kind_sum']:,}".replace(",", " ") in ws["A3"].value and "Отменённых" not in ws["A3"].value, kind
    # the cap counts the rows of the kind
    from src.auth import payments_income

    cap = journal(world, kind="refunded")["total"]
    old = payments_income.MAX_EXPORT_ROWS
    try:
        payments_income.MAX_EXPORT_ROWS = cap
        assert owner(world).get(f"{API}/org/payments/journal/export", params={"from": "2026-10-01", "to": "2026-10-31", "kind": "refunded"}).status_code == 200
        payments_income.MAX_EXPORT_ROWS = cap - 1
        assert owner(world).get(f"{API}/org/payments/journal/export", params={"from": "2026-10-01", "to": "2026-10-31", "kind": "refunded"}).status_code == 422
    finally:
        payments_income.MAX_EXPORT_ROWS = old


def test_journal_rejects_an_unknown_kind(world, seeded):
    r = owner(world).get(f"{API}/org/payments/journal", params={"from": "2026-10-01", "to": "2026-10-31", "kind": "bogus"})
    assert r.status_code == 422


def test_journal_rows_carry_the_correction_and_void_details(world, seeded):
    by = {r["receipt_no"]: r for r in journal(world, size=100)["rows"]}
    new, old, voided, plain = by[11], by[9], by[10], by[13]
    assert (new["replaces_receipt_no"], new["replaces_amount"], new["replaces_reason"], new["corrected_by_name"], new["corrected_at"], new["replaces_method"],
            new["replaces_method_label"]) == (9, 400_000, "Исправлено", "Азиз Хакимов", "2026-10-15T06:30:00", None, None)           # same method card -> null
    assert (old["status"], old["replaced_by_receipt_no"], old["void_reason"], old["voided_by_name"], old["voided_at"]) == (
        "replaced", 11, "Исправлено", "Азиз Хакимов", "2026-10-15T06:30:00")
    assert (voided["status"], voided["void_reason"], voided["voided_by_name"], voided["voided_at"], voided["replaced_by_receipt_no"]) == (
        "voided", "Ошибочный ввод", "Нодира Саидова", "2026-10-20T09:00:00", None)
    for key in ("replaces_receipt_no", "replaces_amount", "replaces_reason", "replaces_method", "replaces_method_label", "corrected_by_name", "corrected_at",
                "void_reason", "voided_by_name", "voided_at", "replaced_by_receipt_no"):
        assert plain[key] is None, key
    # a correction that changed the method says so; the typed reason is kept
    m = seeded["memberships"]["a1"]
    base = dict(org_id=world.org.id, membership_id=m.id, student_id=m.student_id, group_id=m.group_id, paid_on=D(2026, 10, 14))
    wrong = StudentPayment(id=uid(world.org, "s3:wrong"), amount=100_000, method="cash", receipt_no=950, created_by=world.owner.id, created_at=datetime(2026, 10, 14, 5, 0),
                           voided_at=datetime(2026, 10, 14, 7, 0), voided_by=world.admin.id, void_reason="Не тот способ: был перевод", **base)
    world.db.add(wrong)
    world.db.flush()
    world.db.add(StudentPayment(id=uid(world.org, "s3:right"), amount=100_000, method="bank_transfer", receipt_no=951, created_by=world.admin.id,
                                created_at=datetime(2026, 10, 14, 7, 0), replaces_payment_id=wrong.id, **base))
    world.db.commit()
    got = {r["receipt_no"]: r for r in journal(world, size=100)["rows"]}
    right = got[951]
    assert (right["replaces_receipt_no"], right["replaces_reason"], right["replaces_method"], right["replaces_method_label"], right["corrected_by_name"],
            right["corrected_at"], right["method"]) == (950, "Не тот способ: был перевод", "cash", "Наличные", "Нодира Саидова", "2026-10-14T07:00:00", "bank_transfer")
    assert (got[950]["status"], got[950]["replaced_by_receipt_no"]) == ("replaced", 951)


def test_status_active_is_only_active_rows_and_the_counts_ignore_the_status(world, seeded):
    ga = seeded["groups"]["A"]
    for first, last in PERIODS:
        for filters in ({}, {"method": "cash"}, {"group_id": str(ga.id)}):
            act = journal(world, first, last, size=100, status="active", **filters)
            allj = journal(world, first, last, size=100, **filters)                          # the server default stays «all»
            gone = journal(world, first, last, size=100, status="voided", **filters)
            assert {r["status"] for r in act["rows"]} <= {"active"} and act["total"] == len(act["rows"]) == act["totals"]["count"], (first, last, filters)
            assert act["totals"]["sum"] == sum(r["amount"] for r in act["rows"]) == income(world, first, last, **filters)["current"]["received"]
            assert {r["status"] for r in gone["rows"]} <= {"voided", "replaced"} and gone["total"] == len(gone["rows"]) == gone["totals"]["voided_count"]
            assert allj["total"] == act["total"] + gone["total"]
            for key in ("count", "sum", "voided_count", "replaced_count"):
                assert act["totals"][key] == allj["totals"][key] == gone["totals"][key], (key, first, last, filters)
            assert gone["totals"]["replaced_count"] == sum(1 for r in gone["rows"] if r["status"] == "replaced") <= gone["totals"]["voided_count"]
    full = journal(world, size=100)["totals"]
    assert (full["voided_count"], full["replaced_count"]) == (2, 1)                          # «Показать отменённые и исправленные (2)»: one cancelled, one corrected
    assert journal(world, "2026-09-01", "2026-09-30")["totals"]["replaced_count"] == 0


def test_the_correction_details_cost_no_extra_queries(world, seeded):
    """The names of the correctors / voiders ride the one names query: the budget of the journal does not grow with the number of corrections."""
    ms = list(seeded["memberships"].values())
    counter = {"n": 0}

    def more(count):
        for _ in range(count):
            counter["n"] += 1
            n = counter["n"]
            m = ms[n % len(ms)]
            base = dict(org_id=world.org.id, membership_id=m.id, student_id=m.student_id, group_id=m.group_id, paid_on=D(2026, 10, 1 + n % 20))
            who = (world.owner, world.admin, world.teacher, world.teacher2)[n % 4]
            old = StudentPayment(id=uid(world.org, f"q:old{n}"), amount=10_000 + n, method=("cash", "card")[n % 2], receipt_no=2000 + 2 * n, created_by=who.id,
                                 created_at=datetime(2026, 10, 1 + n % 20, 5, 0), voided_at=datetime(2026, 10, 1 + n % 20, 6, 0), voided_by=who.id,
                                 void_reason=f"причина {n}", **base)
            world.db.add(old)
            world.db.flush()
            world.db.add(StudentPayment(id=uid(world.org, f"q:new{n}"), amount=10_000 + n, method="click", receipt_no=2001 + 2 * n, created_by=who.id,
                                        created_at=datetime(2026, 10, 1 + n % 20, 6, 0), replaces_payment_id=old.id, **base))
        world.db.commit()

    c = owner(world)
    probes = {
        "journal": lambda: c.get(f"{API}/org/payments/journal", params={"from": "2026-10-01", "to": "2026-10-31", "size": 100}).json(),
        "journal_active": lambda: c.get(f"{API}/org/payments/journal", params={"from": "2026-10-01", "to": "2026-10-31", "size": 100, "status": "active"}).json(),
        "journal_voided": lambda: c.get(f"{API}/org/payments/journal", params={"from": "2026-10-01", "to": "2026-10-31", "size": 100, "status": "voided"}).json(),
        "journal_kind": lambda: c.get(f"{API}/org/payments/journal", params={"from": "2026-10-01", "to": "2026-10-31", "size": 100, "kind": "current"}).json(),
        "journal_kind_p2": lambda: c.get(f"{API}/org/payments/journal", params={"from": "2026-10-01", "to": "2026-10-31", "size": 5, "page": 2, "kind": "current"}).json(),
        "export_kind": lambda: c.get(f"{API}/org/payments/journal/export", params={"from": "2026-10-01", "to": "2026-10-31", "kind": "current"}).content,
        "export_all": lambda: c.get(f"{API}/org/payments/journal/export", params={"from": "2026-10-01", "to": "2026-10-31"}).content,
    }
    more(3)
    small = {name: count_statements(world, fn)[1] for name, fn in probes.items()}
    more(45)
    big = {name: count_statements(world, fn)[1] for name, fn in probes.items()}
    print("QUERIES S3 small/big:", small, big)
    assert all(n <= 24 for n in big.values()), big
    for name in big:
        assert big[name] <= small[name] + 2, (name, small, big)                              # 15x the corrections: only the optional lookups (the correction chain) may appear
    full = probes["journal"]()
    assert full["totals"]["replaced_count"] >= 48 and sum(1 for r in full["rows"] if r["replaces_reason"]) >= 45
    assert {r["corrected_by_name"] for r in full["rows"] if r["replaces_reason"]} >= {"Азиз Хакимов", "Нодира Саидова", "Мария Ким"}
