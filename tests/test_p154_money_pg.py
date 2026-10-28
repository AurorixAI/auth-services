"""#154: the MONEY of the purposes (PostgreSQL) - the deduction of the recalculation per student, the package, the forecast, the payroll (spec §7: Ц1, Ц2, Ц7, Ц8, Ц9-Ц12, Ц17, Ц18; invariants 2, 3, 4, 7).

Rule 45: every sum is a NUMBER here.  Own month, price 600 000, the period 06.10-05.11 has 13 Mon / Wed / Fri lessons: a lesson the centre cancelled deducts 46 000; the teacher gets
600 000 x 50 % / 12 = 25 000 per student per lesson.  What the answers say is in ``test_p154_create_pg.py``."""
from datetime import date, datetime
from decimal import Decimal

import pytest

from src.models.auth import AttendanceExcuse, GroupAttendance, GroupAudit, GroupScheduleOneOffLesson, PayrollRate, StudentGroup
from tests.ops_helpers import clock, history_row, needs_pg, new_student, pg_session  # noqa: F401  (fixtures)
from tests.payf_helpers import API, D, build_org, cancel, centre, client, pay, pause, rows, statement
from tests.test_payf_recalc_pg import month_world

pytestmark = needs_pg
LESSONS = dict(period_mode="lessons", lessons_per_payment=12)
FIRST = 25_000                      # 600 000 x 50 % / 12


# ============================================================================================================ helpers
def make(w, who, group, day, *, start="12:00", end="13:00", expect=201, **extra):
    r = client(w, who).post(f"{API}/org/schedule/occurrences", json={"groupId": str(group.id), "lessonDate": day, "startTime": start, "endTime": end, **extra})
    assert r.status_code == expect, r.text
    return r


def makeup(w, who, group, day, origin, origin_day, **kw):
    return make(w, who, group, day, purpose="makeup", madeUpLessonId=origin, madeUpDate=origin_day, **kw)


def hold(w, group, lesson_id, day, marks, *, roster=None, start="12:00", end="13:00", confirmed=()):
    """The lesson was held: a completed row with a frozen roster; ``marks`` = {student: mark}; ``confirmed`` = students whose excuse is confirmed."""
    students = list(roster if roster is not None else marks)
    row = history_row(w.db, w.org, group, lesson_id, day, w.A, students, start=start, end=end, records={str(s.id): m for s, m in marks.items()})
    for s in confirmed:
        w.db.add(AttendanceExcuse(org_id=w.org.id, attendance_id=row.id, student_id=s.id, reason="Болел", reason_by=w.admin.id, reason_at=datetime(2026, 10, 1), status="confirmed"))
    w.db.commit()
    return row


def classmates(w, g, n, *, joined=datetime(2026, 10, 6, 6, 0)):
    """``n`` more students of the group, each with a membership; every one pays the period 06.10-05.11."""
    out = []
    for i in range(n):
        s = new_student(w.db, w.org, w.ids, f"x{i}", f"Ученик{i + 1}", "Тест", g, joined=joined)
        out.append(s)
    w.db.commit()
    ms = [w.db.query(StudentGroup).filter_by(group_id=g.id, student_id=s.id).one() for s in out]
    for m in ms:
        pay(w, m, 600_000, D(2026, 10, 6))
    return out, ms


def deducted(w, m):
    w.db.expire_all()
    return statement(w, m)["summary"]["recalculated"]


def rate(w, percent="50"):
    w.db.add(PayrollRate(org_id=w.org.id, teacher_id=w.A.id, lesson_type="group", percent=Decimal(percent), valid_from=date(2020, 1, 1), created_by=w.owner.id))
    w.db.commit()


def sheet(w, month, who="owner"):
    w.db.expire_all()
    r = client(w, who).post(f"{API}/org/payroll/months/{month}/calculate", json={"teacher_ids": [str(w.A.id)]})
    assert r.status_code == 200, r.text
    (s,) = r.json()["sheets"]
    return s


def line_of(s, student):
    return next(x for x in s["lines"] if x["student_id"] == str(student.id))


def mark_of(line, lesson_id):
    return next(m for m in line["marks"] if m["lesson_id"] == lesson_id)


def one_off(w, lesson_id):
    w.db.expire_all()
    return w.db.query(GroupScheduleOneOffLesson).filter_by(id=lesson_id).one()


def decide(w, who, lesson_id, decision="pay", expect=200):
    r = client(w, who).post(f"{API}/org/one-off-lessons/{lesson_id}/money-decision", json={"decision": decision})
    assert r.status_code == expect, r.text
    return r


def five(w, clock):
    """Own month, the setting «вычитать» is on, five students who paid the period 06.10-05.11; the lesson of Monday 26.10 is cancelled by the centre: 46 000 deducted from each.
    ``a b c`` come to the make-up, ``d`` is absent, ``e`` has a confirmed excuse."""
    g, s0, m0 = month_world(w, clock)
    pay(w, m0, 600_000, D(2026, 10, 6))
    more, ms = classmates(w, g, 4)
    students, memberships = [s0, *more], [m0, *ms]
    rate(w)
    cancel(w, g, "lesson-mon", "2026-10-26")
    return g, students, memberships


# ============================================================================================================ Ц1: the deduction per student
def test_c1_a_make_up_lifts_the_deduction_of_those_it_counted_for_and_pays_the_teacher_as_an_ordinary_lesson(pg_session, clock):
    w = build_org(pg_session)
    g, (a, b, c, d, e), ms = five(w, clock)
    assert [deducted(w, m) for m in ms] == [46_000] * 5                                # the lesson was cancelled by the centre: everybody is deducted
    mk = makeup(w, "admin", g, "2026-10-29", "lesson-mon", "2026-10-26").json()["id"]
    assert [deducted(w, m) for m in ms] == [0] * 5                                     # the make-up is ahead: the deduction is paused for everybody of the group
    clock.set("2026-10-29 13:30")
    hold(w, g, mk, "2026-10-29", {a: "present", b: "late", c: "present", d: "absent", e: "excused"}, confirmed=[e])
    # an absence without a reason is a lesson held for the student (the owner's decision); a confirmed excuse keeps the deduction, as on an ordinary lesson
    assert [deducted(w, m) for m in ms] == [0, 0, 0, 0, 46_000]
    # the teacher: an ordinary lesson - the roster who counts (4) x 25 000; the confirmed excuse is not counted
    s = sheet(w, "2026-10")
    assert [mark_of(line_of(s, x), mk)["state"] for x in (a, b, c, d, e)] == ["counted"] * 4 + ["excused_confirmed"]
    assert [line_of(s, x)["amount"] for x in (a, b, c, d, e)] == [FIRST] * 4 + [0] and s["total_amount"] == 100_000
    assert s["lines"][0]["marks"][0]["cover"] == "money"


def test_c2_a_cancelled_make_up_gives_the_deduction_back_and_frees_the_lesson(pg_session, clock):
    w = build_org(pg_session)
    g, students, ms = five(w, clock)
    mk = makeup(w, "admin", g, "2026-10-29", "lesson-mon", "2026-10-26").json()["id"]
    assert [deducted(w, m) for m in ms] == [0] * 5
    r = client(w, "admin").post(f"{API}/org/groups/{g.id}/attendance", json={"lessonDate": "2026-10-29", "lessonId": mk, "status": "cancelled", "records": {}, "cancellationReason": "Не нужно"})
    assert r.status_code == 200, r.text
    assert [deducted(w, m) for m in ms] == [46_000] * 5                                # the deduction came back by itself
    again = makeup(w, "admin", g, "2026-10-30", "lesson-mon", "2026-10-26")            # and a new make-up can be made
    assert again.status_code == 201 and [deducted(w, m) for m in ms] == [0] * 5


def test_c7_a_student_who_left_or_was_on_a_pause_keeps_the_deduction(pg_session, clock):
    w = build_org(pg_session)
    g, (a, b, c, d, e), ms = five(w, clock)
    pause(w, ms[3], D(2026, 10, 28), D(2026, 10, 30))                                  # d is on a pause on the day of the make-up
    clock.set("2026-10-27 09:00")
    w.svc.remove_student_from_group(str(w.org.id), str(g.id), str(e.id))               # e leaves the group between the cancellation and the make-up
    w.db.expire_all()
    clock.set("2026-10-28 09:00")
    before = [deducted(w, m) for m in ms]
    assert before[:3] == [46_000] * 3 and before[3] > 0 and before[4] > 0              # each one's own deduction (a pause / a stay that ended changes the price of his lesson)
    mk = makeup(w, "admin", g, "2026-10-29", "lesson-mon", "2026-10-26").json()["id"]
    assert [deducted(w, m) for m in ms] == [0] * 5                                     # ahead: paused for everybody
    clock.set("2026-10-29 13:30")
    hold(w, g, mk, "2026-10-29", {a: "present", b: "present", c: "present"})           # nobody of d / e is in the roster
    assert [deducted(w, m) for m in ms] == [0, 0, 0, before[3], before[4]]            # d and e keep exactly what they had
    # the payroll pays the three who were there
    s = sheet(w, "2026-10")
    assert [line_of(s, x)["amount"] for x in (a, b, c)] == [FIRST] * 3 and s["total_amount"] == 75_000


def test_c8_an_empty_make_up_waits_for_the_decision_of_the_teachers_money_and_the_students_lesson_is_held(pg_session, clock):
    w = build_org(pg_session)
    g, (a, b, c, d, e), ms = five(w, clock)
    mk = makeup(w, "A", g, "2026-10-29", "lesson-mon", "2026-10-26").json()["id"]
    clock.set("2026-10-29 13:30")
    hold(w, g, mk, "2026-10-29", {a: "absent", b: "absent", c: "absent", d: "absent", e: "absent"})
    assert [deducted(w, m) for m in ms] == [0] * 5                                     # absent = the lesson was held for them
    s = sheet(w, "2026-10")
    assert s["total_amount"] == 0 and [mark_of(line_of(s, x), mk)["state"] for x in (a, b, c, d, e)] == ["undecided"] * 5
    ov = client(w, "owner").get(f"{API}/org/payroll/months/2026-10").json()
    assert next(t for t in ov["teachers"] if t["teacher_id"] == str(w.A.id))["undecided_lessons"] == 1
    # it is in the queue of the decisions with the right title and the counter
    queue = client(w, "admin").get(f"{API}/org/one-off-decisions").json()
    assert [(i["lessonId"], i["purpose"], i["title"], i["studentsPresent"], i["studentsTotal"]) for i in queue["items"]] == [(mk, "makeup", "Отработка без учеников — оплатить преподавателю?", 0, 5)]
    assert client(w, "admin").get(f"{API}/org/substitutions/summary").json()["oneOffDecisionsTotal"] == 1
    decide(w, "owner", mk, "pay")
    s = sheet(w, "2026-10")
    assert s["total_amount"] == 5 * FIRST and one_off(w, mk).money == "paid"           # decided: as an ordinary lesson (absent counts)
    assert [deducted(w, m) for m in ms] == [0] * 5                                     # the decision does not touch the students


def test_c8_an_empty_make_up_decided_not_to_pay_is_without_pay(pg_session, clock):
    w = build_org(pg_session)
    g, (a, b, c, d, e), ms = five(w, clock)
    mk = makeup(w, "A", g, "2026-10-29", "lesson-mon", "2026-10-26").json()["id"]
    clock.set("2026-10-29 13:30")
    hold(w, g, mk, "2026-10-29", {a: "absent", b: "absent", c: "absent", d: "absent", e: "absent"})
    decide(w, "admin", mk, "skip")
    s = sheet(w, "2026-10")
    assert s["total_amount"] == 0 and {mark_of(line_of(s, x), mk)["state"] for x in (a, b, c, d, e)} == {"no_pay"}
    decide(w, "admin", mk, "pay", expect=409)                                          # once


def test_a_make_up_with_students_needs_no_decision(pg_session, clock):
    w = build_org(pg_session)
    g, (a, b, c, d, e), ms = five(w, clock)
    mk = makeup(w, "A", g, "2026-10-29", "lesson-mon", "2026-10-26").json()["id"]
    clock.set("2026-10-29 13:30")
    hold(w, g, mk, "2026-10-29", {a: "present", b: "absent", c: "absent", d: "absent", e: "absent"})
    assert client(w, "admin").get(f"{API}/org/one-off-decisions").json() == {"total": 0, "items": []}
    assert decide(w, "admin", mk, "pay", expect=422).json()["detail"]["code"] == "DECISION_NOT_ALLOWED"


# ============================================================================================================ АД-4: a make-up is covered by the day of the ORIGINAL
def test_ad4_a_make_up_in_the_next_period_is_paid_like_the_lesson_it_makes_up_for(pg_session, clock):
    w = build_org(pg_session)
    clock.set("2026-10-06 09:00")
    (g,), (s,), (m,) = centre(w)
    pay(w, m, 600_000, D(2026, 10, 6))                                                 # paid through 05.11 only; the next period is not paid
    rate(w)
    clock.set("2026-10-30 12:00")
    cancel(w, g, "lesson-mon", "2026-11-02")                                           # a lesson of the PAID period, cancelled by the centre
    mk = makeup(w, "admin", g, "2026-11-09", "lesson-mon", "2026-11-02", start="12:00", end="13:00").json()["id"]
    clock.set("2026-11-09 13:30")
    hold(w, g, mk, "2026-11-09", {s: "present"})
    sh = sheet(w, "2026-11")
    mark = mark_of(line_of(sh, s), mk)
    # by its own date (09.11) the lesson would be «не оплачено» - nothing is paid for the second period; by the date of the original it is covered at the price of the first period
    assert (mark["state"], mark["cover"], mark["price"]) == ("counted", "money", 600_000)
    assert line_of(sh, s)["amount"] == FIRST and sh["total_amount"] == FIRST


# ============================================================================================================ the package and the forecast
def package_world(w, clock, day="2026-10-05"):
    clock.set(f"{day} 12:00")
    (g,), (s,), (m,) = centre(w, joined=datetime(2026, 9, 1, 6, 0), **LESSONS)
    pay(w, m, 600_000, D(2026, 10, 5))
    rate(w)
    return g, s, m


def row_of(w, s):
    (row,) = rows(w, s)
    return row


def test_c1_a_make_up_spends_the_package_like_an_ordinary_lesson_and_the_cancelled_one_spent_nothing(pg_session, clock):
    w = build_org(pg_session)
    g, s, m = package_world(w, clock)
    assert (row_of(w, s)["lessons_used"], row_of(w, s)["paid_through_forecast"]) == (0, "2026-10-30")
    cancel(w, g, "lesson-wed", "2026-10-07")                                           # nothing was spent by the cancelled lesson
    assert row_of(w, s)["lessons_used"] == 0
    mk = makeup(w, "admin", g, "2026-10-08", "lesson-wed", "2026-10-07", start="12:00", end="13:00").json()["id"]
    clock.set("2026-10-08 13:30")
    hold(w, g, mk, "2026-10-08", {s: "present"})
    assert (row_of(w, s)["lessons_used"], row_of(w, s)["lessons_left"]) == (1, 11)     # the package of the student: -1 lesson


def test_c9_an_additional_lesson_never_changes_the_package_the_forecast_or_the_charge_of_a_student(pg_session, clock):
    w = build_org(pg_session)
    g, s, m = package_world(w, clock)
    before = row_of(w, s)
    for money, day in (("paid", "2026-10-06"), ("none", "2026-10-08"), ("wait", "2026-10-13")):
        who = "A" if money == "wait" else "admin"
        made = make(w, who, g, day, start="15:00", end="16:00", purpose="extra", money=money).json()["id"]
        hold(w, g, made, day, {s: "present"}, start="15:00", end="16:00")           # held (the clock stays at 05.10: the forecast is read at the same day)
        now = row_of(w, s)
        assert (now["lessons_used"], now["lessons_left"], now["paid_through_forecast"]) == (before["lessons_used"], before["lessons_left"], before["paid_through_forecast"]), money
    # an absence on it costs nothing either
    made = make(w, "admin", g, "2026-10-14", start="15:00", end="16:00", purpose="extra", money="none").json()["id"]
    hold(w, g, made, "2026-10-14", {s: "absent"}, start="15:00", end="16:00")
    assert row_of(w, s)["lessons_used"] == 0 and row_of(w, s)["paid_through_forecast"] == before["paid_through_forecast"]


def test_inv3_an_additional_lesson_does_not_change_the_charge_in_a_calendar_mode(pg_session, clock):
    w = build_org(pg_session)
    g, s, m = month_world(w, clock)
    pay(w, m, 600_000, D(2026, 10, 6))
    before = statement(w, m)["summary"]
    made = make(w, "admin", g, "2026-10-27", start="15:00", end="16:00", purpose="extra", money="paid").json()["id"]
    clock.set("2026-10-27 17:00")
    hold(w, g, made, "2026-10-27", {s: "present"}, start="15:00", end="16:00")
    assert statement(w, m)["summary"] == before


# ============================================================================================================ Ц9-Ц12: the payroll of an additional lesson
def extra_world(w, clock, money, who):
    """October, five students paid; an additional lesson of Thursday 29.10 15:00-16:00, three came (a, b present, c late), d absent, e excused (pending)."""
    g, (a, b, c, d, e), ms = five(w, clock)
    made = make(w, who, g, "2026-10-29", start="15:00", end="16:00", purpose="extra", money=money).json()["id"]
    clock.set("2026-10-29 17:00")
    hold(w, g, made, "2026-10-29", {a: "present", b: "present", c: "late", d: "absent", e: "excused"}, start="15:00", end="16:00")
    return g, (a, b, c, d, e), ms, made


def states(s, lesson, students):
    return [mark_of(line_of(s, x), lesson)["state"] for x in students]


def test_c9_a_teachers_request_is_not_in_the_sum_until_the_owner_decides_and_then_it_pays_like_a_lesson(pg_session, clock):
    w = build_org(pg_session)
    g, st, ms, made = extra_world(w, clock, "wait", "A")
    s = sheet(w, "2026-10")
    assert states(s, made, st) == ["undecided"] * 5                                                  # everybody who was counted (an excuse that is not confirmed counts)
    assert s["total_amount"] == 0 and s["lines"][0]["undecided"] == 1 and s["undecided_total"] == 5
    assert client(w, "owner").get(f"{API}/org/payroll/months/2026-10").json()["teachers"][0]["undecided_lessons"] == 1
    # the queue shows it with who held it, how many came; the badge counts it with the requests
    queue = client(w, "admin").get(f"{API}/org/one-off-decisions").json()
    (item,) = queue["items"]
    assert (item["lessonId"], item["purpose"], item["title"], item["groupName"], item["lessonDate"], item["studentsPresent"], item["studentsTotal"]) == (
        made, "extra", "Дополнительный урок — оплатить преподавателю?", g.name, "2026-10-29", 3, 5)
    assert item["teacher"]["id"] == str(w.A.id) and item["createdBy"]["role"] == "teacher" and item["payrollSheet"] is None
    summary = client(w, "admin").get(f"{API}/org/substitutions/summary").json()
    assert (summary["oneOffDecisionsTotal"], summary["waitingTotal"], summary["pendingTotal"]) == (1, 1, 0)
    out = decide(w, "admin", made, "pay").json()
    assert (out["lessonId"], out["money"], out["payrollSheet"]) == (made, "paid", None)
    s = sheet(w, "2026-10")
    assert states(s, made, st) == ["counted"] * 5                                                    # the absent student counts like on a lesson
    assert [line_of(s, x)["amount"] for x in st] == [FIRST] * 5                                      # 25 000 each: the price of the day 600 000, the divisor of the month 12, 50 %
    assert s["total_amount"] == 125_000 and s["undecided_total"] == 0
    # the journal of the group: the creation and the decision, with the authors
    w.db.expire_all()
    actions = [(a.action, a.actor_role, a.snapshot.get("to")) for a in w.db.query(GroupAudit).filter(GroupAudit.group_id == g.id, GroupAudit.action.in_(("one_off_created", "one_off_money_decided")))
               .order_by(GroupAudit.created_at)]
    assert actions == [("one_off_created", "teacher", None), ("one_off_money_decided", "admin", "paid")]
    assert client(w, "admin").get(f"{API}/org/one-off-decisions").json()["total"] == 0 and client(w, "admin").get(f"{API}/org/substitutions/summary").json()["waitingTotal"] == 0


def test_c10_not_to_pay_makes_the_lesson_without_pay_and_a_repeat_is_refused(pg_session, clock):
    w = build_org(pg_session)
    g, st, ms, made = extra_world(w, clock, "wait", "A")
    decide(w, "owner", made, "skip")
    s = sheet(w, "2026-10")
    assert states(s, made, st) == ["no_pay"] * 5 and s["total_amount"] == 0 and s["lines"][0]["no_pay"] == 1
    assert one_off(w, made).money == "none"
    for decision in ("pay", "skip"):
        r = decide(w, "owner", made, decision, expect=409)
        assert r.json()["detail"]["code"] == "ALREADY_DECIDED"
    assert one_off(w, made).money == "none"


def test_c12_an_owners_additional_lesson_needs_no_decision(pg_session, clock):
    w = build_org(pg_session)
    g, st, ms, made = extra_world(w, clock, "paid", "owner")
    s = sheet(w, "2026-10")
    assert states(s, made, st) == ["counted"] * 5 and [line_of(s, x)["amount"] for x in st] == [FIRST] * 5 and s["total_amount"] == 125_000
    assert client(w, "admin").get(f"{API}/org/one-off-decisions").json()["total"] == 0
    assert decide(w, "owner", made, "skip", expect=422).json()["detail"]["code"] == "DECISION_NOT_ALLOWED"           # created «paid»: nothing to decide
    w2 = build_org(pg_session)
    g2, st2, ms2, made2 = extra_world(w2, clock, "none", "admin")
    s2 = sheet(w2, "2026-10")
    assert states(s2, made2, st2) == ["no_pay"] * 5 and s2["total_amount"] == 0


def test_inv7_who_may_decide_and_when(pg_session, clock):
    w = build_org(pg_session)
    g, (a, b, c, d, e), ms = five(w, clock)
    made = make(w, "A", g, "2026-10-29", start="15:00", end="16:00", purpose="extra", money="wait").json()["id"]
    assert decide(w, "A", made, "pay", expect=403)                                      # the teacher never decides
    assert decide(w, "admin", made, "pay", expect=409).json()["detail"]["code"] == "LESSON_NOT_COMPLETED"       # not held yet
    assert client(w, "A").get(f"{API}/org/one-off-decisions").status_code == 403
    from tests.pg_world import make_client

    foreign = make_client(w, user=w.foreign_admin, org="other", org_role="admin", role="admin")
    assert foreign.post(f"{API}/org/one-off-lessons/{made}/money-decision", json={"decision": "pay"}).status_code == 404
    assert decide(w, "admin", "one-off-nothing", "pay", expect=404).json()["detail"]["code"] == "ONE_OFF_NOT_FOUND"
    clock.set("2026-10-29 17:00")
    hold(w, g, made, "2026-10-29", {a: "present"}, start="15:00", end="16:00")
    out = decide(w, "owner", made, "pay").json()                                       # held and waiting: the owner decides
    assert out["money"] == "paid"


def test_c11_a_decision_after_the_payout_warns_and_the_sheet_owes_the_difference(pg_session, clock):
    w = build_org(pg_session)
    g, st, ms, made = extra_world(w, clock, "wait", "A")
    hold(w, g, "lesson-wed", "2026-10-21", {x: "present" for x in st}, start="10:00", end="11:00")       # an ordinary lesson of October: 5 x 25 000
    clock.set("2026-11-02 10:00")
    s = sheet(w, "2026-10")
    assert s["total_amount"] == 125_000 and s["undecided_total"] == 5                                    # the additional lesson is «не решено», not in the sum
    saved = client(w, "owner").post(f"{API}/org/payroll/sheets", json={"month": "2026-10", "teacher_id": str(w.A.id), "expected_fingerprint": s["source_fingerprint"]})
    assert saved.status_code == 201, saved.text
    sheet_id = saved.json()["id"]
    paid = client(w, "owner").post(f"{API}/org/payroll/sheets/{sheet_id}/payouts", json={"amount": 125_000, "paid_on": "2026-11-02", "method": "cash"})
    assert paid.status_code == 201 and paid.json()["state"]["state"] == "settled", paid.text
    # the owner sees the month is paid out; an administrator only the month (LESSONS 45)
    assert client(w, "owner").get(f"{API}/org/one-off-decisions").json()["items"][0]["payrollSheet"] == {"month": "2026-10", "status": "settled"}
    assert client(w, "admin").get(f"{API}/org/one-off-decisions").json()["items"][0]["payrollSheet"] == {"month": "2026-10", "status": None}
    out = decide(w, "owner", made, "pay").json()
    assert out["payrollSheet"] == {"month": "2026-10", "status": "settled"}                              # said at the moment of the decision
    got = client(w, "owner").get(f"{API}/org/payroll/sheets/{sheet_id}").json()
    st_ = got["payout"]
    assert got["stale"] is True and (st_["saved_total"], st_["fresh_total"], st_["paid"], st_["balance"], st_["state"]) == (125_000, 250_000, 125_000, 125_000, "due")      # «доплатить 125 000»


def test_c18_a_teacher_sees_no_sums_and_no_decisions(pg_session, clock):
    w = build_org(pg_session)
    g, st, ms, made = extra_world(w, clock, "wait", "A")
    teacher = client(w, "A")
    for path in (f"/org/payroll/months/2026-10", "/org/one-off-decisions", "/org/payroll/sheets"):
        assert teacher.get(f"{API}{path}").status_code == 403, path
    assert teacher.post(f"{API}/org/one-off-lessons/{made}/money-decision", json={"decision": "pay"}).status_code == 403
    # what he reads of his lesson: the purpose and the status, no sum
    listed = teacher.get(f"{API}/org/schedule/occurrences", params={"date_from": "2026-10-29", "date_to": "2026-10-29"}).json()
    item = next(o for o in listed if o["id"] == made)
    assert (item["purposeLabel"], item["money"], item["moneyLabel"]) == ("Дополнительный", "wait", "не решено")
    blob = str(listed)
    assert "amount" not in blob and "25000" not in blob and "payrollSheet" not in blob
    decide(w, "admin", made, "pay")
    item = next(o for o in teacher.get(f"{API}/org/schedule/occurrences", params={"date_from": "2026-10-29", "date_to": "2026-10-29"}).json() if o["id"] == made)
    assert (item["money"], item["moneyLabel"]) == ("paid", "оплачено")


# ============================================================================================================ Ц17: a monthly payment and an additional lesson
def test_c17_in_a_monthly_mode_a_make_up_lifts_the_deduction_and_an_additional_lesson_touches_nothing(pg_session, clock):
    w = build_org(pg_session)
    g, (a, b, c, d, e), ms = five(w, clock)
    before = [statement(w, m)["summary"] for m in ms]
    extra = make(w, "admin", g, "2026-10-28", start="15:00", end="16:00", purpose="extra", money="none").json()["id"]
    mk = makeup(w, "admin", g, "2026-10-29", "lesson-mon", "2026-10-26").json()["id"]
    clock.set("2026-10-29 17:00")
    hold(w, g, extra, "2026-10-28", {a: "present", b: "absent"}, start="15:00", end="16:00")
    hold(w, g, mk, "2026-10-29", {a: "present", b: "present", c: "absent", d: "present", e: "present"})
    after = [statement(w, m)["summary"] for m in ms]
    assert [x["recalculated"] for x in before] == [46_000] * 5 and [x["recalculated"] for x in after] == [0] * 5
    assert [x["charged"] if "charged" in x else None for x in before] == [x["charged"] if "charged" in x else None for x in after]       # the charge itself did not move


# ============================================================================================================ the student card: a neutral label, never the money (LESSONS 51)
def test_the_student_card_says_what_the_lesson_was_and_never_the_money(pg_session, clock):
    w = build_org(pg_session)
    g, (a, b, c, d, e), ms = five(w, clock)
    extra = make(w, "admin", g, "2026-10-28", start="15:00", end="16:00", purpose="extra", money="paid").json()["id"]
    mk = makeup(w, "admin", g, "2026-10-29", "lesson-mon", "2026-10-26").json()["id"]
    clock.set("2026-10-29 17:00")
    hold(w, g, extra, "2026-10-28", {a: "present"}, start="15:00", end="16:00")
    hold(w, g, mk, "2026-10-29", {a: "present"})
    for who in ("owner", "admin"):
        card = client(w, who).get(f"{API}/org/students/{a.id}/attendance", params={"month": "2026-10"})
        assert card.status_code == 200, card.text
        by_id = {x["lesson_id"]: x for x in card.json()["lessons"]}
        assert (by_id[extra]["purpose"], by_id[extra]["purposeLabel"], by_id[extra]["money"], by_id[extra]["moneyLabel"]) == ("extra", "Дополнительный", None, None)
        assert (by_id[mk]["purposeLabel"], by_id[mk]["madeUpFor"]) == ("Отработка за пн 26.10", {"lessonId": "lesson-mon", "date": "2026-10-26"})
        assert "paid" not in str({k: v for k, v in by_id[extra].items() if k in ("money", "moneyLabel")})


# ============================================================================================================ АД-5: the way of payment changed between the cancelled lesson and its make-up
def test_ad5_a_make_up_in_a_package_for_a_lesson_of_a_paid_month_spends_no_lesson(pg_session, clock):
    w = build_org(pg_session)
    history = [{"from": "min", "period_mode": "calendar"}, {"from": "2026-12-01", "period_mode": "lessons", "lessons_per_payment": 12}]
    clock.set("2026-10-06 09:00")
    (g,), (s,), (m,) = centre(w, billing_history=history, period_mode="calendar")
    pay(w, m, 600_000, D(2026, 10, 6))
    pay(w, m, 600_000, D(2026, 11, 6))                                                 # the month 06.11 - 05.12 is paid by the calendar
    rate(w)
    clock.set("2026-11-25 12:00")
    cancel(w, g, "lesson-fri", "2026-11-27")                                           # a lesson of the paid calendar month, cancelled by the centre
    mk = makeup(w, "admin", g, "2026-12-02", "lesson-fri", "2026-11-27").json()["id"]
    pay(w, m, 600_000, D(2026, 12, 1))                                                 # the package of 12 lessons from 01.12
    clock.set("2026-12-02 13:30")
    before = row_of(w, s)
    hold(w, g, mk, "2026-12-02", {s: "present"})
    after = row_of(w, s)
    assert before["period_mode"] == after["period_mode"] == "lessons"
    assert after["lessons_used"] == before["lessons_used"] and after["lessons_left"] == before["lessons_left"]          # spent nothing: the month of the cancelled lesson was paid
    # the teacher is paid for it as for the cancelled lesson would have been: by the calendar day of the original (paid, the price 600 000)
    sh = sheet(w, "2026-12")
    mark = mark_of(line_of(sh, s), mk)
    assert (mark["state"], mark["cover"], mark["price"]) == ("counted", "money", 600_000) and line_of(sh, s)["amount"] == FIRST


# ============================================================================================================ races (LESSONS 23 / 27): the order of the locks is the group first
def test_two_administrators_deciding_at_once_give_one_decision_and_one_409(pg_session, clock):
    from src.auth.one_off_purpose import OneOffPurpose
    from tests.ops_helpers import code_of, run_parallel

    w = build_org(pg_session)
    g, st, ms, made = extra_world(w, clock, "wait", "A")

    def call(decision):
        return lambda svc: OneOffPurpose(svc).decide(str(w.org.id), str(w.admin.id), "admin", made, decision, owner=False)

    results = run_parallel(w, call("pay"), call("skip"))
    ok = [r for r, e in results if e is None]
    errors = [e for r, e in results if e is not None]
    assert len(ok) == 1 and [code_of(e) for e in errors] == ["ALREADY_DECIDED"], results
    w.db.expire_all()
    assert one_off(w, made).money == ok[0]["money"]
    assert w.db.query(GroupAudit).filter_by(group_id=g.id, action="one_off_money_decided").count() == 1


def test_a_make_up_and_the_return_of_the_same_lesson_at_once_never_leave_both(pg_session, clock):
    """The two roads take the group first and then the row of the cancelled lesson: one wins, the other is refused (no deadlock, no lesson held twice)."""
    from datetime import date

    from src.auth.ops_schemas import CreateOneOffLessonRequestB
    from tests.ops_helpers import run_parallel

    w = build_org(pg_session)
    g, st, ms = five(w, clock)

    def create(svc):
        req = CreateOneOffLessonRequestB(groupId=str(g.id), lessonDate=date(2026, 10, 29), startTime="12:00", endTime="13:00", purpose="makeup", madeUpLessonId="lesson-mon", madeUpDate=date(2026, 10, 26))
        return svc.create_one_off_lesson(str(w.org.id), req, actor_id=str(w.admin.id), is_org_admin=True)

    def restore(svc):
        from src.auth.lesson_substitutions import Substitutions

        return Substitutions(svc).restore_lesson(w.org.id, w.admin.id, "admin", g.id, "lesson-mon", date(2026, 10, 26), "Отменили по ошибке")

    results = run_parallel(w, create, restore)
    w.db.expire_all()
    row = w.db.query(GroupAttendance).filter_by(group_id=g.id, lesson_id="lesson-mon", lesson_date="2026-10-26").first()
    made = w.db.query(GroupScheduleOneOffLesson).filter_by(group_id=g.id).count()
    created_ok, restored_ok = results[0][1] is None, results[1][1] is None
    assert created_ok != restored_ok, results                                      # exactly one of the two roads wins
    if created_ok:
        assert row is not None and row.made_up_by and made == 1                  # the lesson stays cancelled, the make-up lives
    else:
        assert row is None and made == 0                                         # the lesson came back, no make-up


def test_the_excel_of_a_sheet_says_the_words_of_the_new_states(pg_session, clock):
    import io

    from openpyxl import load_workbook

    w = build_org(pg_session)
    g, st, ms, made = extra_world(w, clock, "wait", "A")
    s = sheet(w, "2026-10")
    saved = client(w, "owner").post(f"{API}/org/payroll/sheets", json={"month": "2026-10", "teacher_id": str(w.A.id), "expected_fingerprint": s["source_fingerprint"]})
    assert saved.status_code == 201, saved.text
    out = client(w, "owner").get(f"{API}/org/payroll/sheets/{saved.json()['id']}/export.xlsx")
    assert out.status_code == 200
    cells = [c.value for ws in load_workbook(io.BytesIO(out.content)).worksheets for row in ws.iter_rows() for c in row]
    assert cells.count("не решено") == 5                                               # the five students of the lesson
    decide(w, "owner", made, "skip")
    s = sheet(w, "2026-10")
    saved = client(w, "owner").post(f"{API}/org/payroll/sheets", json={"month": "2026-10", "teacher_id": str(w.A.id), "expected_fingerprint": s["source_fingerprint"], "note": "после решения"})
    cells = [c.value for ws in load_workbook(io.BytesIO(client(w, "owner").get(f"{API}/org/payroll/sheets/{saved.json()['id']}/export.xlsx").content)).worksheets for row in ws.iter_rows() for c in row]
    assert cells.count("без оплаты") == 5 and "не решено" not in cells
