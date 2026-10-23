"""The house style of the Excel exports (``src/auth/xlsx_style.py``): palette, blocks, widths, print setup, file names.  No database."""
import io
from datetime import date, datetime

from openpyxl import Workbook, load_workbook

from src.auth import xlsx_style as xs


def sheet():
    wb = Workbook()
    return wb, wb.active


def reread(wb):
    out = io.BytesIO()
    wb.save(out)
    return load_workbook(io.BytesIO(out.getvalue())).active


def test_the_palette_is_the_one_of_the_spec():
    assert (xs.HEAD_FILL, xs.SECTION_FILL, xs.PENDING_FILL, xs.WARN_FILL, xs.EXAMPLE_FILL) == ("D9E2F3", "EDF1F7", "FFF2B3", "FDE2E1", "F2F2F2")
    assert xs.TOTAL_LINE == "8EA9DB" and xs.GRID == "D9D9D9" and xs.CANCELLED_COLOR == "808080" and xs.FONT_NAME == "Calibri"
    assert "2F6FED" not in xs.PALETTE                                                # the bright blue of the old files is gone
    assert (xs.MONEY, xs.DATE, xs.DATETIME, xs.PERCENT) == ("#,##0", "DD.MM.YYYY", "DD.MM.YYYY HH:MM", '0"%"')
    assert xs.HEADER_ROW == 5 and xs.TITLE_ROWS == 3


def test_a_text_is_always_text_and_the_types_stay_real():
    wb, ws = sheet()
    xs.put(ws, 1, 1, "=1+1")
    xs.put(ws, 2, 1, "+7 (999)")
    xs.put(ws, 3, 1, "a\x00b\x0bc￾d")
    xs.put(ws, 4, 1, 1_250_000, number_format=xs.MONEY)
    xs.put(ws, 5, 1, date(2026, 10, 21), number_format=xs.DATE)
    xs.put(ws, 6, 1, datetime(2026, 10, 21, 11, 40), number_format=xs.DATETIME)
    xs.put(ws, 7, 1, "x\ny")
    back = reread(wb)
    assert [back.cell(r, 1).data_type for r in (1, 2, 3)] == ["s", "s", "s"] and back["A1"].value == "=1+1"
    assert back["A3"].value == "abcd" and back["A7"].value == "x\ny"
    assert back["A4"].value == 1_250_000 and back["A4"].number_format == "#,##0" and back["A4"].data_type == "n"
    assert back["A5"].is_date and back["A5"].number_format == "DD.MM.YYYY"
    assert back["A6"].value == datetime(2026, 10, 21, 11, 40) and back["A6"].number_format == "DD.MM.YYYY HH:MM"


def test_title_block_is_three_merged_rows_and_the_header_is_row_five():
    wb, ws = sheet()
    assert xs.title_block(ws, ["Журнал оплат — Центр", "По состоянию на 21.10.2026", None], 6) == 5
    ws2 = reread(wb)
    assert {str(r) for r in ws2.merged_cells.ranges} == {"A1:F1", "A2:F2", "A3:F3"}
    assert ws2["A1"].font.bold and ws2["A1"].font.size == 14 and ws2["A1"].font.name == "Calibri"
    assert not ws2["A2"].font.bold and ws2["A2"].font.size == 11 and ws2["A3"].value is None
    assert ws2["A2"].alignment.wrap_text and ws2["A2"].alignment.horizontal == "left"
    assert ws2.cell(4, 1).value is None


def test_header_section_and_total_rows():
    wb, ws = sheet()
    xs.header_row(ws, 5, ["Дата", "Оплачено", "Подробности"], left=(3,))
    xs.section_row(ws, 6, "Октябрь 2026", 3)
    xs.put(ws, 7, 2, 500_000, number_format=xs.MONEY)
    xs.total_row(ws, 8, "Итого", {2: 500_000, 3: "Аванс 60 000"}, 3)
    back = reread(wb)
    for c in (1, 2, 3):
        head = back.cell(5, c)
        assert head.fill.fgColor.rgb.endswith("D9E2F3") and head.font.bold and head.alignment.wrap_text and head.border.left.style == "thin"
        assert head.border.left.color.rgb.endswith("D9D9D9")
    assert back.cell(5, 1).alignment.horizontal == "center" and back.cell(5, 3).alignment.horizontal == "left"
    assert "A6:C6" in {str(r) for r in back.merged_cells.ranges}
    assert back["A6"].fill.fgColor.rgb.endswith("EDF1F7") and back["A6"].font.bold
    for c in (1, 2, 3):
        total = back.cell(8, c)
        assert total.fill.fgColor.rgb.endswith("D9E2F3") and total.font.bold and total.border.top.style == "medium"
        assert total.border.top.color.rgb.endswith("8EA9DB")
    assert back["A8"].value == "Итого" and back["B8"].value == 500_000 and back["B8"].number_format == "#,##0" and back["C8"].value == "Аванс 60 000"
    assert back["C8"].data_type == "s" and back["B8"].alignment.horizontal == "center"


def test_the_example_row_is_grey_italic_with_its_mark():
    wb, ws = sheet()
    xs.example_row(ws, 2, ["Усмонова", "пример — удалите"])
    back = reread(wb)
    assert back["A2"].fill.fgColor.rgb.endswith("F2F2F2") and back["A2"].font.italic and back["A2"].font.color.rgb.endswith("808080")
    assert xs.EXAMPLE_MARK in back["B2"].value


def test_measuring():
    assert xs.shown_length(date(2026, 1, 1)) == 10 and xs.shown_length(datetime(2026, 1, 1, 1, 1)) == 16
    assert xs.shown_length(1_250_000) == len("1,250,000") and xs.shown_length(None) == 0 and xs.shown_length("a\nlonger line\nb") == len("longer line")
    assert xs.lines_needed("short", 20) == 1 and xs.lines_needed(None, 20) == 1
    assert xs.lines_needed("слово " * 10, 22) == 4                                  # 10 words of 5 chars, 20 characters a line
    assert xs.lines_needed("a\nb\nc", 50) == 3 and xs.lines_needed("x" * 100, 12) == 10         # a word longer than the line is cut by characters


def test_autosize_widths_come_from_the_content_and_nothing_is_cut():
    wb, ws = sheet()
    xs.title_block(ws, ["Журнал оплат — " + "очень длинное название центра " * 4, "второй ряд", None], 4)
    xs.header_row(ws, 5, ["Дата", "Оплачено", "Статус действия платежа", "Подробности"])
    row = [date(2026, 10, 21), 12_345_678_901, "Действует (исправлена)", "за октябрь — 150 000 сум · " * 6]
    for r in (6, 7):
        for c, v in enumerate(row, 1):
            xs.put(ws, r, c, v, number_format=xs.DATE if c == 1 else xs.MONEY if c == 2 else None)
    xs.section_row(ws, 8, "Итого по способам, и это очень длинный заголовок раздела, который не должен раздвигать первую колонку", 4)
    widths = xs.autosize(ws, wrap_cols=(4,), min_width=10)
    back = reread(wb)
    assert widths[0] == 12 and widths[1] >= len("12,345,678,901") + 2                    # a date 10 + 2; a big sum with its separators
    assert widths[2] >= len("Действует (исправлена)") + 2                              # the longest value (the header wraps by words)
    assert widths[3] == xs.WRAP_CAP                                                     # a long text is capped ...
    assert back.cell(6, 4).alignment.wrap_text and back.cell(6, 4).alignment.vertical == "top"        # ... wraps, to the top
    assert back.row_dimensions[6].height >= xs.LINE_HEIGHT * 3
    # the band and the title block are merged: they widen nothing, their ROW gets the height
    assert widths[0] == 12 and back.row_dimensions[1].height and back.row_dimensions[1].height >= 21
    assert [back.column_dimensions[c].width for c in "ABCD"] == widths


def test_autosize_wraps_a_long_header_by_words_and_honours_fixed_widths():
    wb, ws = sheet()
    xs.header_row(ws, 5, ["Не оплачено уроков", "Осталось уроков (пакет)", "Дата"])
    xs.put(ws, 6, 1, 3)
    widths = xs.autosize(ws, fixed={3: 8}, min_width=10)
    back = reread(wb)
    assert widths[2] == 8 and widths[1] < len("Осталось уроков (пакет)") + 2            # the long header wraps instead of widening the column
    assert widths[1] >= len("(пакет)") + 2 and back.row_dimensions[5].height >= 2 * xs.LINE_HEIGHT and back.cell(5, 2).alignment.wrap_text


def test_print_setup():
    wb, ws = sheet()
    xs.print_setup(ws, header=5)
    back = reread(wb)
    assert back.page_setup.orientation == "landscape" and back.page_setup.fitToWidth == 1 and back.page_setup.fitToHeight == 0
    assert back.sheet_properties.pageSetUpPr.fitToPage and back.print_title_rows == "$5:$5" and back.page_setup.paperSize == 9
    assert back.page_margins.left <= 0.5 and back.page_margins.right <= 0.5
    wb, ws = sheet()
    xs.print_setup(ws, header=None)
    assert not reread(wb).print_title_rows


def test_file_names_are_russian_and_safe():
    assert xs.file_name("Выписка", "Алиев Анвар", "октябрь-2026") == "Выписка_Алиев_Анвар_октябрь-2026.xlsx"
    assert xs.file_name("Остатки на начало") == "Остатки_на_начало.xlsx"
    assert xs.file_name("Журнал оплат", None, "x") == "Журнал_оплат_x.xlsx"
    hostile = xs.file_name('a/b\\c:d*e?f"g<h>i|j;k\r\nl', ext="csv")
    assert hostile.endswith(".csv") and not set('/\\:*?"<>|;\r\n') & set(hostile)
    assert xs.file_name("", None) == "Файл.xlsx" and len(xs.file_name("я" * 500)) <= 125
    assert xs.month_slug(date(2026, 10, 3)) == "октябрь-2026" and xs.month_title("2026-10") == "Октябрь 2026"


def test_content_disposition_has_an_ascii_fallback_and_the_utf8_name():
    header = xs.content_disposition("Зарплата_октябрь-2026.xlsx")
    assert header == ("attachment; filename=\"Zarplata_oktyabr-2026.xlsx\"; "
                      "filename*=UTF-8''%D0%97%D0%B0%D1%80%D0%BF%D0%BB%D0%B0%D1%82%D0%B0_%D0%BE%D0%BA%D1%82%D1%8F%D0%B1%D1%80%D1%8C-2026.xlsx")
    header.encode("latin-1")                                                           # an HTTP header: nothing outside latin-1
    assert xs.ascii_fallback("Шаблон_импорта.xlsx") == "Shablon_importa.xlsx" and xs.ascii_fallback("Щука Ёж.xlsx") == "Shchuka_Ezh.xlsx"
    assert xs.ascii_fallback("ъь.xlsx") == "file.xlsx" and xs.ascii_fallback("report.xlsx") == "report.xlsx"
    evil = xs.content_disposition('x"; filename="evil.exe\r\nSet-Cookie: a=b')
    assert "\r" not in evil and "\n" not in evil and evil.count('filename="') == 1
    assert xs.attachment_headers("a.xlsx", **{"Cache-Control": "no-store"}) == {
        "Content-Disposition": 'attachment; filename="a.xlsx"; filename*=UTF-8\'\'a.xlsx', "Cache-Control": "no-store"}
