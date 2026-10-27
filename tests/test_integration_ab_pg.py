"""Integration of pack A (money) with pack B / B2 (operations, substitutions) - cross-pack behaviour, one fresh world (the B2 data set D2).

1. payroll pays the lesson teacher that B resolved (a substitute) with the price A fixed for the lesson (``cover = money``);
2. a student on a pause is not in the roster of a lesson - nor in the roster of a substitute (B2 ``lessonRoster``);
3. a group is archived only when empty (B), the students leave through A's settle flow; the archive never touches money;
4. a lesson cancelled by a B2 decision: A decision 3 (the calendar month costs the same, a package is not consumed);
5. a lesson moved by a B2 decision is neutral for the proration (``moved_one_off_keys`` in ``ledger_service._lesson_days``);
6. the Excel import with an opening balance goes through B's seat rule and A's opening row.
"""
from datetime import date, datetime
from io import BytesIO

import pytest
from openpyxl import Workbook

from src.auth import import_service as imp
from src.models.auth import (
    MembershipOpeningBalance, Organization, PayrollRate, Student, StudentGroup, StudentPayment, StudentWriteoff,
)

from tests.b2_helpers import (  # noqa: F401  (fixtures)
    API, build_d2, client, clock, decide, history_row, needs_pg, new_group, new_slot, new_student, pg_session, req, row_of,
)
from tests.pg_world import give_group_price
from tests.test_ops_b2_decide_pg import Sync

pytestmark = needs_pg
LT = API + "/org/groups/{gid}/lesson-teacher"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture
def w(pg_session, clock):
    clock.set("2026-10-21 14:00")
    return build_d2(pg_session)


@pytest.fixture
def sync(monkeypatch):
    return Sync(monkeypatch)


# ------------------------------------------------------------------------------------------------------------------------ helpers
def money(w, group, price, *, block=None):
    """Calendar fixed_day=1 (accounting from 01.10), the group price from 01.01; ``block`` overrides the payments settings."""
    org = w.db.get(Organization, w.org.id)
    settings = dict(org.settings or {})
    settings["payments"] = {**settings.get("payments", {}), "accounting_start_date": "2026-10-01", "billing_anchor": "fixed_day", "billing_day": 1, **(block or {})}
    org.settings = settings
    w.db.commit()
    give_group_price(w.db, group, price, valid_from=date(2026, 1, 1))


def member(w, student, group):
    w.db.expire_all()
    return w.db.query(StudentGroup).filter_by(student_id=student.id, group_id=group.id).one()


def pay(w, m, amount, on):
    w.db.add(StudentPayment(org_id=w.org.id, membership_id=m.id, student_id=m.student_id, group_id=m.group_id, amount=amount, paid_on=on, created_by=w.owner.id,
                            method="cash", created_at=datetime(on.year, on.month, on.day, 9)))
    w.db.commit()


def joined(w, student, group, day):
    """The student has been in the group since ``day`` (membership and period)."""
    m = member(w, student, group)
    m.joined_at = datetime(day.year, day.month, day.day, 6)
    for period in m.periods:
        period.joined_at = m.joined_at
    w.db.commit()
    return m


def put_teacher(w, group, teacher, *, lesson, day, who="admin"):
    return client(w, who).put(LT.format(gid=group.id), json={"lessonDate": day, "lessonId": lesson, "teacherId": str(getattr(w, teacher).id), "reason": "Болезнь"})


def sheet_of(w, teacher, month="2026-10"):
    r = client(w, "owner").post(f"{API}/org/payroll/months/{month}/calculate", json={"teacher_ids": [str(teacher.id)]})
    assert r.status_code == 200, r.text
    (sheet,) = r.json()["sheets"]
    return sheet


def membership_row(w, student, group):
    r = client(w, "owner").get(f"{API}/org/students/{student.id}/memberships")
    assert r.status_code == 200, r.text
    return next(x for x in r.json() if x["group_id"] == str(group.id))


def request_one(w, who, group, lesson, day):
    r = req(w, who, [(group, lesson, day)])
    assert r.status_code == 201, r.text
    return r.json()["created"][0]


# ======================================================================================================================== 1. payroll x substitute
def test_payroll_pays_the_substitute_that_b_resolved_with_the_price_of_a(w, clock):
    """B: the substitute D leads Wed 21.10 and completes it (row.teacher_id = D).  A: the lesson is paid to D with the price of the period
    (``cover = money``, 600 000 / 12 * 30 % = 15 000 per student); the group teacher A has no mark on that day."""
    money(w, w.G1, 600_000)
    for s in (w.s1, w.s2, w.s3):
        pay(w, member(w, s, w.G1), 600_000, date(2026, 10, 1))
    w.db.add(PayrollRate(org_id=w.org.id, teacher_id=w.D.id, lesson_type="group", percent=30, valid_from=date(2026, 9, 1)))
    w.db.commit()
    assert put_teacher(w, w.G1, "D", lesson="lesson-l2", day="2026-10-21").status_code == 200
    d, base = client(w, "D"), f"{API}/org/groups/{w.G1.id}/attendance"
    clock.set("2026-10-21 16:05")
    assert d.post(base + "/start", json={"lessonDate": "2026-10-21", "lessonId": "lesson-l2"}).status_code == 200
    clock.set("2026-10-21 17:20")
    marks = {str(s.id): "present" for s in (w.s1, w.s2, w.s3)}
    done = d.post(base, json={"lessonDate": "2026-10-21", "lessonId": "lesson-l2", "status": "completed", "records": marks, "expectedVersion": 1})
    assert done.status_code == 200, done.text
    assert row_of(w, w.G1, "lesson-l2", "2026-10-21").teacher_id == w.D.id

    sd = sheet_of(w, w.D)
    assert sorted(line["student_id"] for line in sd["lines"]) == sorted(str(s.id) for s in (w.s1, w.s2, w.s3))
    for line in sd["lines"]:
        (mark,) = line["marks"]
        assert (mark["date"], mark["cover"], mark["price"], mark["state"]) == ("2026-10-21", "money", 600_000, "counted")
        assert line["amount"] == 15_000
    assert sd["total_amount"] == 45_000
    sa = sheet_of(w, w.A)
    dates = {m["date"] for line in sa["lines"] for m in line["marks"]}
    assert "2026-10-21" not in dates and dates >= {"2026-10-12", "2026-10-14", "2026-10-19"}


def test_payroll_hooks_of_a_name_the_sheets_of_both_teachers_and_never_lock(w, clock):
    """A decision 6: a sheet that was paid out is not locked - the change of "who taught" is allowed, the saved sheet is listed with its payout state."""
    from src.auth import payroll_hooks

    money(w, w.G1, 600_000)
    for s in (w.s1, w.s2, w.s3):
        pay(w, member(w, s, w.G1), 600_000, date(2026, 10, 1))
    w.db.add(PayrollRate(org_id=w.org.id, teacher_id=w.D.id, lesson_type="group", percent=30, valid_from=date(2020, 1, 1)))
    w.db.commit()
    own = client(w, "owner")
    fp = sheet_of(w, w.A)["source_fingerprint"]
    saved = own.post(f"{API}/org/payroll/sheets", json={"month": "2026-10", "teacher_id": str(w.A.id), "expected_fingerprint": fp})
    assert saved.status_code == 201, saved.text
    total = saved.json()["total_amount"]
    paid = own.post(f"{API}/org/payroll/sheets/{saved.json()['id']}/payouts", json={"amount": total, "paid_on": "2026-10-21", "method": "cash"})
    assert paid.status_code in (200, 201), paid.text
    payroll_hooks.assert_lesson_teacher_change_allowed(w.db, w.org.id, "2026-10-19", w.A.id, w.D.id)         # does not raise
    touched = payroll_hooks.sheets_touched_by_lesson_teacher_change(w.db, w.org.id, "2026-10-19", w.A.id, w.D.id)
    assert [(t["teacher_id"], t["month"], t["status"]) for t in touched] == [(str(w.A.id), "2026-10", "settled")]
    changed = put_teacher(w, w.G1, "D", lesson="lesson-l1", day="2026-10-19")
    assert changed.status_code == 200 and changed.json()["attendanceUpdated"] is True
    assert [s["status"] for s in changed.json()["payrollSheetsAffected"]] == ["settled"]


# ======================================================================================================================== 2. pause x roster
def test_a_paused_student_is_not_in_the_roster_and_not_in_the_roster_of_a_substitute(w, clock):
    """A pause 26.10-31.10 of s3: B's journal of the substitute D (``lessonRoster``) and the roster frozen when D starts the lesson of 26.10 both leave s3 out."""
    money(w, w.G1, 600_000)
    paused = client(w, "admin").post(f"{API}/org/memberships/{member(w, w.s3, w.G1).id}/pauses", json={"first_day": "2026-10-26", "last_day": "2026-10-31"})
    assert paused.status_code == 201, paused.text
    assert put_teacher(w, w.G1, "D", lesson="lesson-l1", day="2026-10-26").status_code == 200
    d, base = client(w, "D"), f"{API}/org/groups/{w.G1.id}/attendance"
    ahead = d.get(base, params={"month": "2026-10"})
    assert ahead.status_code == 200, ahead.text
    assert {r["id"] for r in ahead.json()["lessonRoster"]} == {str(w.s1.id), str(w.s2.id)}          # the paused s3 is not in the substitute's roster
    clock.set("2026-10-26 10:05")
    assert d.post(base + "/start", json={"lessonDate": "2026-10-26", "lessonId": "lesson-l1"}).status_code == 200
    row = row_of(w, w.G1, "lesson-l1", "2026-10-26")
    assert {r["student_id"] for r in row.roster_snapshot} == {str(w.s1.id), str(w.s2.id)}
    after = d.get(base, params={"month": "2026-10"}).json()
    assert {r["id"] for r in after["lessonRoster"]} == {str(w.s1.id), str(w.s2.id)}
    own = client(w, "A").get(base, params={"month": "2026-10"})                                      # the group teacher's journal agrees
    assert own.status_code == 200 and str(w.s3.id) not in own.json()["rosterStudentIds"]["2026-10-26_lesson-l1"]


# ======================================================================================================================== 3. group archive x leave with money
def debt_group(w, name="Долги", students=2):
    """A group with no schedule (so nothing is missed and the charge is by days) whose students owe."""
    group = new_group(w.db, w.org, w.ids, "D" + name, name, w.C, [])
    people = [new_student(w.db, w.org, w.ids, f"{name}{i}", f"Должник{i}", "Имя", group, joined=datetime(2026, 10, 1, 5, 0)) for i in range(students)]
    w.db.commit()
    money(w, group, 500_000)
    return group, people


def test_group_archive_waits_for_the_students_and_a_settle_flow_closes_their_money(w, clock):
    """B: a group with students cannot be archived (``ACTIVE_STUDENTS``).  A: the students leave with a settle (debt written off, a reason), then the archive passes;
    the archive itself changes no money; the write-offs are in the books and the debtors' tab is empty."""
    group, (a, b) = debt_group(w)
    own = client(w, "owner")
    blocked = client(w, "admin").post(f"{API}/org/groups/{group.id}/archive", json={})
    assert blocked.status_code == 409 and blocked.json()["detail"]["code"] == "GROUP_ARCHIVE_BLOCKED"
    assert [x["code"] for x in blocked.json()["detail"]["blockers"]] == ["ACTIVE_STUDENTS"]
    before = own.get(f"{API}/org/payments/summary").json()
    owed = {s.id: member(w, s, group) for s in (a, b)}
    assert before["students_with_debt"] >= 2 and before["debt_total"] > 0
    total = 0
    for s in (a, b):
        pv = own.get(f"{API}/org/groups/{group.id}/students/{s.id}/leave-preview").json()
        mem = pv["memberships"][0]
        assert "writeoff" in mem["allowed"] and mem["money"]["debt"] > 0
        out = own.post(f"{API}/org/groups/{group.id}/students/{s.id}/leave", json={"settle": [
            {"membership_id": str(owed[s.id].id), "action": "writeoff", "reason": "Группа закрывается"}]})
        assert out.status_code == 200, out.text
        assert out.json()["settled"][0]["action"] == "writeoff" and out.json()["settled"][0]["amount"] > 0
        total += out.json()["settled"][0]["amount"]
    mid = own.get(f"{API}/org/payments/summary").json()
    assert total > 0 and mid["debt_total"] == 0 and mid["students_with_debt"] == 0                  # every debt was written off at leaving
    archived = client(w, "admin").post(f"{API}/org/groups/{group.id}/archive", json={"reason": "Группа закрыта"})
    assert archived.status_code == 200, archived.text
    after = own.get(f"{API}/org/payments/summary").json()
    assert after["debt_total"] == mid["debt_total"] and after["students_with_debt"] == mid["students_with_debt"]      # the archive moved no money
    w.db.expire_all()
    assert w.db.query(StudentWriteoff).filter(StudentWriteoff.membership_id.in_([m.id for m in owed.values()]), StudentWriteoff.voided_at.is_(None)).count() == 2
    debtors = {i["student_id"] for i in own.get(f"{API}/org/payments/status", params={"bucket": "debt", "limit": 100}).json()["items"]}
    assert not debtors & {str(a.id), str(b.id)}


def test_a_leave_without_settle_keeps_the_debt_visible_after_the_group_is_archived(w, clock):
    """The old way out (``DELETE``, no settle) leaves the debt in the books: the debtors' tab keeps the student who left, the group archive does not hide it."""
    group, (a, b) = debt_group(w, "Остаток")
    own = client(w, "owner")
    before = own.get(f"{API}/org/payments/summary").json()
    assert before["debt_total"] > 0
    for s in (a, b):
        assert client(w, "admin").delete(f"{API}/org/groups/{group.id}/students/{s.id}").status_code == 200
    left = own.get(f"{API}/org/payments/summary").json()                                           # the charge stops at the leaving, the rest stays owed
    assert 0 < left["debt_total"] <= before["debt_total"] and left["students_with_debt"] == 2
    assert client(w, "admin").post(f"{API}/org/groups/{group.id}/archive", json={}).status_code == 200
    after = own.get(f"{API}/org/payments/summary").json()
    assert after["debt_total"] == left["debt_total"] and after["students_with_debt"] == 2         # the archive hides nothing
    debtors = {i["student_id"] for i in own.get(f"{API}/org/payments/status", params={"bucket": "debt", "limit": 100}).json()["items"]}
    assert {str(a.id), str(b.id)} <= debtors


# ======================================================================================================================== 4. cancelled / moved lesson x A
def late_joiner(w):
    """s1 joined 14.10 (an incomplete first month): 5 of the 8 lessons of the full October by the pattern (Mon 5, 12, 19, 26 / Wed 7, 14, 21, 28; the slots' validFrom
    does not shrink the month) are his: 14.10, 19.10, 21.10, 26.10, 28.10."""
    money(w, w.G1, 600_000)
    for s in (w.s2, w.s3):
        pay(w, member(w, s, w.G1), 600_000, date(2026, 10, 1))
    joined(w, w.s1, w.G1, date(2026, 10, 14))
    return membership_row(w, w.s1, w.G1)["debt"]


def test_a_lesson_cancelled_by_a_b2_decision_leaves_the_calendar_month_price_unchanged(w, clock, sync):
    """A decision 3 meets B2 ``cancel``: the cancelled row stays a lesson of the month for the proration (the price of the month is the same)."""
    debt = late_joiner(w)
    assert debt == 375_000                                                         # 5 of 8 lessons of 600 000
    item = request_one(w, "A", w.G1, "lesson-l1", "2026-10-26")
    clock.set("2026-10-21 14:02")
    out = decide(w, "admin", item["id"], decision="cancel", reason="Преподаватель болен, замены нет")
    assert out.status_code == 200, out.text
    assert row_of(w, w.G1, "lesson-l1", "2026-10-26").status == "cancelled"
    assert membership_row(w, w.s1, w.G1)["debt"] == debt


def test_a_lesson_moved_by_a_b2_decision_is_neutral_for_the_proration(w, clock, sync):
    """B2 ``reschedule`` = the cancelled original (counted, decision 3) + a one-off copy; ``moved_one_off_keys`` keeps the copy out of ``_lesson_days``: the price is the same."""
    debt = late_joiner(w)
    item = request_one(w, "A", w.G1, "lesson-l2", "2026-10-28")
    clock.set("2026-10-21 14:02")
    out = decide(w, "admin", item["id"], decision="reschedule", newDate="2026-10-29", startTime="16:00", endTime="17:30", reason="Перенос")
    assert out.status_code == 200, out.text
    assert out.json()["result"]["newLesson"]["id"]
    assert row_of(w, w.G1, "lesson-l2", "2026-10-28").status == "cancelled"
    assert membership_row(w, w.s1, w.G1)["debt"] == debt                            # 375 000 - not 6/9 of the month


def test_a_lesson_cancelled_by_a_b2_decision_does_not_consume_a_package(w, clock, sync):
    """Packages of 12: the three held lessons of G1 are used, the cancelled one of today is not."""
    money(w, w.G1, 600_000, block={"period_mode": "lessons", "lessons_per_payment": 12, "billing_history": [{"from": "min", "period_mode": "lessons", "lessons_per_payment": 12}]})
    pay(w, member(w, w.s1, w.G1), 600_000, date(2026, 10, 1))

    def row():
        r = client(w, "owner").get(f"{API}/org/payments/status", params={"limit": 100})
        assert r.status_code == 200, r.text
        return next(i for i in r.json()["items"] if i["student_id"] == str(w.s1.id))

    before = row()
    assert before["lessons_used"] == 3
    item = request_one(w, "A", w.G1, "lesson-l2", "2026-10-21")                     # today 16:00, two hours ahead
    clock.set("2026-10-21 14:02")
    assert decide(w, "admin", item["id"], decision="cancel", reason="Преподаватель болен, замены нет").status_code == 200
    clock.set("2026-10-21 23:00")
    after = row()
    assert after["lessons_used"] == 3 and after["lessons_left"] == before["lessons_left"] and after["debt"] == before["debt"]


# ======================================================================================================================== 6. import x seat rule x opening
def book(rows):
    wb = Workbook()
    wb.remove(wb.active)
    g = wb.create_sheet(imp.SHEET_GROUPS)
    g.append(["Название группы*", "Класс (1-11)", "Email учителя", "Макс. учеников"])
    for name in ("9-А", "7-Б"):
        g.append([name, 9, None, None])
    s = wb.create_sheet(imp.SHEET_STUDENTS)
    s.append(imp.COLS_STUDENTS)
    for r in rows:
        s.append(r)
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def test_import_with_an_opening_balance_obeys_the_seat_rule_of_b_and_writes_the_opening_of_a(w, clock):
    """9-А is full (3/3): the row is skipped ``GROUP_FULL`` and no opening is written for it; 7-Б has room: the student is enrolled by B's path and gets A's opening row;
    an existing student gets no opening; the preview counts the openings that WILL be written; an admin without the right to money enrols without an opening."""
    money(w, w.G2, 600_000)
    data = book([
        ["Полный", "Ученик", 9, None, "9-А", None, date(2026, 10, 31), None, None],                  # a full group
        ["Новый", "Ученик", 7, None, "7-Б", None, date(2026, 10, 31), None, None],                    # paid until 31.10
        ["Должник", "Ученик", 7, None, "7-Б", None, None, "250 000", "30.09.2026"],                   # a debt
        [w.s4.last_name, w.s4.first_name, 7, None, "7-Б", None, date(2026, 10, 31), None, None],       # already in 7-Б: unchanged, no opening
    ])
    files = {"file": ("x.xlsx", data, XLSX)}
    own = client(w, "owner")
    pv = own.post(f"{API}/org/import/preview", files=files).json()
    assert pv["ok"], pv
    assert pv["counts"]["openings"] == 2 and pv["counts"]["students_enrolled"] == 2
    assert sorted(s["code"] for s in pv["skipped"]) == ["ALREADY_IN_GROUP", "GROUP_FULL"]
    done = own.post(f"{API}/org/import/commit", files=files).json()
    assert done["ok"] and done["counts"]["openings"] == 2 and done["counts"]["students_enrolled"] == 2
    w.db.expire_all()
    assert w.db.query(Student).filter_by(org_id=w.org.id, last_name="Полный").count() == 0
    openings = w.db.query(MembershipOpeningBalance).filter_by(org_id=w.org.id, voided_at=None).all()
    assert {(o.paid_until, o.debt_amount, o.debt_on, o.source) for o in openings} == {
        (date(2026, 10, 31), None, None, "import"), (None, 250_000, date(2026, 9, 30), "import")}
    for o in openings:
        m = w.db.get(StudentGroup, o.membership_id)
        assert m.group_id == w.G2.id and m.is_active
    data2 = book([["Без прав", "Ученик", 7, None, "7-Б", None, date(2026, 10, 31), None, None]])
    adm = client(w, "admin").post(f"{API}/org/import/commit", files={"file": ("x.xlsx", data2, XLSX)}).json()
    assert adm["ok"] and adm["counts"]["students_enrolled"] == 1 and adm["counts"]["openings"] == 0
    assert any("Нет прав вносить остатки" in i["message"] for i in adm["issues"])


def test_import_row_whose_opening_is_refused_is_skipped_whole(w, clock, monkeypatch):
    """The opening is written in the row's transaction: when A refuses it, the student of the row is not left behind (B's all-or-nothing row)."""
    from src.auth import logica_service

    money(w, w.G2, 600_000)

    def refuse(*a, **k):
        raise ValueError("отказ остатка")

    monkeypatch.setattr(logica_service.OpeningService, "apply", refuse)
    data = book([["Отказ", "Ученик", 7, None, "7-Б", None, date(2026, 10, 31), None, None]])
    done = client(w, "owner").post(f"{API}/org/import/commit", files={"file": ("x.xlsx", data, XLSX)}).json()
    assert done["ok"] and done["counts"]["students_enrolled"] == 0
    assert [s["code"] for s in done["skipped"]] == ["ROW_ERROR"]
    w.db.expire_all()
    assert w.db.query(Student).filter_by(org_id=w.org.id, last_name="Отказ").count() == 0


# ======================================================================================================================== 7. migration chain
def test_the_linear_chain_composes_the_check_constraints_of_a_and_b(monkeypatch):
    """e1f2 -> f2a3 (A) -> lb1 -> lb2 -> lb3 (B / B2): one head; the three CHECKs that both packs extend from their LIVE definition hold the values of both after the
    upgrade; a downgrade of B's revisions keeps A's values (and the other way round); the audit rows of a package go with its downgrade only."""
    import os
    from pathlib import Path

    from alembic import command
    from alembic.config import Config
    from alembic.script import ScriptDirectory
    from sqlalchemy import create_engine, text
    from sqlalchemy.engine import make_url

    from src.core.db_timezone import utc_connect_args

    base = os.environ.get("A5_AUTH_MIGRATION_TEST_DATABASE_URL")
    if not base:
        pytest.skip("Requires isolated a5_auth_migrations PostgreSQL")
    root = Path(__file__).resolve().parents[1]
    config = Config(str(root / "alembic.ini"))
    config.set_main_option("script_location", str(root / "alembic"))
    script = ScriptDirectory.from_config(config)
    assert script.get_heads() == ["w1a2b3c4d5e"]                              # W1 on top of PAY-F, on top of the console pack, the IA fix pack (student journal) and LOGIC-B2
    assert [r.revision for r in script.walk_revisions()][:9] == ["w1a2b3c4d5e", "pf1a2b3c4d5", "cn1a2b3c4d5", "sa1b2c3d4e5", "lb3c4d5e6f7", "lb2b3c4d5e6", "lb1a2b3c4d5", "f2a3b4c5d6e", "e1f2a3b4c5d"]
    name = "a5_auth_chain_migrations"
    url_base = make_url(base)
    admin = create_engine(url_base.set(database="postgres"), isolation_level="AUTOCOMMIT", connect_args=utc_connect_args(base))
    with admin.connect() as c:
        c.execute(text(f"DROP DATABASE IF EXISTS {name} WITH (FORCE)"))
        c.execute(text(f"CREATE DATABASE {name}"))
    url = url_base.set(database=name).render_as_string(hide_password=False)
    monkeypatch.setenv("DATABASE_URL", url)
    eng = create_engine(url, connect_args=utc_connect_args(url))
    checks = (("group_audit", "ck_group_audit_action"), ("group_attendance_audit", "ck_group_attendance_audit_action"),
              ("attendance_excuses", "ck_attendance_excuse_status"))
    A_VALUES = {"group_audit": ["membership_pause_set", "student_left_group", "group_lesson_type_set"],
                "group_attendance_audit": ["excuse_closed_pause"], "attendance_excuses": ["not_required"]}
    B_VALUES = {"group_audit": ["group_archived", "group_restored", "lesson_teacher_set", "lesson_teacher_cleared"],
                "group_attendance_audit": ["teacher_changed"], "attendance_excuses": []}

    def definition(table, constraint):
        with eng.connect() as c:
            return c.execute(text("SELECT pg_get_constraintdef(c.oid) FROM pg_constraint c JOIN pg_class t ON t.oid=c.conrelid WHERE t.relname=:t AND c.conname=:n"),
                             {"t": table, "n": constraint}).scalar() or ""

    try:
        command.upgrade(config, "head")
        for table, constraint in checks:
            d = definition(table, constraint)
            assert all(f"'{v}'" in d for v in A_VALUES[table] + B_VALUES[table]), (table, d)
        command.downgrade(config, "f2a3b4c5d6e")                                    # B / B2 go, A stays
        for table, constraint in checks:
            d = definition(table, constraint)
            assert all(f"'{v}'" in d for v in A_VALUES[table]) and not any(f"'{v}'" in d for v in B_VALUES[table]), (table, d)
        command.upgrade(config, "head")
        command.downgrade(config, "e1f2a3b4c5d")                                    # everything goes
        for table, constraint in checks:
            d = definition(table, constraint)
            assert not any(f"'{v}'" in d for v in A_VALUES[table] + B_VALUES[table]), (table, d)
        command.upgrade(config, "head")
        with eng.connect() as c:
            assert c.execute(text("SELECT version_num FROM alembic_version")).scalar_one() == "w1a2b3c4d5e"
    finally:
        eng.dispose()
        with admin.connect() as c:
            c.execute(text(f"DROP DATABASE IF EXISTS {name} WITH (FORCE)"))
        admin.dispose()


# ======================================================================================================================== 8. group page x pause
def test_the_group_page_flags_the_students_on_pause_today(w, clock):
    """``GET /org/groups/{id}``: every student carries ``on_pause`` / ``paused_until`` (the last day of the pause that covers TODAY), so the lesson panel hides him BEFORE the
    start; a pause of the future or of the past does not flag him; a voided pause neither."""
    money(w, w.G1, 600_000)
    adm = client(w, "admin")
    now = adm.post(f"{API}/org/memberships/{member(w, w.s3, w.G1).id}/pauses", json={"first_day": "2026-10-20", "last_day": "2026-10-24"})
    later = adm.post(f"{API}/org/memberships/{member(w, w.s2, w.G1).id}/pauses", json={"first_day": "2026-10-26", "last_day": "2026-10-31"})
    assert now.status_code == 201 and later.status_code == 201, (now.text, later.text)
    group = client(w, "A").get(f"{API}/org/groups/{w.G1.id}")
    assert group.status_code == 200, group.text
    by_id = {s["id"]: s for s in group.json()["students"]}
    assert (by_id[str(w.s3.id)]["on_pause"], by_id[str(w.s3.id)]["paused_until"]) == (True, "2026-10-24")
    assert (by_id[str(w.s2.id)]["on_pause"], by_id[str(w.s2.id)]["paused_until"]) == (False, None)
    assert (by_id[str(w.s1.id)]["on_pause"], by_id[str(w.s1.id)]["paused_until"]) == (False, None)
    voided = adm.post(f"{API}/org/pauses/{now.json()['pause']['id']}/void", json={"reason": "ошибся датой"})
    assert voided.status_code == 200, voided.text
    again = {s["id"]: s for s in client(w, "A").get(f"{API}/org/groups/{w.G1.id}").json()["students"]}
    assert again[str(w.s3.id)]["on_pause"] is False
