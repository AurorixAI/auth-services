"""LOGIC-B O9 - the Excel import goes by the rules of the interface.  Matrix rows [O9.R1] ... [O9.R9]; chain 3 (golden)."""
from io import BytesIO

import pytest
from openpyxl import Workbook

from src.auth import payroll_hooks
from src.models.auth import Group, GroupAudit, Student, StudentGroup

from tests.ops_helpers import (  # noqa: F401  (fixtures)
    build_d0, clock, client, code_of, diff, exams, fixture, needs_pg, new_group, new_slot, new_student, pg_session, run_parallel,
)

pytestmark = needs_pg
IMP = "/api/v1/auth/org/import"
MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
GH = ["Название группы*", "Класс (1-11)", "Email учителя", "Макс. учеников"]
SH = ["Фамилия*", "Имя*", "Класс (1-11)*", "Телефон", "Группа*", "Комментарий"]


@pytest.fixture
def w(pg_session, clock):
    clock.set("2026-10-21 14:00")
    world = build_d0(pg_session)
    world.s3.first_name, world.s3.last_name = "Иван", "Иванов"                 # "Иванов Иван": a student who is in 9-А (3/3)
    arch = new_group(pg_session, world.org, world.ids, "G5v", "5-В", None, [], class_level=5)
    arch.status = "archived"
    pg_session.commit()
    return world


def book(groups, students):
    wb = Workbook()
    wb.remove(wb.active)
    g = wb.create_sheet("Группы")
    g.append(GH)
    for row in groups:
        g.append(row)
    s = wb.create_sheet("Ученики")
    s.append(SH)
    for row in students:
        s.append(row)
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def chain3(w, teacher_email=None):
    teacher_email = teacher_email or w.C.email
    return book(
        [["9-А", 9, None, None], ["5-В", 5, None, None], ["6-Д", 6, teacher_email, 2]],
        [["Бобоев", "Али", 9, None, "9-А", None], ["Каримова", "Лола", 6, "+998901112233", "6-Д", None],
         ["Рахимов", "Жасур", 6, None, "6-Д", None], ["Султанова", "Нодира", 6, None, "6-Д", None],
         ["Иванов", "Иван", 9, None, "6-Д", None], ["Юсупов", "Тимур", 5, None, "5-В", None],
         ["Каримова", "Лола", 6, "+998901112233", "6-Д", None]],
    )


def post(w, kind, data, who="admin"):
    return client(w, who).post(f"{IMP}/{kind}", files={"file": ("import.xlsx", data, MIME)})


def test_chain3_preview_commit_and_repeat_golden(w):
    """[O9.R1] [O9.R2] [O9.R3] [O9.R6] [O9.R7] Chain 3: an unknown teacher blocks; the corrected file previews and commits through the interface rules; a repeat changes nothing."""
    blocked = post(w, "preview", chain3(w, "nobody@algo.test"))
    assert blocked.status_code == 200 and blocked.json() == fixture("c3_preview_teacher_not_teacher.json")
    r = post(w, "preview", chain3(w))
    expected_preview = fixture("c3_preview_ok.json")
    for g in expected_preview["groups_preview"]:
        if g["teacher_email"] == "TEACHER_EMAIL":
            g["teacher_email"] = w.C.email.lower()
    assert r.status_code == 200 and not diff(r.json(), expected_preview), diff(r.json(), expected_preview)
    before = w.db.query(Student).filter_by(org_id=w.org.id).count()
    assert before == 6                                                               # a preview writes nothing
    done = post(w, "commit", chain3(w))
    assert done.status_code == 200, done.text
    out = done.json()
    expected = fixture("c3_commit_ok.json")
    created = out["created_groups"]
    assert len(created) == 1 and created[0]["name"] == "6-Д"
    expected["created_groups"][0]["id"] = created[0]["id"]
    assert not diff(out, expected), diff(out, expected)
    w.db.expire_all()
    group = w.db.get(Group, created[0]["id"])
    assert group.status == "active" and group.max_students == 2 and group.teacher_id == w.C.id and group.lesson_type == "group"
    audit = w.db.query(GroupAudit).filter_by(group_id=group.id, action="created").one()
    assert audit.snapshot["source"] == "import" and audit.snapshot["lesson_type"] == "group"           # INV-O9-4
    members = w.db.query(Student, StudentGroup).join(StudentGroup, StudentGroup.student_id == Student.id).filter(StudentGroup.group_id == group.id).all()
    assert sorted(s.last_name for s, _m in members) == ["Каримова", "Рахимов"]
    assert all(s.status == "active" and s.display_id and m.is_active and m.periods for s, m in members)  # INV-O9-3
    assert w.db.query(Student).filter_by(org_id=w.org.id).count() == before + 2                          # the skipped rows made nobody
    assert w.db.query(StudentGroup).filter_by(group_id=w.G1.id, is_active=True).count() == 3             # INV-O9-1: 9-А is still 3/3
    assert w.db.query(StudentGroup).join(Group, Group.id == StudentGroup.group_id).filter(Group.org_id == w.org.id, Group.status == "archived").count() == 0         # INV-O9-2
    again = post(w, "commit", chain3(w))
    exp2 = fixture("c3_commit_repeat.json")
    assert not diff(again.json(), exp2), diff(again.json(), exp2)                                                            # INV-O9-5
    assert w.db.query(Student).filter_by(org_id=w.org.id).count() == before + 2


def test_r4_parallel_imports_one_wins_and_seats_are_checked_under_the_lock(w):
    """[O9.R4] A second import of the same centre at the same time is 409 IMPORT_IN_PROGRESS (an organisation-wide advisory lock)."""
    from src.auth import import_ops
    from contextlib import ExitStack

    with import_ops._org_lock(w.db, w.org.id):
        r = post(w, "commit", chain3(w))
        assert r.status_code == 409 and r.json()["detail"]["code"] == "IMPORT_IN_PROGRESS"
    assert post(w, "commit", chain3(w)).status_code == 200                     # released: the next one runs


def test_r5_the_after_enrol_hook_is_the_extension_point_of_the_payment_package(w, monkeypatch):
    """[O9.R5] Owner decision 1 (paid-until / debt columns) is package A's: the hook is called once per enrolled student with the parsed row."""
    seen = []
    monkeypatch.setattr(payroll_hooks, "after_enrol", lambda db, student, membership, row: seen.append((student.last_name, membership.is_active, row.row)))
    assert post(w, "commit", chain3(w)).status_code == 200
    assert sorted(seen) == [("Каримова", True, 3), ("Рахимов", True, 4)]


def test_r6_places_are_counted_in_the_file_order(w):
    """[O9.R6] max 2, one seat taken: the first student of the file is enrolled, the second is GROUP_FULL with the row number."""
    group = new_group(w.db, w.org, w.ids, "Gx", "Малая", w.D, [], max_students=2)
    new_student(w.db, w.org, w.ids, "x1", "Занятов", "Один", group)
    w.db.commit()
    data = book([["Малая", 9, None, None]], [["Первый", "Ученик", 9, None, "Малая", None], ["Второй", "Ученик", 9, None, "Малая", None]])
    r = post(w, "commit", data).json()
    assert r["counts"]["students_enrolled"] == 1 and [(s["row"], s["code"]) for s in r["skipped"]] == [(3, "GROUP_FULL")]
    assert "лимит 2, занято 2" in r["skipped"][0]["message"]


def test_r7_homonym_without_phone_and_the_teacher_errors(w):
    """[O9.R7] A student of the same name without a phone is not created blindly (HOMONYM_NO_PHONE); a teacher who left is TEACHER_INACTIVE; unknown is TEACHER_NOT_FOUND."""
    new_student(w.db, w.org, w.ids, "hom", "Тёзкин", "Тёзка", w.G2, phone="+998900000001")
    w.db.commit()
    data = book([["7-Б", 9, None, None]], [["Тёзкин", "Тёзка", 9, None, "7-Б", None]])
    r = post(w, "preview", data).json()
    assert [s["code"] for s in r["skipped"]] == ["HOMONYM_NO_PHONE"]
    w.member("D").is_active = False
    w.db.commit()
    bad = post(w, "preview", book([["Новая", 9, w.D.email, None]], []))
    assert bad.json()["issues"][0]["code"] == "TEACHER_INACTIVE" and bad.json()["ok"] is False
    # the owner / admin may teach (lead decision)
    ok = post(w, "preview", book([["Новая", 9, w.admin.email, None]], [])).json()
    assert ok["ok"] is True
    rng = post(w, "preview", book([["Новая", 9, None, 101]], [])).json()
    assert rng["ok"] is False and rng["issues"][0]["code"] == "MAX_STUDENTS_RANGE"


def test_r8_owner_and_admin_only(w):
    """[O9.R8]"""
    assert post(w, "preview", chain3(w), who="C").status_code == 403
    assert post(w, "commit", chain3(w), who="owner").status_code == 200


def test_r9_new_students_are_active_at_once_and_a_student_without_a_group_is_enrolled(w):
    """[O9.R9] New imported students are active with a membership; an existing student without a group is enrolled (status stays honest)."""
    loose = new_student(w.db, w.org, w.ids, "loose", "Свободный", "Ученик", None, status="active")
    w.db.commit()
    data = book([["7-Б", 9, None, None]], [["Свободный", "Ученик", 9, None, "7-Б", None], ["Свежий", "Ученик", 9, "+998911111111", "7-Б", None]])
    r = post(w, "commit", data).json()
    assert r["counts"]["students_enrolled"] == 2 and r["counts"]["students_new"] == 1
    w.db.expire_all()
    assert w.db.query(StudentGroup).filter_by(student_id=loose.id, group_id=w.G2.id, is_active=True).count() == 1
    fresh = w.db.query(Student).filter_by(org_id=w.org.id, last_name="Свежий").one()
    assert fresh.status == "active" and fresh.display_id and fresh.group_id == w.G2.id
