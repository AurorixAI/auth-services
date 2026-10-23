"""PAY-2 stream B: the payment API on PostgreSQL - golden fixtures through the endpoints, write-off, refund, replace / void,
idempotency, transfers, reports."""
import json
import os
from datetime import date, datetime, timedelta
from pathlib import Path
from uuid import uuid4

import pytest

from src.auth import payment_status as ps
from src.models.auth import (
    GroupAudit, MembershipTransfer, Student, StudentGroup, StudentPayment, StudentRefund, StudentWriteoff,
)

from tests.pg_world import (  # noqa: F401  (fixtures)
    give_group_price, make_client, make_group, make_student, needs_pg, pg_session, world,
)

pytestmark = needs_pg
API = "/api/v1/auth"
D = date
TODAY = D(2026, 10, 21)
FIX = Path(__file__).resolve().parent / "fixtures" / "pay2"


def fx(name):
    return json.loads((FIX / f"{name}.json").read_text(encoding="utf-8"))


_CLOCKED_MODULES = ("service", "leave_service", "membership_periods", "payments_service", "student_archive", "logica_service", "payroll_service")


@pytest.fixture
def fixed_today(monkeypatch):
    """The business day is ``day`` (``ps.business_today``) AND the clock of the code that stamps rows (``datetime.now`` in the modules below, ``business_clock``)
    stands at 12:00 of that day (Tashkent): a test no longer depends on the date it runs at (a leave stamps ``left_at`` with the clock, the ledger reads the day)."""
    import importlib
    from datetime import datetime as real_datetime, timezone

    from src.auth import business_clock as bc

    state = {"day": TODAY}

    def moment():
        day = state["day"]
        hour, minute = (int(x) for x in os.environ.get("ALGO_TEST_CLOCK", "12:00").split(":"))        # Tashkent wall time; a stress run sets other values
        return real_datetime(day.year, day.month, day.day, hour, minute, tzinfo=bc.BUSINESS_TIMEZONE)

    class Frozen(real_datetime):
        @classmethod
        def now(cls, tz=None):
            m = moment()
            return m.astimezone(tz) if tz is not None else m.astimezone(timezone.utc).replace(tzinfo=None)

        @classmethod
        def utcnow(cls):
            return moment().astimezone(timezone.utc).replace(tzinfo=None)

    for name in _CLOCKED_MODULES:
        module = importlib.import_module(f"src.auth.{name}")
        if getattr(module, "datetime", None) is real_datetime:
            monkeypatch.setattr(module, "datetime", Frozen)
    bc.set_clock(moment)

    def fix(day=TODAY):
        state["day"] = day
        monkeypatch.setattr(ps, "business_today", lambda tz=None: day)
    fix()
    yield fix
    bc.set_clock(None)


def owner(w):
    return make_client(w, user=w.owner)


def settings_block(w, **block):
    org = w.org
    settings = dict(org.settings or {})
    settings["payments"] = {**settings.get("payments", {}), **block}
    org.settings = settings
    w.db.commit()


def membership(w, student=None, group=None):
    w.db.expire_all()
    return w.db.query(StudentGroup).filter_by(student_id=(student or w.students[0]).id,
                                              group_id=(group or w.group).id).one()


def stay(w, m, first, last=None):
    m.joined_at = datetime(first.year, first.month, first.day, 6)
    if last is not None:
        m.left_at = datetime(last.year, last.month, last.day, 6)
        m.is_active = False
    (period,) = m.periods
    period.joined_at, period.left_at = m.joined_at, m.left_at
    w.db.commit()


def seed_payment(w, m, amount, paid_on, receipt_no=None, method="cash"):
    p = StudentPayment(org_id=w.org.id, membership_id=m.id, student_id=m.student_id, group_id=m.group_id, amount=amount,
                       paid_on=paid_on, created_by=w.owner.id, method=method, receipt_no=receipt_no,
                       created_at=datetime(paid_on.year, paid_on.month, paid_on.day, 9))
    w.db.add(p)
    w.db.commit()
    return p


def e1_world(w):
    """The owner's example: payment_date, price 500 000 from 01.01, joined 15.09, 231 000 paid on 20.09, today 21.10."""
    settings_block(w, accounting_start_date="2026-09-01")
    give_group_price(w.db, w.group, 500_000, valid_from=D(2026, 1, 1))
    m = membership(w)
    stay(w, m, D(2026, 9, 15))
    seed_payment(w, m, 231_000, D(2026, 9, 20), receipt_no=17)
    return m


def preview(w, m, **params):
    r = owner(w).get(f"{API}/org/memberships/{m.id}/payments/preview", params=params)
    assert r.status_code == 200, r.text
    return r.json()


def post_payment(w, m, expect=201, headers=None, **body):
    body.setdefault("paid_on", TODAY.isoformat())
    body.setdefault("method", "cash")
    r = owner(w).post(f"{API}/org/memberships/{m.id}/payments", json=body, headers=headers or {})
    assert r.status_code == expect, r.text
    return r.json()


# ---- golden fixtures through the endpoints -------------------------------------------------------------------------


def test_e1_preview_endpoint_equals_the_fixtures(world, fixed_today):
    m = e1_world(world)
    assert preview(world, m) == fx("e1_preview_due")
    for amount in (250_000, 300_000, 769_000, 790_000):
        assert preview(world, m, amount=amount) == fx(f"e1_preview_{amount}"), amount


def test_e9_price_ahead_endpoint_equals_the_fixture(world, fixed_today):
    m = e1_world(world)
    give_group_price(world.db, world.group, 600_000, valid_from=D(2026, 11, 1))
    assert preview(world, m, amount=1_069_000) == fx("e9_preview_price_ahead_1069000")


def test_e2_and_e3_preview_endpoints(world, fixed_today):
    settings_block(world, accounting_start_date="2026-09-01", billing_anchor="fixed_day", billing_day=1)
    give_group_price(world.db, world.group, 600_000, valid_from=D(2026, 1, 1))
    # the group meets Mon / Wed / Fri (the test group meets daily: replace the schedule)
    world.group.schedule = [{"id": f"s{d}", "dayOfWeek": d, "startTime": "10:00", "endTime": "11:00", "format": "offline"}
                            for d in (1, 3, 5)]
    world.db.commit()
    m = membership(world)
    stay(world, m, D(2026, 10, 16))
    for amount in (200_000, 323_000, 900_000):
        assert preview(world, m, amount=amount) == fx(f"e2_preview_{amount}"), amount
    assert preview(world, m) == fx("e2_preview_due")


def test_e3_gap(world, fixed_today):
    fixed_today(D(2026, 10, 25))
    settings_block(world, accounting_start_date="2026-09-01", billing_anchor="fixed_day", billing_day=1)
    give_group_price(world.db, world.group, 300_000, valid_from=D(2026, 1, 1))
    world.group.schedule = []
    world.db.commit()
    m = membership(world)
    stay(world, m, D(2026, 10, 1), D(2026, 10, 10))
    world.db.add(__import__("src.models.auth", fromlist=["StudentGroupPeriod"]).StudentGroupPeriod(
        membership_id=m.id, joined_at=datetime(2026, 10, 20, 6)))
    m.is_active = True
    m.left_at = None
    world.db.commit()
    for amount in (100_000, 213_000):
        assert preview(world, m, amount=amount) == fx(f"e3_preview_{amount}"), amount


def test_e5_lessons_preview_endpoint(world, fixed_today):
    from tests.pg_world import reconcile, save
    fixed_today(D(2026, 10, 27))
    settings_block(world, period_mode="lessons", lessons_per_payment=12, accounting_start_date="2026-09-01",
                   billing_history=[{"from": "min", "period_mode": "lessons", "lessons_per_payment": 12}])
    give_group_price(world.db, world.group, 600_000, valid_from=D(2026, 1, 1))
    # 12 held lessons for the first student are made by the helper below (the journal), dated by the test calendar
    m = membership(world)
    stay(world, m, D(2026, 9, 25))
    from src.models.auth import GroupAttendance
    s1 = world.students[0]
    dates = [D(2026, 9, 25), D(2026, 9, 28), D(2026, 9, 30), D(2026, 10, 2), D(2026, 10, 5), D(2026, 10, 7), D(2026, 10, 9),
             D(2026, 10, 12), D(2026, 10, 14), D(2026, 10, 16), D(2026, 10, 19), D(2026, 10, 21)]
    for i, d in enumerate(dates):
        world.db.add(GroupAttendance(org_id=world.org.id, group_id=world.group.id, lesson_date=d.isoformat(),
                                     lesson_id=f"s{d.isoweekday()}", status="completed", records={str(s1.id): "present"},
                                     roster_snapshot=[{"student_id": str(s1.id)}]))
    world.db.commit()
    fixed_today(D(2026, 10, 21))
    got = preview(world, m, amount=630_000)
    want = fx("e5_preview_630000_held12")
    assert got["result"]["lessons"] == want["result"]["lessons"]
    assert got["result"]["summary_text"] == want["result"]["summary_text"]
    assert [(a["amount"], a["days"], a["bucket"]) for a in got["result"]["allocations"]] == [
        (a["amount"], a["days"], a["bucket"]) for a in want["result"]["allocations"]]
    assert got["due"]["debt"] == 600_000 and got["suggested_amount"] == 600_000


# ---- create / replace / void -----------------------------------------------------------------------------------


def test_create_payment_is_money_with_a_distribution(world, fixed_today):
    m = e1_world(world)
    body = post_payment(world, m, amount=300_000, note="за октябрь")
    want = fx("e1_create_300000")
    assert body["allocation"] == want["allocation"] and body["warnings"] == want["warnings"]
    p = body["payment"]
    assert (p["amount"], p["method"], p["kind"], p["receipt_no"], p["legacy_period"]) == (300_000, "cash", "payment", 18, False)
    assert p["valid_from"] == "2026-09-15" and p["valid_to"] == "2026-10-15"            # the shim for the old front
    assert p["allocation"] == want["payment"]["allocation"]
    for key in ("status", "paid_until", "days_overdue", "debt", "past_debt", "current_due", "advance", "partial_credit",
                "refundable", "writeoff_allowed", "price_missing", "lessons_unpaid", "period_mode"):
        assert body["row"][key] == want["row"][key], key
    # the snapshot of the moment is what the receipt prints
    r = owner(world).get(f"{API}/org/payments/{p['id']}/receipt")
    assert r.status_code == 200
    rec, want_rec = r.json(), fx("e1_receipt_300000")
    for key in ("allocations", "debt_before", "debt_after", "advance_after", "paid_until_before", "paid_until_after",
                "price_snapshot", "method", "mode", "valid_from", "valid_to", "amount_words"):
        assert rec[key] == want_rec[key], key
    # ... and it does not change when the facts change later (a later payment), while the payment's own allocation does
    post_payment(world, m, amount=469_000)
    assert owner(world).get(f"{API}/org/payments/{p['id']}/receipt").json()["allocations"] == want_rec["allocations"]


def test_payment_validation_codes(world, fixed_today):
    m = e1_world(world)
    for body, code in (({"amount": 0}, "AMOUNT_RANGE"), ({"amount": 10_000_000_001}, "AMOUNT_RANGE"),
                       ({"amount": 1000, "paid_on": "2026-10-22"}, "PAID_ON_RANGE"),
                       ({"amount": 1000, "paid_on": "2023-10-20"}, "PAID_ON_RANGE"),
                       ({"amount": 1000, "method": "unknown"}, "METHOD_REQUIRED")):
        r = owner(world).post(f"{API}/org/memberships/{m.id}/payments", json={"paid_on": TODAY.isoformat(), "method": "cash", **body})
        assert r.status_code == 422 and r.json()["detail"]["code"] == code, (body, r.text)
    # no method at all: the old front - stored as "other"
    r = owner(world).post(f"{API}/org/memberships/{m.id}/payments", json={"amount": 1000, "paid_on": TODAY.isoformat()})
    assert r.status_code == 201 and r.json()["payment"]["method"] == "other"
    # the legacy period fields are accepted and ignored
    r = owner(world).post(f"{API}/org/memberships/{m.id}/payments", json={
        "amount": 2000, "paid_on": TODAY.isoformat(), "valid_from": "2026-01-01", "valid_to": "2026-01-31", "lessons_count": 5})
    assert r.status_code == 201
    row = world.db.get(StudentPayment, r.json()["payment"]["id"])
    assert (row.valid_from, row.valid_to, row.lessons_count) == (None, None, None)


def test_idempotency_key_and_soft_double_click(world, fixed_today, monkeypatch):
    m = e1_world(world)
    key = str(uuid4())
    first = post_payment(world, m, amount=100_000, headers={"Idempotency-Key": key})
    r = owner(world).post(f"{API}/org/memberships/{m.id}/payments", json={"amount": 100_000, "paid_on": TODAY.isoformat(),
                                                                         "method": "cash"}, headers={"Idempotency-Key": key})
    assert r.status_code == 409 and r.json()["detail"]["code"] == "DUPLICATE_PAYMENT"
    assert r.json()["detail"]["payment_id"] == first["payment"]["id"]
    count = world.db.query(StudentPayment).filter_by(membership_id=m.id).count()
    assert count == 2                                                                           # the seeded one + one
    # without a key: the same student + date + sum within a minute is a double click
    import src.auth.payments_service as svc
    monkeypatch.setattr(svc, "SOFT_DUPLICATE_SECONDS", 60)
    second = post_payment(world, m, amount=7_000)
    r = owner(world).post(f"{API}/org/memberships/{m.id}/payments", json={"amount": 7_000, "paid_on": TODAY.isoformat(), "method": "cash"})
    assert r.status_code == 409 and r.json()["detail"]["code"] == "DUPLICATE_PAYMENT"
    assert r.json()["detail"]["payment_id"] == second["payment"]["id"]
    # a different sum is another payment; with the guard off two equal payments are legal
    post_payment(world, m, amount=7_001)
    monkeypatch.setattr(svc, "SOFT_DUPLICATE_SECONDS", 0)
    post_payment(world, m, amount=7_000)


def test_archived_student_pays_the_debt_only_and_can_be_refunded_and_written_off(world, fixed_today):
    """S-5 (LOGIC-A): a student who left pays his DEBT (not more); an advance cannot grow; the preview says so."""
    m = e1_world(world)                                                   # debt 769 000
    world.students[0].status = "archived"
    world.db.commit()
    r = owner(world).post(f"{API}/org/memberships/{m.id}/payments", json={"amount": 800_000, "paid_on": TODAY.isoformat(), "method": "cash"})
    assert r.status_code == 422 and r.json()["detail"]["code"] == "PAYMENT_EXCEEDS_DEBT" and r.json()["detail"]["debt"] == 769_000
    r = owner(world).post(f"{API}/org/memberships/{m.id}/payments", json={"amount": 1000, "paid_on": TODAY.isoformat(), "method": "cash"})
    assert r.status_code == 201
    r = owner(world).post(f"{API}/org/memberships/{m.id}/writeoffs", json={"amount": 1000, "reason": "льгота"})
    assert r.status_code == 201
    pv = preview(world, m)
    assert pv["warnings"][0]["code"] == "left_debt_only" and pv["suggested_amount"] == 767_000


def test_e8_replace_and_void(world, fixed_today):
    m = e1_world(world)
    p = post_payment(world, m, amount=300_000)["payment"]
    r = owner(world).post(f"{API}/org/payments/{p['id']}/replace", json={
        "amount": 250_000, "paid_on": TODAY.isoformat(), "method": "cash", "reason": "Неверная сумма"})
    assert r.status_code == 201, r.text
    body, want = r.json(), fx("e8_replace_300000_by_250000")
    assert body["allocation"] == want["allocation"]                       # the "before" excludes the replaced payment
    assert body["payment"]["replaces_payment_id"] == p["id"] and body["payment"]["receipt_no"] == 19
    for key in ("status", "paid_until", "days_overdue", "debt", "past_debt", "current_due", "partial_credit"):
        assert body["row"][key] == want["row"][key], key
    old = world.db.get(StudentPayment, p["id"])
    world.db.refresh(old)
    assert old.voided_at is not None and old.void_reason == "Неверная сумма"
    # void the replacement: the row returns to the state before it
    r = owner(world).post(f"{API}/org/payments/{body['payment']['id']}/void", json={"reason": "ошибка оператора"})
    assert r.status_code == 200
    assert r.json()["row"]["debt"] == 769_000 and r.json()["payment"]["voided_at"]
    r = owner(world).post(f"{API}/org/payments/{body['payment']['id']}/void", json={"reason": "ещё раз"})
    assert r.status_code == 409 and r.json()["detail"]["code"] == "ALREADY_VOIDED"
    r = owner(world).post(f"{API}/org/payments/{p['id']}/void", json={"reason": "а"})
    assert r.status_code == 422                                                             # the reason is 3..500


# ---- write-off (E6) and refund (E7) ----------------------------------------------------------------------------


def e6_world(w):
    settings_block(w, accounting_start_date="2026-09-01")
    give_group_price(w.db, w.group, 500_000, valid_from=D(2026, 1, 1))
    m = membership(w, w.students[1])
    stay(w, m, D(2026, 10, 15))
    seed_payment(w, m, 400_000, D(2026, 10, 15), receipt_no=31)
    return m


def test_e6_writeoff_flow(world, fixed_today):
    m = e6_world(world)
    c = owner(world)
    r = c.post(f"{API}/org/memberships/{m.id}/writeoffs", json={"amount": 40_000, "reason": "болезнь, договорённость с родителем",
                                                               "effective_on": TODAY.isoformat()})
    assert r.status_code == 201, r.text
    body, want = r.json(), fx("e6_writeoff_40000")
    assert (body["debt_before"], body["debt_after"]) == (100_000, 60_000)
    assert body["row"]["paid_until"] == want["row"]["paid_until"] and body["row"]["partial_credit"] == want["row"]["partial_credit"]
    assert body["writeoff"]["reason"] == "болезнь, договорённость с родителем" and body["writeoff"]["amount"] == 40_000
    # limits and codes
    r = c.post(f"{API}/org/memberships/{m.id}/writeoffs", json={"amount": 60_001, "reason": "много"})
    assert r.status_code == 422 and r.json()["detail"]["code"] == "WRITEOFF_EXCEEDS_DEBT" and r.json()["detail"]["debt"] == 60_000
    r = c.post(f"{API}/org/memberships/{m.id}/writeoffs", json={"amount": 100, "reason": "ab"})
    assert r.status_code == 422 and r.json()["detail"]["code"] == "REASON_REQUIRED"
    r = c.post(f"{API}/org/memberships/{m.id}/writeoffs", json={"amount": 100})
    assert r.status_code == 422
    # the student pays the rest: debt 0
    assert post_payment(world, m, amount=60_000)["row"]["debt"] == 0
    # "all": nothing to write off now
    r = c.post(f"{API}/org/memberships/{m.id}/writeoffs", json={"all": True, "reason": "всё"})
    assert r.status_code == 422 and r.json()["detail"]["debt"] == 0
    # void the write-off: the debt comes back; the audit has both rows with debt before / after
    wid = body["writeoff"]["id"]
    r = c.post(f"{API}/org/writeoffs/{wid}/void", json={"reason": "ошибка"})
    assert r.status_code == 200 and r.json()["row"]["advance"] == 0 and r.json()["writeoff"]["voided_at"]
    assert c.post(f"{API}/org/writeoffs/{wid}/void", json={"reason": "ещё"}).status_code == 409
    # scoped to THIS world: the test database is shared by every test (a foreign row of another test used to win the dict below, depending on the heap order)
    rows = world.db.query(GroupAudit).filter(GroupAudit.org_id == world.org.id, GroupAudit.action.in_(["payment_writeoff", "payment_writeoff_void"])).all()
    snaps = {a.action: a.snapshot for a in rows}
    assert snaps["payment_writeoff"]["debt_before"] == 100_000 and snaps["payment_writeoff"]["debt_after"] == 60_000
    assert snaps["payment_writeoff_void"]["amount"] == 40_000
    # whole debt
    m2 = membership(world, world.students[2])
    stay(world, m2, D(2026, 10, 15))
    r = c.post(f"{API}/org/memberships/{m2.id}/writeoffs", json={"all": True, "reason": "льгота"})
    assert r.status_code == 201 and r.json()["debt_after"] == 0 and r.json()["writeoff"]["amount"] == 500_000


def test_e7_refund_flow(world, fixed_today):
    fixed_today(D(2026, 10, 25))
    m = e1_world(world)
    seed_payment(world, m, 790_000, TODAY, receipt_no=18)
    stay(world, m, D(2026, 9, 15), D(2026, 10, 25))
    c = owner(world)
    row = c.get(f"{API}/org/students/{world.students[0].id}/memberships").json()[0]
    assert (row["advance"], row["refundable"]) == (344_000, 344_000)
    r = c.post(f"{API}/org/memberships/{m.id}/refunds", json={"amount": 344_001, "method": "cash"})
    assert r.status_code == 422 and r.json()["detail"]["code"] == "REFUND_EXCEEDS_ADVANCE" and r.json()["detail"]["refundable"] == 344_000
    r = c.post(f"{API}/org/memberships/{m.id}/refunds", json={"amount": 344_000})
    assert r.status_code == 422 and r.json()["detail"]["code"] == "METHOD_REQUIRED"
    r = c.post(f"{API}/org/memberships/{m.id}/refunds", json={"amount": 344_000, "method": "cash", "reason": "ученик ушёл"})
    assert r.status_code == 201, r.text
    body = r.json()
    assert (body["advance_before"], body["advance_after"]) == (344_000, 0) and body["row"]["advance"] == 0
    # voiding the payment now would leave the refund without money: refused
    pid = world.db.query(StudentPayment).filter_by(membership_id=m.id, amount=790_000).one().id
    r = c.post(f"{API}/org/payments/{pid}/void", json={"reason": "ошибка"})
    assert r.status_code == 409 and r.json()["detail"]["code"] == "VOID_BREAKS_REFUND"
    assert world.db.get(StudentPayment, pid).voided_at is None
    # cancel the refund: the advance is back; then the payment can go
    assert c.post(f"{API}/org/refunds/{body['refund']['id']}/void", json={"reason": "ошибка"}).json()["row"]["advance"] == 344_000
    assert c.post(f"{API}/org/payments/{pid}/void", json={"reason": "ошибка"}).status_code == 200
    assert {a.action for a in world.db.query(GroupAudit).filter(GroupAudit.org_id == world.org.id, GroupAudit.action.like("payment_refund%"))} == {
        "payment_refund", "payment_refund_void"}


# ---- statement, export ------------------------------------------------------------------------------------------


def test_e1_statement_endpoint_equals_the_fixture(world, fixed_today):
    m = e1_world(world)
    post_payment(world, m, amount=790_000, note="аванс")
    r = owner(world).get(f"{API}/org/memberships/{m.id}/statement")
    assert r.status_code == 200, r.text
    got, want = r.json(), fx("e1_statement_790000")
    assert got["summary"] == want["summary"] and got["upcoming"] == want["upcoming"] and got["accounting_start"] == "2026-09-01"
    strip = lambda rows: [{k: v for k, v in row.items() if k != "details"} | {"alloc": [(a["amount"], a["days"], a["bucket"], a["closes"]) for a in row["details"]["allocations"]]}
                          for row in rows]
    assert strip(got["rows"]) == strip(want["rows"])
    assert got["rows"][-1]["balance_after"] == got["summary"]["advance"] - got["summary"]["debt"]
    # credit of a payment row is the gross amount even after a partial refund (the refund is its own row)
    cut = owner(world).get(f"{API}/org/memberships/{m.id}/statement", params={"from": "2026-10-01", "to": "2026-10-31"}).json()
    assert [r["seq"] for r in cut["rows"]] == [3, 4] and cut["rows"][-1]["balance_after"] == 21_000
    assert cut["summary"] == got["summary"]


def test_statement_with_refund_keeps_the_gross_credit_and_voided_rows(world, fixed_today):
    m = e6_world(world)
    c = owner(world)
    post_payment(world, m, amount=300_000)                       # advance 200 000 on top of the period
    r = c.post(f"{API}/org/memberships/{m.id}/refunds", json={"amount": 50_000, "method": "card"})
    assert r.status_code == 201, r.text
    pay_id = r.json()["row"]["last_payment"]["id"]
    wo = c.post(f"{API}/org/memberships/{m.id}/writeoffs", json={"amount": 1_000, "reason": "тест"})
    st = c.get(f"{API}/org/memberships/{m.id}/statement", params={"include_voided": True}).json()
    credits = [(r["type"], r["credit"], r["outflow"]) for r in st["rows"]]
    assert ("payment", 300_000, 0) in credits and ("refund", 0, 50_000) in credits
    assert st["summary"]["refunded"] == 50_000 and st["summary"]["reconciles"] is True
    c.post(f"{API}/org/writeoffs/{wo.json()['writeoff']['id'] if wo.status_code == 201 else uuid4()}/void", json={"reason": "отмена"})
    st = c.get(f"{API}/org/memberships/{m.id}/statement", params={"include_voided": True}).json()
    assert st["voided"] == [] or st["voided"][0]["voided_at"] and st["voided"][0]["balance_after"] is None


def test_statement_export_filename(world, fixed_today):
    m = e1_world(world)
    for fmt, ctype in (("xlsx", "spreadsheetml"), ("csv", "text/csv")):
        r = owner(world).get(f"{API}/org/memberships/{m.id}/statement/export", params={"format": fmt})
        assert r.status_code == 200 and ctype in r.headers["content-type"]
        cd = r.headers["content-disposition"]
        assert cd.startswith("attachment;") and f".{fmt}" in cd and "statement-" in cd
    import io
    from openpyxl import load_workbook
    r = owner(world).get(f"{API}/org/memberships/{m.id}/statement/export", params={"format": "xlsx"})
    ws = load_workbook(io.BytesIO(r.content)).active
    flat = [c.value for row in ws.iter_rows() for c in row if c.value is not None]
    assert 231000 in flat and ws["A1"].value.startswith("Выписка — ") and "Оплата № 17 · наличные" in " ".join(str(x) for x in flat)
    head = next(r for r in range(1, 30) if ws.cell(r, 1).value == "Дата")
    assert [ws.cell(head, c).value for c in range(1, 7)] == ["Дата", "Операция", "Стоимость", "Оплачено", "Баланс", "Подробности"]
    assert ws.freeze_panes == f"A{head + 1}"


# ---- transfers ---------------------------------------------------------------------------------------------------


def test_transfer_through_add_student_to_group_makes_a_link(world, fixed_today):
    """AuthService.add_student_to_group(mode=transfer) records the link; money follows by itself; a payment voided in the old
    group changes what arrived in the new one."""
    db = world.db
    settings_block(world, accounting_start_date="2026-09-01")
    group_b = make_group(db, world.org, world.teacher, "Группа B")
    group_b.schedule = world.group.schedule = []
    give_group_price(db, world.group, 500_000, valid_from=D(2026, 1, 1))
    give_group_price(db, group_b, 400_000, valid_from=D(2026, 1, 1))
    db.commit()
    a = membership(world)
    stay(world, a, D(2026, 10, 1))
    seed_payment(world, a, 500_000, D(2026, 10, 1), receipt_no=1)
    # the real code path: the transfer happens "now" (the test moves the clock of the business date only)
    now = datetime.utcnow()
    result = world.svc.add_student_to_group(
        str(world.org.id), str(group_b.id), str(world.students[0].id), mode="transfer", actor_id=str(world.owner.id),
        payments_detail=True)
    links = db.query(MembershipTransfer).filter_by(student_id=world.students[0].id).all()
    assert len(links) == 1 and links[0].from_membership_id == a.id
    codes = [w.get("code") for w in result["warnings"]]
    assert "payment_carried_over" in codes
    assert any("Остаток денег" in w["message"] for w in result["warnings"])


# ---- reports -----------------------------------------------------------------------------------------------------


def report_world(w):
    """The centre of the report fixture: B (E1 student) 790 000 cash on 21.10, E (student 2) 200 000 card on 12.10 and a
    write-off of 30 000 on 20.10, F (student 3) 550 000 cash on 01.10 refunded 324 000 on 15.10 after leaving on 14.10."""
    settings_block(w, accounting_start_date="2026-09-01")
    w.group.name = "Математика 7А"
    give_group_price(w.db, w.group, 500_000, valid_from=D(2026, 1, 1))
    g2 = make_group(w.db, w.org, w.teacher2, "Алгебра 9Б")
    w.db.commit()
    give_group_price(w.db, g2, 300_000, valid_from=D(2026, 1, 1))
    m1 = membership(w)
    stay(w, m1, D(2026, 9, 15))
    seed_payment(w, m1, 231_000, D(2026, 9, 20), receipt_no=1)
    seed_payment(w, m1, 790_000, D(2026, 10, 21), receipt_no=4)
    s2 = make_student(w.db, w.org, g2, "Алия2", "Каримова2")
    w.db.commit()
    m2 = membership(w, s2, g2)
    stay(w, m2, D(2026, 9, 10))
    seed_payment(w, m2, 300_000, D(2026, 9, 12), receipt_no=2)
    seed_payment(w, m2, 200_000, D(2026, 10, 12), receipt_no=5, method="card")
    w.db.add(StudentWriteoff(org_id=w.org.id, membership_id=m2.id, student_id=m2.student_id, group_id=m2.group_id,
                             amount=30_000, reason="льгота", effective_on=D(2026, 10, 20), created_by=w.owner.id))
    m3 = membership(w, w.students[2])
    stay(w, m3, D(2026, 10, 1), D(2026, 10, 14))
    seed_payment(w, m3, 550_000, D(2026, 10, 1), receipt_no=3)
    w.db.add(StudentRefund(org_id=w.org.id, membership_id=m3.id, student_id=m3.student_id, group_id=m3.group_id,
                           amount=324_000, method="cash", effective_on=D(2026, 10, 15), created_by=w.owner.id))
    w.db.commit()
    return m1, m2, m3


def test_receipts_report_matches_the_fixture(world, fixed_today):
    report_world(world)
    r = owner(world).get(f"{API}/org/payments/receipts", params={"from": "2026-10-01", "to": "2026-10-31", "group_by": "day"})
    assert r.status_code == 200, r.text
    got, want = r.json(), fx("receipts_report_2026-10")
    # the fixture's centre has only these three memberships; the test world has more students (nobody paid) - the position
    # differs by their debts, everything about money received is the same
    for key in ("as_of", "from", "to", "group_by", "totals", "by_method", "series"):
        assert got[key] == want[key], key
    assert [(g["group_name"], g["received"], g["to_past_debt"], g["to_current"], g["to_advance"], g["written_off"],
             g["refunded"], g["count"]) for g in got["by_group"]] == [
        (g["group_name"], g["received"], g["to_past_debt"], g["to_current"], g["to_advance"], g["written_off"],
         g["refunded"], g["count"]) for g in want["by_group"]]
    assert got["reconciliation"]["ok"] is True
    assert {c["code"]: c["diff"] for c in got["reconciliation"]["checks"]} == {
        "received_eq_allocated_plus_refunded": 0, "charged_minus_paid_eq_debt_minus_advance": 0, "transfers_net_zero": 0}
    # drill-down: the past-debt column is the payment of the E1 student
    r = owner(world).get(f"{API}/org/payments/receipts/items", params={"from": "2026-10-01", "to": "2026-10-31", "column": "to_past_debt"})
    items = r.json()
    assert items["total"] == 1 and items["items"][0]["amount"] == 790_000 and items["items"][0]["column_amount"] == 269_000
    assert items["items"][0]["receipt_no"] == 4 and len(items["items"][0]["allocations"]) == 3
    r = owner(world).get(f"{API}/org/payments/receipts/items", params={"from": "2026-10-01", "to": "2026-10-31", "column": "received", "bucket": "2026-10-12"})
    assert [i["amount"] for i in r.json()["items"]] == [200_000]
    r = owner(world).get(f"{API}/org/payments/receipts/items", params={"from": "2026-10-01", "to": "2026-10-31", "column": "refunded"})
    assert [(i["type"], i["amount"]) for i in r.json()["items"]] == [("refund", 324_000)]
    # ranges
    assert owner(world).get(f"{API}/org/payments/receipts", params={"from": "2026-01-01", "to": "2026-10-31", "group_by": "day"}).json()["detail"]["code"] == "RANGE_TOO_LONG"
    assert owner(world).get(f"{API}/org/payments/receipts", params={"from": "2025-01-01", "to": "2026-10-31", "group_by": "month"}).status_code == 422
    r = owner(world).get(f"{API}/org/payments/receipts", params={"from": "2026-09-01", "to": "2026-10-31", "group_by": "month"})
    assert r.status_code == 200 and [p["bucket"] for p in r.json()["series"]] == ["2026-09", "2026-10"]
    # filters
    r = owner(world).get(f"{API}/org/payments/receipts", params={"from": "2026-10-01", "to": "2026-10-31", "method": "card"})
    assert r.json()["totals"]["received"] == 200_000


def test_balances_and_summary_and_status_debt_tab(world, fixed_today):
    m1, m2, m3 = report_world(world)
    c = owner(world)
    r = c.get(f"{API}/org/payments/balances", params={"kind": "debt"})
    items = {i["membership_id"]: i for i in r.json()["items"]}
    assert str(m2.id) in items and items[str(m2.id)]["debt"] == 70_000 and items[str(m2.id)]["past_debt"] == 0
    assert str(m1.id) not in items
    adv = {i["membership_id"]: i for i in c.get(f"{API}/org/payments/balances", params={"kind": "advance"}).json()["items"]}
    assert adv[str(m1.id)]["advance"] == 21_000
    s = c.get(f"{API}/org/payments/summary").json()
    assert s["accounting_start_date"] == "2026-09-01" and s["debt_total"] >= 70_000 and s["advance_total"] == 21_000
    assert s["memberships_with_debt"] >= 1 and s["students_with_debt"] >= 1
    page = c.get(f"{API}/org/payments/status", params={"bucket": "debt"}).json()
    assert str(m2.id) in {i["membership_id"] for i in page["items"]} and all(i["debt"] > 0 for i in page["items"])
    assert [i["debt"] for i in page["items"]] == sorted((i["debt"] for i in page["items"]), reverse=True)
    sorted_all = c.get(f"{API}/org/payments/status", params={"bucket": "all", "sort": "debt"}).json()["items"]
    assert [i["debt"] for i in sorted_all] == sorted((i["debt"] for i in sorted_all), reverse=True)


# ---- settings ----------------------------------------------------------------------------------------------------


def test_settings_accounting_start_history_and_dry_run(world, fixed_today):
    c = owner(world)
    got = c.get(f"{API}/org/settings/payments").json()
    assert got["accounting_start_date"] is not None
    r = c.put(f"{API}/org/settings/payments", json={"accounting_start_date": "2026-12-01"})
    assert r.status_code == 422 and r.json()["detail"]["code"] == "ACCOUNTING_START_RANGE"
    r = c.put(f"{API}/org/settings/payments", json={"accounting_start_date": "2026-09-01"})
    assert r.status_code == 200 and r.json()["accounting_start_date"] == "2026-09-01" and r.json()["impact"] is None
    # a change of the billing day is appended to the history from today
    r = c.put(f"{API}/org/settings/payments", json={"billing_anchor": "fixed_day", "billing_day": 10})
    assert r.status_code == 200
    world.db.expire_all()
    history = world.org.settings["payments"]["billing_history"]
    assert [h["from"] for h in history] == ["min", TODAY.isoformat()] and history[0]["billing_anchor"] == "payment_date"
    assert history[1]["billing_day"] == 10
    r = c.put(f"{API}/org/settings/payments", json={"billing_day": 12})
    world.db.expire_all()
    assert [h["from"] for h in world.org.settings["payments"]["billing_history"]] == ["min", TODAY.isoformat()]  # same day: replaced
    assert world.org.settings["payments"]["billing_history"][1]["billing_day"] == 12
    # dry run: nothing is written, the impact is returned
    give_group_price(world.db, world.group, 500_000, valid_from=D(2026, 1, 1))
    for i, s in enumerate(world.students):
        m = membership(world, s)
        stay(world, m, D(2026, 8, 1))
    r = c.put(f"{API}/org/settings/payments", params={"dry_run": "true"}, json={"accounting_start_date": "2026-08-01"})
    assert r.status_code == 200
    body = r.json()
    assert body["dry_run"] is True and body["accounting_start_date"] == "2026-09-01"
    assert body["impact"]["memberships_with_new_debt"] == 3 and body["impact"]["debt_delta"] > 0
    world.db.expire_all()
    assert world.org.settings["payments"]["accounting_start_date"] == "2026-09-01"


def test_permissions_unchanged_for_the_new_endpoints(world, fixed_today):
    m = e1_world(world)
    teacher = make_client(world, user=world.teacher)
    admin = make_client(world, user=world.admin)
    for client, expect in ((teacher, 403), (admin, 403)):
        for method, url, body in (
            ("get", f"/org/memberships/{m.id}/payments/preview", None),
            ("get", f"/org/memberships/{m.id}/statement", None),
            ("get", "/org/payments/receipts?from=2026-10-01&to=2026-10-31", None),
            ("get", "/org/payments/balances", None),
            ("post", f"/org/memberships/{m.id}/writeoffs", {"amount": 1, "reason": "abc"}),
            ("post", f"/org/memberships/{m.id}/refunds", {"amount": 1, "method": "cash"}),
        ):
            r = getattr(client, method)(API + url, **({"json": body} if body else {}))
            assert r.status_code == expect, (method, url, r.status_code)
    # an administrator the owner allowed money: all of it works (the owner may switch it on)
    settings_block(world, admin_can_manage=True)
    assert admin.get(f"{API}/org/memberships/{m.id}/payments/preview").status_code == 200
    assert admin.post(f"{API}/org/memberships/{m.id}/writeoffs", json={"amount": 1000, "reason": "льгота"}).status_code == 201
    roles = {a.actor_role for a in world.db.query(GroupAudit).filter(GroupAudit.action == "payment_writeoff", GroupAudit.org_id == world.org.id)}
    assert roles == {"admin"}


def test_the_old_quote_endpoint_is_a_shim_for_one_release(world, fixed_today):
    m = e1_world(world)
    r = owner(world).get(f"{API}/org/memberships/{m.id}/payments/quote", params={"paid_on": TODAY.isoformat(),
                                                                                  "valid_from": "", "valid_to": ""})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["suggested_amount"] == 769_000 == preview(world, m)["suggested_amount"]
    assert body["valid_from"] == "2026-09-15" and body["valid_to"] == "2026-10-14" and body["price"]["amount"] == 500_000
    assert body["mode"] == "calendar" and body["explanation"] == "К оплате 769 000 сум"


def test_a_student_without_any_money_event_stays_none_with_the_debt_visible(world, fixed_today):
    m = e1_world(world)
    world.db.query(StudentPayment).filter_by(membership_id=m.id).delete()
    world.db.commit()
    row = owner(world).get(f"{API}/org/students/{world.students[0].id}/memberships").json()[0]
    assert (row["status"], row["paid_until"], row["debt"], row["past_debt"], row["current_due"]) == ("none", None, 1_000_000, 500_000, 500_000)
    page = owner(world).get(f"{API}/org/payments/status", params={"bucket": "none"}).json()
    assert str(m.id) in {i["membership_id"] for i in page["items"]}


def test_the_soft_double_click_guard_does_not_apply_to_requests_with_a_key(world, fixed_today, monkeypatch):
    import src.auth.payments_service as svc

    monkeypatch.setattr(svc, "SOFT_DUPLICATE_SECONDS", 60)
    m = e1_world(world)
    a = post_payment(world, m, amount=250_000, headers={"Idempotency-Key": str(uuid4())})
    b = post_payment(world, m, amount=250_000, headers={"Idempotency-Key": str(uuid4())})
    assert a["payment"]["id"] != b["payment"]["id"]
    # without a key the same request within the minute is a double click
    c = post_payment(world, m, amount=111_000)
    r = owner(world).post(f"{API}/org/memberships/{m.id}/payments", json={"amount": 111_000, "paid_on": TODAY.isoformat(), "method": "cash"})
    assert r.status_code == 409 and r.json()["detail"]["payment_id"] == c["payment"]["id"]


def test_replaying_a_replace_with_the_same_key_is_the_duplicate_answer_not_already_voided(world, fixed_today):
    m = e1_world(world)
    old = post_payment(world, m, amount=300_000)["payment"]
    key = str(uuid4())
    body = {"amount": 250_000, "paid_on": TODAY.isoformat(), "method": "cash", "reason": "исправление"}
    first = owner(world).post(f"{API}/org/payments/{old['id']}/replace", json=body, headers={"Idempotency-Key": key})
    assert first.status_code == 201
    again = owner(world).post(f"{API}/org/payments/{old['id']}/replace", json=body, headers={"Idempotency-Key": key})
    assert again.status_code == 409 and again.json()["detail"]["code"] == "DUPLICATE_PAYMENT"
    assert again.json()["detail"]["payment_id"] == first.json()["payment"]["id"]
    assert world.db.query(StudentPayment).filter_by(membership_id=m.id, voided_at=None).count() == 2   # seeded + the replacement


def test_accounting_start_bounds_and_the_earliest_money_event(world, fixed_today):
    m = e1_world(world)                                       # a payment of 231 000 on 20.09; start 01.09
    c = owner(world)
    r = c.put(f"{API}/org/settings/payments", json={"accounting_start_date": "2019-12-31"})
    assert r.status_code == 422 and r.json()["detail"]["code"] == "ACCOUNTING_START_RANGE" and "2020" in r.json()["detail"]["message"]
    r = c.put(f"{API}/org/settings/payments", json={"accounting_start_date": "2026-10-01"})
    assert r.status_code == 422 and r.json()["detail"]["message"] == "Есть оплаты до этой даты: 20.09.2026"
    r = c.put(f"{API}/org/settings/payments", params={"dry_run": "true"}, json={"accounting_start_date": "2026-09-21"})
    assert r.status_code == 422                                # the dry run says the same
    # not later than the payment, earlier is fine; a write-off / refund dated earlier counts too
    assert c.put(f"{API}/org/settings/payments", json={"accounting_start_date": "2026-09-20"}).status_code == 200
    assert c.post(f"{API}/org/memberships/{m.id}/writeoffs", json={"amount": 1000, "reason": "льгота", "effective_on": "2026-09-19"}).status_code == 201
    r = c.put(f"{API}/org/settings/payments", json={"accounting_start_date": "2026-09-20"})
    assert r.status_code == 422 and "19.09.2026" in r.json()["detail"]["message"]


def test_the_dry_run_impact_reports_the_advance_too(world, fixed_today):
    settings_block(world, accounting_start_date="2026-10-01")
    give_group_price(world.db, world.group, 500_000, valid_from=D(2026, 1, 1))
    m = membership(world)
    stay(world, m, D(2026, 8, 15))
    seed_payment(world, m, 2_000_000, D(2026, 10, 2))
    r = owner(world).put(f"{API}/org/settings/payments", params={"dry_run": "true"}, json={"accounting_start_date": "2026-08-15"})
    assert r.status_code == 200, r.text
    impact = r.json()["impact"]
    assert impact["advance_delta"] < 0 and impact["refundable_delta"] < 0 and impact["memberships_with_changed_advance"] == 1
    assert impact["debt_delta"] >= 0


def test_receipts_report_position_follows_the_filters_and_transfers_are_split_from_refunds(world, fixed_today):
    m1, m2, m3 = report_world(world)
    c = owner(world)
    whole = c.get(f"{API}/org/payments/receipts", params={"from": "2026-10-01", "to": "2026-10-31"}).json()
    only_g2 = c.get(f"{API}/org/payments/receipts", params={"from": "2026-10-01", "to": "2026-10-31", "group_id": str(m2.group_id)}).json()
    assert only_g2["position"]["debt_total"] == 70_000 and only_g2["position"]["advance_total"] == 0
    assert whole["position"]["advance_total"] == 21_000 and only_g2["totals"]["received"] == 200_000
    assert whole["totals"]["refunded_from_received"] == 324_000 and whole["totals"]["transferred_from_received"] == 0
    assert whole["reconciliation"]["ok"] and only_g2["reconciliation"]["ok"]


def test_the_statement_export_never_makes_a_formula(world, fixed_today):
    import io
    from openpyxl import load_workbook

    m = e1_world(world)
    post_payment(world, m, amount=1000, note="=HYPERLINK(\"http://evil\")")
    r = owner(world).get(f"{API}/org/memberships/{m.id}/statement/export", params={"format": "xlsx"})
    ws = load_workbook(io.BytesIO(r.content)).active
    cells = [c for row in ws.iter_rows() for c in row if isinstance(c.value, str) and "=HYPERLINK" in c.value]
    assert cells and all(c.data_type == "s" for c in cells)             # the note travels into «Подробности» and stays a text
    assert not any(c.data_type == "f" for row in ws.iter_rows() for c in row)
    r = owner(world).get(f"{API}/org/memberships/{m.id}/statement/export", params={"format": "csv"})
    assert not any(part.strip().startswith("=") for line in r.text.splitlines() for part in line.split(";"))


def export_sheet(w, m, **params):
    import io
    from openpyxl import load_workbook
    r = owner(w).get(f"{API}/org/memberships/{m.id}/statement/export", params={"format": "xlsx", **params})
    assert r.status_code == 200, r.text
    return load_workbook(io.BytesIO(r.content)).active


def statement_xlsx_world(w, fix):
    """Calendar months, the student joins on 18.08 (a partial month) with an individual price; three months, payments with
    allocation lines, a write-off, a refund, a cancelled payment; today 21.10."""
    settings_block(w, accounting_start_date="2026-08-01", billing_anchor="fixed_day", billing_day=1)
    give_group_price(w.db, w.group, 500_000, valid_from=D(2026, 1, 1))
    w.group.schedule = []
    w.students[0].display_id = "A-0042"
    w.db.commit()
    m = membership(w)
    stay(w, m, D(2026, 8, 18))
    c = owner(w)
    r = c.put(f"{API}/org/memberships/{m.id}/price-override", json={"amount": 450_000, "valid_from": "2026-08-18", "reason": "скидка"})
    assert r.status_code == 200, r.text
    post_payment(w, m, amount=200_000, paid_on="2026-08-20", note="первый взнос")
    post_payment(w, m, amount=700_000, paid_on="2026-09-05", method="card")
    post_payment(w, m, amount=153_000, paid_on="2026-10-15", method="click")
    fix(D(2026, 10, 19))                                                      # every record on its own day: the order of a day is not fixed
    r = c.post(f"{API}/org/memberships/{m.id}/writeoffs", json={"amount": 50_000, "reason": "болел"})
    assert r.status_code == 201, r.text
    fix(D(2026, 10, 20))
    post_payment(w, m, amount=100_000, method="payme", paid_on="2026-10-20")  # an advance for November
    fix(TODAY)
    r = c.post(f"{API}/org/memberships/{m.id}/refunds", json={"amount": 40_000, "method": "card", "reason": "переехал"})
    assert r.status_code == 201, r.text
    wrong = post_payment(w, m, amount=10_000, paid_on="2026-10-19")["payment"]["id"]
    assert c.post(f"{API}/org/payments/{wrong}/void", json={"reason": "ошибка ввода"}).status_code == 200
    return m


def sheet_grid(ws, first):
    return [[c.value for c in row] for row in ws.iter_rows(min_row=first, max_col=6)]


def test_the_statement_xlsx_reads_like_the_page(world, fixed_today):
    """Golden: the file of a real ledger - a partial month with an individual price, three months, an allocation over three
    periods, a write-off, a refund, a cancelled payment; all months, then a single month with the carry line."""
    from datetime import datetime as dt
    m = statement_xlsx_world(world, fixed_today)
    ws = export_sheet(world, m, include_voided=True)
    head = 5
    assert [ws.cell(r, 1).value for r in (1, 2, 3, 4)] == [
        f"Выписка — Каримова Алия · ID A-0042 · {world.group.name} · Пётр Teacher",
        "По состоянию на 21.10.2026 · Учёт ведётся с 01.08.2026 · Период: Все месяцы",
        "Стоимость обучения 1 103 000 · Оплачено 1 153 000 · Прощено 50 000 · Возвращено 40 000 · Баланс: Аванс 60 000", None]
    assert ws["A1"].font.bold and all(ws.cell(r, 1).data_type == "s" for r in (1, 2, 3))
    assert [ws.cell(head, c).value for c in range(1, 7)] == ["Дата", "Операция", "Стоимость", "Оплачено", "Баланс", "Подробности"]
    assert all(ws.cell(head, c).font.bold for c in range(1, 7)) and ws.freeze_panes == "A6"
    who = "Иван Owner"
    assert sheet_grid(ws, head + 1) == [
        ["Август 2026", None, None, None, None, None],
        [dt(2026, 8, 18), "Август 2026 — неполный месяц (14 дней из 31)", 203_000, None, "Долг 203 000", "цена 450 000 сум в месяц · индивидуальная цена"],
        [dt(2026, 8, 20), f"Оплата № 1 · наличные · внёс {who}", None, 200_000, "Долг 3 000", "комментарий: первый взнос\nза август — 200 000 (осталось 3 000)"],
        ["Сентябрь 2026", None, None, None, None, None],
        [dt(2026, 9, 1), "Сентябрь 2026", 450_000, None, "Долг 453 000", "цена 450 000 сум в месяц · индивидуальная цена"],
        [dt(2026, 9, 5), f"Оплата № 2 · карта · внёс {who}", None, 700_000, "Аванс 247 000",
         "за август — 3 000 (август закрыт)\nза сентябрь — 450 000 (сентябрь закрыт)\nаванс на октябрь — 247 000"],
        ["Октябрь 2026", None, None, None, None, None],
        [dt(2026, 10, 1), "Октябрь 2026", 450_000, None, "Долг 203 000", "цена 450 000 сум в месяц · индивидуальная цена"],
        [dt(2026, 10, 15), f"Оплата № 3 · Click · внёс {who}", None, 153_000, "Долг 50 000", "за октябрь — 153 000 (осталось 50 000)"],
        [dt(2026, 10, 19), f"Долг прощён — причина: болел · {who}", None, 50_000, "0", None],
        [dt(2026, 10, 20), f"Оплата № 4 · Payme · внёс {who}", None, 100_000, "Аванс 100 000", "аванс на ноябрь — 60 000"],
        [dt(2026, 10, 21), f"Возврат аванса — карта · {who}", None, -40_000, "Аванс 60 000", "причина: переехал"],
        [None] * 6,
        ["Следующие периоды (ещё не начались)", None, None, None, None, None],
        ["01.11–30.11: оплачено 60 000 сум из 450 000 сум", None, None, None, None, None],
        [None] * 6,
        ["Отменённые записи", None, None, None, None, None],
        [dt(2026, 10, 19), "Оплата № 5", None, 10_000, None, f"Отменено: ошибка ввода · {who}"],
    ]
    # types and formats: a real date, whole sums with a thousands format, texts stay texts, month bands are merged
    first = head + 2
    assert ws.cell(first, 1).is_date and ws.cell(first, 1).number_format == "DD.MM.YYYY"
    assert isinstance(ws.cell(first, 3).value, int) and ws.cell(first, 3).number_format == "#,##0" and ws.cell(first + 1, 4).number_format == "#,##0"
    assert all(ws.cell(r, c).data_type == "s" for r in range(first, ws.max_row + 1) for c in (2, 5, 6) if ws.cell(r, c).value)
    merged = {str(x) for x in ws.merged_cells.ranges}
    assert {f"A{head + 1}:F{head + 1}", f"A{head + 4}:F{head + 4}", f"A{head + 7}:F{head + 7}"} <= merged
    assert ws.cell(head + 1, 1).font.bold and ws.cell(head + 1, 1).fill.fgColor.rgb not in (None, "00000000")
    flat = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
    assert not any("Начисл" in x for x in flat) and not any(x in flat for x in ("cash", "card", "click", "payme", "bank_transfer"))
    assert ws.cell(ws.max_row, 2).font.strike and ws.cell(ws.max_row, 6).value.startswith("Отменено:")

    # one month: the carry line comes from the WHOLE ledger (the balance before October), the title says the month
    ws = export_sheet(world, m, **{"from": "2026-10-01", "to": "2026-10-31"})
    head = 5
    assert ws.cell(2, 1).value.endswith("Период: Октябрь 2026")
    got = sheet_grid(ws, head + 1)
    assert got[0][0] == "Октябрь 2026" and got[1][0] == "С сентября перенесён аванс 247 000"
    assert [g[1] for g in got[2:]] == ["Октябрь 2026", f"Оплата № 3 · Click · внёс {who}", f"Долг прощён — причина: болел · {who}",
                                         f"Оплата № 4 · Payme · внёс {who}", f"Возврат аванса — карта · {who}"]
    assert ws.max_row == head + 7                                            # a month: no «Следующие периоды», no cancelled block
    assert got[2][4] == "Долг 203 000"                                       # the ledger's balance, not a sum over the cut rows
    assert any(str(x) == f"A{head + 2}:F{head + 2}" for x in ws.merged_cells.ranges)
    # all months: no carry line anywhere
    ws = export_sheet(world, m)
    assert not any(str(ws.cell(r, 1).value).startswith("С ") and "перенесён" in str(ws.cell(r, 1).value) for r in range(1, ws.max_row + 1))

