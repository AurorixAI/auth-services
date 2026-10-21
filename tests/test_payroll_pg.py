"""N13 P3: payroll - the month, the coverage by payments, sheets, revisions, diff, Excel, permissions (PostgreSQL)."""
import calendar
import io
from datetime import date, datetime, timedelta
from decimal import Decimal
from types import SimpleNamespace
from uuid import uuid4

import pytest
from sqlalchemy import event

from src.models.auth import (
    AttendanceExcuse, Group, GroupAttendance, GroupAudit, Organization, PayrollRate, PayrollSheet, PayrollSheetRevision,
    Student, StudentGroup, StudentPayment,
)

from tests.pg_world import (  # noqa: F401  (fixtures)
    business_today, give_group_price, give_price, make_client, make_group, make_member, make_student, make_user,
    needs_pg, pg_session, world,
)

pytestmark = needs_pg

API = "/api/v1/auth"
FEE = 450_000


# ---- the month under test -------------------------------------------------------------------------


def previous_month():
    """The whole previous calendar month (Tashkent): ``.d(n)`` is its n-th day."""
    today = business_today()
    last = today.replace(day=1) - timedelta(days=1)
    first = last.replace(day=1)
    return SimpleNamespace(first=first, last=last, label=first.strftime("%Y-%m"), d=lambda n: first.replace(day=n))


def current_month():
    today = business_today()
    first = today.replace(day=1)
    return SimpleNamespace(
        first=first, last=first.replace(day=calendar.monthrange(first.year, first.month)[1]),
        label=first.strftime("%Y-%m"), d=lambda n: first.replace(day=n),
    )


# ---- helpers: data ---------------------------------------------------------------------------------


def student_group(w, student, group=None):
    w.db.expire_all()
    return w.db.query(StudentGroup).filter_by(student_id=student.id, group_id=(group or w.group).id).one()


def set_fee(w, student, fee, group=None):
    """The price of the student (an individual price from always; ``None`` - no price of his own)."""
    give_price(w.db, student_group(w, student, group), fee)


def lesson(w, day_, marks, *, group=None, teacher=None, roster=None, status="completed", lesson_id=None):
    """A lesson row: ``marks`` = {student: mark}; ``roster`` (students) defaults to the marked ones."""
    group = group or w.group
    members = list(marks) if roster is None else list(roster)
    row = GroupAttendance(
        org_id=w.org.id, group_id=group.id, lesson_date=day_.isoformat(), lesson_id=lesson_id or f"s{day_.isoweekday()}",
        status=status, records={str(s.id): m for s, m in marks.items()},
        roster_snapshot=[{"student_id": str(s.id)} for s in members],
        teacher_id=(teacher or w.teacher).id if teacher is not False else None,
    )
    w.db.add(row)
    w.db.commit()
    return row


def pay_calendar(w, student, first, last, *, group=None, fee=FEE, paid_on=None):
    """PAY-2: the student is paid THROUGH ``last``; the accounting of the centre starts on ``first`` (a lesson before it is
    not charged, so not owed).  The old test said «a payment for first .. last»; money is one account now, a prefix of days."""
    from tests.pg_world import ensure_billing, money_for, payments_block

    membership = student_group(w, student, group)
    if fee is not None:
        give_price(w.db, membership, fee)
    org = w.db.get(type(w.org), w.org.id)
    current = ((org.settings or {}).get("payments") or {}).get("accounting_start_date")
    if current is None or date.fromisoformat(current) > first:
        payments_block(w.db, org, accounting_start_date=first.isoformat())
    ensure_billing(w, group or w.group, price=None)
    on = paid_on or first
    amount = money_for(w, membership.id, through=last, on=on)
    w.db.add(StudentPayment(
        org_id=w.org.id, membership_id=membership.id, student_id=student.id, group_id=membership.group_id,
        amount=amount, paid_on=on, created_by=w.owner.id, method="cash",
    ))
    w.db.commit()


def pay_package(w, student, valid_from, size, *, group=None, fee=FEE, paid_on=None):
    """PAY-2: the centre bills by packages of ``size`` lessons from ``valid_from``; one package of the price ``fee`` is paid."""
    from tests.pg_world import ensure_billing, payments_block

    membership = student_group(w, student, group)
    if fee is not None:
        give_price(w.db, membership, fee)
    ensure_billing(w, group or w.group, price=None)
    payments_block(w.db, w.org, period_mode="lessons", lessons_per_payment=size, accounting_start_date=valid_from.isoformat())
    from src.auth import pricing

    price = pricing.load_books(w.db, [membership.group_id], [membership.id]).price(membership.id, membership.group_id, valid_from)
    payment = StudentPayment(
        org_id=w.org.id, membership_id=membership.id, student_id=student.id, group_id=membership.group_id,
        amount=fee or price.amount or 1, paid_on=paid_on or valid_from, created_by=w.owner.id, method="cash",
    )
    w.db.add(payment)
    w.db.commit()
    return payment


def excuse(w, att, student, status="pending"):
    w.db.add(AttendanceExcuse(
        org_id=w.org.id, attendance_id=att.id, student_id=student.id, reason="Болел", reason_by=w.admin.id,
        reason_at=datetime.utcnow(), status=status,
    ))
    w.db.commit()


# ---- helpers: HTTP ------------------------------------------------------------------------------------


def owner(w):
    return make_client(w, user=w.owner)


def pct(w, value="50", lesson_type="group", teacher=None):
    return {"teacher_id": str((teacher or w.teacher).id), "lesson_type": lesson_type, "percent": value}


EPOCH_MONTH = date(2020, 1, 1)


def set_rate(w, teacher, lesson_type="group", percent="50", valid_from=EPOCH_MONTH, by=None):
    """A row of ``payroll_rates`` (UX-2): the percent of the teacher from a month (default: always)."""
    w.db.add(PayrollRate(
        org_id=w.org.id, teacher_id=teacher.id, lesson_type=lesson_type, percent=Decimal(str(percent)),
        valid_from=valid_from, created_by=(by or w.owner).id,
    ))
    w.db.commit()


def seed_rates(w, percents):
    """The tests of the release before sent ``percents`` with every request; now they are rates: set the ones that
    are not in force yet (from always)."""
    for item in percents or ():
        teacher = w.teacher if item.get("teacher_id") == str(w.teacher.id) else w.teacher2
        w.db.expire_all()
        now = w.db.query(PayrollRate).filter_by(org_id=w.org.id, teacher_id=teacher.id, lesson_type=item["lesson_type"],
                                                valid_from=EPOCH_MONTH).order_by(PayrollRate.created_at.desc()).first()
        if now is None or now.percent != Decimal(str(item["percent"])):
            set_rate(w, teacher, item["lesson_type"], item["percent"])


def calc(w, m, percents=None, teacher_ids=None, expect=200, client=None):
    seed_rates(w, percents if percents is not None else [pct(w)])
    body = {}
    if teacher_ids is not None:
        body["teacher_ids"] = teacher_ids
    r = (client or owner(w)).post(f"{API}/org/payroll/months/{m.label}/calculate", json=body)
    assert r.status_code == expect, r.text
    return r.json()


def preview(w, m, teacher=None, value="50"):
    """The preview sheet of one teacher."""
    teacher = teacher or w.teacher
    body = calc(w, m, [pct(w, value, teacher=teacher)], teacher_ids=[str(teacher.id)])
    (sheet,) = body["sheets"]
    return sheet


def save(w, m, teacher=None, percents=None, expect=201, **extra):
    teacher = teacher or w.teacher
    seed_rates(w, percents or [pct(w, teacher=teacher)])
    body = {"month": m.label, "teacher_id": str(teacher.id), **extra}
    r = owner(w).post(f"{API}/org/payroll/sheets", json=body)
    assert r.status_code == expect, r.text
    return r.json()


def patch(w, sheet, expect=200, **body):
    body.setdefault("revision", sheet["revision"])
    r = owner(w).patch(f"{API}/org/payroll/sheets/{sheet['id']}", json=body)
    assert r.status_code == expect, r.text
    return r.json()


def by_student(sheet):
    return {line["student_id"]: line for line in sheet["lines"]}


def line_of(sheet, student):
    return by_student(sheet)[str(student.id)]


def overview(w, m, client=None):
    r = (client or owner(w)).get(f"{API}/org/payroll/months/{m.label}")
    assert r.status_code == 200, r.text
    return r.json()


def teacher_row(ov, teacher):
    return next(t for t in ov["teachers"] if t["teacher_id"] == str(teacher.id))


def full_month(w, m, students=None, days=range(1, 14), fee=FEE, **kw):
    """Everybody present on ``days`` of the month, everybody paid for the whole month."""
    students = students or w.students
    for n in days:
        lesson(w, m.d(n), {s: "present" for s in students}, **kw)
    for s in students:
        pay_calendar(w, s, m.first, m.last, fee=fee)


# ============================================================================================
# the partner's table, the arithmetic end to end
# ============================================================================================


def test_partner_table_through_the_api(world):
    """450 000, 50 %: 11 / 13 / 5 lessons -> 206 250 / 243 750 / 93 750."""
    m = previous_month()
    s0, s1, s2 = world.students
    for n in range(1, 14):
        roster = [s for s, (a, b) in ((s0, (1, 11)), (s1, (1, 13)), (s2, (9, 13))) if a <= n <= b]
        lesson(world, m.d(n), {s: "present" for s in roster}, roster=roster)
    for s in world.students:
        pay_calendar(world, s, m.first, m.last)

    sheet = preview(world, m)
    assert [line_of(sheet, s)["amount"] for s in (s0, s1, s2)] == [206_250, 243_750, 93_750]
    assert [line_of(sheet, s)["counted"] for s in (s0, s1, s2)] == [11, 13, 5]
    assert sheet["total_amount"] == 543_750 and sheet["counted_total"] == 29 and sheet["unpaid_total"] == 0
    assert sheet["saved"] is False and sheet["id"] is None and sheet["revision"] is None
    assert sheet["month"] == m.label and sheet["period_start"] == m.first.isoformat() and sheet["period_end"] == m.last.isoformat()
    assert sheet["month_open"] is False and sheet["percents"] == {"group": "50"} and sheet["divisor_default"] == 12
    line = line_of(sheet, s2)
    assert line["student_name"] == "Орлова Дана" and line["display_id"] == "N13-003"      # «Фамилия Имя»
    assert line["group_name"] == "Группа" and line["lesson_type"] == "group" and line["lesson_time"] == "10:00–11:00"
    assert (line["monthly_fee"], line["source_fee"], line["fee_overridden"], line["percent"], line["divisor"]) == (FEE, FEE, False, "50", 12)
    assert len(line["marks"]) == 13 and [x["mark"] for x in line["marks"]] == [None] * 8 + [1] * 5
    assert [x["state"] for x in line["marks"]][:2] == ["not_in_group", "not_in_group"] and line["marks"][8]["state"] == "counted"
    (block,) = sheet["groups"]
    assert block["group_name"] == "Группа" and block["line_count"] == 3 and len(block["dates"]) == 13
    assert block["total_amount"] == 543_750 and block["percent"] == "50"

    saved = save(world, m)                                                  # the saved sheet is the same sheet
    assert saved["total_amount"] == 543_750 and saved["revision"] == 1 and saved["saved"] is True and saved["stale"] is False
    got = owner(world).get(f"{API}/org/payroll/sheets/{saved['id']}").json()
    assert [line_of(got, s)["amount"] for s in (s0, s1, s2)] == [206_250, 243_750, 93_750]
    assert got["created_by_name"] == "Иван Owner" and got["teacher_name"] == "Пётр Teacher"


def test_the_divisor_is_a_centre_setting(world):
    m = previous_month()
    full_month(world, m, days=range(1, 7))
    c = owner(world)
    assert c.get(f"{API}/org/payroll/settings").json()["divisor"] == 12
    base = preview(world, m)
    assert line_of(base, world.students[0])["amount"] == FEE * 50 * 6 // 1200
    r = c.put(f"{API}/org/payroll/settings", json={"divisor": 10})
    assert r.status_code == 200 and r.json()["divisor"] == 10 and r.json()["updated_by"] == str(world.owner.id)
    ten = preview(world, m)
    assert line_of(ten, world.students[0])["amount"] == FEE * 50 * 6 // 1000 and ten["divisor_default"] == 10
    assert line_of(ten, world.students[0])["divisor"] == 10
    for bad in (0, 32, -1, "x"):
        assert c.put(f"{API}/org/payroll/settings", json={"divisor": bad}).status_code == 422
    # the payments settings and the payroll settings share one JSON column and do not erase each other
    assert c.put(f"{API}/org/settings/payments", json={"period_mode": "lessons"}).status_code == 200
    assert c.get(f"{API}/org/payroll/settings").json()["divisor"] == 10
    assert c.put(f"{API}/org/payroll/settings", json={"divisor": 31}).status_code == 200
    assert c.get(f"{API}/org/settings/payments").json()["period_mode"] == "lessons"
    # garbage in the column falls back to the default, the page still opens
    world.org.settings = {"payroll": {"divisor": 99}}
    world.db.commit()
    assert c.get(f"{API}/org/payroll/settings").json()["divisor"] == 12


# ============================================================================================
# who teaches, which lessons count
# ============================================================================================


def test_only_completed_lessons_of_the_teacher_by_the_snapshot_count(world):
    m = previous_month()
    s0 = world.students[0]
    pay_calendar(world, s0, m.first, m.last)
    lesson(world, m.d(1), {s0: "present"})
    lesson(world, m.d(2), {s0: "present"}, status="cancelled")
    lesson(world, m.d(3), {s0: "present"}, status="in_progress")
    lesson(world, m.d(4), {s0: "late"}, teacher=world.teacher2)          # taught by somebody else
    lesson(world, m.d(5), {s0: "present"}, teacher=False)                # nobody was the teacher
    outside = lesson(world, m.d(6), {s0: "present"})
    outside.lesson_date = (m.first - timedelta(days=1)).isoformat()      # last day of the month before
    after = lesson(world, m.d(7), {s0: "present"}, lesson_id="late-one")
    after.lesson_date = (m.last + timedelta(days=1)).isoformat()
    world.db.commit()
    sheet = preview(world, m)
    line = line_of(sheet, s0)
    assert line["counted"] == 1 and [x["date"] for x in line["marks"]] == [m.d(1).isoformat()]
    other = preview(world, m, teacher=world.teacher2)
    assert line_of(other, s0)["counted"] == 1 and line_of(other, s0)["marks"][0]["date"] == m.d(4).isoformat()
    ov = overview(world, m)
    assert ov["lessons_without_teacher"] == 1
    assert teacher_row(ov, world.teacher)["lessons_count"] == 1 and teacher_row(ov, world.teacher2)["lessons_count"] == 1


def test_marks_present_late_absent_count_and_a_student_outside_the_roster_does_not(world):
    m = previous_month()
    s0, s1, s2 = world.students
    for s in world.students:
        pay_calendar(world, s, m.first, m.last)
    lesson(world, m.d(1), {s0: "present", s1: "present", s2: "present"}, roster=[s0, s1])     # s2 not in the roster
    lesson(world, m.d(2), {s0: "late", s1: "absent"})
    lesson(world, m.d(3), {s0: "absent"}, roster=[s0, s1])                                    # s1 in the roster, never marked
    sheet = preview(world, m)
    assert (line_of(sheet, s0)["counted"], line_of(sheet, s1)["counted"]) == (3, 2)
    assert [x["state"] for x in line_of(sheet, s1)["marks"]] == ["counted", "counted", "not_in_group"]
    assert str(s2.id) not in by_student(sheet)                  # in the records of a lesson but never on a roster: no line


def test_excuses_confirmed_pending_rejected_and_without_a_decision(world):
    m = previous_month()
    s0, s1, s2 = world.students
    for s in world.students:
        pay_calendar(world, s, m.first, m.last)
    att = lesson(world, m.d(1), {s0: "excused", s1: "excused", s2: "excused"})
    excuse(world, att, s0, "confirmed")
    excuse(world, att, s1, "pending")
    excuse(world, att, s2, "rejected")
    att2 = lesson(world, m.d(2), {s0: "excused", s1: "present", s2: "present"})                 # excused, no row at all
    sheet = preview(world, m)
    l0, l1, l2 = (line_of(sheet, s) for s in (s0, s1, s2))
    assert [x["state"] for x in l0["marks"]] == ["excused_confirmed", "counted"] and [x["mark"] for x in l0["marks"]] == [0, 1]
    assert l0["counted"] == 1 and l0["amount"] == FEE * 50 // 1200 and l0["pending_excused"] == 0
    assert [x["state"] for x in l1["marks"]] == ["excused_pending", "counted"] and l1["counted"] == 2 and l1["pending_excused"] == 1
    assert [w["code"] for w in l1["warnings"]] == ["EXCUSE_PENDING"]
    assert [x["state"] for x in l2["marks"]] == ["counted", "counted"] and l2["counted"] == 2 and l2["pending_excused"] == 0
    assert sheet["pending_excused_total"] == 1 and sheet["warnings"]["pending_excused"] == 1
    ov = overview(world, m)
    assert teacher_row(ov, world.teacher)["warnings"]["pending_excused"] == 1
    # the owner confirms it: the lesson is no longer counted, and the saved sheet notices
    saved = save(world, m)
    world.db.query(AttendanceExcuse).filter_by(attendance_id=att.id, student_id=s1.id).update({"status": "confirmed"})
    world.db.commit()
    again = preview(world, m)
    assert line_of(again, s1)["counted"] == 1 and [x["mark"] for x in line_of(again, s1)["marks"]] == [0, 1]
    assert owner(world).get(f"{API}/org/payroll/sheets/{saved['id']}").json()["stale"] is True


def test_a_student_who_left_keeps_the_lessons_he_attended(world):
    m = previous_month()
    s0, s1, _ = world.students
    full_month(world, m, days=range(1, 6))
    membership = student_group(world, s1)
    membership.is_active = False
    membership.left_at = datetime.utcnow()
    world.db.commit()
    sheet = preview(world, m)
    assert line_of(sheet, s1)["counted"] == 5 and line_of(sheet, s1)["amount"] == FEE * 50 * 5 // 1200


def test_two_teachers_in_one_group_split_the_month_by_the_snapshot(world):
    m = previous_month()
    s0 = world.students[0]
    pay_calendar(world, s0, m.first, m.last)
    for n in range(1, 7):
        lesson(world, m.d(n), {s0: "present"})
    for n in range(7, 13):
        lesson(world, m.d(n), {s0: "present"}, teacher=world.teacher2)
    a, b = preview(world, m), preview(world, m, teacher=world.teacher2, value="40")
    assert line_of(a, s0)["counted"] == 6 and line_of(a, s0)["amount"] == FEE * 50 * 6 // 1200
    assert line_of(b, s0)["counted"] == 6 and line_of(b, s0)["amount"] == FEE * 40 * 6 // 1200
    assert line_of(a, s0)["membership_id"] == line_of(b, s0)["membership_id"]       # one membership, a line per teacher
    ov = overview(world, m)
    assert {t["teacher_id"] for t in ov["teachers"]} == {str(world.teacher.id), str(world.teacher2.id)}


def test_lines_per_group_and_the_kind_of_lessons_set_the_percent(world):
    m = previous_month()
    s0 = world.students[0]
    second = make_group(world.db, world.org, world.teacher, "Парные")
    second.lesson_type = "pair"
    world.db.add(StudentGroup(student_id=s0.id, group_id=second.id, is_active=True, joined_at=datetime.utcnow() - timedelta(days=60)))
    world.db.commit()
    for n in range(1, 5):
        lesson(world, m.d(n), {s0: "present"})
        lesson(world, m.d(n), {s0: "present"}, group=second)
    pay_calendar(world, s0, m.first, m.last)
    pay_calendar(world, s0, m.first, m.last, group=second, fee=600_000)
    # both kinds need a percent: one missing is refused and says which
    set_rate(world, world.teacher, "group", "50")
    r = owner(world).post(f"{API}/org/payroll/months/{m.label}/calculate", json={})
    assert r.status_code == 422 and r.json()["detail"]["code"] == "PERCENT_REQUIRED"
    (missing,) = r.json()["detail"]["missing"]
    assert missing["lesson_type"] == "pair" and missing["teacher_id"] == str(world.teacher.id) and missing["lesson_type_label"] == "Парный"
    body = calc(world, m, [pct(world, "50"), pct(world, "70", "pair")])
    (sheet,) = body["sheets"]
    lines = {x["group_name"]: x for x in sheet["lines"]}
    assert lines["Группа"]["amount"] == FEE * 50 * 4 // 1200 and lines["Парные"]["amount"] == 600_000 * 70 * 4 // 1200
    assert sheet["percents"] == {"group": "50", "pair": "70"} and [b["lesson_type"] for b in sheet["groups"]] == ["group", "pair"]


# ============================================================================================
# coverage by payments
# ============================================================================================


def test_calendar_lessons_outside_the_paid_period_are_unpaid_and_can_be_included(world):
    m = previous_month()
    s0 = world.students[0]
    for n in range(1, 11):
        lesson(world, m.d(n), {s0: "present"})
    pay_calendar(world, s0, m.first, m.d(6))                       # the first six lessons are paid
    sheet = preview(world, m)
    line = line_of(sheet, s0)
    assert (line["counted"], line["unpaid"], line["amount"]) == (6, 4, FEE * 50 * 6 // 1200)
    assert [x["state"] for x in line["marks"]] == ["counted"] * 6 + ["unpaid"] * 4
    assert [w["code"] for w in line["warnings"]] == ["UNPAID_LESSONS"] and sheet["warnings"]["unpaid"] == 4
    assert teacher_row(overview(world, m), world.teacher)["warnings"]["unpaid"] == 4
    saved = save(world, m)
    assert line_of(saved, s0)["include_unpaid"] is False
    ticked = patch(world, saved, line_overrides=[{"membership_id": line["membership_id"], "include_unpaid": True}])
    tl = line_of(ticked, s0)
    assert tl["include_unpaid"] is True and tl["amount"] == FEE * 50 * 10 // 1200 and (tl["counted"], tl["unpaid"]) == (6, 4)
    assert ticked["revision"] == 2 and ticked["total_amount"] == tl["amount"]
    # a save with the tick from the start
    again = save(world, m, revision=2, line_overrides=[{"membership_id": line["membership_id"], "include_unpaid": True}])
    assert line_of(again, s0)["amount"] == FEE * 50 * 10 // 1200 and again["revision"] == 3



def test_no_fee_gives_zero_and_a_warning(world):
    m = previous_month()
    s0, s1, _ = world.students
    for n in range(1, 4):
        lesson(world, m.d(n), {s0: "present", s1: "present"})
    pay_calendar(world, s1, m.first, m.last)                                     # s0 has no price and so no way to pay
    sheet = preview(world, m)
    zero = line_of(sheet, s0)
    assert zero["monthly_fee"] is None and zero["amount"] == 0 and [w["code"] for w in zero["warnings"]] == ["NO_FEE", "UNPAID_LESSONS"]
    assert sheet["warnings"]["no_fee"] == 1 and teacher_row(overview(world, m), world.teacher)["warnings"]["no_fee"] == 1
    # the price is set where prices live (the card of the student), never on the sheet (UX-2)
    set_fee(world, s0, 300_000)
    pay_calendar(world, s0, m.first, m.last, fee=None)                           # now he has a price, so he can pay
    saved = save(world, m)
    fixed = line_of(saved, s0)
    assert fixed["monthly_fee"] == 300_000 and fixed["source_fee"] == 300_000 and fixed["fee_overridden"] is False
    assert fixed["price_source"] == "override" and fixed["amount"] == 300_000 * 50 * 3 // 1200 and "NO_FEE" not in [w["code"] for w in fixed["warnings"]]


# ============================================================================================
# transfers and carried-over payments (the real transfer code)
# ============================================================================================


from tests.test_payment_carry_over_pg import clock, transfer, two  # noqa: E402,F401


def test_transfer_in_the_middle_of_the_month_splits_the_lessons_and_the_payment_follows(two, monkeypatch):
    w, m = two, previous_month()
    s = w.student
    moved = m.d(15)
    give_group_price(w.db, w.group, FEE)
    give_group_price(w.db, w.second, 500_000)                                           # the new group has ITS price
    pay_calendar(w, s, m.first, m.last, fee=None)
    for n in range(10, 15):
        lesson(w, m.d(n), {s: "present"})                                               # first group, teacher 1
    with clock(monkeypatch, moved, 10):
        transfer(w, w.second)
    for n in range(16, 21):
        lesson(w, m.d(n), {s: "present"}, group=w.second, teacher=w.teacher2)           # second group, teacher 2
    first = preview(w, m)
    second = preview(w, m, teacher=w.teacher2)
    a, b = line_of(first, s), line_of(second, s)
    assert a["group_name"] == "Группа" and b["group_name"] == "Вторая" and a["membership_id"] != b["membership_id"]
    assert (a["counted"], a["unpaid"], a["amount"]) == (5, 0, FEE * 50 * 5 // 1200)
    assert (b["counted"], b["unpaid"]) == (5, 0), b["marks"]                           # covered by the carry-over
    # the price did NOT travel with the student (#105): in the new group he is priced by the price of that group
    assert b["monthly_fee"] == 500_000 and b["amount"] == 104_167 and b["price_source"] == "group"        # 500 000 * 50 % * 5 / 12 = 104 166.67
    # the same student is in two lines of two teachers: nothing is counted twice
    assert a["membership_id"] != b["membership_id"] and len(first["lines"]) == 1 and len(second["lines"]) == 1


def test_a_lesson_after_the_transfer_in_the_old_group_is_not_the_old_payments(two, monkeypatch):
    """The old payment is cut at the transfer day; the carry-over covers the new group from there."""
    w, m = two, previous_month()
    s = w.student
    give_group_price(w.db, w.second, FEE)
    pay_calendar(w, s, m.first, m.d(20), fee=FEE)                                    # paid until the 20th
    with clock(monkeypatch, m.d(12), 10):
        transfer(w, w.second)
    lesson(w, m.d(14), {s: "present"}, group=w.second, teacher=w.teacher2)
    lesson(w, m.d(18), {s: "present"}, group=w.second, teacher=w.teacher2)
    lesson(w, m.d(22), {s: "present"}, group=w.second, teacher=w.teacher2)           # after the paid period
    b = line_of(preview(w, m, teacher=w.teacher2), s)
    assert [x["state"] for x in b["marks"]] == ["counted", "counted", "unpaid"]


def test_package_carry_over_keeps_the_size_of_the_original_package(two, monkeypatch):
    w, m = two, previous_month()
    s = w.student
    give_group_price(w.db, w.group, 480_000)
    give_group_price(w.db, w.second, 480_000)
    pay_package(w, s, m.first, 8, fee=None)
    for n in range(2, 5):
        lesson(w, m.d(n), {s: "present"})                                                # 3 lessons in group 1
    with clock(monkeypatch, m.d(10), 10):
        transfer(w, w.second)
    for n in range(12, 17):
        lesson(w, m.d(n), {s: "present"}, group=w.second, teacher=w.teacher2)           # 5 left: exactly the rest
    lesson(w, m.d(18), {s: "present"}, group=w.second, teacher=w.teacher2)              # the 9th lesson: no package
    a = line_of(preview(w, m), s)
    b = line_of(preview(w, m, teacher=w.teacher2), s)
    assert (a["counted"], a["divisor"], a["amount"]) == (3, 8, 3 * 480_000 * 50 // 800)
    assert (b["counted"], b["unpaid"], b["divisor"]) == (5, 1, 8), [x["state"] for x in b["marks"]]
    assert b["amount"] == 5 * 480_000 * 50 // 800
    # the whole package is paid out exactly once: 8 lessons * a lesson price * p %
    assert a["amount"] + b["amount"] == 480_000 * 50 // 100


# ============================================================================================
# the month page
# ============================================================================================


def test_the_month_page_lists_teachers_types_and_suggestions(world):
    m = previous_month()
    s0 = world.students[0]
    pay_calendar(world, s0, m.first, m.last)
    lesson(world, m.d(1), {s0: "present"})
    lesson(world, m.d(2), {s0: "present"}, teacher=world.teacher2)
    ov = overview(world, m)
    assert ov["month"] == m.label and ov["month_open"] is False and ov["divisor"] == 12
    assert [t["teacher_name"] for t in ov["teachers"]] == ["Мария Teacher", "Пётр Teacher"]       # by name
    row = teacher_row(ov, world.teacher)
    assert row["lesson_types"] == [{"lesson_type": "group", "groups": 1, "suggested_percent": None}]
    assert row["sheet"] is None
    assert {k: v for k, v in row["warnings"].items() if k != "open_lessons"} == {"pending_excused": 0, "no_fee": 0, "unpaid": 0, "unpaid_students": 0}
    save(world, m, percents=[pct(world, "55.5")])
    ov = overview(world, m)
    row = teacher_row(ov, world.teacher)
    assert row["sheet"]["revision"] == 1 and row["sheet"]["stale"] is False and row["sheet"]["total_amount"] > 0
    assert row["lesson_types"][0]["suggested_percent"] == "55.5"
    # the suggestion is the last saved percent of THAT teacher for that kind; the other teacher has none
    assert teacher_row(ov, world.teacher2)["lesson_types"][0]["suggested_percent"] is None
    save(world, m, percents=[pct(world, "60")], revision=1)
    assert teacher_row(overview(world, m), world.teacher)["lesson_types"][0]["suggested_percent"] == "60"


def test_unfinished_lessons_are_counted_per_teacher_in_the_month(world):
    """The very list of «Незакрытые уроки»: a started lesson left open and a scheduled one never started."""
    m = previous_month()
    s0 = world.students[0]
    world.group.schedule = [{"id": "mon", "dayOfWeek": 1, "startTime": "10:00", "endTime": "11:00", "format": "offline"}]
    world.db.query(StudentGroup).filter_by(group_id=world.group.id).update({"joined_at": datetime(2020, 1, 1)})
    world.db.commit()
    pay_calendar(world, s0, m.first, m.last)
    lesson(world, m.d(1), {s0: "present"})
    started = lesson(world, m.d(3), {}, status="in_progress")
    started.started_at = datetime.utcnow()
    lesson(world, m.d(4), {}, status="in_progress", teacher=world.teacher2)
    world.db.commit()
    mondays = [n for n in range(1, m.last.day + 1) if m.d(n).isoweekday() == 1]
    never_started = len([n for n in mondays if n not in (1, 3, 4)])           # a Monday with no row at all
    ov = overview(world, m)
    assert teacher_row(ov, world.teacher)["warnings"]["open_lessons"] == 1 + never_started
    # a teacher that only has an unfinished lesson is listed too: nothing to pay yet, but the owner sees what blocks
    only = teacher_row(ov, world.teacher2)
    assert only["lessons_count"] == 0 and only["sheet"] is None and only["lesson_types"] == []
    assert only["warnings"] == {"open_lessons": 1, "pending_excused": 0, "no_fee": 0, "unpaid": 0, "unpaid_students": 0}
    assert preview(world, m)["warnings"]["open_lessons"] == 1 + never_started


def test_months_in_the_future_or_malformed_are_refused_and_the_current_month_is_a_preview(world):
    c = owner(world)
    for bad in ("2026-13", "2026-00", "26-09", "2026-9", "abcd-ef", "1999-01", "2026-09-01"):
        r = c.get(f"{API}/org/payroll/months/{bad}")
        assert r.status_code == 422 and r.json()["detail"]["code"] == "MONTH_INVALID", bad
    future = business_today().replace(day=1) + timedelta(days=40)
    r = c.get(f"{API}/org/payroll/months/{future:%Y-%m}")
    assert r.status_code == 422 and r.json()["detail"]["code"] == "MONTH_INVALID"
    assert c.post(f"{API}/org/payroll/months/2026-99/calculate", json={"percents": []}).status_code == 422
    assert c.get(f"{API}/org/payroll/sheets", params={"month": "nope"}).status_code == 422
    cur = current_month()
    s0 = world.students[0]
    lesson(world, business_today() - timedelta(days=1) if business_today().day > 1 else business_today(), {s0: "present"})
    pay_calendar(world, s0, cur.first, cur.last)
    ov = overview(world, cur)
    assert ov["month_open"] is True
    sheet = preview(world, cur)
    assert sheet["month_open"] is True
    saved = save(world, cur)                                                  # saving an open month is allowed
    assert saved["month_open"] is True and saved["saved"] is True


# ============================================================================================
# sheets: saving, revisions, conflicts, the "data changed" banner
# ============================================================================================


def test_percent_is_required_before_anything_is_calculated_or_saved(world):
    m = previous_month()
    full_month(world, m, days=range(1, 4))
    c = owner(world)
    r = c.post(f"{API}/org/payroll/months/{m.label}/calculate", json={})
    assert r.status_code == 422 and r.json()["detail"]["code"] == "PERCENT_REQUIRED"
    assert r.json()["detail"]["message"] == "Не задана ставка преподавателя: укажите её в «Сотрудниках»"
    r = c.post(f"{API}/org/payroll/sheets", json={"month": m.label, "teacher_id": str(world.teacher.id)})
    assert r.status_code == 422 and r.json()["detail"]["code"] == "PERCENT_REQUIRED"
    assert r.json()["detail"]["message"] == "Не задана ставка преподавателя: укажите её в «Сотрудниках»"
    assert world.db.query(PayrollSheet).filter_by(org_id=world.org.id).count() == 0
    # the percents of the request (the front of the release before) are ignored - they do not make a rate
    r = c.post(f"{API}/org/payroll/months/{m.label}/calculate", json={"percents": [pct(world)]})
    assert r.status_code == 422 and r.json()["detail"]["code"] == "PERCENT_REQUIRED"
    # a teacher without lessons needs no percent and gets an empty sheet
    empty = calc(world, m, [], teacher_ids=[str(world.teacher2.id)])
    assert empty["sheets"][0]["lines"] == [] and empty["sheets"][0]["total_amount"] == 0
    assert c.post(f"{API}/org/payroll/months/{m.label}/calculate",
                  json={"teacher_ids": [str(uuid4())]}).status_code == 404


def test_calculate_writes_nothing(world):
    m = previous_month()
    full_month(world, m, days=range(1, 4))
    calc(world, m)
    world.db.expire_all()
    assert world.db.query(PayrollSheet).filter_by(org_id=world.org.id).count() == 0
    assert owner(world).get(f"{API}/org/payroll/sheets", params={"month": m.label}).json() == []
    assert owner(world).get(f"{API}/org/payroll/months/{m.label}/export.xlsx").status_code == 404      # nothing saved


def test_save_edit_history_and_conflicts(world):
    m = previous_month()
    s0 = world.students[0]
    full_month(world, m, days=range(1, 7))
    first = save(world, m, note="Первый расчёт")
    assert first["revision"] == 1 and first["saved"] and first["note"] == "Первый расчёт"
    sid = first["id"]
    # a second POST without the revision is a conflict (and so is a stale one) - nothing is overwritten
    r = owner(world).post(f"{API}/org/payroll/sheets", json={"month": m.label, "teacher_id": str(world.teacher.id)})
    assert r.status_code == 409 and r.json()["detail"]["code"] == "SHEET_REVISION_CONFLICT" and r.json()["detail"]["current_revision"] == 1
    r = owner(world).post(f"{API}/org/payroll/sheets", json={"month": m.label, "teacher_id": str(world.teacher.id), "revision": 7})
    assert r.status_code == 409
    # a note, a new rate, a tick: new revisions (the rate is read again only by "recalculate")
    line = line_of(first, s0)
    second = patch(world, first, note="Скидка")
    assert second["revision"] == 2 and line_of(second, s0)["amount"] == FEE * 50 * 6 // 1200 and second["total_amount"] == first["total_amount"]
    set_rate(world, world.teacher, "group", "60", valid_from=m.first)
    assert owner(world).get(f"{API}/org/payroll/sheets/{sid}").json()["stale"] is True            # the rate of the month moved
    third = patch(world, second, recalculate=True)
    assert third["revision"] == 3 and third["percents"] == {"group": "60"} and third["stale"] is False
    assert line_of(third, s0)["amount"] == FEE * 60 * 6 // 1200
    back = patch(world, third, line_overrides=[{"membership_id": line["membership_id"], "include_unpaid": True}])
    assert back["revision"] == 4 and back["percents"] == {"group": "60"} and line_of(back, s0)["include_unpaid"] is True
    # the revision protects against somebody else's change: a stale revision is a 409 and changes nothing
    r = owner(world).patch(f"{API}/org/payroll/sheets/{sid}", json={"revision": 2, "note": "x"})
    assert r.status_code == 409 and r.json()["detail"]["code"] == "SHEET_REVISION_CONFLICT" and r.json()["detail"]["current_revision"] == 4
    assert owner(world).get(f"{API}/org/payroll/sheets/{sid}").json()["percents"] == {"group": "60"}
    # history: newest first, with the author, the note and the total of each revision
    history = owner(world).get(f"{API}/org/payroll/sheets/{sid}/revisions").json()
    assert [h["revision"] for h in history] == [4, 3, 2, 1]
    assert history[3]["change_note"] == "Первый расчёт" and history[2]["change_note"] == "Скидка"
    assert history[3]["total_amount"] == first["total_amount"] and history[3]["percents"] == {"group": "50"}
    assert [h["line_count"] for h in history] == [3, 3, 3, 3]
    assert history[0]["changed_by"] == str(world.owner.id) and history[0]["changed_by_name"] == "Иван Owner" and history[0]["changed_at"]
    assert world.db.query(PayrollSheetRevision).filter_by(sheet_id=sid).count() == 4
    # nothing to change is not a revision (the percents of the body are ignored, so they are nothing either)
    r = owner(world).patch(f"{API}/org/payroll/sheets/{sid}", json={"revision": 4})
    assert r.status_code == 422 and r.json()["detail"]["code"] == "NOTHING_TO_CHANGE"
    r = owner(world).patch(f"{API}/org/payroll/sheets/{sid}", json={"revision": 4, "percents": [pct(world, "10")]})
    assert r.status_code == 422 and r.json()["detail"]["code"] == "NOTHING_TO_CHANGE"
    # unknown line; the price is not edited on a sheet
    r = owner(world).patch(f"{API}/org/payroll/sheets/{sid}", json={"revision": 4, "line_overrides": [{"membership_id": str(uuid4()), "include_unpaid": True}]})
    assert r.status_code == 422 and r.json()["detail"]["code"] == "LINE_NOT_FOUND"
    r = owner(world).patch(f"{API}/org/payroll/sheets/{sid}", json={"revision": 4, "line_overrides": [{"membership_id": line["membership_id"], "monthly_fee": 1}]})
    assert r.status_code == 422 and r.json()["detail"]["code"] == "PRICE_NOT_EDITABLE"
    assert owner(world).get(f"{API}/org/payroll/sheets/{sid}").json()["revision"] == 4
    # the list and the single sheet
    listed = owner(world).get(f"{API}/org/payroll/sheets", params={"month": m.label}).json()
    assert [(x["id"], x["revision"], x["month"], x["teacher_name"]) for x in listed] == [(sid, 4, m.label, "Пётр Teacher")]
    assert listed[0]["stale"] is False and listed[0]["total_amount"] == back["total_amount"]
    assert len(owner(world).get(f"{API}/org/payroll/sheets").json()) == 1                     # no month: every month


def test_a_failed_first_save_leaves_nothing_behind_and_a_lost_race_is_a_409(world, monkeypatch):
    from sqlalchemy.exc import IntegrityError

    from src.auth.payroll_service import PayrollService

    m = previous_month()
    full_month(world, m, days=range(1, 3))

    def boom(self, *a, **k):
        raise IntegrityError("insert", {}, Exception("duplicate key value violates unique constraint"))

    set_rate(world, world.teacher)
    with monkeypatch.context() as patched:
        patched.setattr(PayrollService, "_add_revision", boom)
        r = owner(world).post(f"{API}/org/payroll/sheets", json={"month": m.label, "teacher_id": str(world.teacher.id)})
    assert r.status_code == 409 and r.json()["detail"]["code"] == "SHEET_REVISION_CONFLICT"
    world.db.expire_all()
    assert world.db.query(PayrollSheet).filter_by(org_id=world.org.id).count() == 0
    assert save(world, m)["revision"] == 1


def test_revisions_are_unique_per_sheet_and_cascade(world):
    from sqlalchemy.exc import IntegrityError

    m = previous_month()
    full_month(world, m, days=range(1, 3))
    sheet = save(world, m)
    row = world.db.query(PayrollSheetRevision).filter_by(sheet_id=sheet["id"], revision=1).one()
    world.db.add(PayrollSheetRevision(sheet_id=row.sheet_id, revision=1, changed_by=world.owner.id, percent_inputs={}, lines=[], total_amount=0))
    with pytest.raises(IntegrityError):
        world.db.commit()
    world.db.rollback()
    world.db.query(PayrollSheet).filter_by(id=sheet["id"]).delete()
    world.db.commit()
    assert world.db.query(PayrollSheetRevision).filter_by(sheet_id=sheet["id"]).count() == 0


def test_a_saved_sheet_never_changes_by_itself_and_the_diff_says_what_would(world):
    m = previous_month()
    s0, s1, s2 = world.students
    full_month(world, m, days=range(1, 7))
    saved = save(world, m)
    before = {x["student_id"]: x["amount"] for x in saved["lines"]}
    fp = saved["source_fingerprint"]
    assert len(fp) == 64 and owner(world).get(f"{API}/org/payroll/sheets/{saved['id']}").json()["stale"] is False
    clean = owner(world).get(f"{API}/org/payroll/sheets/{saved['id']}/diff").json()
    assert clean["stale"] is False and clean["lines"] == [] and clean["delta_amount"] == 0 and clean["saved_total"] == clean["fresh_total"]

    # the attendance is corrected afterwards: a lesson of s0 is found, a mark of s1 becomes "excused"
    lesson(world, m.d(20), {s0: "present"})
    day2 = world.db.query(GroupAttendance).filter_by(group_id=world.group.id, lesson_date=m.d(2).isoformat()).one()
    day2.records = {str(s0.id): "present", str(s1.id): "excused", str(s2.id): "present"}
    world.db.commit()
    excuse(world, day2, s1, "confirmed")
    stale = owner(world).get(f"{API}/org/payroll/sheets/{saved['id']}").json()
    assert stale["stale"] is True and stale["revision"] == 1
    assert {x["student_id"]: x["amount"] for x in stale["lines"]} == before                  # the sheet itself did not move
    ov_row = teacher_row(overview(world, m), world.teacher)
    assert ov_row["sheet"]["stale"] is True and ov_row["sheet"]["revision"] == 1
    assert owner(world).get(f"{API}/org/payroll/sheets", params={"month": m.label}).json()[0]["stale"] is True
    diff = owner(world).get(f"{API}/org/payroll/sheets/{saved['id']}/diff").json()
    assert diff["stale"] is True and diff["revision"] == 1
    rows = {x["student_name"]: x for x in diff["lines"]}
    assert set(rows) == {"Каримова Алия", "Юсупов Бек"}                                       # Орлова is untouched
    a, b = rows["Каримова Алия"], rows["Юсупов Бек"]
    assert a["status"] == "changed" and a["saved"]["counted"] == 6 and a["fresh"]["counted"] == 7
    assert a["delta_amount"] == FEE * 50 // 1200 and "засчитано" in a["changes"]
    assert b["fresh"]["counted"] == 5 and b["delta_amount"] == -(FEE * 50 // 1200)
    assert diff["delta_amount"] == 0 and diff["fresh_total"] == diff["saved_total"]
    # the sheet is a snapshot: recalculating is an explicit new revision that keeps the owner's ticks
    line = line_of(saved, s2)
    keep = patch(world, saved, line_overrides=[{"membership_id": line["membership_id"], "include_unpaid": True}])
    assert keep["stale"] is True and keep["source_fingerprint"] == fp                          # a tick reads no data
    fresh = patch(world, keep, recalculate=True, note="Пересчёт")
    assert fresh["revision"] == 3 and fresh["stale"] is False and fresh["source_fingerprint"] != fp
    assert line_of(fresh, s0)["counted"] == 7 and line_of(fresh, s1)["counted"] == 5
    assert line_of(fresh, s2)["include_unpaid"] is True                                          # the tick survived
    assert owner(world).get(f"{API}/org/payroll/sheets/{saved['id']}/diff").json()["lines"] == []


def test_the_fingerprint_is_stable_and_follows_the_inputs(world):
    m = previous_month()
    s0 = world.students[0]
    full_month(world, m, days=range(1, 4))
    a = preview(world, m)["source_fingerprint"]
    assert preview(world, m)["source_fingerprint"] == a                       # no state, no clock, no order
    assert preview(world, m, value="77")["source_fingerprint"] == a            # percents are not data
    s0.first_name = "Иная"
    world.db.commit()
    assert preview(world, m)["source_fingerprint"] == a                       # a name is not data
    set_fee(world, s0, FEE + 1)
    b = preview(world, m)["source_fingerprint"]
    assert b != a                                                             # a fee is
    lesson(world, m.d(9), {s0: "present"})
    c = preview(world, m)["source_fingerprint"]
    assert c not in (a, b)
    world.db.query(GroupAttendance).filter_by(lesson_date=m.d(9).isoformat()).update({"teacher_id": world.teacher2.id})
    world.db.commit()
    assert preview(world, m)["source_fingerprint"] == b                       # the lesson moved to another teacher
    assert owner(world).put(f"{API}/org/payroll/settings", json={"divisor": 11}).status_code == 200
    assert preview(world, m)["source_fingerprint"] != b                       # the divisor is data too


def test_a_fee_change_on_the_membership_after_saving_is_noticed_but_the_sheet_keeps_its_fee(world):
    m = previous_month()
    s0 = world.students[0]
    full_month(world, m, days=range(1, 4))
    saved = save(world, m)
    set_fee(world, s0, 500_000)
    now = owner(world).get(f"{API}/org/payroll/sheets/{saved['id']}").json()
    assert now["stale"] is True and line_of(now, s0)["monthly_fee"] == FEE
    (row,) = [x for x in owner(world).get(f"{API}/org/payroll/sheets/{saved['id']}/diff").json()["lines"] if x["student_name"] == "Каримова Алия"]
    assert row["saved"]["monthly_fee"] == FEE and row["fresh"]["monthly_fee"] == 500_000


def test_edge_cases_of_percents_overrides_and_missing_data(world, monkeypatch):
    from src.auth.service import AuthService

    m = previous_month()
    s0, s1, _ = world.students
    full_month(world, m, days=range(1, 4))
    c = owner(world)
    # the percents of a request are ignored whatever they hold - another teacher's, an invalid one - and make no rate
    r = c.post(f"{API}/org/payroll/sheets", json={"month": m.label, "teacher_id": str(world.teacher.id),
                                                  "percents": [pct(world, teacher=world.teacher2)]})
    assert r.status_code == 422 and r.json()["detail"]["code"] == "PERCENT_REQUIRED"
    saved = save(world, m)
    r = c.patch(f"{API}/org/payroll/sheets/{saved['id']}", json={"revision": 1, "percents": [pct(world, "9", teacher=world.teacher2)], "note": "n"})
    assert r.status_code == 200 and r.json()["percents"] == {"group": "50"}
    saved = r.json()
    r = c.post(f"{API}/org/payroll/months/{m.label}/calculate", json={"percents": "junk"})
    assert r.status_code == 200
    # a save for a revision that does not exist yet
    other = previous_month()
    r = c.post(f"{API}/org/payroll/sheets", json={"month": other.label, "teacher_id": str(world.teacher2.id), "revision": 3})
    assert r.status_code == 409 and r.json()["detail"]["code"] == "SHEET_REVISION_CONFLICT"
    # a recalculation that finds a new kind of lessons needs its percent
    second = make_group(world.db, world.org, world.teacher, "Индивидуальная")
    second.lesson_type = "individual"
    world.db.add(StudentGroup(student_id=s0.id, group_id=second.id, is_active=True, joined_at=datetime(2020, 1, 1)))
    world.db.commit()
    lesson(world, m.d(5), {s0: "present"}, group=second)
    r = c.patch(f"{API}/org/payroll/sheets/{saved['id']}", json={"revision": 2, "recalculate": True})
    assert r.status_code == 422 and r.json()["detail"]["code"] == "PERCENT_REQUIRED"
    assert r.json()["detail"]["missing"][0]["lesson_type"] == "individual"
    set_rate(world, world.teacher, "individual", "70")
    grown = patch(world, saved, recalculate=True)
    assert grown["percents"] == {"group": "50", "individual": "70"} and len(grown["groups"]) == 2
    # ... and an override for a line the fresh data does not have
    r = c.patch(f"{API}/org/payroll/sheets/{saved['id']}", json={
        "revision": 3, "recalculate": True, "line_overrides": [{"membership_id": str(uuid4()), "include_unpaid": True}]})
    assert r.status_code == 422 and r.json()["detail"]["code"] == "LINE_NOT_FOUND"
    # a student that no longer exists has nobody to take a fee from: the lesson stays, the line is not made
    lesson(world, m.d(6), {s1: "present"})
    ghost = Student(org_id=world.org.id, first_name="Призрак", last_name="Удалённый", class_level=9)
    world.db.add(ghost)
    world.db.commit()
    lesson(world, m.d(7), {ghost: "present"})
    world.db.delete(ghost)
    world.db.commit()
    both = calc(world, m, [pct(world, "50"), pct(world, "70", "individual")], teacher_ids=[str(world.teacher.id)])
    assert all(x["student_name"] != "Удалённый Призрак" for x in both["sheets"][0]["lines"])
    assert len(both["sheets"][0]["lines"]) == 4                         # three students in the group + the one in the other
    # a hand-edited percent in an old sheet must not break the suggestions
    world.db.query(PayrollSheet).filter_by(id=saved["id"]).update({"percent_inputs": {"group": "abc", "pair": "40"}})
    world.db.commit()
    row = teacher_row(overview(world, m), world.teacher)
    assert {t["lesson_type"]: t["suggested_percent"] for t in row["lesson_types"]}["group"] is None
    # the list of unfinished lessons failing must not take the page down
    def broken(self, *args, **kwargs):
        raise RuntimeError("boom")

    monkeypatch.setattr(AuthService, "get_open_attendance_actions", broken)
    ov = overview(world, m)
    assert teacher_row(ov, world.teacher)["warnings"]["open_lessons"] == 0


# ============================================================================================
# permissions and tenancy
# ============================================================================================


def endpoints(w, sheet_id):
    m = previous_month()
    return [
        ("get", f"{API}/org/payroll/settings", None),
        ("put", f"{API}/org/payroll/settings", {"divisor": 12}),
        ("get", f"{API}/org/payroll/months/{m.label}", None),
        ("post", f"{API}/org/payroll/months/{m.label}/calculate", {"percents": []}),
        ("get", f"{API}/org/payroll/months/{m.label}/export.xlsx", None),
        ("post", f"{API}/org/payroll/sheets", {"month": m.label, "teacher_id": str(w.teacher.id), "percents": []}),
        ("get", f"{API}/org/payroll/sheets", None),
        ("get", f"{API}/org/payroll/sheets/{sheet_id}", None),
        ("patch", f"{API}/org/payroll/sheets/{sheet_id}", {"revision": 1, "note": "x"}),
        ("get", f"{API}/org/payroll/sheets/{sheet_id}/revisions", None),
        ("get", f"{API}/org/payroll/sheets/{sheet_id}/diff", None),
        ("get", f"{API}/org/payroll/sheets/{sheet_id}/export.xlsx", None),
        ("get", f"{API}/org/payroll/rates", None),
        ("put", f"{API}/org/payroll/rates/{w.teacher.id}", {"valid_from": m.label, "percents": {"group": "50"}}),
        ("get", f"{API}/org/payroll/rates/{w.teacher.id}/history", None),
    ]


def call(client, method, url, body):
    return getattr(client, method)(url, json=body) if body is not None else getattr(client, method)(url)


def test_an_admin_is_refused_even_with_the_right_to_payments_and_so_is_a_teacher(world):
    m = previous_month()
    full_month(world, m, days=range(1, 3))
    saved = save(world, m)
    # the owner gives the administrator the right to payments - payroll is not part of it
    assert owner(world).put(f"{API}/org/settings/payments", json={"admin_can_manage": True}).status_code == 200
    assert make_client(world, user=world.admin).get(f"{API}/org/payments/summary").status_code == 200
    for who in (world.admin, world.teacher):
        client = make_client(world, user=who)
        for method, url, body in endpoints(world, saved["id"]):
            r = call(client, method, url, body)
            assert r.status_code == 403, (who.role, method, url, r.text)
    assert owner(world).get(f"{API}/org/payroll/sheets/{saved['id']}").status_code == 200
    # nothing was changed by the refused writes
    world.db.expire_all()
    assert world.db.query(PayrollSheet).filter_by(id=saved["id"]).one().revision == 1


def test_no_organisation_and_a_platform_account_are_refused(world):
    for client in (
        make_client(world, user=world.owner, org=None),                                  # a token without an organisation
        make_client(world, user=world.owner, org_role="admin"),                           # the live membership says admin
        make_client(world, user=world.owner, role="superadmin", org_role="superadmin"),   # superadmin passes the role gate
        make_client(world, user=world.owner, role="superadmin", org=None),                # ... even without an organisation
    ):
        for method, url, body in endpoints(world, uuid4()):
            r = call(client, method, url, body)
            assert r.status_code == 403, (method, url, r.status_code)


def test_another_centre_sees_nothing_and_gets_404(world):
    m = previous_month()
    full_month(world, m, days=range(1, 4))
    saved = save(world, m)
    foreign_owner = make_user(world.db, "owner")
    make_member(world.db, world.other_org, foreign_owner, "owner")
    world.db.commit()
    foreign = make_client(world, user=foreign_owner, org="other")
    base = f"{API}/org/payroll/sheets/{saved['id']}"
    for suffix in ("", "/revisions", "/diff", "/export.xlsx"):
        assert foreign.get(base + suffix).status_code == 404, suffix
    assert foreign.patch(base, json={"revision": 1, "note": "x"}).status_code == 404
    assert foreign.get(f"{API}/org/payroll/sheets", params={"month": m.label}).json() == []
    theirs = foreign.get(f"{API}/org/payroll/months/{m.label}").json()["teachers"]
    assert str(world.teacher.id) not in {t["teacher_id"] for t in theirs} and all(t["sheet"] is None for t in theirs)
    assert foreign.get(f"{API}/org/payroll/months/{m.label}/export.xlsx").status_code == 404
    # our teacher is not theirs, ours cannot be saved or previewed from there
    r = foreign.post(f"{API}/org/payroll/sheets", json={"month": m.label, "teacher_id": str(world.teacher.id), "percents": []})
    assert r.status_code == 404
    assert foreign.post(f"{API}/org/payroll/months/{m.label}/calculate", json={"percents": [], "teacher_ids": [str(world.teacher.id)]}).status_code == 404
    # their settings are theirs
    assert foreign.put(f"{API}/org/payroll/settings", json={"divisor": 20}).status_code == 200
    assert owner(world).get(f"{API}/org/payroll/settings").json()["divisor"] == 12
    # and the lessons of another centre never reach our sheet
    foreign_student = world.foreign_student
    lesson(world, m.d(1), {foreign_student: "present"}, group=world.foreign_group, teacher=world.foreign_admin)
    assert preview(world, m)["lines"] and all(x["student_id"] != str(foreign_student.id) for x in preview(world, m)["lines"])


# ============================================================================================
# Excel
# ============================================================================================


def workbook(response):
    from openpyxl import load_workbook

    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert response.headers["content-disposition"].startswith("attachment;") and "no-store" in response.headers["cache-control"]
    return load_workbook(io.BytesIO(response.content))


def cells(ws):
    return [c for row in ws.iter_rows() for c in row if c.value is not None]


def test_excel_follows_the_partner_table_with_red_zeros_and_a_total(world):
    m = previous_month()
    s0, s1, s2 = world.students
    att = lesson(world, m.d(1), {s0: "excused", s1: "excused", s2: "present"})
    excuse(world, att, s0, "confirmed")
    excuse(world, att, s1, "pending")
    for n in (2, 3, 4):
        lesson(world, m.d(n), {s0: "present", s1: "present", s2: "present"})
    for s in world.students:
        pay_calendar(world, s, m.first, m.d(3))                                   # the 4th lesson is unpaid
    saved = save(world, m)
    r = owner(world).get(f"{API}/org/payroll/sheets/{saved['id']}/export.xlsx")
    wb = workbook(r)
    assert wb.sheetnames == ["Пётр Teacher"]
    ws = wb.active
    values = {c.coordinate: c.value for c in cells(ws)}
    assert "Пётр Teacher" in values["A1"] and m.label.split("-")[0] in values["A1"]
    assert "Редакция 1" in values["A2"]
    header = [c.value for c in ws[5] if c.value is not None]
    assert header[:6] == ["Группа", "Тип занятий", "Время", "Ученик", "Оплата в месяц", "%"]
    assert header[6:10] == [f"{m.d(n):%d.%m}" for n in (1, 2, 3, 4)] and header[10:] == ["Засчитано", "Не оплачено", "Включено", "Начислено"]
    rows = {ws.cell(row=r_, column=4).value: r_ for r_ in range(6, ws.max_row + 1) if ws.cell(row=r_, column=4).value}
    karimova, yusupov, orlova = (rows[n] for n in ("Каримова Алия", "Юсупов Бек", "Орлова Дана"))
    # zero (a confirmed excuse) is red; the pending excuse is yellow text; the unpaid lesson is grey
    zero = ws.cell(row=karimova, column=7)
    assert zero.value == 0 and zero.font.color.rgb.endswith("C00000")
    pending = ws.cell(row=yusupov, column=7)
    assert pending.value == "Ув.?" and pending.fill.start_color.rgb.endswith("FFF2B3")
    unpaid = ws.cell(row=orlova, column=10)
    assert unpaid.value == 1 and unpaid.fill.start_color.rgb.endswith("E5E7EB")
    assert ws.cell(row=orlova, column=7).value == 1
    assert ws.cell(row=orlova, column=5).value == FEE and ws.cell(row=orlova, column=6).value == 50
    # only the zero in the date columns is red
    reds = [c for c in cells(ws) if c.value == 0 and c.font.color is not None and c.font.color.rgb.endswith("C00000")]
    assert len(reds) == 1
    # the totals: counted / unpaid / amount; the sheet's total
    total_row = max(r_ for r_ in range(6, ws.max_row + 1) if ws.cell(row=r_, column=1).value == "ИТОГО")
    assert ws.cell(row=total_row, column=14).value == saved["total_amount"]
    assert ws.cell(row=total_row, column=11).value == saved["counted_total"] and ws.cell(row=total_row, column=12).value == 3
    assert ws.cell(row=orlova, column=14).value == line_of(saved, s2)["amount"]
    assert any(ws.cell(row=r_, column=1).value == "Итого: Группа" for r_ in range(6, ws.max_row + 1))


def test_excel_of_a_month_has_a_sheet_per_saved_teacher(world):
    m = previous_month()
    s0 = world.students[0]
    pay_calendar(world, s0, m.first, m.last)
    lesson(world, m.d(1), {s0: "present"})
    lesson(world, m.d(2), {s0: "present"}, teacher=world.teacher2)
    save(world, m)
    only = workbook(owner(world).get(f"{API}/org/payroll/months/{m.label}/export.xlsx"))
    assert only.sheetnames == ["Пётр Teacher"]
    save(world, m, teacher=world.teacher2, percents=[pct(world, "40", teacher=world.teacher2)])
    both = owner(world).get(f"{API}/org/payroll/months/{m.label}/export.xlsx")
    wb = workbook(both)
    assert wb.sheetnames == ["Мария Teacher", "Пётр Teacher"]                      # by the name of the teacher
    assert "filename*=UTF-8''" in both.headers["content-disposition"] and ".xlsx" in both.headers["content-disposition"]


def test_excel_cells_from_names_are_never_formulas(world):
    m = previous_month()
    s0, s1, s2 = world.students
    s0.last_name, s0.first_name = "=HYPERLINK(\"http://evil\",\"x\")", "+1"
    s1.last_name, s1.first_name = "@SUM(1+1)", "-2"
    world.group.name = "=cmd|' /C calc'!A0"
    world.teacher.first_name, world.teacher.last_name = "=1+1", "x"
    world.db.commit()
    for n in (1, 2):
        lesson(world, m.d(n), {s0: "present", s1: "present", s2: "present"})
    for s in world.students:
        pay_calendar(world, s, m.first, m.last)
    saved = save(world, m, note="=2+2")
    ws = workbook(owner(world).get(f"{API}/org/payroll/sheets/{saved['id']}/export.xlsx")).active
    for cell in cells(ws):
        assert cell.data_type != "f", (cell.coordinate, cell.value)
    dangerous = [c for c in cells(ws) if isinstance(c.value, str) and c.value.startswith("'") and c.value[1:2] in "=+-@"]
    assert {c.value[:8] for c in dangerous} >= {"'=HYPERL", "'@SUM(1+", "'=cmd|' "}, [c.value for c in dangerous]
    assert all(c.data_type == "s" for c in dangerous)
    assert ws.title == "=1+1 x"                                                   # a worksheet name is not evaluated
    note = [c.value for c in cells(ws) if isinstance(c.value, str) and "2+2" in c.value]
    assert note == ["Примечание: =2+2"]


def test_excel_helpers():
    from src.auth import payroll_export as px

    assert px.safe_text("=A1") == "'=A1" and px.safe_text("+1") == "'+1" and px.safe_text("-1") == "'-1"
    assert px.safe_text("a\x00b\x1fc\x0bd") == "abcd" and px.safe_text("\x00=1") == "'=1"            # control characters are dropped
    assert px.safe_text("@x") == "'@x" and px.safe_text("\t=1") == "' =1" and px.safe_text(" +1") == "' +1" and px.safe_text("Алия") == "Алия" and px.safe_text(None) == ""
    assert px.sheet_title("a/b:c*d?e[f]g\\h", []) == "a b c d e f g h"
    assert px.sheet_title("x" * 50, []) == "x" * 31
    assert px.sheet_title("Иван", ["иван"]) == "Иван (2)" and px.sheet_title("'quoted'", []) == "quoted"
    assert px.sheet_title(None, []) == "Преподаватель" and px.sheet_title("", ["Преподаватель"]) == "Преподаватель (2)"
    assert px.export_filename("2026-09", 'Иван "Эй"/\\') == "Зарплата 2026-09 Иван Эй.xlsx"


# ============================================================================================
# lessons_type of a group
# ============================================================================================


def test_group_lesson_type_is_stored_audited_and_returned(world):
    c = owner(world)
    listed = {g["id"]: g for g in c.get(f"{API}/org/groups").json()}
    assert listed[str(world.group.id)]["lesson_type"] == "group"                    # every old group is "group"
    r = c.post(f"{API}/org/groups", json={"name": "Парная", "class_level": 9, "teacher_id": str(world.teacher.id), "lesson_type": "pair"})
    assert r.status_code == 201 and r.json()["lesson_type"] == "pair"
    gid = r.json()["id"]
    assert c.post(f"{API}/org/groups", json={"name": "Без типа", "class_level": 9}).json()["lesson_type"] == "group"
    assert c.post(f"{API}/org/groups", json={"name": "Пустой тип", "lesson_type": None}).json()["lesson_type"] == "group"
    assert c.post(f"{API}/org/groups", json={"name": "Плохая", "lesson_type": "huge"}).status_code == 422
    created = world.db.query(GroupAudit).filter_by(group_id=gid, action="created").one()
    assert created.snapshot["lesson_type"] == "pair"
    # an administrator changes it: the change is in the audit log
    admin = make_client(world, user=world.admin)
    r = admin.put(f"{API}/org/groups/{gid}", json={"lesson_type": "individual"})
    assert r.status_code == 200 and r.json()["lesson_type"] == "individual"
    audit = world.db.query(GroupAudit).filter_by(group_id=gid, action="updated").one()
    assert audit.snapshot == {"lesson_type": ["pair", "individual"]} and audit.actor_id == world.admin.id
    assert admin.put(f"{API}/org/groups/{gid}", json={"lesson_type": "individual"}).status_code == 200
    assert world.db.query(GroupAudit).filter_by(group_id=gid, action="updated").count() == 1    # no change, no row
    assert admin.put(f"{API}/org/groups/{gid}", json={"lesson_type": "nope"}).status_code == 422
    assert admin.put(f"{API}/org/groups/{gid}", json={"name": "Другое имя"}).json()["lesson_type"] == "individual"   # untouched
    assert c.get(f"{API}/org/groups/{gid}").json()["lesson_type"] == "individual"


def test_a_teacher_cannot_pick_the_kind_of_lessons_that_sets_the_own_percent(world):
    teacher = make_client(world, user=world.teacher)
    r = teacher.post(f"{API}/org/groups", json={"name": "Моя", "class_level": 9, "lesson_type": "individual"})
    assert r.status_code == 403 and "администратор" in r.json()["detail"]
    r = teacher.post(f"{API}/org/groups", json={"name": "Моя", "class_level": 9})
    assert r.status_code == 201 and r.json()["lesson_type"] == "group"
    r = teacher.post(f"{API}/org/groups", json={"name": "Моя 2", "class_level": 9, "lesson_type": "group"})
    assert r.status_code == 201
    assert teacher.put(f"{API}/org/groups/{r.json()['id']}", json={"lesson_type": "pair"}).status_code == 403   # edit is for managers


# ============================================================================================
# queries do not grow with the size of the centre
# ============================================================================================


def count_queries(w, fn):
    from src.auth.service import AuthService  # noqa: F401

    engine = w.db.get_bind()
    n = {"queries": 0}

    def before(conn, cursor, statement, parameters, context, executemany):
        n["queries"] += 1

    event.listen(engine, "before_cursor_execute", before)
    try:
        result = fn()
    finally:
        event.remove(engine, "before_cursor_execute", before)
    return n["queries"], result


def build_centre(w, teachers, students_per, lessons=8):
    """``teachers`` teachers, each with a group of ``students_per`` students and ``lessons`` lessons."""
    m = previous_month()
    percents = []
    for t in range(teachers):
        teacher = make_user(w.db, "teacher", f"Учитель{t:02d}")
        make_member(w.db, w.org, teacher, "teacher")
        group = make_group(w.db, w.org, teacher, f"Группа {t:02d}")
        group.lesson_type = "group" if t % 2 else "pair"
        pupils = []
        for i in range(students_per):
            pupils.append(make_student(w.db, w.org, group, f"Имя{i:03d}", f"Фам{t:02d}{i:03d}"))
        w.db.flush()
        memberships = {x.student_id: x for x in w.db.query(StudentGroup).filter(StudentGroup.group_id == group.id)}
        for n in range(1, lessons + 1):
            w.db.add(GroupAttendance(
                org_id=w.org.id, group_id=group.id, lesson_date=m.d(n).isoformat(), lesson_id=f"s{m.d(n).isoweekday()}",
                status="completed", records={str(s.id): ("present" if (i + n) % 5 else "absent") for i, s in enumerate(pupils)},
                roster_snapshot=[{"student_id": str(s.id)} for s in pupils], teacher_id=teacher.id))
        for i, s in enumerate(pupils):
            membership = memberships[s.id]
            give_price(w.db, membership, 300_000 + 1000 * i)
            if i % 3 == 0:
                w.db.add(StudentPayment(org_id=w.org.id, membership_id=membership.id, student_id=s.id, group_id=group.id,
                                        amount=1, paid_on=m.first, valid_from=m.first, valid_to=m.last, created_by=w.owner.id))
            elif i % 3 == 1:
                w.db.add(StudentPayment(org_id=w.org.id, membership_id=membership.id, student_id=s.id, group_id=group.id,
                                        amount=1, paid_on=m.first, valid_from=m.first, lessons_count=6, created_by=w.owner.id))
        percents.append({"teacher_id": str(teacher.id), "lesson_type": group.lesson_type, "percent": "40"})
        w.db.add(PayrollRate(org_id=w.org.id, teacher_id=teacher.id, lesson_type=group.lesson_type, percent=Decimal("40"),
                             valid_from=EPOCH_MONTH, created_by=w.owner.id))
    w.db.commit()
    return m, percents


def test_a_month_of_ten_teachers_and_two_hundred_students_costs_a_fixed_number_of_queries(world):
    small_m, small_percents = build_centre(world, 2, 3)
    small_client = owner(world)
    q_small, small = count_queries(world, lambda: small_client.post(
        f"{API}/org/payroll/months/{small_m.label}/calculate", json={}))
    assert small.status_code == 200 and len(small.json()["sheets"]) == 2
    q_overview_small, _ = count_queries(world, lambda: small_client.get(f"{API}/org/payroll/months/{small_m.label}"))

    # a second centre, five times the teachers and ~ 30 times the students
    org = Organization(name="Big centre", status="active")
    world.db.add(org)
    world.db.flush()
    owner2 = make_user(world.db, "owner")
    make_member(world.db, org, owner2, "owner")
    world.db.commit()
    wb = SimpleNamespace(**{**vars(world), "org": org, "owner": owner2})
    m, percents = build_centre(wb, 10, 20)
    client = make_client(wb, user=owner2)
    q_big, resp = count_queries(world, lambda: client.post(f"{API}/org/payroll/months/{m.label}/calculate", json={}))
    assert resp.status_code == 200, resp.text
    sheets = resp.json()["sheets"]
    assert len(sheets) == 10 and sum(len(s["lines"]) for s in sheets) == 200
    assert all(s["total_amount"] > 0 for s in sheets)
    q_overview_big, ov = count_queries(world, lambda: client.get(f"{API}/org/payroll/months/{m.label}"))
    assert ov.status_code == 200 and len(ov.json()["teachers"]) == 10
    # bounded: the same number of queries for 2 teachers x 3 students and for 10 teachers x 20 students
    assert q_big == q_small and q_overview_big == q_overview_small, (q_small, q_big, q_overview_small, q_overview_big)
    assert q_big <= 45 and q_overview_big <= 45, (q_big, q_overview_big)
    # saving one sheet and opening it is bounded as well
    teacher_id = sheets[0]["teacher_id"]
    q_save, saved = count_queries(world, lambda: client.post(f"{API}/org/payroll/sheets", json={
        "month": m.label, "teacher_id": teacher_id}))
    assert saved.status_code == 201 and q_save <= 60, q_save


# ============================================================================================
# races (separate sessions, real row locks)
# ============================================================================================


def threaded_client(world):
    """A client whose every request gets its OWN session, like production (the shared session of ``make_client`` would serialise)."""
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from sqlalchemy.orm import sessionmaker

    from src.auth.dependencies import get_current_active_user_payload
    from src.auth.payroll_router import router
    from src.core.database import get_db

    factory = sessionmaker(world.db.get_bind())

    def own_session():
        db = factory()
        try:
            yield db
        finally:
            db.close()

    payload = {"sub": str(world.owner.id), "role": "owner", "type": "access", "org_id": str(world.org.id), "org_role": "owner"}
    app = FastAPI()
    app.include_router(router, prefix="/api/v1")
    app.dependency_overrides[get_db] = own_session
    app.dependency_overrides[get_current_active_user_payload] = lambda: payload
    return lambda: TestClient(app)


def test_simultaneous_edits_with_the_same_revision_have_exactly_one_winner(world):
    from concurrent.futures import ThreadPoolExecutor

    m = previous_month()
    full_month(world, m, days=range(1, 5))
    saved = save(world, m)
    new_client = threaded_client(world)
    url = f"{API}/org/payroll/sheets/{saved['id']}"

    def edit(n):
        return new_client().patch(url, json={"revision": 1, "note": f"правка {n}"})

    with ThreadPoolExecutor(6) as pool:
        results = list(pool.map(edit, range(6)))
    codes = sorted(r.status_code for r in results)
    assert codes == [200] + [409] * 5, [(r.status_code, r.text[:200]) for r in results]
    assert all(r.json()["detail"]["code"] == "SHEET_REVISION_CONFLICT" for r in results if r.status_code == 409)
    world.db.expire_all()
    sheet = world.db.query(PayrollSheet).filter_by(id=saved["id"]).one()
    revisions = world.db.query(PayrollSheetRevision).filter_by(sheet_id=sheet.id).order_by(PayrollSheetRevision.revision).all()
    assert sheet.revision == 2 and [r.revision for r in revisions] == [1, 2]            # no gap, no duplicate
    winner = next(r for r in results if r.status_code == 200).json()
    assert revisions[1].change_note == winner["note"] and sheet.percent_inputs == winner["percents"]


def test_simultaneous_first_saves_make_one_sheet(world):
    from concurrent.futures import ThreadPoolExecutor

    m = previous_month()
    full_month(world, m, days=range(1, 4))
    set_rate(world, world.teacher)
    new_client = threaded_client(world)
    body = {"month": m.label, "teacher_id": str(world.teacher.id)}

    def create(_):
        return new_client().post(f"{API}/org/payroll/sheets", json=body)

    with ThreadPoolExecutor(6) as pool:
        results = list(pool.map(create, range(6)))
    assert sorted(r.status_code for r in results) == [201] + [409] * 5, [(r.status_code, r.text[:200]) for r in results]
    assert all(r.json()["detail"]["current_revision"] == 1 for r in results if r.status_code == 409)   # the real revision
    world.db.expire_all()
    assert world.db.query(PayrollSheet).filter_by(org_id=world.org.id).count() == 1
    assert world.db.query(PayrollSheetRevision).count() >= 1
    sheet = world.db.query(PayrollSheet).filter_by(org_id=world.org.id).one()
    assert sheet.revision == 1 and world.db.query(PayrollSheetRevision).filter_by(sheet_id=sheet.id).count() == 1


# ============================================================================================
# R3 fixes: sanitising, save what was seen, local time, boundary order, carried decisions
# ============================================================================================


def test_text_xml_cannot_hold_never_breaks_the_workbook(world):
    from src.auth import payroll_export as px

    m = previous_month()
    s0, s1, s2 = world.students
    world.teacher.first_name, world.teacher.last_name = "Пё\x0bтр\ufffe", "Te\tach\ner\uffff"
    s0.first_name = "А\x0cл\x01ия\x1f"
    world.group.name = "Груп\x0cпа\r\nA"
    world.db.commit()
    for n in (1, 2):
        lesson(world, m.d(n), {s0: "present", s1: "present", s2: "present"})
    for s in world.students:
        pay_calendar(world, s, m.first, m.last)
    save(world, m, note="При\x0bмечание\n2")
    wb = workbook(owner(world).get(f"{API}/org/payroll/months/{m.label}/export.xlsx"))
    assert wb.sheetnames == ["Пётр Te ach er"]                                   # \x0b, U+FFFE / U+FFFF dropped, tab / LF -> space
    ws = wb.active
    texts = [c.value for c in cells(ws) if isinstance(c.value, str)]
    assert "Группа  A" in texts                                                    # \x0c dropped, CR LF -> two spaces
    assert not any(ch in t for t in texts for ch in "\x00\x0b\x0c\x01\x1f\ufffe\uffff")
    assert any(t == "Примечание: Примечание 2" for t in texts)
    # the single-sheet export and the file name survive the same names
    one = owner(world).get(f"{API}/org/payroll/sheets/{world.db.query(PayrollSheet).filter_by(org_id=world.org.id).one().id}/export.xlsx")
    assert workbook(one).sheetnames == ["Пётр Te ach er"] and "\x0b" not in one.headers["content-disposition"]
    assert px.sheet_title("a\x0bb\tc\nd\ufffe\uffff", []) == "ab c d"
    assert px.clean_text("x\r\ny\x00") == "x  y"


def test_save_what_the_owner_saw(world):
    m = previous_month()
    s0 = world.students[0]
    full_month(world, m, days=range(1, 5))
    seen = preview(world, m)
    fp = seen["source_fingerprint"]
    lesson(world, m.d(20), {s0: "present"})                                      # the data moves after the preview
    r = owner(world).post(f"{API}/org/payroll/sheets", json={
        "month": m.label, "teacher_id": str(world.teacher.id), "expected_fingerprint": fp})
    assert r.status_code == 409
    detail = r.json()["detail"]
    assert detail["code"] == "SHEET_DATA_CHANGED" and detail["message"] == "Данные изменились после расчёта — пересчитайте и проверьте суммы"
    fresh = preview(world, m)["source_fingerprint"]
    assert detail["source_fingerprint"] == fresh and fresh != fp
    world.db.expire_all()
    assert world.db.query(PayrollSheet).filter_by(org_id=world.org.id).count() == 0           # nothing was saved
    saved = save(world, m, expected_fingerprint=fresh)                                          # the new preview can be saved
    assert saved["revision"] == 1 and saved["source_fingerprint"] == fresh
    assert save(world, m, revision=1)["revision"] == 2                                           # without it: no check

    # PATCH with recalculate: the same protection
    lesson(world, m.d(21), {s0: "present"})
    r = owner(world).patch(f"{API}/org/payroll/sheets/{saved['id']}", json={"revision": 2, "recalculate": True, "expected_fingerprint": fresh})
    assert r.status_code == 409 and r.json()["detail"]["code"] == "SHEET_DATA_CHANGED"
    newest = r.json()["detail"]["source_fingerprint"]
    assert owner(world).get(f"{API}/org/payroll/sheets/{saved['id']}").json()["revision"] == 2
    ok = owner(world).patch(f"{API}/org/payroll/sheets/{saved['id']}", json={"revision": 2, "recalculate": True, "expected_fingerprint": newest})
    assert ok.status_code == 200 and ok.json()["revision"] == 3 and ok.json()["source_fingerprint"] == newest
    # without recalculation nothing is read again, so there is nothing to compare
    lesson(world, m.d(22), {s0: "present"})
    plain = owner(world).patch(f"{API}/org/payroll/sheets/{saved['id']}", json={"revision": 3, "note": "x", "expected_fingerprint": "0" * 64})
    assert plain.status_code == 200 and plain.json()["revision"] == 4


def test_the_export_shows_the_save_time_in_tashkent(world):
    m = previous_month()
    full_month(world, m, days=range(1, 3))
    saved = save(world, m)
    world.db.query(PayrollSheet).filter_by(id=saved["id"]).update({"updated_at": datetime(2026, 9, 30, 20, 30)})   # UTC
    world.db.commit()
    ws = workbook(owner(world).get(f"{API}/org/payroll/sheets/{saved['id']}/export.xlsx")).active
    assert ws["A2"].value == "Редакция 1 · сохранено 01.10.2026 01:30"                         # UTC+5, next day
    from src.auth import payroll_export as px
    assert px.local_time("2026-01-15T22:05:59") == "16.01.2026 03:05" and px.local_time(None) == ""


def test_an_unpaid_lesson_is_priced_by_the_latest_package_not_by_a_payment_that_covers_nothing(world):
    m = previous_month()
    s0 = world.students[0]
    for n in range(1, 11):
        lesson(world, m.d(n), {s0: "present"})
    pay_package(world, s0, m.first, 8, fee=480_000)
    membership = student_group(world, s0)
    # a calendar payment, paid LATER, that moved away before its period began: it covers nothing at all
    world.db.add(StudentPayment(
        org_id=world.org.id, membership_id=membership.id, student_id=s0.id, group_id=world.group.id, amount=1,
        paid_on=m.d(25), valid_from=m.d(20), valid_to=m.last, created_by=world.owner.id,
        carried_out_at=datetime(m.first.year, m.first.month, 5, 10, 0)))
    world.db.commit()
    line = line_of(preview(world, m), s0)
    assert (line["counted"], line["unpaid"]) == (8, 2)
    assert [x["divisor"] for x in line["marks"]] == [8] * 10                      # not 12: the package is the latest payment


def test_two_lessons_of_one_day_at_the_package_boundary_go_by_start_time(world):
    """The package ends in the middle of a day: the EARLIER lesson (teacher 1) is paid, the later one (teacher 2) is not."""
    m = previous_month()
    s0 = world.students[0]
    day3 = m.d(3)
    world.group.schedule = [
        {"id": "z-early", "dayOfWeek": day3.isoweekday(), "startTime": "10:00", "endTime": "11:00", "format": "offline"},
        {"id": "a-late", "dayOfWeek": day3.isoweekday(), "startTime": "15:00", "endTime": "16:00", "format": "offline"},
    ]
    world.db.commit()
    lesson(world, m.d(1), {s0: "present"}, lesson_id="x1")
    lesson(world, m.d(2), {s0: "present"}, lesson_id="x2")
    lesson(world, day3, {s0: "present"}, lesson_id="z-early")                    # 10:00, teacher 1 ("z" sorts AFTER "a")
    lesson(world, day3, {s0: "present"}, lesson_id="a-late", teacher=world.teacher2)   # 15:00, teacher 2
    pay_package(world, s0, m.first, 3)
    first, second = line_of(preview(world, m), s0), line_of(preview(world, m, teacher=world.teacher2), s0)
    assert [x["state"] for x in first["marks"]] == ["counted", "counted", "counted"] and first["unpaid"] == 0
    assert [x["state"] for x in second["marks"]] == ["unpaid"] and second["counted"] == 0
    page = {r["student_id"]: r for r in owner(world).get(f"{API}/org/payments/status", params={"limit": 50}).json()["items"]}
    assert (page[str(s0.id)]["lessons_used"], page[str(s0.id)]["lessons_total"], page[str(s0.id)]["lessons_left"]) == (4, 3, 0)
    # the same day with the times the other way round: now teacher 2's lesson is the earlier one
    world.group.schedule = [
        {"id": "z-early", "dayOfWeek": day3.isoweekday(), "startTime": "16:00", "endTime": "17:00", "format": "offline"},
        {"id": "a-late", "dayOfWeek": day3.isoweekday(), "startTime": "09:00", "endTime": "10:00", "format": "offline"},
    ]
    world.db.commit()
    first, second = line_of(preview(world, m), s0), line_of(preview(world, m, teacher=world.teacher2), s0)
    assert [x["state"] for x in first["marks"]] == ["counted", "counted", "unpaid"]
    assert [x["state"] for x in second["marks"]] == ["counted"]


def test_a_new_revision_by_post_keeps_the_decisions_of_the_last_one(world):
    m = previous_month()
    s0, s1, _ = world.students
    for n in range(1, 7):
        lesson(world, m.d(n), {s0: "present", s1: "present"})
    pay_calendar(world, s0, m.first, m.d(4))                                       # two unpaid lessons each
    pay_calendar(world, s1, m.first, m.d(4))
    base = preview(world, m)
    mid0, mid1 = line_of(base, s0)["membership_id"], line_of(base, s1)["membership_id"]
    first = save(world, m, line_overrides=[{"membership_id": mid0, "include_unpaid": True},
                                           {"membership_id": mid1, "include_unpaid": True}])
    assert line_of(first, s0)["include_unpaid"] and line_of(first, s1)["include_unpaid"]
    # a plain re-save (the rate of the month changed): the ticks stay
    second = save(world, m, revision=1, percents=[pct(world, "60")])
    a, b = line_of(second, s0), line_of(second, s1)
    assert (a["monthly_fee"], a["fee_overridden"], a["include_unpaid"]) == (FEE, False, True)
    assert (b["monthly_fee"], b["include_unpaid"]) == (FEE, True)
    assert a["amount"] == FEE * 60 * 6 // 1200 and b["amount"] == FEE * 60 * 6 // 1200
    # what the request says wins; a field it leaves out stays
    third = save(world, m, revision=2, line_overrides=[{"membership_id": mid1, "include_unpaid": False}])
    a, b = line_of(third, s0), line_of(third, s1)
    assert (a["monthly_fee"], a["include_unpaid"]) == (FEE, True) and (b["monthly_fee"], b["include_unpaid"]) == (FEE, False)
    fourth = save(world, m, revision=3, line_overrides=[{"membership_id": mid1, "include_unpaid": True}])
    assert line_of(fourth, s0)["include_unpaid"] is True and line_of(fourth, s1)["include_unpaid"] is True


def test_revisions_and_the_month_export_use_a_fixed_number_of_queries(world):
    m = previous_month()
    s0 = world.students[0]
    pay_calendar(world, s0, m.first, m.last)
    lesson(world, m.d(1), {s0: "present"})
    lesson(world, m.d(2), {s0: "present"}, teacher=world.teacher2)
    a = save(world, m, note="первая")
    save(world, m, teacher=world.teacher2, percents=[pct(world, "40", teacher=world.teacher2)], note="вторая")
    patch(world, a, note="ещё")
    client = owner(world)
    statements = []

    def grab(conn, cursor, statement, parameters, context, executemany):
        statements.append(statement)

    engine = world.db.get_bind()
    event.listen(engine, "before_cursor_execute", grab)
    try:
        revisions = client.get(f"{API}/org/payroll/sheets/{a['id']}/revisions")
        statements_revisions, statements[:] = list(statements), []
        exported = client.get(f"{API}/org/payroll/months/{m.label}/export.xlsx")
    finally:
        event.remove(engine, "before_cursor_execute", grab)
    assert revisions.status_code == 200 and [h["line_count"] for h in revisions.json()] == [1, 1]
    assert any("jsonb_array_length" in q for q in statements_revisions)                       # counted by the database
    assert not any("payroll_sheet_revisions.lines AS" in q for q in statements_revisions)      # the lines are not loaded
    assert workbook(exported).sheetnames == ["Мария Teacher", "Пётр Teacher"]
    assert len([q for q in statements if "FROM payroll_sheet_revisions" in q]) == 1            # the notes of all sheets: one query


def test_the_diff_carries_the_fingerprint_to_save_exactly_what_it_showed(world):
    m = previous_month()
    s0 = world.students[0]
    full_month(world, m, days=range(1, 5))
    saved = save(world, m)
    url = f"{API}/org/payroll/sheets/{saved['id']}"
    clean = owner(world).get(f"{url}/diff").json()
    assert clean["source_fingerprint"] == saved["source_fingerprint"] and clean["stale"] is False
    lesson(world, m.d(20), {s0: "present"})                                       # the data moves
    diff = owner(world).get(f"{url}/diff").json()
    assert diff["stale"] is True and diff["source_fingerprint"] != saved["source_fingerprint"]
    assert diff["source_fingerprint"] == preview(world, m)["source_fingerprint"]      # the fresh one
    lesson(world, m.d(21), {s0: "present"})                                       # ... and moves again after the owner looked
    r = owner(world).patch(url, json={"revision": 1, "recalculate": True, "expected_fingerprint": diff["source_fingerprint"]})
    assert r.status_code == 409 and r.json()["detail"]["code"] == "SHEET_DATA_CHANGED"
    assert owner(world).get(url).json()["revision"] == 1
    again = owner(world).get(f"{url}/diff").json()
    ok = owner(world).patch(url, json={"revision": 1, "recalculate": True, "expected_fingerprint": again["source_fingerprint"]})
    assert ok.status_code == 200 and ok.json()["revision"] == 2 and ok.json()["source_fingerprint"] == again["source_fingerprint"]
    assert ok.json()["stale"] is False and line_of(ok.json(), s0)["counted"] == 6


# ---- PAY-2: the lessons of the account ------------------------------------------------------------------------------


def pay_money(w, student, amount, paid_on, *, group=None):
    """Plain money on the account (the sum, the date)."""
    membership = student_group(w, student, group)
    w.db.add(StudentPayment(org_id=w.org.id, membership_id=membership.id, student_id=student.id, group_id=membership.group_id,
                            amount=amount, paid_on=paid_on, created_by=w.owner.id, method="cash"))
    w.db.commit()


def test_money_pays_the_oldest_days_first_and_a_voided_payment_covers_nothing(world):
    m = previous_month()
    s0, s1, _ = world.students
    for n in range(1, 4):
        lesson(world, m.d(n), {s0: "present", s1: "present"})
    pay_calendar(world, s0, m.d(2), m.last)
    pay_calendar(world, s1, m.first, m.last)
    world.db.query(StudentPayment).filter_by(student_id=s1.id).update(
        {"voided_at": datetime.utcnow(), "voided_by": world.owner.id, "void_reason": "ошибка"})
    world.db.commit()
    sheet = preview(world, m)
    # the account is paid THROUGH the end of the month: the oldest days are paid first, nothing is "before the payment"
    assert [x["state"] for x in line_of(sheet, s0)["marks"]] == ["counted"] * 3
    assert [x["state"] for x in line_of(sheet, s1)["marks"]] == ["unpaid"] * 3 and line_of(sheet, s1)["amount"] == 0


def test_lessons_before_the_start_of_the_accounting_are_not_owed(world):
    """Grace: a day before the accounting began (or before the student's first charged day) is neither charged nor unpaid."""
    from tests.pg_world import ensure_billing, payments_block

    m = previous_month()
    s0 = world.students[0]
    for n in (1, 2, 20):
        lesson(world, m.d(n), {s0: "present"})
    ensure_billing(world, price=FEE)
    payments_block(world.db, world.org, accounting_start_date=m.d(10).isoformat())
    assert [x["state"] for x in line_of(preview(world, m), s0)["marks"]] == ["counted", "counted", "unpaid"]


def test_a_package_gives_its_size_as_the_divisor(world):
    """8 lessons of an 8-lesson package: the teacher gets exactly p % of the price, not 8/12."""
    m = previous_month()
    s0, s1, _ = world.students
    for n in range(1, 9):
        lesson(world, m.d(n), {s0: "present", s1: "present"})
    pay_package(world, s0, m.first, 8, fee=480_000)
    pay_package(world, s1, m.first, 8, fee=480_000)
    sheet = preview(world, m)
    for student in (s0, s1):
        line = line_of(sheet, student)
        assert (line["counted"], line["divisor"], line["amount"]) == (8, 8, 240_000)
        assert {x["divisor"] for x in line["marks"]} == {8}


def test_package_runs_out_the_rest_is_unpaid_and_the_status_page_agrees(world):
    """Parity: the lessons the payments page says are covered are exactly the ones payroll counts as paid."""
    m = previous_month()
    s0, s1, _ = world.students
    for n in range(1, 9):
        lesson(world, m.d(n), {s0: "present", s1: "present"})
    held = lesson(world, m.d(9), {s0: "excused", s1: "present"})
    excuse(world, held, s0, "confirmed")                                        # not counted, not consumed
    pay_package(world, s0, m.first, 5)
    pay_package(world, s1, m.first, 5)
    pay_money(world, s1, FEE * 4 // 5, m.d(7))                                    # four more lessons of the same account
    page = {r["student_id"]: r for r in owner(world).get(f"{API}/org/payments/status", params={"limit": 50}).json()["items"]}
    sheet = preview(world, m)
    l0, l1 = line_of(sheet, s0), line_of(sheet, s1)
    assert (l0["counted"], l0["unpaid"]) == (5, 3)
    assert [x["state"] for x in l0["marks"]] == ["counted"] * 5 + ["unpaid"] * 3 + ["excused_confirmed"]
    assert [x["state"] for x in l1["marks"]] == ["counted"] * 8 + ["counted"]    # 5 + 4 paid lessons cover all nine
    assert [x["divisor"] for x in l0["marks"][:8]] == [5] * 8 and l0["divisor"] == 5
    for line, student in ((l0, s0), (l1, s1)):
        row = page[str(student.id)]
        assert row["lessons_total"] - 0 >= min(row["lessons_used"], row["lessons_total"])
        assert line["counted"] == min(row["lessons_used"], row["lessons_total"]), student.first_name   # the page's number = ours
        assert line["counted"] + line["unpaid"] == row["lessons_used"]
    assert l0["amount"] == FEE * 50 * 5 // (100 * 5)                              # five lessons of a 5-lesson package


def test_an_included_unpaid_lesson_of_a_package_student_is_priced_by_his_package(world):
    m = previous_month()
    s0 = world.students[0]
    for n in range(1, 11):
        lesson(world, m.d(n), {s0: "present"})
    pay_package(world, s0, m.first, 8, fee=480_000)
    sheet = preview(world, m)
    line = line_of(sheet, s0)
    assert (line["counted"], line["unpaid"], line["amount"]) == (8, 2, 240_000)
    assert [x["divisor"] for x in line["marks"]] == [8] * 10 and line["divisor"] == 8
    saved = save(world, m)
    ticked = patch(world, saved, line_overrides=[{"membership_id": line["membership_id"], "include_unpaid": True}])
    assert line_of(ticked, s0)["amount"] == 10 * 480_000 * 50 // 800                    # ten lessons at 1/8 of the price


@pytest.mark.parametrize("seed", range(10))
def test_covered_lessons_equal_what_the_status_page_covers_on_random_data(world, seed):
    """Random lessons, marks, excuse decisions and a random number of paid lessons per student: payroll counts as paid
    exactly the lessons the payments page says are covered."""
    import random

    from src.auth import payment_status as ps
    from tests.pg_world import ensure_billing, money_for, payments_block

    rnd = random.Random(seed)
    m = previous_month()
    students = world.students
    size = rnd.randint(1, 6)
    for n in sorted(rnd.sample(range(1, 26), rnd.randint(6, 16))):
        roster = [s for s in students if rnd.random() > 0.15]
        marks = {s: rnd.choice(["present", "present", "late", "absent", "excused"]) for s in roster}
        att = lesson(world, m.d(n), marks, roster=roster)
        for s, mark in marks.items():
            decision = rnd.choice(["confirmed", "pending", "rejected", None]) if mark == "excused" else None
            if decision:
                excuse(world, att, s, decision)
    ensure_billing(world, price=FEE)
    payments_block(world.db, world.org, period_mode="lessons", lessons_per_payment=size, accounting_start_date=m.first.isoformat())
    for s in students:
        membership = student_group(world, s)
        amount = money_for(world, membership.id, lessons=rnd.randint(1, 12))
        pay_money(world, s, amount, m.first)
    world.db.expire_all()
    rows = {r["student_id"]: r for r in owner(world).get(f"{API}/org/payments/status", params={"limit": 50}).json()["items"]}
    held = world.db.query(GroupAttendance).filter_by(group_id=world.group.id, status="completed").all()
    excuses = {(e.attendance_id, str(e.student_id)): e.status for e in world.db.query(AttendanceExcuse).filter_by(org_id=world.org.id)}
    sheet = preview(world, m)
    for s in students:
        sid = str(s.id)
        counts = sum(
            1 for a in held
            if sid in {i["student_id"] for i in a.roster_snapshot}
            and ps.lesson_counts(mark=a.records.get(sid), excuse_status=excuses.get((a.id, sid)))
        )
        line = by_student(sheet).get(sid)
        counted, unpaid = (line["counted"], line["unpaid"]) if line else (0, 0)
        assert counted + unpaid == counts, (seed, sid)                                  # the same lessons are counted
        assert rows[sid]["lessons_used"] == counts, (seed, sid)
        assert counted == min(counts, rows[sid]["lessons_total"]), (seed, sid, counted, rows[sid])   # and paid as on the page
        if line:
            assert all(x["divisor"] for x in line["marks"] if x["state"] in ("counted", "excused_pending", "unpaid"))



def test_lessons_before_the_first_price_do_not_stop_the_packages_of_the_account(world):
    """Lessons mode, a price only from the 15th: the lessons of the 10th and 12th are skipped (no price, nothing owed and
    nothing paid), the package starts with the lesson of the 15th and the paid money counts it and the next one."""
    from tests.pg_world import ensure_billing, payments_block

    m = previous_month()
    s0 = world.students[0]
    for n in (10, 12, 15, 17):
        lesson(world, m.d(n), {s0: "present"})
    ensure_billing(world, price=None)
    give_group_price(world.db, world.group, 400_000, m.d(15))
    payments_block(world.db, world.org, period_mode="lessons", lessons_per_payment=2, accounting_start_date=m.first.isoformat())
    pay_money(world, s0, 400_000, m.d(15))
    line = line_of(preview(world, m), s0)
    assert [x["state"] for x in line["marks"]] == ["unpaid", "unpaid", "counted", "counted"]
    assert [x["price"] for x in line["marks"]] == [None, None, 400_000, 400_000] and line["no_price_lessons"] == 2
    assert line["amount"] == 2 * 400_000 * 50 // 200
    page = {r["student_id"]: r for r in owner(world).get(f"{API}/org/payments/status", params={"limit": 50}).json()["items"]}
    assert page[str(s0.id)]["lessons_used"] == 2 and page[str(s0.id)]["price_missing"] is True
